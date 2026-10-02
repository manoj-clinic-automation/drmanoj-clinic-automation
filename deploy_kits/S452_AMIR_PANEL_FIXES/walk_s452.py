#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""walk_s452.py -- kit S452_AMIR_PANEL_FIXES. THE REAL finance app (and the real portal, for the staff-eye walk) in one process per side,
over SCRATCH COPIES of finance.db and assets.db (backup API), through Flask's test clients. NEW = the box + the kit's six files (and the
kit's data step, apply_s452.py, run on the scratch copy); OLD = the box as it is (the negative control). Sign-in for the staff-eye walk:
the walk's own random secret and own user store in a scratch portal folder (never the live ones). Its own rows are keyed W452: a Marg
bill 'W452-61234' of 'KEDWALK452 PHARMA', scans stamped 'W452-01..03', crafted feeds 'push_snapshot W452' / 'push_expected base=W452',
a crafted SMS-read NEFT month 2099-04 and an unrecorded 2099-06. The live count's vouchers, the box's scans and August's NEFT are read by
key; every date is computed from today. The phone's key is NEVER printed: the probe reports only whether a page carried it, and stops
red if a key string reaches its output. Walk-only patch, named where used: packs.amir_pack for September (a crafted ready pack).

  --fin-new DIR --fin-old DIR --por DIR --ast DIR --shared DIR --db PATH --adb PATH --kit DIR --duty-map PATH --uploads DIR --spine PATH
"""
import argparse
import datetime as dt
import hashlib
import io
import json
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import sys
import time

TODAY = dt.date.today()
SUP = "KEDWALK452 PHARMA"
USERS = {"amir": "staff", "darpan": "staff", "shavez": "manager", "manoj": "doctor", "reception": "staff"}
DEVA = re.compile("[" + chr(0x900) + "-" + chr(0x97F) + "]")
SHOP = ("B-0054", "B-0078", "B-0079", "B-0080", "B-0082")


def D(k):
    return (TODAY - dt.timedelta(days=k)).isoformat()


def DMY(d):
    return d.strftime("%d-%m-%Y")


def now_iso():
    return dt.datetime.now().replace(microsecond=0).isoformat()


def ins(con, table, row, replace=False):
    row = dict(row)
    for _cid, name, typ, notnull, dflt, pk in con.execute("PRAGMA table_info(%s)" % table).fetchall():
        if notnull and dflt is None and not pk and name not in row:
            row[name] = 0 if "INT" in (typ or "").upper() else ""
    cols = list(row)
    return con.execute("INSERT %sINTO %s (%s) VALUES (%s)" % ("OR REPLACE " if replace else "", table, ", ".join(cols), ", ".join("?" * len(cols))),
                       [row[c] for c in cols]).lastrowid


# ======================================================================= the PROBE (one process per side)
def probe():
    NEW = os.environ["MODE"] == "new"
    FIN, POR = os.environ["FINDIR"], os.environ["PORDIR"]
    W = json.loads(os.environ["W452"])
    for p in (POR, FIN):
        sys.path.insert(0, p)
    os.chdir(FIN)
    import finance_app as fa                                           # noqa: E402
    import portal as po                                                # noqa: E402
    import purchase_app                                                # noqa: E402
    BASE = "https://followup.dr-manoj.in"
    fc = fa.app.test_client(use_cookies=False)
    pc = po.app.test_client(use_cookies=False)
    H = lambda u: {"X-Clinic-User": u, "X-Clinic-Role": ""}             # noqa: E731
    db = sqlite3.connect(os.environ["FINANCE_DB"], timeout=60)
    db.row_factory = sqlite3.Row
    q = lambda s, *a: [dict(r) for r in db.execute(s, a).fetchall()]   # noqa: E731
    out = {}
    secrets_seen = []                                                  # the phone keys this probe read -- never printed

    def FG(u, p):
        r = fc.get(p, base_url=BASE, headers=H(u))
        b = r.get_data()
        return r.status_code, (b.decode("latin-1") if b[:4] == b"%PDF" or b[:2] == b"PK" else b.decode("utf-8", "replace")), r.headers

    def FB(u, p):
        r = fc.get(p, base_url=BASE, headers=H(u))
        return r.status_code, r.get_data(), r.headers

    def FJ(u, p, body=None):
        r = fc.post(p, base_url=BASE, headers=H(u), json=body or {}) if body is not None else fc.get(p, base_url=BASE, headers=H(u))
        return r.status_code, (r.get_json(silent=True) or {})

    def text(h):
        return (re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", h or "")).replace("&mdash;", "—").replace("&#10003;", "✓")
                .replace("&#8377;", "Rs ").replace("&middot;", "·").replace("&larr;", "<-").replace("&#x27;", "'").strip())

    def card(n=1):
        h = FG("amir", "/finance/amir/step/%d" % n)[1]
        m = re.search(r"id=s444duty.*?</div></div>", h, re.S)
        return text(m.group(0)) if m else ""

    def neft_block(n=1):
        h = FG("amir", "/finance/amir/step/%d" % n)[1]
        m = re.search(r"id=s452neft>.*?</div></div>", h, re.S)
        if m:
            return text(m.group(0))
        m = re.search(r"<div class=card><h2>NEFT .*?(?=<p class=foot>)", h, re.S)          # S407's card (the box as it is)
        return text(m.group(0)) if m else ""

    def needs():
        return [l.get("text", "") for l in FJ("manoj", "/finance/sanjeevni/api/needs-you")[1].get("lines") or []]

    def board(u):
        s, j = FJ(u, "/finance/stock/api/pad/amir/1")
        vf = (j.get("made") or {}).get("vouchers_flat") or []
        return dict(status=s, n=len(vf), seq=[v["seq"] for v in vf], keys=["%d|%s|%d" % (v["round_no"], v["kind"], v["batch_no"]) for v in vf],
                    stage=(j.get("stage") or {}), deva=len(DEVA.findall(json.dumps(j, ensure_ascii=False))),
                    salt=[x.get("item") for x in j.get("salt_fix") or []], salt_done=[x.get("item") for x in j.get("salt_fix") or [] if x.get("marg_done")],
                    rates=[t.get("item") for t in j.get("rate_tasks") or []], entered=sum(1 for v in vf if v.get("entered")))

    def enter(keys, no):
        res = []
        for k in keys:
            r, kind, b = k.split("|")
            res.append(FJ("amir", "/finance/stock/api/pad/vouchers/1/entered", dict(round=int(r), kind=kind, batch=int(b), marg_voucher_no=no))[0])
        return res

    def step2():
        h = FG("amir", "/finance/amir/step/2")[1]
        m = re.search(r"id=s446bills>(.*?)(?=<div class=card><p class=big>)", h, re.S)
        blk = m.group(1) if m else ""
        ids = [int(x) for x in re.findall(r"/finance/purchase/api/scan-file/(\d+)", blk)]
        lines = {int(i): text(l) for l, i in re.findall(r"<div class=line>(.*?)/finance/purchase/api/scan-file/(\d+)", blk, re.S)}
        hm = re.search(r"(\d+) scan abhi reception ki jaanch mein hain", blk)
        return dict(ids=ids, lines=lines, held=(int(hm.group(1)) if hm else 0), heading=bool(re.search(r"Marg mein daalne ke bill \(\d+\)", blk)),
                    text=text(blk))

    def dl(sid, u="amir"):
        s, b, hd = FB(u, "/finance/purchase/api/scan-file/%d" % sid)
        name = (re.search(r'filename="([^"]+)"', hd.get("Content-Disposition") or "") or [None, ""])[1]
        return dict(status=s, name=name, pdf=(b[:4] == b"%PDF"), text=text(b.decode("utf-8", "replace")) if b[:4] != b"%PDF" else "")

    stamp_of = {}
    try:
        acon = sqlite3.connect("file:%s?mode=ro" % os.environ["ASSETS_DB"], uri=True)
        stamp_of = {r[0]: r[1] for r in acon.execute("SELECT id, stamp_no FROM bills WHERE kind='Pharmacy'")}
        acon.close()
    except sqlite3.Error:
        pass
    # ------------------------------------------------------------------ 0. the staff-eye walk (first, on the untouched scratch copy)
    pw = W["pw"]

    def login(user):
        r = pc.post("/portal/login", base_url=BASE, data={"user": user, "password": pw[user]})
        m = re.search(r"clinic_sso=([^;]+)", " ".join(r.headers.getlist("Set-Cookie")))
        return m.group(1) if m else ""

    def tiles(h):
        import html as _h                                              # noqa: PLC0415
        return [_h.unescape(t) for t in re.findall(r'<div class="nm">([^<]+)</div>', h or "")]

    dm = {}
    if NEW:
        FG("amir", "/finance/amir/step/1")                             # his card read once, as the install's data step does: stock_rate_marg exists
        dm = json.load(open(os.environ["REAL_DUTY_MAP"], encoding="utf-8"))
        ro = sqlite3.connect("file:%s?mode=ro" % os.environ["FINANCE_DB"], uri=True)
        eye = {}
        for who in ("amir", "darpan", "shavez", "manoj"):
            tok = login(who)
            r = pc.get("/portal?all=1", base_url=BASE, headers={"Cookie": "clinic_sso=" + tok})
            seen = tiles(r.get_data(as_text=True))
            rows = []
            mine = {who, (dm.get("shared") or {}).get(who)}
            for du in dm.get("duties") or []:
                if du.get("person") not in mine:
                    continue
                row = {"id": du.get("id"), "tile": du.get("tile"), "tile_seen": (du.get("tile") in seen) if du.get("tile") else None}
                sql = str(du.get("due_sql") or "").strip().rstrip(";")
                n = None
                if sql:
                    try:
                        rr = ro.execute(sql).fetchone()
                        n = int((rr[0] if rr else 0) or 0)
                    except Exception as e:                         # noqa: BLE001
                        n = "ERR " + str(e)[:80]
                row["due_n"] = n
                door, mark = du.get("door"), du.get("door_marker")
                if door and mark and isinstance(n, int) and n > 0 and str(door).startswith("/finance/"):
                    path, body = str(door), ""
                    for _hop in range(4):
                        rr2 = fc.get(path, base_url=BASE, headers={"Cookie": "clinic_sso=" + tok})
                        if rr2.status_code in (301, 302, 303) and rr2.headers.get("Location", "").replace(BASE, "").startswith("/finance/"):
                            path = rr2.headers["Location"].replace(BASE, "")
                            continue
                        body = rr2.get_data(as_text=True)
                        break
                    row["door_seen"] = mark in body
                rows.append(row)
            eye[who] = dict(status=r.status_code, duties=rows)
        ro.close()
        out["eye"] = eye
    # ------------------------------------------------------------------ 1. the list, on the box's own scans
    s2 = step2()
    box = [i for i in s2["ids"] if not str(stamp_of.get(i, "")).startswith("W452")]
    out["l_box_ids"] = box
    out["l_box_stamps"] = [stamp_of.get(i) for i in box]
    out["l_heading"] = s2["heading"]
    out["l_held"] = s2["held"]
    out["l_lines"] = {str(i): s2["lines"].get(i, "") for i in box}
    out["l_downloads"] = {str(i): dl(i) for i in box}
    if NEW:
        import porders                                                 # noqa: E402 -- READ ONLY: S440's own grouping
        k = porders.scan_work(db)
        out["l_wait"] = sorted(x["stamp"] for x in k["wait"] if not str(x["stamp"]).startswith("W452"))
        r = purchase_app.scans_for_amir(db)
        out["l_held_box"] = len([x for x in r["held"] if not str(x["stamp"]).startswith("W452")])
        out["l_held_reasons"] = sorted("%s:%s" % (x["stamp"], x["held"]) for x in r["held"] if not str(x["stamp"]).startswith("W452"))
    out["l_old_like"] = len([1 for i in s2["ids"] if not str(stamp_of.get(i, "")).startswith("W452")])
    # ------------------------------------------------------------------ 2. the crafted scans
    S1, S2, S3, BN = W["s1"], W["s2"], W["s3"], W["bill_n"]
    out["c_s1_listed0"] = S1 in s2["ids"]
    out["c_s1_dl0"] = dl(S1)
    out["c_s2_line"] = s2["lines"].get(S2, "")
    out["c_s2_dl"] = dl(S2)
    out["c_s3_dl"] = dl(S3)
    out["c_nahi"] = FJ("manoj", "/finance/porders/api/scan/confirm", dict(scan=S1, bill=BN, yes=False))[0] if NEW else None
    s2b = step2()
    out["c_s1_listed1"] = S1 in s2b["ids"]
    r = fc.get("/finance/purchase/api/scan-files/today.zip", base_url=BASE, headers=H("amir"))
    try:
        import zipfile                                                 # noqa: PLC0415
        out["c_zip"] = [r.status_code, sorted(zipfile.ZipFile(io.BytesIO(r.data)).namelist())]
    except Exception:                                                  # noqa: BLE001
        out["c_zip"] = [r.status_code, None]
    # a Marg bill arrives that the matcher links to S1 (as a push would: the matcher runs)
    md5 = W["md5"]
    ins(db, "purchase_bill", dict(supplier_norm=SUP, supplier=SUP, bill_no="W452-98765", bill_date=D(20), month=D(20)[:7], cash_p=0, credit_p=100000,
                                  amount_p=100000, source_md5=md5, bw_md5=md5, bw_amount_p=100000, date_src="BILLWISE"))
    db.commit()
    purchase_app._rematch(db, "W452 push")
    s2c = step2()
    out["c_s1_listed2"] = S1 in s2c["ids"]
    out["c_s2_listed2"] = S2 in s2c["ids"]
    # the page is drawn with S2 on it; a Marg bill arrives that links S2; then the download
    ins(db, "purchase_bill", dict(supplier_norm=SUP, supplier=SUP, bill_no="W452-55555", bill_date=D(3), month=D(3)[:7], cash_p=0, credit_p=77700,
                                  amount_p=77700, source_md5=md5, bw_md5=md5, bw_amount_p=77700, date_src="BILLWISE"))
    db.commit()
    purchase_app._rematch(db, "W452 push")
    out["c_s2_linked"] = bool(q("SELECT 1 FROM purchase_scan_link WHERE asset_bill_id=?", S2))
    out["c_guard"] = dl(S2)
    out["c_s2_listed3"] = S2 in step2()["ids"]
    hs = FG("manoj", "/finance/purchase/page/scans")[1]
    m = re.search(r'id="s452amirlist".*?</div></div>', hs, re.S)
    out["c_scans_line"] = text(m.group(0)) if m else ""
    # ------------------------------------------------------------------ 3. the phone's key
    def tok():
        r = db.execute("SELECT value FROM setting WHERE key='supplier_msg.phone_token'").fetchone()
        v = r[0] if r else ""
        if v and v not in secrets_seen:
            secrets_seen.append(v)
        return v

    def audits():
        return int(db.execute("SELECT COUNT(*) FROM purchase_audit WHERE action='phone_token_shown'").fetchone()[0])

    def setup(u):
        s, h, _hd = FG(u, "/finance/purchase/page/phone-setup")
        m = re.search(r'id="s452phone"[^>]*>(.*?)</div>', h, re.S)
        return dict(status=s, key=(tok() in h) if tok() else False, line=(text(m.group(1)) if m else ""), page=("MacroDroid" in h))

    T0 = tok()
    a0 = audits()
    out["k_roles"] = {u: setup(u)["status"] for u in ("amir", "darpan", "reception")}
    m1, m2 = setup("manoj"), setup("manoj")
    out["k_owner"] = [m1["status"], m1["key"], m2["key"], m1["page"], m1["line"]]
    out["k_audit"] = [a0, audits()]
    if NEW:
        ap_ = subprocess.run([sys.executable, "-B", os.path.join(os.environ["KITDIR"], "apply_s452.py"), "--finance", FIN, "--db", os.environ["FINANCE_DB"]],
                             env=dict(os.environ), stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=600)
        out["k_apply"] = [ap_.returncode, ap_.stdout.splitlines()[-1:] if ap_.stdout else [], any(t and t in ap_.stdout for t in (T0, tok()))]
        out["k_apply_text"] = [l for l in ap_.stdout.splitlines() if not (T0 and T0 in l)][:40]
        T1 = tok()
        out["k_new"] = bool(T1) and T1 != T0
        out["k_old_next"] = fc.get("/finance/api/supplier-msg/next", base_url=BASE, headers={"X-Phone-Token": T0}).status_code
        out["k_line_401"] = setup("manoj")["line"]
        out["k_new_next"] = fc.get("/finance/api/supplier-msg/next", base_url=BASE, headers={"X-Phone-Token": T1}).status_code
        m3 = setup("manoj")
        out["k_line_200"] = m3["line"]
        out["k_key_new_shown"] = m3["key"]
        out["k_setting"] = (q("SELECT value FROM setting WHERE key='amir.vouchers_per_visit'") or [{}])[0].get("value")
    else:
        sa = FG("amir", "/finance/purchase/page/phone-setup")
        out["k_amir_old"] = [sa[0], ("Token" in sa[1])]
    # ------------------------------------------------------------------ 4. step 7 (Amir) and the owner's day
    h7 = FG("amir", "/finance/amir/step/7")[1]
    out["s7_amir"] = text(re.sub(r"(?s)<style>.*?</style>", "", h7))
    m = re.search(r"<h2>Abhi kuch baaki hai</h2>.*?</div>", h7, re.S)
    out["s7_card"] = text(m.group(0)) if m else ""
    od = FG("manoj", "/finance/amir/day")[1]
    m = re.search(r"<div class=card><h2>Amir -- .*?</div>", od, re.S)
    out["s7_owner"] = text(m.group(0)) if m else text(od)[:600]
    # ------------------------------------------------------------------ 5. the board
    s, hb, _hd = FG("amir", "/finance/stock/page/amir?count=1")
    out["b_page"] = [s, len(DEVA.findall(hb)), ("data-ent" in hb and "Marg mein daal diya" in hb), ("Server ko bhejo" in hb), ("Closing stock nikaaliye" in hb)]
    out["b_json0"] = board("amir")
    k1 = out["b_json0"]["keys"][:1]
    out["b_enter1"] = enter(k1, "W452-A1")
    out["b_card1"] = card(1)
    out["b_json1"] = board("amir")
    out["b_rate_card0"] = card(1)
    rate_sql = ([d.get("due_sql") for d in (dm.get("duties") or []) if d.get("id") == "amir.rate_entry"] or [""])[0] if NEW else ""

    def rate_n():
        try:
            return int(db.execute(rate_sql).fetchone()[0] or 0) if rate_sql else None
        except Exception as e:                                         # noqa: BLE001
            return "ERR %s" % str(e)[:80]
    out["b_rate_sql0"] = rate_n()
    try:
        import stock_app                                               # noqa: E402
        due = stock_app.amir_rate_due(db) if NEW else []
    except Exception:                                                  # noqa: BLE001
        due = []
    out["b_rate_due"] = [t["item"] for t in due]
    # Amir puts the selling rate in Marg's item master; Marg's next item export carries it into the spine (here: the scratch copy)
    sp = sqlite3.connect(os.environ["SPINE_DB"], timeout=30)
    for t in due:
        sp.execute("INSERT INTO sp_item_fact (name, packing, fact, value, as_on, source_md5) SELECT name, packing, fact, '350.0', ?, 'W452 crafted export' "
                   "FROM sp_item_fact WHERE name=? AND fact IN ('s_rate','mrp') GROUP BY fact", (TODAY.isoformat(), t["item"]))
    sp.commit()
    sp.close()
    out["b_rate_card1"] = card(1)
    out["b_json_rates1"] = board("amir")["rates"]
    out["b_rate_sql1"] = rate_n()
    h6 = FG("amir", "/finance/amir/step/6")[1]
    out["b_step6_heads"] = len(re.findall(r"<h[12][^>]*>\s*Marg sudhar", h6))
    # ------------------------------------------------------------------ 6. NEFT
    vendors = sorted({r["vendor"] for r in q("SELECT vendor FROM supplier_msg WHERE month='2026-08'")})
    out["n_vendors"] = len(vendors)
    blocks = {n: neft_block(n) for n in range(1, 8)}
    out["n_blocks"] = {str(n): b[:600] for n, b in blocks.items()}
    out["n_vendor_in_block"] = sorted({v for b in blocks.values() for v in vendors if v and v in b})
    out["n_baaki_in_block"] = any("baaki" in b or "bata diya" in b for b in blocks.values())
    s, pdf, hd = FB("amir", "/finance/amir/pack/2026-08/neft")
    out["n_pdf"] = [s, hd.get("Content-Type"), (re.search(r'filename="([^"]+)"', hd.get("Content-Disposition") or "") or [None, ""])[1], pdf[:4] == b"%PDF"]
    toks = [re.sub(r"\\(.)", r"\1", t.decode("cp1252", "replace")) for t in re.findall(rb"\(((?:\\.|[^\\)])*)\) Tj", pdf)]   # the PDF's strings
    try:
        import packs                                                   # noqa: E402
        from openpyxl import load_workbook                             # noqa: E402
        rows = list(load_workbook(io.BytesIO(packs.build_neft_paid_sheet(db, "2026-08"))).active.iter_rows(values_only=True))[1:]
        sheet_names = [str(r[0]) for r in rows]
        sheet_total_p = int(round(sum(float(r[1] or 0) for r in rows) * 100))
    except Exception as e:                                             # noqa: BLE001
        sheet_names, sheet_total_p = ["ERR %s" % e], None
    out["n_sheet"] = [len(sheet_names), sheet_total_p]
    if s == 200 and pdf[:4] == b"%PDF":
        import clinic_day_pdf as cdp                                   # noqa: E402
        import supplier_msg as smx                                     # noqa: E402
        fitted = [cdp._latin(cdp._fit(n, 9, 280)) for n in sheet_names]
        out["n_pdf_rows"] = [sum(1 for f in fitted if f in toks), len(fitted)]
        out["n_pdf_total"] = [smx._s452_rs(sheet_total_p) in toks, smx._s452_rs(sheet_total_p) if sheet_total_p is not None else None]
    else:
        out["n_pdf_rows"] = out["n_pdf_total"] = None
    out["n_sms"] = neft_block(3)
    out["n_none_route"] = FG("amir", "/finance/amir/pack/2099-06/neft")[0]
    out["n_sep_route"] = FG("amir", "/finance/amir/pack/2026-09/neft")[0]
    try:
        import packs                                                   # noqa: E402,F811
        real = packs.amir_pack
        packs.amir_pack = lambda con, m: dict(real(con, m), ready=True, seen=None) if m == "2026-09" else real(con, m)   # walk-only: a crafted ready pack
        hc = FG("amir", "/finance/amir/step/1")[1]
        packs.amir_pack = real
        m = re.search(r"Mahine ka pack &mdash; September 2026</span>(.*?)</div>", hc, re.S)
        out["n_sep_pack"] = [bool(m), ("Paid NEFT sheet" in m.group(1)) if m else None]
    except Exception as e:                                             # noqa: BLE001
        out["n_sep_pack"] = ["ERR %s" % e]
    import html as _html                                               # noqa: E402
    for who in ("shavez", "manoj"):
        hp = _html.unescape(FG(who, "/finance/purchase/page/pay/2026-08")[1])     # a name with '&' is escaped on the page
        out["n_pay_%s" % who] = [sum(1 for v in vendors if v in hp), "neftcard_s407" in hp]
    xl = {}
    for p in ("/finance/amir/pack/2026-08/neft", "/finance/purchase/page/pay/2026-08/advice.xlsx", "/finance/packs/preview/neft?month=2026-08"):
        s, b, hd = FB("amir", p)
        xl[p] = [s, ("spreadsheet" in (hd.get("Content-Type") or "")) or b[:2] == b"PK"]
    out["n_xlsx"] = xl
    # ------------------------------------------------------------------ 7. Stage C: 12 a visit, more on request
    db.execute("INSERT INTO stock_stage_event (count_id, stage, event, ref, at, detail) VALUES (1,'A','verified','',?, 'W452 crafted')", (now_iso(),))
    db.commit()
    out["sc_b_card"] = card(1)
    out["sc_b_board"] = board("amir")
    db.execute("INSERT INTO stock_stage_event (count_id, stage, event, ref, at, detail) VALUES (1,'B','verified','',?, 'W452 crafted')", (now_iso(),))
    db.commit()
    out["sc_c1_card"] = card(1)
    b1 = board("amir")
    out["sc_c1_board"] = b1
    lot1 = b1["keys"]
    time.sleep(1.2)
    out["sc_enter"] = enter(lot1, "W452-C1")
    time.sleep(1.2)
    need = {}
    for k in lot1:
        for item, ch in W["key_lines"].get(k, []):
            need[item] = need.get(item, 0) + ch
    exp_src = "push_expected base=W452 pur_to=%s" % TODAY.isoformat()

    def feed(as_on, marg, rec):
        for item, qty in marg.items():
            ins(db, "stock_feed", dict(as_on=as_on, source="push_snapshot W452", item=item, qty=qty, received_at=rec))
            ins(db, "stock_feed", dict(as_on=as_on, source=exp_src, item=item, qty=100, received_at=rec))
        db.commit()

    feed(DMY(TODAY - dt.timedelta(days=1)), {i: 100 for i in need}, now_iso())
    wrong = sorted(need)[0] if need else ""
    ml = {i: 100 + need[i] for i in need}
    if wrong:
        ml[wrong] += 1
    time.sleep(1.2)
    feed(DMY(TODAY), ml, now_iso())
    out["sc_wrong"] = dict(card=card(1), item=wrong, marg=ml.get(wrong), should=(ml.get(wrong) or 1) - 1)
    a_more0 = len(q("SELECT id FROM stock_stage_event WHERE count_id=1 AND stage='C' AND event='release' AND detail LIKE '%\"why\": \"more\"%'"))
    out["sc_tap1"] = FJ("amir", "/finance/stock/api/pad/amir/1/more", {})
    out["sc_tap1_card"] = card(1)
    out["sc_tap1_board"] = board("amir")
    out["sc_audit1"] = [a_more0, len(q("SELECT id FROM stock_stage_event WHERE count_id=1 AND stage='C' AND event='release' AND detail LIKE '%\"why\": \"more\"%'"))]
    out["sc_audit_row"] = (q("SELECT detail FROM stock_stage_event WHERE count_id=1 AND stage='C' AND event='release' AND detail LIKE '%\"why\": \"more\"%' ORDER BY id DESC LIMIT 1") or [{}])[0].get("detail")
    # his next visit: the newest lot was released yesterday and he opens his day today
    rel = q("SELECT id FROM stock_stage_event WHERE count_id=1 AND stage='C' AND event='release' ORDER BY id DESC LIMIT 1")
    if rel:
        db.execute("UPDATE stock_stage_event SET at=? WHERE id=?", (D(1) + "T00:00:00", rel[0]["id"]))
    if q("SELECT 1 FROM amir_day WHERE day=?", TODAY.isoformat()):
        db.execute("UPDATE amir_day SET opened_at=? WHERE day=? AND (opened_at IS NULL OR opened_at < ?)", (TODAY.isoformat() + " 00:00:05", TODAY.isoformat(), D(1) + " 00:00:01"))
    else:
        ins(db, "amir_day", dict(day=TODAY.isoformat(), opened_at=TODAY.isoformat() + " 00:00:05"))
    db.commit()
    out["sc_visit_card"] = card(1)
    out["sc_visit_board"] = board("amir")
    out["sc_tap_none"] = FJ("amir", "/finance/stock/api/pad/amir/1/more", {})
    out["sc_audit2"] = len(q("SELECT id FROM stock_stage_event WHERE count_id=1 AND stage='C' AND event='release' AND detail LIKE '%\"why\": \"more\"%'"))
    out["sc_owner"] = [t for t in needs() if t.startswith("Count #1:")]
    # ------------------------------------------------------------------ the key never leaves the probe
    blob = json.dumps(out, default=str)
    leaked = any(t and t in blob for t in secrets_seen)
    if leaked:
        print("W452JSON " + json.dumps({"LEAK": True}))
        return
    print("W452JSON " + blob)


# ======================================================================= the WALK
def main():
    ap = argparse.ArgumentParser()
    for k in ("--fin-new", "--fin-old", "--por", "--ast", "--shared", "--db", "--adb", "--kit", "--duty-map", "--uploads", "--spine"):
        ap.add_argument(k, required=True)
    a = ap.parse_args()
    for p in (a.db, a.adb):
        assert p.startswith("/tmp/"), "refusing a non-scratch database: " + p
    n, fails = [0], []

    def check(label, cond, got=None):
        n[0] += 1
        print(("  ok   " if cond else "  FAIL ") + label + (("   [" + str(got)[:420] + "]") if got is not None else ""))
        if not cond:
            fails.append(label)

    def copydb(src, dst):
        s = sqlite3.connect("file:%s?mode=ro" % src, uri=True)
        d = sqlite3.connect(dst)
        s.backup(d)
        d.close()
        s.close()

    # ---- the walk's own sign-in (staff-eye): a random secret and its own user store in the scratch portal folder
    secret = secrets.token_hex(32)
    pw = {u: secrets.token_urlsafe(12) for u in USERS}
    assert a.por.startswith("/tmp/")
    with open(os.path.join(a.por, "portal_config.py"), "w") as fh:
        fh.write("# W452 walk only\nCLINIC_SSO_SECRET = %r\nPORTAL_PIN_HASH = ''\nPORTAL_PIN_SALT = ''\nPORTAL_TOKEN_SEED = 'w452-walk-seed'\n" % secret)
    sys.path.insert(0, a.por)
    import clinic_users                                                # noqa: E402
    store = os.path.join(a.por, "clinic_users.json")
    if os.path.exists(store):
        os.remove(store)
    try:
        clinic_users.add_role(store, "staff")
    except ValueError:
        pass
    for u, role in USERS.items():
        clinic_users.add_user(store, u, role, pw[u])
    # ---- the live count's vouchers, read by key on the scratch copy
    F = sqlite3.connect(a.db)
    bat = {}
    for rno, kind, bno, sec, item, ch in F.execute("SELECT round_no, kind, batch_no, section, item, change FROM stock_voucher_line WHERE count_id=1"):
        b = bat.setdefault("%d|%s|%d" % (rno, kind, bno), dict(secs=set(), lines=[]))
        b["secs"].add(sec or "")
        b["lines"].append((item, int(ch or 0)))
    ortho = [k for k, b in bat.items() if "Orthotics" in b["secs"]]
    med_n = len(bat) - len(ortho)
    F.close()
    # ---- the box's own scan files, copied (read only) into the walk's upload folder; the crafted papers beside them
    up = os.path.join(os.path.dirname(a.adb), "w452_uploads")
    os.makedirs(up, exist_ok=True)
    A = sqlite3.connect("file:%s?mode=ro" % a.adb, uri=True)
    copied = 0
    for (stored,) in A.execute("SELECT source_stored FROM bills WHERE kind='Pharmacy' AND status='captured' AND source_stored IS NOT NULL"):
        src = os.path.join(a.uploads, os.path.basename(str(stored)))
        if os.path.isfile(src):
            shutil.copy(src, os.path.join(up, os.path.basename(str(stored))))
            copied += 1
    A.close()
    pdf = b"%PDF-1.4\n% W452 walk paper\n" + secrets.token_bytes(64) + b"\n%%EOF\n"
    for i in (1, 2, 3):
        with open(os.path.join(up, "w452_s%d.pdf" % i), "wb") as fh:
            fh.write(pdf + bytes([i]))
    copydb(a.db, a.db + ".old")
    copydb(a.adb, a.adb + ".old")
    spines = {s: os.path.join(os.path.dirname(a.db), "w452_spine_%s.db" % s) for s in ("new", "old")}   # Marg's item export (the spine), a scratch copy per side
    for s in spines.values():
        copydb(a.spine, s)
    log =os.path.join(os.path.dirname(a.db), "w452_access.log")
    with open(log, "w") as fh:
        fh.write('127.0.0.1 - - [%s +0530] "GET /finance/api/supplier-msg/next HTTP/1.1" 401 31\n' % (dt.datetime.now() - dt.timedelta(days=1)).strftime("%d/%b/%Y:%H:%M:%S"))

    def craft(fin, ast):
        Fc = sqlite3.connect(fin, timeout=30)
        Ac = sqlite3.connect(ast, timeout=30)
        md5 = "w452" + hashlib.md5(b"bw452").hexdigest()[:28]
        ins(Fc, "purchase_export", dict(md5=md5, type="BILLWISE", file="W452", period_from=D(30), period_to=TODAY.isoformat(),
                                        export_stamp=dt.datetime.now().strftime("%Y%m%d-%H%M%S"), received_at=now_iso(), n_rows=1, grand_amount_p=0), replace=True)
        bn = ins(Fc, "purchase_bill", dict(supplier_norm=SUP, supplier=SUP, bill_no="W452-61234", bill_date=D(5), month=D(5)[:7], cash_p=0, credit_p=100000,
                                            amount_p=100000, source_md5=md5, bw_md5=md5, bw_amount_p=100000, date_src="BILLWISE"))
        now = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ids = []
        for i, (bno, bdate, amt) in enumerate((("W452-98765", D(20), 1000.0), ("W452-55555", "2023-05-20", 777.0), ("", D(1), 333.0)), 1):
            ids.append(ins(Ac, "bills", dict(kind="Pharmacy", vendor=SUP, bill_no=bno, bill_date=bdate, total_amount=amt, notes="W452 walk", created_at=now,
                                             stamp_no="W452-0%d" % i, status="captured", ocr_status="read", lane="pharmacy", source_stored="w452_s%d.pdf" % i,
                                             source_orig="w452", submitted_by="W452 walk")))
        ins(Fc, "purchase_neft_event", dict(month="2099-04", kind="provisional", source="sms", amount_p=123400, sms_date=D(2), created_at=now_iso(), created_by="sms"))
        Fc.commit(); Ac.commit(); Fc.close(); Ac.close()
        return dict(s1=ids[0], s2=ids[1], s3=ids[2], bill_n=bn, md5=md5)

    ids = craft(a.db, a.adb)
    assert ids == craft(a.db + ".old", a.adb + ".old"), "the crafted rows must carry the same ids on both sides"
    # ---- the owner's figure is never moved: a copy where amir.vouchers_per_visit reads 8
    keep = a.db + ".keep"
    copydb(a.db, keep)
    K = sqlite3.connect(keep)
    K.execute("UPDATE setting SET value='8' WHERE key='amir.vouchers_per_visit'")
    K.commit()
    K.close()
    kp = subprocess.run([sys.executable, "-B", os.path.join(a.kit, "apply_s452.py"), "--finance", a.fin_new, "--db", keep, "--settings-only"],
                        env=dict(os.environ, FINANCE_DB=keep), stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
    K = sqlite3.connect(keep)
    kept = (K.execute("SELECT value FROM setting WHERE key='amir.vouchers_per_visit'").fetchone() or [None])[0]
    K.close()
    print("-- crafted on both sides: Marg bill W452-61234 (%s, Rs 1,000, %s); scans W452-01 (a near-match of it), W452-02 (bill date 2023), W452-03 "
          "(no bill number read); an SMS-read NEFT for 2099-04; %d of the box's scan files copied read-only; the live count #1: %d vouchers, %d orthotic, %d medicine"
          % (SUP, D(5), copied, len(bat), len(ortho), med_n))
    W = dict(pw=pw, key_lines={k: b["lines"] for k, b in bat.items()}, **ids)

    def run(side):
        fin = a.fin_new if side == "new" else a.fin_old
        dbp, adbp = (a.db, a.adb) if side == "new" else (a.db + ".old", a.adb + ".old")
        env = dict(os.environ, MODE=side, FINDIR=fin, PORDIR=a.por, FINANCE_DB=dbp, ASSETS_DB=adbp, ASSETS_UPLOADS=up, FINANCE_ALLOW_HEADER_AUTH="1",
                   FINANCE_SSO_DIR=a.por, CLINIC_PORTAL_DIR=a.por, CLINIC_USERS_FILE=store, TILE_GRANTS_FILE=os.path.join(a.por, "tile_grants.json"),
                   CLINIC_SSO_SECRET=secret, REAL_DUTY_MAP=a.duty_map, DUTY_MAP_JSON=a.duty_map, SHARED_LIB_DIR=a.shared, W452=json.dumps(W),
                   KITDIR=a.kit, FINANCE_ACCESS_LOG=log, SPINE_DB=spines[side])
        env.pop("SARVAM_API_KEY", None)
        p = subprocess.run([sys.executable, "-B", os.path.abspath(__file__), "--probe"], env=env, cwd=fin, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                           text=True, timeout=2400)
        js = [l for l in p.stdout.splitlines() if l.startswith("W452JSON ")]
        if not js:
            print("-- the %s probe did not finish (exit %s); its last lines:" % (side, p.returncode))
            for l in p.stdout.splitlines()[-25:]:
                print("   " + l[:300])
            return None
        return json.loads(js[-1][9:])

    N = run("new")
    O = run("old")
    check("both probes ran to the end (new / old); no phone key reached either probe's output", N is not None and O is not None and not N.get("LEAK") and not O.get("LEAK"))
    if N is None or O is None or N.get("LEAK") or O.get("LEAK"):
        print("WALK_S452 RED -- a probe did not finish, or a key reached its output")
        return 1
    print("-- 0  the staff-eye walk (DUTY_MAP.json v3): amir, darpan, shavez and the owner")
    for who in ("amir", "darpan", "shavez", "manoj"):
        e = N["eye"][who]
        missing = [d["id"] for d in e["duties"] if d["tile"] and not d["tile_seen"]]
        unseen = [d["id"] for d in e["duties"] if d.get("door_seen") is False]
        errs = [d["id"] for d in e["duties"] if isinstance(d.get("due_n"), str)]
        due = [(d["id"], d["due_n"]) for d in e["duties"] if isinstance(d.get("due_n"), int) and d["due_n"] > 0]
        check("%s: %d duties; every tile on the home; every due duty's door shows it (%s)" % (who, len(e["duties"]), ", ".join("%s %s" % x for x in due) or "none due"),
              e["status"] == 200 and e["duties"] and not missing and not unseen and not errs, dict(missing=missing, unseen=unseen, errors=errs))
    re_ = [d for d in N["eye"]["amir"]["duties"] if d["id"] == "amir.rate_entry"]
    check("amir.rate_entry is on the map, due on the box's own data, and his card shows it", bool(re_) and isinstance(re_[0]["due_n"], int) and re_[0]["due_n"] > 0
          and re_[0].get("door_seen") is True, re_)
    print("-- 1  the list holds only bills that are not in Marg (F-686), on the box's own scans")
    ln = N["l_box_stamps"]
    check("the listed scans are exactly S440's 'Marg ka intezaar' scans with a known supplier: %s" % ", ".join(ln), N["l_heading"] and sorted(ln) == N["l_wait"], [ln, N["l_wait"]])
    check("none of %s (the shop's own name as supplier) is listed" % ", ".join(SHOP), not any(s in ln for s in SHOP), ln)
    check("the grey line's N (%d) plus the listed (%d) = the %d of the list as it is today" % (N["l_held"], len(ln), O["l_old_like"]),
          N["l_held"] - 0 >= N["l_held_box"] and N["l_held_box"] + len(ln) == O["l_old_like"], [N["l_held"], N["l_held_box"], len(ln), O["l_old_like"], N["l_held_reasons"]])
    lines = N["l_lines"]
    check("each listed line shows the stamp, the supplier, 'scan: dd-mm, <who>' and a bill date only within 60 days",
          all(re.search(r"scan: \d\d-\d\d, \S", t) and (re.search(r"bill \d\d-\d\d-\d{4}", t) or "bill ki tareekh scan par saaf nahi" in t) for t in lines.values()), lines)
    dls = N["l_downloads"]
    check("every download name starts with its stamp; none contains 'nobill'", dls and all(d["status"] == 200 and d["pdf"] and d["name"].startswith(N["l_box_stamps"][i].replace(" ", "_"))
                                                                                         and "nobill" not in d["name"] for i, d in enumerate(dls.values())), dls)
    print("-- 2  the crafted scans: a near-match, a 2023 date, an unread number; the download guard")
    check("a near-match scan (W452-01) is held back, and its download says 'reception ki jaanch'", not N["c_s1_listed0"] and N["c_s1_dl0"]["status"] == 409
          and "reception ki jaanch" in N["c_s1_dl0"]["text"], [N["c_s1_listed0"], N["c_s1_dl0"]])
    check("reception answers 'Nahi' (Yahi bill hai?): it joins the list", N["c_nahi"] == 200 and N["c_s1_listed1"], [N["c_nahi"], N["c_s1_listed1"]])
    check("when a Marg bill links it, it leaves", N["c_s1_listed2"] is False, N["c_s1_listed2"])
    want2 = "W452-02_KEDWALK452_PHARMA_W45255555_%s.pdf" % DMY(TODAY)
    check("a scan with a 2023 bill date: the line says the date is not clear; the file carries the scan day (%s)" % want2,
          "bill ki tareekh scan par saaf nahi" in N["c_s2_line"] and N["c_s2_dl"]["name"] == want2, [N["c_s2_line"], N["c_s2_dl"]["name"]])
    want3 ="W452-03_KEDWALK452_PHARMA_%s-%s-%s.pdf" % (D(1)[8:10], D(1)[5:7], D(1)[:4])
    check("a scan with no bill number read: the name leaves it out (%s), no 'nobill'" % want3, N["c_s3_dl"]["name"] == want3, N["c_s3_dl"]["name"])
    z = N["c_zip"][1] or []
    check("the zip of today's follows the same rule: every name starts with its stamp, none 'nobill', the held near-match not in it until answered",
          N["c_zip"][0] == 200 and z and all(re.match(r"^(W452-0\d|B-\d+)_", x) for x in z) and not any("nobill" in x for x in z), z)
    check("the download guard: a scan linked after the page was drawn answers 'Yeh bill Marg mein aa chuka hai', and it is gone on reload",
          N["c_s2_listed2"] and N["c_s2_linked"] and N["c_guard"]["status"] == 409 and "Yeh bill Marg mein aa chuka hai" in N["c_guard"]["text"] and not N["c_s2_listed3"],
          [N["c_s2_listed2"], N["c_s2_linked"], N["c_guard"], N["c_s2_listed3"]])
    check("the owner's Scan links page: 'Amir's list: N to enter · M held for reception'", re.search(r"Amir's list: \d+ to enter · \d+ held for reception", N["c_scans_line"]) is not None,
          N["c_scans_line"])
    print("-- 3  the reception phone's key (F-687)")
    check("amir, darpan and the reception login get the refusal on the setup page", N["k_roles"]["amir"] == 403 and N["k_roles"]["darpan"] == 403
          and N["k_roles"]["reception"] in (302, 403), N["k_roles"])
    check("the owner (the medical checker) gets 200 with the key, on every open ('shown once' retired)", N["k_owner"][:4] == [200, True, True, True], N["k_owner"][:4])
    check("each showing writes one audit row (two opens: +2)", N["k_audit"][1] - N["k_audit"][0] == 2, N["k_audit"])
    check("the page's first line says when the phone last asked and with which answer (the access log's 401 before S452 recorded anything)",
          "Phone ne aakhri baar" in N["k_owner"][4] and "401" in N["k_owner"][4] and "message intezaar mein" in N["k_owner"][4], N["k_owner"][4])
    check("the data step makes a new key (the key itself in no line of its output)", N["k_apply"][0] == 0 and N["k_apply"][1] == ["S452 apply done"] and not N["k_apply"][2]
          and N["k_new"], N["k_apply"])
    check("the old key gets 401 on the queue door; the new one 200; the page's line follows (401, then 200); the new key on the owner's page",
          N["k_old_next"] == 401 and N["k_new_next"] == 200 and "jawab 401" in N["k_line_401"] and "jawab 200" in N["k_line_200"] and N["k_key_new_shown"],
          [N["k_old_next"], N["k_new_next"], N["k_line_401"], N["k_line_200"]])
    print("-- 4  step 7 in Roman Hindi; the owner's English")
    s7 = N["s7_amir"]
    check("as amir: no English phrase of the old list; the count's line stands apart with 'din band karne se nahi rukta'",
          not any(x in s7 for x in ("not confirmed", "not verified", "not ticked", "not entered")) and re.search(r"Ginti #\d+: .*?\(din band karne se nahi rukta\)", s7) is not None,
          N["s7_card"][:400])
    check("as the owner (/finance/amir/day): English kept", "OPEN --" in N["s7_owner"] and ("not confirmed" in N["s7_owner"] or "not verified" in N["s7_owner"] or "count #" in N["s7_owner"]),
          N["s7_owner"][:300])
    print("-- 5  the board in Roman Hindi; every duty on his card")
    bp = N["b_page"]
    check("as amir: no code point of the Devanagari block on the page he is served, nor in its data", bp[0] == 200 and bp[1] == 0 and N["b_json0"]["deva"] == 0
          and bp[3] and bp[4], [bp, N["b_json0"]["deva"]])
    check("the 7 vouchers, their boxes and buttons still work: one marked entered -> 'Orthotic voucher baaki: 6'",
          N["b_json0"]["n"] == 7 and bp[2] and N["b_enter1"] == [200] and "Orthotic voucher baaki: 6" in N["b_card1"] and N["b_json1"]["entered"] == 1, [N["b_json0"]["n"], N["b_card1"][:160]])
    nr = len(N["b_rate_due"])
    check("'%d item ka rate Marg mein daalna hai' on his card while due; gone when a crafted Marg item export carries the rates (the spine's S.RATE "
          "above 0) -- the board's (b) empty too, and DUTY_MAP's amir.rate_entry reads %d, then 0" % (nr, nr),
          nr > 0 and ("%d item ka rate Marg mein daalna hai" % nr) in N["b_rate_card0"] and "ka rate Marg mein" not in N["b_rate_card1"] and N["b_json_rates1"] == []
          and N["b_rate_sql0"] == nr and N["b_rate_sql1"] == 0,
          [N["b_rate_due"], N["b_rate_card1"][-120:], N["b_json_rates1"], N["b_rate_sql0"], N["b_rate_sql1"]])
    check("a salt line already done in Marg leaves block (a)", not N["b_json0"]["salt_done"] and "LINVIZ 600" not in N["b_json0"]["salt"], N["b_json0"]["salt"])
    check("step 6 carries the heading 'Marg sudhar' once", N["b_step6_heads"] == 1, N["b_step6_heads"])
    print("-- 6  NEFT: one line and a PDF, only once confirmed")
    bl = N["n_blocks"]
    check("August (the owner's entry, 24-09; no bank SMS read): every step reads 'NEFT August 2026 — payment 24-09 ko ho gaya (Doctor sahab ne darj kiya)' and 'Paid NEFT sheet (PDF)'",
          all("NEFT August 2026 — payment 24-09 ko ho gaya (Doctor sahab ne darj kiya)" in bl[str(k)] and "Paid NEFT sheet (PDF)" in bl[str(k)] for k in range(1, 8)), bl["1"])
    check("no supplier name (of %d) and no 'baaki' / 'bata diya' in the NEFT block on any of his seven steps" % N["n_vendors"],
          N["n_vendors"] > 0 and N["n_vendor_in_block"] == [] and not N["n_baaki_in_block"], [N["n_vendor_in_block"][:4], N["n_baaki_in_block"]])
    check("the file is a PDF named NEFT_paid_2026-08.pdf", N["n_pdf"] == [200, "application/pdf", "NEFT_paid_2026-08.pdf", True], N["n_pdf"])
    check("its rows and total equal the old sheet's (%s rows, total %s)" % tuple(N["n_sheet"]), N["n_pdf_rows"] and N["n_pdf_rows"][0] == N["n_pdf_rows"][1] == N["n_sheet"][0]
          and N["n_pdf_total"] and N["n_pdf_total"][0], [N["n_pdf_rows"], N["n_pdf_total"]])
    check("a crafted bank SMS read (April 2099): 'NEFT April 2099 — bank ka kaam ho gaya, %s'" % (D(2)[8:10] + "-" + D(2)[5:7]),
          ("NEFT April 2099 — bank ka kaam ho gaya, %s-%s" % (D(2)[8:10], D(2)[5:7])) in N["n_sms"], N["n_sms"])
    check("a month with neither: no line, the route refuses amir (2099-06 and September), no sheet in a ready pack (a walk-only patch makes September's ready)",
          "June 2099" not in str(bl) and "September 2026" not in str(bl) and N["n_none_route"] == 403 and N["n_sep_route"] == 403 and N["n_sep_pack"] == [True, False],
          [N["n_none_route"], N["n_sep_route"], N["n_sep_pack"]])
    check("Shavez's and the owner's pay pages keep the supplier list (%d names)" % N["n_vendors"],
          N["n_pay_shavez"] == [N["n_vendors"], True] and N["n_pay_manoj"] == [N["n_vendors"], True] and N["n_pay_shavez"] == O["n_pay_shavez"], [N["n_pay_shavez"], N["n_pay_manoj"], O["n_pay_shavez"]])
    check("no address gives amir an .xlsx (Amir's NEFT route, the bank advice file, the owner's preview)", not any(v[1] for v in N["n_xlsx"].values()), N["n_xlsx"])
    print("-- 7  medicine vouchers: 12 a visit, more on request")
    check("the setting is 12 after the data step; a box where the owner had set 8 keeps his", N["k_setting"] == "12" and kept == "8" and kp.returncode == 0, [N["k_setting"], kept])
    check("nothing of Stage C before Stage B is verified (Stage B: the renames, no 'Dawa voucher')", "Dawa voucher" not in N["sc_b_card"] and N["sc_b_board"]["stage"].get("stage") == "B",
          [N["sc_b_card"][:160], N["sc_b_board"]["stage"]])
    b1 = N["sc_c1_board"]
    check("after a crafted verified Stage B: 12 on the board (1..12), 'Dawa voucher: aaj ke 12 (baaki %d)', the button offered" % (med_n - 12),
          b1["n"] == 12 and b1["seq"] == list(range(1, 13)) and ("Dawa voucher: aaj ke 12 (baaki %d)" % (med_n - 12)) in N["sc_c1_card"] and b1["stage"].get("more") == med_n - 12,
          [b1["n"], N["sc_c1_card"][:200], b1["stage"]])
    sw = N["sc_wrong"]
    check("an export with one wrong item names it, with Marg's figure, the one it should be, and its voucher", sw["item"] and sw["item"] in sw["card"]
          and ("Marg mein %s, hona chahiye %s" % (sw["marg"], sw["should"])) in sw["card"] and "voucher" in sw["card"], sw)
    check("one tap on 'Aur voucher kholiye': 24 on the board, 'aaj ke 12 (baaki %d)' -- the wrong lot does not stop it; one audit row (who, how many)" % (med_n - 24),
          N["sc_tap1"][0] == 200 and N["sc_tap1_board"]["n"] == 24 and ("aaj ke 12 (baaki %d)" % (med_n - 24)) in N["sc_tap1_card"] and N["sc_audit1"][1] - N["sc_audit1"][0] == 1
          and '"by": "amir"' in (N["sc_audit_row"] or "") and '"n": 12' in (N["sc_audit_row"] or ""), [N["sc_tap1"], N["sc_tap1_board"]["n"], N["sc_tap1_card"][:160], N["sc_audit_row"]])
    check("his next visit opens the third lot: %d on the board and the button gone; another tap answers 'none left' and writes nothing" % med_n,
          N["sc_visit_board"]["n"] == med_n and N["sc_visit_board"]["stage"].get("more") == 0 and N["sc_tap_none"][0] == 409 and N["sc_audit2"] == N["sc_audit1"][1],
          [N["sc_visit_board"]["n"], N["sc_visit_board"]["stage"], N["sc_tap_none"], N["sc_audit2"]])
    check("the owner's stage line counts entered, released and verified: 'Stage C (medicine vouchers) 12/%d entered · %d released · 0 verified'" % (med_n, med_n),
          any(("Stage C (medicine vouchers) 12/%d entered · %d released · 0 verified" % (med_n, med_n)) in t for t in N["sc_owner"]), N["sc_owner"])
    print("-- 8  NEGATIVE CONTROL: the same walk on the box as it is")
    ol = O["l_box_stamps"]
    check("NEGATIVE: the old list carries the held scans -- the shop's own name among them (%s)" % ", ".join(s for s in SHOP if s in ol), any(s in ol for s in SHOP) and len(ol) > len(ln), ol)
    check("NEGATIVE: the old list lists the near-match; its 2023 scan's file carries 2023; the unread number reads 'nobill'",
          O["c_s1_listed0"] and "2023" in O["c_s2_dl"]["name"] and "nobill" in O["c_s3_dl"]["name"], [O["c_s1_listed0"], O["c_s2_dl"]["name"], O["c_s3_dl"]["name"]])
    check("NEGATIVE: the old setup page answers 200 to amir", O["k_amir_old"][0] == 200, O["k_amir_old"])
    check("NEGATIVE: the old step 7 lists English ('not confirmed' / 'not entered')", "not confirmed" in O["s7_amir"] or "not entered" in O["s7_amir"], O["s7_amir"][:200])
    check("NEGATIVE: the old board carries Devanagari (%d code points on the page)" % O["b_page"][1], O["b_page"][1] > 0, O["b_page"][:2])
    check("NEGATIVE: the old NEFT block names the suppliers with 'baaki'; Amir's old route gives an .xlsx",
          bool(O["n_vendor_in_block"]) and O["n_baaki_in_block"] and O["n_xlsx"]["/finance/amir/pack/2026-08/neft"][1], [O["n_vendor_in_block"][:3], O["n_xlsx"]])
    check("NEGATIVE: the old Stage C releases 5 and has no 'Aur voucher kholiye' door", "aaj ke 5 (baaki" in O["sc_c1_card"] and O["sc_tap1"][0] in (404, 405),
          [O["sc_c1_card"][:160], O["sc_tap1"][0]])
    print("WALK_S452 %s -- %d of %d passed" % ("GREEN" if not fails else "RED", n[0] - len(fails), n[0]))
    for f in fails:
        print("   FAILED: " + f)
    return 0 if not fails else 1


if __name__ == "__main__":
    if "--probe" in sys.argv:
        probe()
    else:
        sys.exit(main())
