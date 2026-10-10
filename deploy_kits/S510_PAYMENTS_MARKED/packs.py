#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""packs.py -- S408 (26-Sep-2026, D624). Month-end packs: the statement shelf, the accountant pack, Amir's pack, Shavez's
monthly checklist. OWNER: PARENT (clinic). The three Sanjeevni pieces (Amir's pack, the NEFT details, the Sanjeevni statement
matching) are marked below. Everything is server-side; the accountants receive nothing until the owner taps Send.

  the shelf     stmt_slot (one per account/card the owner listed; the tail is LEARNED by his one tap) · stmt_file (every file
                stmt_shelf.py fetched, identified BY CONTENT through pdftotext: the bank's name, the holder's words, the
                account / card tail, the period). Bank statements are read into bank_statement_period/_line by finance_yesbank
                (Yes Bank) or finance_icici (ICICI, new); a card statement is filed as arrived. Every account's month is one
                cell: arrived / read / matched. A month with any cell empty by the 10th -> a Needs-you line + a checklist line.
  the pack      /finance/packs (owner, English; unit 'packs', the doctor checker): every item of the owner's formal list as a
                row -- ready / missing (why) / late (belongs to <month>); previews; the two paper items as ticks; ONE button
                'Send to accountants' -> one mail (part 2 over packs.mail_limit_mb) to setting packs.accountant_to through the
                box's SMTP (the health report's .env), receipt in pack_send; a second send is a re-send, recorded as such.
  Amir's pack   (Sanjeevni) when both Sanjeevni statements of a month are on the shelf: the paid NEFT sheet (S407's marks) +
                the two PDFs, on his board as 'Pichle mahine ka pack' with tap-to-open and a 'dekh liya' tick. No email to Amir.
  the checklist /finance/packs/checklist (staff, Hindi; unit 'packs', shavez maker): packs_item rows -- 'system se ho gaya'
                where automatic (linked to the pack row), a tick for the manual ones (who, when, 10-minute guard); the Lab
                items ticks only; what is still open on the 10th -> the owner's Needs you.
NO patient data, no account number beyond a tail, no password anywhere in this file. The mail credentials are read from
/root/wa/.env at send time and never printed. PACKS_TODAY / PACKS_MAIL_STUB / PACKS_ENV / STMT_INBOX serve the walk.

S411 (the shelf's first real run, 26-Sep-2026): the identifier reads the HEADER of a statement -- the bank, the holder line ('Your
Details With Us'), the account tail, the period -- never a narration (a Yes Bank UPI narration inside an ICICI statement had flipped
the bank; 'NK PATHOLOGY' inside a transfer narration had placed Dr Bhawna's statement on NK's slot). Holders match as word sets and the
longest match wins (MANOJ KUMAR AGARWAL HUF beats MANOJ KUMAR AGARWAL); a tie is unplaced. Three ICICI slots the run showed were added
(NK Pathology current, Dr Bhawna savings, the HUF savings). A by-words placement learns the account tail. ICICI's pipe-delimited .txt
statements are read; a password-locked PDF is named as such and waits for a decrypted copy; a statement covering only part of a month is
'partial' and never the month's statement. ICICI lines live in icici_statement_period/_line (finance_icici v1.1), never in the Yes Bank
tables the owner's Bank card reads.

S412 (Yes Bank unlock, 26-Sep-2026): the shelf opens Yes Bank's password-locked e-statements ITSELF. The owner types the bank's statement
password once (the 'Statement passwords' card; candidates one per line) -> stmt_secret, kept in finance.db only, never printed, never logged,
never sent back to any page (the page shows 'set on <date>' and Replace). process_inbox runs an unlock step: the venv python (the only one
with pypdf) opens every locked PDF the candidates fit into <inbox>/<id>.unlocked.pdf (stmt_shelf.py unlock), the locked original stays,
and the file is then identified and read as any other; a file nothing opens reads 'locked -- no password' and is retried on every run.
The branch's own statements (not locked) are the PRIMARY Yes Bank source: when a branch copy of the same account-month is on the shelf,
the locked twin is 'duplicate of branch copy' and never read into the tables. The four non-Sanjeevni Yes Bank accounts go to
yesbank_account_statement_period/_line (their own tables) -- never bank_statement_*, which the owner's Bank card, the statement-covers
test and the NEFT bank check read as 'the Yes Bank statement' with no account filter. The ICICI anchor moves only on a statement with a
real header period and a PRINTED closing (never the pipe .txt). No secret, no account number beyond a tail, anywhere in this file.

S417 (F-636, day one of the packs page, 26-Sep-2026): the CARD statements are READ -- the statement date, the billing period and the card
number from the decrypted PDF's own text (HDFC's two layouts, ICICI's one), never the 'linked savings account XX…' line that had given the
old identifier a bank account's tail. A card's month is the month its statement is DATED in. The slot learns every card number its
statements print (the ICICI Amazon card changed number in Feb 2026: both forms, one card). The password-locked original cannot be read; its
twin is the file of the SAME NAME in the card's Decrypted folder (the owner's script keeps the name): with a twin it is 'duplicate of
decrypted'; without one its month reads 'decrypted copy not yet made' on the pack row. An .xlsx in the Credit Card Statements root is the
running Excel (folder all_txn, pack row 4's attachment) and is never asked about.

S433 (F-652, 28-Sep-2026): the Yes Bank BRANCH statement is read. The branch sends each account's month unlocked in its own layout
('STATEMENT OF ACCOUNT' ... 'Period :   01-AUG-2026 To 31-AUG-2026' ... 'A/C Number:'); the shelf could give it no month and the Yes Bank
reader refused it, so every Yes Bank row of the pack stayed 'missing'. yes_branch.py (new, parent's) identifies it -- the holder is the
left column of the 'OD Limit' line, never the proprietor line under it ('PROP MANOJ KUMAR AGARWAL HUF' on Sanjeevni's statement had put it
on the HUF slot) -- and reads it with the same four-way proof as S360; on the shared Sanjeevni tables a month already held is CHECKED, not
written twice. finance_yesbank.py is not touched. 'M/S' is dropped and single letters join ('N K PATHOLOGY' is NK PATHOLOGY). Two owner
rulings as settings (data, not code): packs.off_shelf_tails -- an account the shelf leaves alone (the owner, 28-Sep: Yes Bank current
...0460, to be closed); packs.retired_slots -- a slot that is no longer a pack row (yes_cur_clinic: that ...0460 account was the list's
third Yes Bank current). packs.hide_locked=1 -- the locked e-statement copies are counted on the passwords card, not asked about one by
one, since the branch copies are the source.

S434 (28-Sep-2026, the owner's walk of the August pack): the clinic's Yes Bank current row is back (S433 had retired it on a wrong
reading -- the account exists; the branch did not send August); the payment digest keeps only payments (receipts, invoices, bills,
renewals) -- never an OTP, a sign-in link, a failed or paused payment, an expiry or storage notice, a trial reminder, a report mail,
and NEVER a Docterz appointment line (a patient's name); an exact repeat of one mail is one row; an amount the mail did not state reads
'see the receipt'; the pharmacy's cash/UPI corrections carry their real amounts (finance_app's own keys); electricity counts the
ICICI statements AND the card sheet (All_Transactions.xlsx -- the owner: the second bill is auto-paid by the ICICI Amazon Pay card) and
reads ready only when packs.electricity_expected bills are found; the income summary names the split takings; the NEFT letter goes as
a PDF; every attachment carries a readable name; the page is folded into sections that open on a tap.

S476 (F-723, F-725, 04-Oct-2026): the statement road is WATCHED BY WHAT IT PRODUCES. The relay that files the banks' statements into
Drive runs under the owner's personal Google authorisation; when that lapsed (1-Oct-2026) its failure mail reached his personal inbox
only and nothing here knew. statement_road(): from the 3rd of a month, amber while NO bank statement has reached the shelf this month
and last month's cells are still empty or partial (it stays amber until one arrives); from the 3rd to the 10th a grey note when some
came and then none for three days while a cell is still empty; quiet otherwise (settings packs.road_from_day / packs.road_to_day /
packs.road_quiet_days). The packs page prints the line above the summary card and
/finance/health carries it as 'Bank statements reaching Drive' (statement_road_row -- one guarded call from finance_app). And one
door, the owner's only: /finance/packs/api/text/<id> answers the text the readers see of one shelf file (pdftotext -layout of the
unlocked copy when the shelf opened it) -- a statement is read in its own bytes before its reader is judged (F-720). Nothing stored.

S479 (F-725, 04-Oct-2026): the Yes Bank MONTHLY E-MAILED statement is read. A third layout ('YOUR ACCOUNT STATEMENT FROM .. TO ..',
'Primary Holder : ..', one foot line with the opening, the two totals and the closing): the shelf opened it with the owner's password,
the identifier took its bank from a narration and gave it no period, and the net-banking reader refused it. yes_monthly.py (new,
parent's) names it and reads it with the same proofs as S360 / S433, and never writes a row another file already holds (date +
amounts): the net-banking download of 1 Aug - 20 Sep sits beside September's statement on the shared Sanjeevni tables. And a person is
handed the copy he can open: the accountants' attachment, the owner's preview and Amir's pack take the UNLOCKED copy of a statement
the shelf opened (_best_path) -- the locked original is something only the owner can read.

S503 (10-Oct-2026, the owner: "passwords for ICICI and the cards on our own page, read both kinds, name what is locked"):
  * Yes Bank's CONSOLIDATED statement is read (yes_monthly.py 1.3); a file a reader REFUSED is offered again by itself whenever a
    reader is newer than the refusal (_retry_refused) -- the one-time correction is part of the design, not a step for anyone.
  * the passwords card has a box per ICICI account and per CARD as well as per Yes Bank account; a locked card original with no
    decrypted twin is opened by the shelf itself and stands as the month's card statement (nothing waits on an outside script).
  * the card names every locked file still closed, in words ('ICICI e-statement, Sep 2026'), and the owner may set one aside
    ('set aside': never tried, never counted, put back with one tap).
  * a second copy of an ICICI account-month (the bank's e-statement beside the branch's or the net-banking one) is CHECKED against
    the copy already read and never written twice; each cell shows every copy of its month side by side, and the pack uses one.

S504 (10-Oct-2026, the owner: the card sheet "is too crude ... properly categorised for me and accountant use ... like COCO is the
main petrol pump where fuel is filled"):
  * every card statement on the shelf is read LINE BY LINE here (card_lines.py) and PROVED against the statement's own printed
    debits and credits; a statement that does not prove is shown as such and never put in a pack as if whole.
  * each line carries two words: a plain one for the owner and the accountant's ledger; a merchant is named ONCE on this page
    ('Name these merchants') and every line of it, past and future, follows.
  * the pack's card sheet is built from these lines (By ledger . Every line . Statements) when every card statement dated in the
    month is proved; until then the running All_Transactions.xlsx still goes, as before.
  * the electricity bill a card pays is found in these lines first.
"""
import datetime as dt
import glob
import hashlib
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile

from flask import Blueprint, jsonify, render_template_string, request, send_file

bp = Blueprint("packs", __name__)
_db = _require = None
_unit = "packs"
HERE = os.path.dirname(os.path.abspath(__file__))
INBOX = os.environ.get("STMT_INBOX", os.path.join(HERE, "statements", "inbox"))
VPY = os.environ.get("STMT_VENV_PYTHON", "/root/wa/venv/bin/python3")      # S412: the unlock step runs under the venv (pypdf lives there only)
PACK_DIR = os.environ.get("PACKS_DIR", os.path.join(HERE, "statements", "packs"))
UPLOADS = os.environ.get("ASSETS_UPLOADS", "/root/assetapp/uploads")
ENV_PATH = os.environ.get("PACKS_ENV", "/root/wa/.env")
REPEAT_MIN = 10
DUE_DAY = 10

DDL = (
    "CREATE TABLE IF NOT EXISTS stmt_slot (id INTEGER PRIMARY KEY, key TEXT NOT NULL UNIQUE, bank TEXT NOT NULL, holder_label TEXT NOT NULL,"
    " kind TEXT NOT NULL CHECK (kind IN ('current','savings','card')), ident_tail TEXT, ident_words TEXT, owner_set TEXT, sanjeevni INTEGER NOT NULL DEFAULT 0,"
    " sort INTEGER NOT NULL DEFAULT 0)",
    "CREATE TABLE IF NOT EXISTS stmt_file ("
    " id INTEGER PRIMARY KEY, drive_id TEXT NOT NULL UNIQUE, name TEXT NOT NULL, mtime TEXT, size INTEGER, folder TEXT NOT NULL,"
    " subfolder TEXT, slot_id INTEGER, period_from TEXT, period_to TEXT, read_status TEXT, matched_status TEXT, sha256 TEXT,"
    " fetched_at TEXT, local_path TEXT, bank TEXT, holder TEXT, tail TEXT, kind TEXT, ident_how TEXT, note TEXT, identified_at TEXT,"
    " locked INTEGER NOT NULL DEFAULT 0, unlocked_path TEXT, unlocked_at TEXT)",
    # S412: the owner's statement passwords -- candidates as a JSON list, per bank; never selected into any page or log
    "CREATE TABLE IF NOT EXISTS stmt_secret (bank TEXT PRIMARY KEY, secret TEXT NOT NULL, set_by TEXT, set_at TEXT)",
    "CREATE TABLE IF NOT EXISTS stmt_secret_log (id INTEGER PRIMARY KEY, bank TEXT NOT NULL, action TEXT NOT NULL, who TEXT, at TEXT)",
    # S412: the non-Sanjeevni Yes Bank accounts' own tables (the shape of the ICICI ones + slot_key); the shared bank_statement_* stay Sanjeevni's
    "CREATE TABLE IF NOT EXISTS yesbank_account_statement_period (id INTEGER PRIMARY KEY, account_ref TEXT NOT NULL, slot_key TEXT, period_from TEXT NOT NULL,"
    " period_to TEXT NOT NULL, opening_p INTEGER, closing_p INTEGER, source_file TEXT, sha256 TEXT, ingested_at TEXT, UNIQUE (account_ref, period_from, period_to))",
    "CREATE TABLE IF NOT EXISTS yesbank_account_statement_line (id INTEGER PRIMARY KEY, account_ref TEXT NOT NULL, slot_key TEXT, txn_date TEXT NOT NULL,"
    " value_date TEXT, description TEXT NOT NULL, reference TEXT, withdrawal_p INTEGER NOT NULL DEFAULT 0, deposit_p INTEGER NOT NULL DEFAULT 0,"
    " balance_p INTEGER, is_cash_deposit INTEGER NOT NULL DEFAULT 0, source_file TEXT, sha256 TEXT, ingested_at TEXT,"
    " UNIQUE (account_ref, txn_date, reference, deposit_p, withdrawal_p))",
    "CREATE TABLE IF NOT EXISTS pack_send (id INTEGER PRIMARY KEY, month TEXT NOT NULL, sent_at TEXT NOT NULL, sent_by TEXT, what TEXT,"
    " message_id TEXT, to_list TEXT, parts INTEGER, bytes INTEGER, resend INTEGER NOT NULL DEFAULT 0)",
    "CREATE TABLE IF NOT EXISTS pack_tick (month TEXT NOT NULL, item TEXT NOT NULL, done_by TEXT, done_at TEXT, PRIMARY KEY (month, item))",
    "CREATE TABLE IF NOT EXISTS packs_item (id INTEGER PRIMARY KEY, grp TEXT NOT NULL, item TEXT NOT NULL, auto_key TEXT, sort INTEGER NOT NULL DEFAULT 0,"
    " active INTEGER NOT NULL DEFAULT 1, edited_by TEXT, edited_at TEXT, source TEXT)",
    "CREATE TABLE IF NOT EXISTS packs_done (month TEXT NOT NULL, item_id INTEGER NOT NULL, done_by TEXT, done_at TEXT, note TEXT, PRIMARY KEY (month, item_id))",
)

# the owner's §1 list: the slots. bank · holder label · kind · the words the statement must carry (ALL of them, upper-case) · Sanjeevni?
SLOTS = [
    ("yes_cur_nk",       "YES",   "NK Pathology (Yes Bank current)",           "current", "NK PATHOLOGY",                 0),
    ("yes_cur_clinic",   "YES",   "Dr Manoj Agarwal Clinic (Yes Bank current)", "current", "MANOJ AGARWAL CLINIC",         0),
    ("yes_cur_sanj",     "YES",   "Sanjeevni Medicos (Yes Bank current)",      "current", "SANJEEVNI MEDICOS",            1),
    ("yes_sav_manoj",    "YES",   "Dr Manoj Agarwal (Yes Bank savings)",       "savings", "MANOJ KUMAR AGARWAL",          0),
    ("yes_sav_bhawna",   "YES",   "Dr Bhawna Agarwal (Yes Bank savings)",      "savings", "BHAWNA AGARWAL",               0),
    ("yes_sav_huf",      "YES",   "Manoj Kumar Agarwal HUF (Yes Bank savings)", "savings", "MANOJ KUMAR AGARWAL HUF",      0),
    ("icici_sanj",       "ICICI", "Sanjeevni Medicos (ICICI)",                 "current", "SANJEEVNI MEDICOS",            1),
    ("icici_clinic",     "ICICI", "Dr Manoj Agarwal Clinic (ICICI)",           "current", "MANOJ AGARWAL CLINIC",         0),
    ("icici_nk",         "ICICI", "NK Pathology (ICICI current)",              "current", "NK PATHOLOGY",                 0),   # S411
    ("icici_personal",   "ICICI", "Dr Manoj Agarwal (ICICI)",                  "savings", "MANOJ KUMAR AGARWAL",          0),
    ("icici_bhawna",     "ICICI", "Dr Bhawna Agarwal (ICICI savings)",         "savings", "BHAWNA AGARWAL",               0),   # S411
    ("icici_huf",        "ICICI", "Manoj Kumar Agarwal HUF (ICICI savings)",   "savings", "MANOJ KUMAR AGARWAL HUF",      0),   # S411
    ("card_hdfc_regalia", "HDFC", "HDFC Business Regalia (card)",              "card",    "REGALIA",                      0),
    ("card_icici_amazon", "ICICI", "ICICI Amazon Pay (card)",                  "card",    "AMAZON",                       0),
    ("card_icici_coral",  "ICICI", "ICICI Coral (card)",                       "card",    "CORAL",                        0),
]
# S411: the words the S408 seed wrote, by key -- the seed of S411 replaces exactly these (an owner-edited value is left alone)
S408_WORDS = {"yes_cur_sanj": "SANJEEVNI", "yes_sav_manoj": "MANOJ AGARWAL", "yes_sav_bhawna": "BHAWNA", "yes_sav_huf": "HUF",
              "icici_sanj": "SANJEEVNI", "icici_personal": "MANOJ AGARWAL"}
LOCKED_NOTE = "password-protected PDF -- type the bank's statement password once under 'Statement passwords' on this page; the shelf opens it itself"
LOCKED_NO_PW_NOTE = "password-protected PDF -- none of the stored passwords opens it (replace them under 'Statement passwords')"
LOCKED_NO_PW = "locked -- no password"
DUP_OF_BRANCH = "duplicate of branch copy"
DUP_OF_DECRYPTED = "duplicate of decrypted"                    # S417
NO_TWIN = "decrypted copy not yet made"                        # S417
SECRET_BANKS = (("YES", "Yes Bank"), ("ICICI", "ICICI Bank"), ("HDFC", "HDFC Bank"))       # the banks that send locked PDFs (Yes Bank now)
SET_ASIDE = "set aside"                                              # S503: the owner's tap -- a locked file never tried, never counted
READERS = ("yes_monthly.py", "yes_branch.py", "finance_yesbank.py", "finance_icici.py", "packs.py")   # S503: a newer reader retries a refusal
YES_ACCOUNT_TABLES = ("yesbank_account_statement_period", "yesbank_account_statement_line")
STMT_FILE_COLS = (("locked", "INTEGER NOT NULL DEFAULT 0"), ("unlocked_path", "TEXT"), ("unlocked_at", "TEXT"))
PACKS_ITEM_COLS = (("due_day", "INTEGER"),)                       # S472: the day of the FOLLOWING month an item is due by
# S472: the default due day by the item's kind (the owner edits any item's day on his page; a NULL keeps this rule)
DUE_BY_KEY = {"clinic_days": 1, "sanj_stmts": 5, "clinic_stmts": 5, "scans_clear": 5, "pm_final": 5, "neft_done": 7, "sent": 10}
DUE_BY_GROUP = {"Clinic": 3, "Lab": 10, "Sanjeevni": 10, "Accountant ka samaan": 10}
HANDOVER_WORDS = ("register", "receipt book", "cheque register", "petty book", "kaagaz")   # a manual item that hands a thing over: photo
HANDOVER_DIR = os.environ.get("PACKS_HANDOVER_DIR", os.path.join(PACK_DIR, "handover"))
NOISE_TOKENS = {"M", "S", "MS", "DR", "MR", "MRS", "SHRI", "SMT", "AND", "THE", "OF"}
CARD_FOLDERS = {"HDFC Business Regalia": "card_hdfc_regalia", "ICICI Amazon Pay": "card_icici_amazon", "ICICI Card 5007": "card_icici_coral", "ICICI Coral": "card_icici_coral"}
SETTINGS = {
    "packs.accountant_to": ("", "S408 D624 -- the accountants' addresses (read from the filer script at install); a data edit changes it"),
    "packs.electricity_words": ("PVVNL,UPPCL,ELECTRICITY,BIJLI,PASCHIMANCHAL,POWER CORP", "S408 D624 -- ICICI narration words that mean an electricity auto-pay; seeded from typical words (no real ICICI statement was on the box) -- edit when the real one shows"),
    "packs.mail_limit_mb": ("20", "S408 D624 -- a pack heavier than this goes as part 1 / part 2"),
    "packs.off_shelf_tails": ("", "S433 -- account tails the shelf leaves alone, comma-separated (owner 28-Sep: 0460, to be closed)"),
    "packs.retired_slots": ("", "S433 -- slot keys that are no longer pack rows, comma-separated (owner 28-Sep: yes_cur_clinic)"),
    "packs.hide_locked": ("1", "S433 -- 1: locked e-statement copies are counted on the passwords card, not listed as questions"),
    "packs.electricity_expected": ("2", "S434 -- how many electricity bills a month has (the owner: two, one from ICICI savings, one by the Amazon Pay card)"),
    "packs.road_from_day": ("3", "S476 F-723 -- the statement road is judged from this day of the month (amber while nothing has come this month)"),
    "packs.road_to_day": ("10", "S476 F-723 -- the grey 'quiet for N days' note is said from road_from_day to this day only"),
    "packs.road_quiet_days": ("3", "S476 F-723 -- inside that window, this many days without a new bank file is said out loud (a grey note)"),
    "packs.digest_drop_words": ("OTP,SECURE LINK,UNSUCCESSFUL,FAILED,PAUSED,TURNED OFF,EXPIRING,STORAGE IS FULL,TRIAL,ACTIVATED,ENROLLMENT,ENROLMENT,"
                                "REPORT FOR,STATEMENT FROM,APPOINTMENT,REMINDER,WELCOME,VERIFY,PASSWORD",
                                "S434 -- a payment-register row whose subject carries any of these words is not a payment and never goes to the accountants"),
}
DIGEST_KEEP = ("RECEIPT", "INVOICE", "BILL", "RENEWAL", "PAYMENT SUCCESSFUL", "ORDER", "PAID", "PAYMENT RECEIVED", "TAX INVOICE")   # S434
OFF_STATUS = "off the shelf"                                      # S433
OFF_NOTE = "account ...%s is off the shelf (the owner's ruling)"   # S433


def _csv_setting(con, key):
    return {x.strip() for x in str(_setting(con, key, SETTINGS[key][0]) or "").split(",") if x.strip()}
# Shavez's checklist: the workbook could not be read by the box (Drive 404 to the service account), so the four groups are seeded
# from the brief's own words; the owner edits them on his page. auto_key ties an item to a pack row / a system fact.
ITEMS = [
    ("Sanjeevni", "Vendor payments sheet lock (mahine ka purchase final)", "pm_final"),
    ("Sanjeevni", "NEFT done aur supplier ko bata diya", "neft_done"),
    ("Sanjeevni", "Cheque register update (handed over)", None),
    ("Sanjeevni", "Pharmacy bill scan poore (Bill scan karo khaali)", "scans_clear"),
    ("Sanjeevni", "Yes Bank aur ICICI Sanjeevni statement shelf par", "sanj_stmts"),
    ("Lab", "Lab register accounts (kaagaz) accountant ko", None),
    ("Lab", "Lab receipt book (kaagaz) accountant ko", None),
    ("Lab", "NK Pathology ka Yes Bank statement shelf par", "slot:yes_cur_nk"),
    ("Clinic", "Docterz day revenue har din aa gaya", "clinic_days"),
    ("Clinic", "Clinic ka Yes Bank / ICICI statement shelf par", "clinic_stmts"),
    ("Clinic", "Petty book aur attendance close", None),
    ("Accountant ka samaan", "Income sheet + clinic ledger tayyar", "row:income"),
    ("Accountant ka samaan", "UPI totals + correction report tayyar", "row:upi"),
    ("Accountant ka samaan", "Saare bank statement shelf par", "row:stmts"),
    ("Accountant ka samaan", "Card statements (password hataye hue) + All_Transactions", "row:cards"),
    ("Accountant ka samaan", "Bijli ke do bill (ICICI se)", "row:electricity"),
    ("Accountant ka samaan", "Sanjeevni NEFT advice + letter", "row:neft"),
    ("Accountant ka samaan", "Payment digest", "row:digest"),
    ("Accountant ka samaan", "Scan kiye bill bundle (lab + expense)", "row:bundles"),
    ("Accountant ka samaan", "Pack accountant ko bheja", "sent"),
]


def init(app, db_getter, require_fn, unit="packs"):
    global _db, _require, _unit
    _db, _require, _unit = db_getter, require_fn, unit
    app.register_blueprint(bp)
    return bp


# ---------------------------------------------------------------- small things
def _today():
    v = os.environ.get("PACKS_TODAY", "")
    try:
        return dt.date.fromisoformat(v) if v else dt.date.today()
    except ValueError:
        return dt.date.today()


def _now():
    return dt.datetime.now().replace(microsecond=0).isoformat()


def _month_name(ym):
    try:
        return dt.date(int(ym[:4]), int(ym[5:7]), 1).strftime("%B %Y")
    except (TypeError, ValueError):
        return ym


def _prev_month(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    return "%04d-%02d" % ((y - 1, 12) if m == 1 else (y, m - 1))


def _month_end(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    nxt = dt.date(y + (m == 12), (m % 12) + 1, 1)
    return (nxt - dt.timedelta(days=1)).isoformat()


def _good_month(m):
    return bool(re.match(r"^\d{4}-\d{2}$", str(m or "")))


def _dmy(iso):
    try:
        return dt.date.fromisoformat(str(iso)[:10]).strftime("%d-%b-%Y")
    except (TypeError, ValueError):
        return str(iso or "")


def _inr(p):
    v = str(int(round(abs(int(p or 0)) / 100.0)))
    if len(v) > 3:
        head, tail = v[:-3], v[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        v = ",".join(parts) + "," + tail
    return "₹" + v


def _esc(s):
    return (str("" if s is None else s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


def _has(con, t):
    return con.execute("SELECT 1 FROM sqlite_master WHERE name=?", (t,)).fetchone() is not None


def _setting(con, key, default=""):
    try:
        r = con.execute("SELECT value FROM setting WHERE key=?", (key,)).fetchone()
        return (r[0] if r and r[0] is not None else default)
    except Exception:                                    # noqa: BLE001
        return default


def ensure(con):
    for d in DDL:
        con.execute(d)
    have = {r[1] for r in con.execute("PRAGMA table_info(stmt_file)")}          # S412: the columns on a shelf made before this kit
    for c, d in STMT_FILE_COLS:
        if c not in have:
            con.execute("ALTER TABLE stmt_file ADD COLUMN %s %s" % (c, d))
    have_i = {r[1] for r in con.execute("PRAGMA table_info(packs_item)")}          # S472
    for c, d in PACKS_ITEM_COLS:
        if c not in have_i:
            con.execute("ALTER TABLE packs_item ADD COLUMN %s %s" % (c, d))
    for i, (key, bank, label, kind, words, sanj) in enumerate(SLOTS):
        con.execute("INSERT OR IGNORE INTO stmt_slot (key, bank, holder_label, kind, ident_words, sanjeevni, sort) VALUES (?,?,?,?,?,?,?)",
                    (key, bank, label, kind, words, sanj, i))
    if not con.execute("SELECT 1 FROM packs_item LIMIT 1").fetchone():
        for i, (grp, item, key) in enumerate(ITEMS):
            con.execute("INSERT INTO packs_item (grp, item, auto_key, sort, source) VALUES (?,?,?,?,?)", (grp, item, key, i, "seed S408 (the brief's four groups)"))
    con.commit()


def _who(u):
    return (u or {}).get("user") or (u or {}).get("username") or ""


def _con():
    con = sqlite3.connect(os.environ.get("FINANCE_DB", os.path.join(HERE, "finance.db")), timeout=60)
    con.row_factory = sqlite3.Row
    return con


# ---------------------------------------------------------------- the shelf: identification by content
def pdf_text(path):
    try:
        r = subprocess.run(["pdftotext", "-layout", path, "-"], stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=60)
        return r.stdout.decode("utf-8", "replace")
    except Exception:                                    # noqa: BLE001
        return ""


TAIL_RES = (re.compile(r"(?:card|account|a/c)\s*(?:no\.?|number|#)?\s*[:\-]?\s*(?:[Xx*]{2,}[\s\-]*)+(\d{4})\b", re.I),
            re.compile(r"(?:account|a/c)\s*(?:no\.?|number|#)?\s*[:\-]?\s*(\d{9,18})\b", re.I),
            re.compile(r"\b\d{4}[\s\-]?[Xx*]{4}[\s\-]?[Xx*]{4}[\s\-]?(\d{4})\b"),
            re.compile(r"\b(?:[Xx*]{4}[\s\-]?){3}(\d{4})\b"))
PERIOD_RES = (re.compile(r"(?:statement\s+period|period|from)\s*[:\-]?\s*(\d{2}[/-]\d{2}[/-]\d{4})\s*(?:to|-|–)\s*(\d{2}[/-]\d{2}[/-]\d{4})", re.I),
              re.compile(r"statement\s+date\s*[:\-]?\s*(\d{2}[/-]\d{2}[/-]\d{4})", re.I))


def _iso_any(s):
    for fmt in ("%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


PIPE_ACCT_RE = re.compile(r"^\s*[Xx*\d]{2,16}?(\d{4})\|", re.M)     # S411: ICICI's pipe .txt prints the account at the start of every row


def _tokens(s):
    """Holder words as a set: upper-case, the salutations and 'M/S' dropped, SANJEEVANI folded to SANJEEVNI."""
    up = str(s or "").upper().replace("SANJEEVANI", "SANJEEVNI")
    up = re.sub(r"\bM\s*/\s*S\b\.?", " ", up)                         # S433: 'M/S.' is not a word of the name
    toks = []
    for t in re.split(r"[^A-Z0-9]+", up):
        if not t:
            continue
        if len(t) == 1 and toks and toks[-1][1]:                      # S433: 'N K PATHOLOGY' -> NK PATHOLOGY
            toks[-1] = (toks[-1][0] + t, True)
        else:
            toks.append((t, len(t) == 1))
    return {t for t, _ in toks if t not in NOISE_TOKENS}


def _header(text, n=40):
    """The statement's head: the non-empty lines BEFORE the first transaction row (a dated row, a B/F row or a pipe row) -- where the
    bank, the holder, the account and the period are printed; a narration never reaches it."""
    out = []
    for ln in (text or "").splitlines():
        if not ln.strip():
            continue
        if re.match(r"^\s*\d{2}[-–/]\d{2}[-–/]\d{4}\s", ln) or re.search(r"\bB/F\b", ln) or re.match(r"^\s*[Xx*\d]{6,}\|", ln):
            break
        out.append(ln)
        if len(out) >= n:
            break
    return "\n".join(out)


def _bank_of(up):
    """The bank named FIRST in the text (the head names the issuing bank before anything else)."""
    hits = []
    for word, bank in (("YES BANK", "YES"), ("YESBANK", "YES"), ("YES FIRST", "YES"), ("ICICI", "ICICI"), ("HDFC", "HDFC")):
        i = up.find(word)
        if i >= 0:
            hits.append((i, bank))
    return min(hits)[1] if hits else ""


def identify_text(text):
    """S411: the bank, the kind, the tail, the period and the HOLDER from the statement's head -- never from a narration.
    'Your Details With Us:' (ICICI iCRM) names the holder on the next line; otherwise the head's own words are the holder text."""
    try:                                                              # S433: the Yes Bank branch statement names itself
        import yes_branch                                             # noqa: PLC0415
        if yes_branch.is_branch(text):
            b = yes_branch.ident(text)
            return dict(bank="YES", kind=b["kind"], tail=b["tail"], period_from=b["period_from"], period_to=b["period_to"],
                        up=" ".join(b["holder"].upper().split()), holder=b["holder"], tokens=sorted(_tokens(b["holder"])))
    except ImportError:
        pass
    try:                                                              # S479 (F-725): Yes Bank's monthly e-mailed statement names itself
        import yes_monthly                                            # noqa: PLC0415
        if yes_monthly.is_monthly(text):
            b = yes_monthly.ident(text)
            return dict(bank="YES", kind=b["kind"], tail=b["tail"], period_from=b["period_from"], period_to=b["period_to"],
                        up=" ".join(b["holder"].upper().split()), holder=b["holder"], tokens=sorted(_tokens(b["holder"])))
    except ImportError:
        pass
    head = _header(text)
    hup = " ".join(head.upper().split())
    up = " ".join(text.upper().split())
    bank = _bank_of(hup) or _bank_of(up)
    kind = "card" if ("CREDIT CARD" in hup or "MINIMUM AMOUNT DUE" in hup or "CARD NUMBER" in hup) else ("savings" if "SAVINGS" in hup else ("current" if "CURRENT" in hup else ""))
    holder = ""
    hl = [ln.strip() for ln in head.splitlines()]
    for i, ln in enumerate(hl):
        if "YOUR DETAILS WITH US" in ln.upper() and i + 1 < len(hl):
            holder = hl[i + 1]
            break
    tail = ""
    for scope in (head, text):
        for rx in TAIL_RES + (PIPE_ACCT_RE,):
            m = rx.search(scope)
            if m:
                tail = m.group(1)[-4:]
                break
        if tail:
            break
    pf = pt = None
    m = PERIOD_RES[0].search(text)
    if m:
        pf, pt = _iso_any(m.group(1)), _iso_any(m.group(2))
    else:
        m = PERIOD_RES[1].search(text)
        if m:
            d = _iso_any(m.group(1))
            if d:
                pt = d
                pf = d[:8] + "01"
    return dict(bank=bank, kind=kind, tail=tail, period_from=pf, period_to=pt, up=hup, holder=holder,
                tokens=sorted(_tokens(holder) if holder else _tokens(hup)))


# ---------------------------------------------------------------- S417: the card statements, read (F-636)
_MON = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?"
_CD1 = r"(\d{1,2}\s+%s,?\s+\d{4})" % _MON          # 15 Aug, 2026        (HDFC)
_CD2 = r"(%s\s+\d{1,2},\s*\d{4})" % _MON            # August 12, 2026     (ICICI)
_CD3 = r"(\d{2}[/-]\d{2}[/-]\d{4})"                  # 15/03/2025          (HDFC's older layout)
_CARD_SD = (r"(?i)statement\s+date\s*:?\s*" + _CD1, r"(?i)statement\s+date\s*:?\s*" + _CD3, r"(?i)statement\s+date[^\n]*\n\s*" + _CD2,
            r"(?i)statement\s+date\s*:?\s*" + _CD2)
_CARD_PER = (r"(?i)billing\s+period\s*:?\s*" + _CD1 + r"\s*(?:-|–|to)\s*" + _CD1, r"(?i)statement\s+period\s*:?\s*" + _CD2 + r"\s*(?:-|–|to)\s*" + _CD2,
             r"(?i)(?:billing|statement)\s+period\s*:?\s*" + _CD3 + r"\s*(?:-|–|to)\s*" + _CD3)
_CARD_NO = (re.compile(r"(?i)card\s*(?:no\.?|number)\s*:?\s*([0-9Xx*][0-9Xx* \-]{6,24}?\d{4})\b"), re.compile(r"\b[1-9]\d{3}[Xx*]{6,10}(\d{4})\b"))


def _card_date(s):
    s = " ".join(str(s or "").replace(",", " ").replace(".", " ").split())
    for fmt in ("%d %b %Y", "%d %B %Y", "%B %d %Y", "%b %d %Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def card_read(text):
    """S417: a card statement's statement date, billing period and card number(s), from its own text. The period, when the statement
    does not print one (HDFC's older layout), is the month up to the statement date. Returns dict(statement_date, period_from,
    period_to, tails) -- dates ISO or None; tails the last four of every card number printed, in order, once each."""
    t = text or ""
    sd = pf = pt = None
    for rx in _CARD_SD:
        m = re.search(rx, t)
        if m and _card_date(m.group(1)):
            sd = _card_date(m.group(1))
            break
    for rx in _CARD_PER:
        m = re.search(rx, t)
        if m and _card_date(m.group(1)) and _card_date(m.group(2)):
            pf, pt = _card_date(m.group(1)), _card_date(m.group(2))
            break
    sd = sd or pt
    pt = pt or sd
    if sd and not pf:
        y, mo = (sd.year - 1, 12) if sd.month == 1 else (sd.year, sd.month - 1)
        last = (dt.date(sd.year, sd.month, 1) - dt.timedelta(days=1)).day
        pf = dt.date(y, mo, min(sd.day, last)) + dt.timedelta(days=1)
    tails = []
    for m in _CARD_NO[0].finditer(t):
        if m.group(1)[-4:] not in tails:
            tails.append(m.group(1)[-4:])
    for m in _CARD_NO[1].finditer(t):
        if m.group(1) not in tails:
            tails.append(m.group(1))
    iso = lambda d: d.isoformat() if d else None      # noqa: E731
    return dict(statement_date=iso(sd), period_from=iso(pf), period_to=iso(pt), tails=tails)


def _tailset(t):
    """S417: a slot may know several tails ('9012,9004' -- one card re-issued); a file's tail may carry several."""
    return {x for x in re.split(r"[^0-9]+", str(t or "")) if x}


def _running_excel(f):
    """S417: an .xlsx / .xls in the Credit Card Statements ROOT is the running Excel (All_Transactions), never a statement."""
    return f.get("folder") == "cards" and not (f.get("subfolder") or "") and str(f.get("name") or "").lower().endswith((".xlsx", ".xls"))


def _name_date(name):
    """S417: the date a card file's NAME starts with, less a day (the bank's mail comes the day after the statement) -- used only for a
    locked original that has no decrypted twin, to say which month waits for its decrypted copy. Never used to place a file."""
    m = re.match(r"^(\d{4}-\d{2}-\d{2})", str(name or ""))
    try:
        return (dt.date.fromisoformat(m.group(1)) - dt.timedelta(days=1)).isoformat() if m else None
    except ValueError:
        return None


def _original_status(con, f, slot, card=None):
    """S417: a card folder's (password-locked) original. Its twin is the file of the same name in the same card's Decrypted folder; with a
    twin it is a duplicate and takes the twin's month, without one its month waits for the decrypted copy. Returns (read_status, matched)."""
    tw = con.execute("SELECT id, period_from, period_to FROM stmt_file WHERE folder='decrypted' AND id<>? AND LOWER(name)=LOWER(?) "
                     "AND (COALESCE(subfolder,'')=COALESCE(?,'') OR slot_id=?) ORDER BY id DESC LIMIT 1",
                     (f["id"], f["name"], f.get("subfolder"), slot["id"] if slot else -1)).fetchone()
    if tw:
        if tw[2]:
            con.execute("UPDATE stmt_file SET period_from=?, period_to=? WHERE id=?", (tw[1], tw[2], f["id"]))
        return DUP_OF_DECRYPTED, "file %d is the decrypted copy" % tw[0]
    if f.get("unlocked_path") and card and card.get("statement_date"):          # S503: opened by the shelf with the stored password
        con.execute("UPDATE stmt_file SET period_from=?, period_to=? WHERE id=?", (card.get("period_from"), card.get("period_to"), f["id"]))
        return "read", "statement dated %s · opened by the shelf" % _dmy(card["statement_date"])
    sd = (card or {}).get("statement_date") or _name_date(f.get("name"))
    if sd:
        con.execute("UPDATE stmt_file SET period_to=?, period_from=COALESCE(period_from, ?) WHERE id=?", (sd, sd, f["id"]))
    return "locked original", NO_TWIN


def _has_twin(con, f):
    """S503: a card original's decrypted twin (the same name in the same card's Decrypted folder) is on the shelf."""
    return con.execute("SELECT 1 FROM stmt_file WHERE folder='decrypted' AND id<>? AND LOWER(name)=LOWER(?) AND COALESCE(subfolder,'')=COALESCE(?,'') LIMIT 1",
                       (f["id"], f["name"], f.get("subfolder"))).fetchone() is not None


def _mark_twin(con, f, slot, card):
    """S417: a decrypted card statement makes its locked original (same name, same card folder) a duplicate and gives it its month."""
    con.execute("UPDATE stmt_file SET read_status=?, matched_status=?, period_from=?, period_to=? WHERE folder='cards' AND id<>? AND LOWER(name)=LOWER(?) "
                "AND (COALESCE(subfolder,'')=COALESCE(?,'') OR slot_id=?)",
                (DUP_OF_DECRYPTED, "file %d is the decrypted copy" % f["id"], (card or {}).get("period_from"), (card or {}).get("period_to"),
                 f["id"], f["name"], f.get("subfolder"), slot["id"]))


def _learn_card_tails(con):
    """S417: every card slot learns the card numbers its decrypted statements print, the newest first ('9012,9004': one card, two numbers);
    a tail the owner set himself (his tap) is never overwritten."""
    for s in [dict(r) for r in con.execute("SELECT id, ident_tail, owner_set FROM stmt_slot WHERE kind='card'")]:
        if s["ident_tail"] and s["owner_set"] and not str(s["owner_set"]).startswith("learned"):
            continue
        seen, n = {}, 0
        for tail, pt in con.execute("SELECT tail, period_to FROM stmt_file WHERE slot_id=? AND folder='decrypted' AND read_status='read' "
                                    "AND tail IS NOT NULL AND tail<>''", (s["id"],)).fetchall():
            n += 1
            for t in [x.strip() for x in str(tail).split(",") if x.strip()]:
                if str(pt or "") >= seen.get(t, ""):
                    seen[t] = str(pt or "")
        if not seen:
            continue
        want = ",".join(sorted(seen, key=lambda t: (seen[t], t), reverse=True))
        if want != (s["ident_tail"] or ""):
            con.execute("UPDATE stmt_slot SET ident_tail=?, owner_set=? WHERE id=?", (want, "learned from the card statements (%d files) %s" % (n, _now()), s["id"]))


def place(con, ident):
    """(slot row, how) or (None, why). By the learned tail first; else by the holder's words + bank (+ kind when the text says);
    a tie is unplaced -- the owner's one tap decides."""
    if ident.get("tail") and ident["tail"] in _csv_setting(con, "packs.off_shelf_tails"):      # S433
        return None, OFF_NOTE % ident["tail"]
    slots = [dict(r) for r in con.execute("SELECT * FROM stmt_slot ORDER BY sort")]
    if ident["tail"]:
        for s in slots:
            if s["ident_tail"] and (_tailset(s["ident_tail"]) & _tailset(ident["tail"])) and (not ident["bank"] or s["bank"] == ident["bank"]):   # S417: sets
                return s, "by tail"
    have = set(ident.get("tokens") or _tokens(ident.get("holder") or ident.get("up") or ""))
    text_up = " ".join(str(ident.get("holder") or ident.get("up") or "").upper().replace("SANJEEVANI", "SANJEEVNI").split())
    cands = []
    for s in slots:
        if ident["bank"] and s["bank"] != ident["bank"]:
            continue
        if (ident["kind"] == "card") != (s["kind"] == "card"):
            continue
        raw = (s["ident_words"] or "").upper()
        words = _tokens(raw.replace(",", " "))
        if not words or not words <= have:
            continue
        # the phrase as printed ranks above a word set: 'HUF' printed beats 'MANOJ AGARWAL' assembled from 'MANOJ KUMAR AGARWAL HUF'
        phrase = 1 if all(" ".join(g.split()) in text_up for g in raw.split(",") if g.strip()) else 0
        cands.append((phrase, len(words), 1 if (ident["kind"] and s["kind"] == ident["kind"]) else 0, s))
    if not cands:
        return None, "no slot's words are in the holder text%s" % ((" '%s'" % ident["holder"]) if ident.get("holder") else "")
    cands.sort(key=lambda x: (-x[0], -x[1], -x[2]))
    if len(cands) > 1 and cands[0][:3] == cands[1][:3]:
        return None, "two slots fit the words (%s / %s)" % (cands[0][3]["holder_label"], cands[1][3]["holder_label"])
    best = cands[0][3]
    if best["ident_tail"] and ident["tail"] and not (_tailset(best["ident_tail"]) & _tailset(ident["tail"])):
        return None, "the words say %s but the tail %s is not that account's (%s)" % (best["holder_label"], ident["tail"], best["ident_tail"])
    return best, "by words"


def read_file(con, f, slot):
    """Read a placed bank statement into the bank tables; a card is filed as arrived. Returns (read_status, matched_status).
    S412: a locked file is read from its unlocked copy; a Yes Bank statement of a NON-Sanjeevni account goes to the per-account tables;
    when a branch copy (not locked) of the same account-month is on the shelf, the locked twin is 'duplicate of branch copy' and never
    read into the tables; the period the reader proved is written back to the shelf row (the pipe .txt's widened one too)."""
    if not slot or slot["kind"] == "card":
        return ("n/a" if slot and slot["kind"] == "card" else "no file"), ""
    path = f.get("unlocked_path") or f.get("local_path")
    if not path or not os.path.exists(path):
        return "no file", ""
    with open(path, "rb") as fh:
        blob = fh.read()
    try:
        if slot["bank"] == "YES":
            import finance_yesbank                        # noqa: PLC0415
            import yes_branch                             # noqa: PLC0415  S433
            import yes_monthly                            # noqa: PLC0415  S479 (F-725)
            rdr = yes_branch if yes_branch.is_branch_blob(blob) else (yes_monthly if yes_monthly.is_monthly_blob(blob) else finance_yesbank)
            parsed = rdr.parse_statement(blob)
            pf, pt = parsed["period_from"], parsed["period_to"]
            if f.get("locked"):
                dup = con.execute("SELECT id FROM stmt_file WHERE id<>? AND slot_id=? AND period_from=? AND period_to=? AND read_status='read' AND COALESCE(locked,0)=0",
                                  (f["id"], slot["id"], pf, pt)).fetchone()
                if dup:
                    con.execute("UPDATE stmt_file SET period_from=?, period_to=? WHERE id=?", (pf, pt, f["id"]))
                    return DUP_OF_BRANCH, "file %d is the branch copy" % dup[0]
            else:
                con.execute("UPDATE stmt_file SET read_status=?, matched_status=? WHERE id<>? AND slot_id=? AND period_from=? AND period_to=? AND locked=1 AND read_status='read'",
                            (DUP_OF_BRANCH, "file %d is the branch copy" % f["id"], f["id"], slot["id"], pf, pt))
            tables = None if slot.get("sanjeevni") else YES_ACCOUNT_TABLES
            res = rdr.ingest_statement(con, f["name"], blob, None, now=_now(), tables=tables)
            if tables:
                for t in tables:
                    con.execute("UPDATE %s SET slot_key=? WHERE account_ref=? AND (slot_key IS NULL OR slot_key='')" % t, (slot["key"], res["account_ref"]))
        elif slot["bank"] == "ICICI":
            import finance_icici                          # noqa: PLC0415
            # S503: the bank's own e-statement and the branch / net-banking copy of an account are BOTH kept and shown. When another
            # read copy of the same account overlaps this one's dates, a line is written only if no line of that account already has its
            # date, amounts and running balance (the two layouts print references differently) -- so nothing is ever counted twice, and
            # the days only this copy covers are still written. Its closing is checked against a held statement ending the same day.
            parsed = finance_icici.parse_statement(blob)
            pf, pt = parsed["period_from"], parsed["period_to"]
            other = con.execute("SELECT id FROM stmt_file WHERE id<>? AND slot_id=? AND read_status='read' AND period_from<=? AND period_to>=? ORDER BY id LIMIT 1",
                                (f["id"], slot["id"], pt, pf)).fetchone()
            if other:
                finance_icici.ensure_tables(con)
                fname = re.sub(r"\d{6,}", lambda m: "x" + m.group(0)[-4:], str(f["name"] or "icici.pdf"))[-120:]
                sha = hashlib.sha256(blob if isinstance(blob, bytes) else blob.encode()).hexdigest()
                held = con.execute("SELECT closing_p FROM %s WHERE account_ref=? AND period_to=? AND COALESCE(sha256,'')<>? ORDER BY id DESC LIMIT 1"
                                   % finance_icici.PERIOD_TABLE, (parsed["account_ref"], pt, sha)).fetchone()
                con.execute("INSERT OR IGNORE INTO %s (account_ref, period_from, period_to, opening_p, closing_p, source_file, sha256, ingested_at, layout, "
                            "closing_printed) VALUES (?,?,?,?,?,?,?,?,?,?)" % finance_icici.PERIOD_TABLE,
                            (parsed["account_ref"], pf, pt, parsed["opening_p"], parsed["closing_p"], fname, sha, _now(), parsed["layout"],
                             1 if parsed["closing_printed"] else 0))
                new = 0
                for ln in parsed["lines"]:
                    if con.execute("SELECT 1 FROM %s WHERE account_ref=? AND txn_date=? AND withdrawal_p=? AND deposit_p=? AND COALESCE(balance_p,-1)=COALESCE(?,-1) LIMIT 1"
                                   % finance_icici.LINE_TABLE, (parsed["account_ref"], ln["txn_date"], ln["withdrawal_p"], ln["deposit_p"], ln["balance_p"])).fetchone():
                        continue
                    cur = con.execute("INSERT OR IGNORE INTO %s (account_ref, txn_date, value_date, description, reference, withdrawal_p, deposit_p, balance_p, "
                                      "is_cash_deposit, source_file, sha256, ingested_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)" % finance_icici.LINE_TABLE,
                                      (parsed["account_ref"], ln["txn_date"], ln["value_date"], ln["description"], ln["reference"], ln["withdrawal_p"],
                                       ln["deposit_p"], ln["balance_p"], ln["is_cash_deposit"], fname, sha, _now()))
                    new += cur.rowcount if cur.rowcount and cur.rowcount > 0 else 0
                con.execute("UPDATE stmt_file SET period_from=?, period_to=? WHERE id=?", (pf, pt, f["id"]))
                close = ("closing agrees (%s)" % _inr(parsed["closing_p"]) if held and parsed.get("closing_p") is not None and int(held[0] or 0) == int(parsed["closing_p"] or 0)
                         else "CLOSING DIFFERS from the statement held to the same day -- look at both" if held and held[0] is not None
                         else "no other statement ends the same day")
                return "read", "second copy beside file %d -- %d lines, %d new, %d already held; %s" % (other[0], len(parsed["lines"]), new, len(parsed["lines"]) - new, close)
            res = finance_icici.ingest_statement(con, f["name"], blob, now=_now())
            if slot["key"] == "icici_sanj":
                _anchor_refresh(con, res)
        else:
            return "n/a", ""
    except Exception as ex:                              # noqa: BLE001
        return "refused: %s" % str(ex)[:160], ""
    if res.get("period") and res["period"][0] and res["period"][1]:
        con.execute("UPDATE stmt_file SET period_from=?, period_to=? WHERE id=?", (res["period"][0], res["period"][1], f["id"]))
    matched = "ingested %d lines" % res.get("lines", 0)
    if res.get("already_held"):                                       # S479: one line once -- what was written, and what another file held
        matched = "%d lines: %d new, %d already held from another file" % (res.get("lines", 0), res.get("new_lines", 0), res["already_held"])
    if res.get("checked_against"):                                    # S433: the shared Sanjeevni tables already held the month
        matched = ("agrees with the statement held (%s), %d lines" % (res["checked_against"], res.get("lines", 0)) if res.get("agrees")
                   else "DIFFERS from the statement held (%s): %s" % (res["checked_against"], "; ".join(res.get("differs") or [])))
    return "read", matched


def _anchor_refresh(con, res):
    """(Sanjeevni) bank_anchor for ICICI Sanjeevni follows the newest statement's closing balance (S377's row, refreshed)."""
    if not _has(con, "bank_anchor"):
        return
    # S412 (3a): only a statement whose header period is REAL (from < to) and whose closing balance is PRINTED and proven may move the
    # anchor -- never the pipe .txt (no closing printed, a degenerate header): S411's first run had moved it to 10-Sep on exactly that.
    per = res.get("period") or (None, None)
    if not res.get("closing_printed") or res.get("layout") == "pipe" or not (per[0] and per[1] and per[0] < per[1]):
        return
    r = con.execute("SELECT as_on FROM bank_anchor WHERE unit='medical' AND account='icici'").fetchone()
    if r and str(r[0] or "") >= res["period"][1]:
        return
    con.execute("INSERT INTO bank_anchor (unit, account, as_on, balance_p, source, entered_by, entered_at) VALUES ('medical','icici',?,?,?,?,?) "
                "ON CONFLICT(unit, account) DO UPDATE SET as_on=excluded.as_on, balance_p=excluded.balance_p, source=excluded.source, "
                "entered_by=excluded.entered_by, entered_at=excluded.entered_at",
                (res["period"][1], res["closing_p"], "ICICI Sanjeevni statement %s..%s, the bank's own closing balance (S408 shelf)" % tuple(res["period"]), "shelf", _now()))


def _best_path(con, fid):
    """S479: the file a person can open -- the unlocked copy when the shelf opened a locked e-statement, else the file as fetched.
    (path, name); (None, name) when neither is on disk; (None, None) when the shelf has no such file."""
    r = con.execute("SELECT local_path, name, unlocked_path FROM stmt_file WHERE id=?", (fid,)).fetchone()
    if not r:
        return None, None
    for p in (r[2], r[0]):
        if p and os.path.exists(p):
            return p, r[1]
    return None, r[1]


def _file_text(lp):
    """(text, locked): pdftotext -layout for a PDF (locked=True when it is password-protected), the file itself for .txt / .csv."""
    if not lp or not os.path.exists(lp):
        return "", False
    low = lp.lower()
    if low.endswith(".pdf"):
        try:
            r = subprocess.run(["pdftotext", "-layout", lp, "-"], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
        except Exception:                                # noqa: BLE001
            return "", False
        err = r.stderr.decode("utf-8", "replace")
        if "Incorrect password" in err or "Encrypted" in err:
            return "", True
        return r.stdout.decode("utf-8", "replace"), False
    if low.endswith((".txt", ".csv")):
        try:
            with open(lp, "rb") as fh:
                return fh.read(400000).decode("utf-8", "replace"), False
        except OSError:
            return "", False
    return "", False


def _secret_count(con):
    try:
        return con.execute("SELECT COUNT(*) FROM stmt_secret").fetchone()[0]
    except sqlite3.Error:
        return 0


def _db_path(con):
    try:
        return next((r[2] for r in con.execute("PRAGMA database_list") if r[1] == "main"), "") or ""
    except sqlite3.Error:
        return ""


def _unlock_locked(con, retry_all=True):
    """S412: the venv python opens every locked PDF the stored candidates fit (stmt_shelf.py unlock -- pypdf lives in the venv only).
    Returns how many opened now. Nothing runs without a stored secret AND a locked file still closed; the secret never leaves the database
    (the subprocess reads it there) and nothing of it is printed. retry_all=False tries only the files not yet marked 'no password'
    (the later passes of one run); every run's first pass tries every closed file again."""
    if not _secret_count(con):
        return 0
    extra = "" if retry_all else " AND COALESCE(read_status,'')<>'%s'" % LOCKED_NO_PW
    if not con.execute("SELECT 1 FROM stmt_file WHERE locked=1 AND (unlocked_path IS NULL OR unlocked_path='') AND COALESCE(read_status,'')<>'%s'%s LIMIT 1"
                       % (SET_ASIDE, extra)).fetchone():
        return 0
    con.commit()
    dbp = _db_path(con)
    if not dbp or not os.path.exists(VPY):
        return 0
    try:
        r = subprocess.run([VPY, "-B", os.path.join(HERE, "stmt_shelf.py"), "unlock"], cwd=HERE,
                           env=dict(os.environ, FINANCE_DB=dbp, STMT_INBOX=INBOX, STMT_UNLOCK_FRESH=("" if retry_all else "1")),
                           stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=600)
    except Exception:                                    # noqa: BLE001
        return 0
    m = re.search(r"opened (\d+)", r.stdout.decode("utf-8", "replace"))
    return int(m.group(1)) if m else 0


def process_inbox(con=None, verbose=False):
    """S411: identify in passes -- a file that prints only its account number (ICICI's .txt) may arrive before the statement that
    teaches the account's tail; while a pass learns a new tail, the still-unplaced files get another look (at most three passes).
    S412: after a pass, the locked PDFs get the unlock step; every file it opened gets its look in the next pass."""
    own = con is None
    con = con or _con()
    ensure(con)
    _retry_refused(con)                                               # S503
    tails = lambda: con.execute("SELECT COUNT(*) FROM stmt_slot WHERE ident_tail IS NOT NULL AND ident_tail<>''").fetchone()[0]   # noqa: E731
    out = None
    first = True
    for _ in range(4):
        t0 = tails()
        r = _process_once(con, verbose)
        out = r if out is None else dict(placed=out["placed"] + r["placed"], read=out["read"] + r["read"], refused=out["refused"] + r["refused"],
                                         skipped=out["skipped"] + r["skipped"], unplaced=r["unplaced"], unlocked=out.get("unlocked", 0))
        out.setdefault("unlocked", 0)
        opened = _unlock_locked(con, retry_all=first)
        first = False
        if opened:
            out["unlocked"] += opened
            continue
        if tails() == t0 or not r["unplaced"]:
            break
    out["cards_read"] = _read_card_lines(con)                          # S504
    if own:
        con.close()
    return out


def _process_once(con=None, verbose=False):
    """Identify every fetched file not yet identified (or re-identify the unplaced ones after the owner's tap). Idempotent."""
    own = con is None
    con = con or _con()
    ensure(con)
    out = dict(placed=0, unplaced=0, read=0, refused=0, skipped=0)
    for f in [dict(r) for r in con.execute("SELECT * FROM stmt_file WHERE (slot_id IS NULL AND COALESCE(read_status,'') NOT IN (?, ?)) OR read_status IS NULL ORDER BY id", (OFF_STATUS, SET_ASIDE))]:
        if f["folder"] == "all_txn" or _running_excel(f):         # S417: an .xlsx in the Credit Card Statements root IS the running Excel
            con.execute("UPDATE stmt_file SET folder='all_txn', read_status='n/a', matched_status='', note=NULL, identified_at=? WHERE id=?", (_now(), f["id"]))
            out["skipped"] += 1
            continue
        lp = f.get("unlocked_path") or f["local_path"] or ""          # S412: the unlocked copy when the venv opened it
        text, locked = _file_text(lp)
        if locked and not f.get("locked") and (f["folder"] != "cards" or not _has_twin(con, f)):   # S503: a card original with no twin is opened by the shelf
            con.execute("UPDATE stmt_file SET locked=1 WHERE id=?", (f["id"],))
            f["locked"] = 1
        ident = identify_text(text) if text else dict(bank="", kind="", tail="", period_from=None, period_to=None, up="", holder="")
        card = card_read(text) if (text and (f["folder"] in ("cards", "decrypted") or ident.get("kind") == "card")) else None     # S417
        if card and card["statement_date"]:
            ident.update(kind="card", tail=",".join(card["tails"]), period_from=card["period_from"], period_to=card["period_to"])
        slot, how = place(con, ident) if text else (None, ((LOCKED_NO_PW_NOTE if _secret_count(con) else LOCKED_NOTE) if locked else ("not a readable PDF" if lp else "no file")))
        if f["folder"] in ("cards", "decrypted") and not slot and f["subfolder"] in CARD_FOLDERS:
            slot = dict(con.execute("SELECT * FROM stmt_slot WHERE key=?", (CARD_FOLDERS[f["subfolder"]],)).fetchone() or {}) or None
            how = "by the card folder" if slot else how
        con.execute("UPDATE stmt_file SET bank=?, kind=?, tail=?, period_from=?, period_to=?, holder=?, slot_id=?, ident_how=?, note=?, identified_at=? WHERE id=?",
                    (ident["bank"], ident["kind"] or (slot["kind"] if slot else ""), ident["tail"], ident["period_from"], ident["period_to"],
                     (slot["holder_label"] if slot else (ident.get("holder") or "")), (slot["id"] if slot else None), (how if slot else None), (None if slot else how), _now(), f["id"]))
        if slot:
            out["placed"] += 1
            if how == "by words" and ident["tail"] and not (slot.get("ident_tail") or "") and f["folder"] == "bank":   # S411: the tail is learned
                con.execute("UPDATE stmt_slot SET ident_tail=?, owner_set=? WHERE id=? AND (ident_tail IS NULL OR ident_tail='')",
                            (ident["tail"], "learned from the statement (file %d) %s" % (f["id"], _now()), slot["id"]))
            if f["folder"] == "cards":
                rs, ms = _original_status(con, f, slot, card)                       # S417
            elif slot.get("kind") == "card":
                rs, ms = (("read", "statement dated %s" % _dmy(card["statement_date"])) if (card and card["statement_date"])
                          else ("n/a", "the statement date could not be read from the PDF"))
                _mark_twin(con, f, slot, card)
            else:
                rs, ms = read_file(con, f, slot)
            con.execute("UPDATE stmt_file SET read_status=?, matched_status=? WHERE id=?", (rs, ms, f["id"]))
            out["read" if rs == "read" else ("refused" if rs.startswith("refused") else "skipped")] += 1
        else:
            if str(how).startswith(OFF_NOTE.split("%s")[0]):           # S433: the owner's ruling, not a question
                con.execute("UPDATE stmt_file SET read_status=? WHERE id=?", (OFF_STATUS, f["id"]))
                out["skipped"] += 1
                continue
            out["unplaced"] += 1
            # S412: 'locked -- no password' is the unlock step's verdict (stmt_shelf.py unlock); it stays until a new password resets it
            con.execute("UPDATE stmt_file SET read_status=? WHERE id=?", ((LOCKED_NO_PW if (locked and f.get("read_status") == LOCKED_NO_PW) else None), f["id"]))
    _learn_card_tails(con)                                                 # S417
    con.commit()
    if own:
        con.close()
    return out


def assign(con, file_id, slot_id, who):
    """The owner's one tap: this unplaced file is that account -> the slot learns the tail for ever; the file is read."""
    ensure(con)
    f = con.execute("SELECT * FROM stmt_file WHERE id=?", (file_id,)).fetchone()
    s = con.execute("SELECT * FROM stmt_slot WHERE id=?", (slot_id,)).fetchone()
    if not f or not s:
        return dict(ok=False, error="no_such"), 404
    f, s = dict(f), dict(s)
    if f["tail"]:
        con.execute("UPDATE stmt_slot SET ident_tail=?, owner_set=? WHERE id=?", (f["tail"], "%s %s" % (who, _now()), s["id"]))
    locked = str(f.get("note") or "").startswith("password-protected") and not f.get("unlocked_path")     # S412: an opened one reads as any other
    con.execute("UPDATE stmt_file SET slot_id=?, holder=?, ident_how=?, note=NULL, identified_at=? WHERE id=?", (s["id"], s["holder_label"], "the owner's tap", _now(), f["id"]))
    rs, ms = ("locked original", "") if (f["folder"] == "cards" or locked) else read_file(con, f, s)
    con.execute("UPDATE stmt_file SET read_status=?, matched_status=? WHERE id=?", (rs, ms, f["id"]))
    con.commit()
    return dict(ok=True, file=f["id"], slot=s["key"], tail_learned=f["tail"] or None, read_status=rs), 200


# ---------------------------------------------------------------- the month cells
def _stitch(fs, lo, hi):
    """S471: the statements of one account that TOGETHER cover the month lo..hi, each starting on or before the day after the
    previous one ends (ICICI's cycle statements: 11th -> 10th, 15th -> 14th). Returns the pieces in date order, or []."""
    cand = sorted([x for x in fs if x.get("period_from") and x.get("period_to")], key=lambda x: (x["period_from"], x["period_to"]))
    for i, first in enumerate(cand):
        if first["period_from"] > lo:
            break
        chain, end = [first], first["period_to"]
        for nxt in cand[i + 1:]:
            if end >= hi:
                break
            if nxt["period_to"] <= end:
                continue
            if nxt["period_from"] <= _iso_plus(end, 1):
                chain.append(nxt)
                end = nxt["period_to"]
        if end >= hi and len(chain) > 1:
            return chain
    return []


def _iso_plus(iso, days):
    try:
        return (dt.date.fromisoformat(iso) + dt.timedelta(days=days)).isoformat()
    except (TypeError, ValueError):
        return iso


def _partial_words(f, lo, hi):
    """S471: what a single cycle statement leaves uncovered, in the owner's words, with the rough date the next one is due."""
    pf, pt = f.get("period_from") or "", f.get("period_to") or ""
    if pf <= lo and pt < hi:
        try:
            d0 = dt.date.fromisoformat(pt)
            nxt_from, nxt_to = d0 + dt.timedelta(days=1), (d0 + dt.timedelta(days=32)).replace(day=min(d0.day, 28))
            return "%s → %s on the shelf; the %s → %s statement completes the month (due from the bank about %s)" % (
                _dmy(pf), _dmy(pt), _dmy(nxt_from.isoformat()), _dmy(nxt_to.isoformat()), _dmy((nxt_to + dt.timedelta(days=1)).isoformat()))
        except ValueError:
            pass
    if pf > lo:
        return "%s → %s on the shelf; the earlier statement (up to %s) is not on the shelf" % (_dmy(pf), _dmy(pt), _dmy(_iso_plus(pf, -1)))
    return "only %s → %s on the shelf, not the whole month" % (_dmy(pf), _dmy(pt))


def cells(con, month):
    """One cell per slot for the month: the file (the one whose period covers the month, or falls in it) and its state."""
    ensure(con)
    lo, hi = month + "-01", _month_end(month)
    out = []
    retired = _csv_setting(con, "packs.retired_slots")                                   # S433
    for s in [dict(r) for r in con.execute("SELECT * FROM stmt_slot ORDER BY sort")]:
        if s["key"] in retired:
            continue
        fs = [dict(r) for r in con.execute(
            "SELECT * FROM stmt_file WHERE slot_id=? AND folder<>'cards' AND period_to IS NOT NULL AND period_from<=? AND period_to>=? "
            "AND COALESCE(read_status,'')<>'duplicate of branch copy' ORDER BY period_to DESC, COALESCE(locked,0) ASC, id DESC", (s["id"], hi, lo))]   # S412: the branch copy first
        full = [x for x in fs if x["period_from"] <= lo and x["period_to"] >= hi]   # S411: the month's statement covers the whole month
        pieces = _stitch(fs, lo, hi) if (not full and s["kind"] != "card") else []   # S471: two cycle statements together cover it
        if s["kind"] == "card":        # S417: a card's month is the month its statement is DATED in (the statement date = period_to)
            fs = [dict(r) for r in con.execute("SELECT * FROM stmt_file WHERE slot_id=? AND (folder='decrypted' OR (folder='cards' AND "
                                               "COALESCE(unlocked_path,'')<>'' AND read_status='read')) AND period_to BETWEEN ? AND ? "   # S503
                                               "ORDER BY period_to DESC, id DESC", (s["id"], lo, hi))]
            full = fs[:1]
        f = full[0] if full else (pieces[-1] if pieces else (fs[0] if fs else None))
        partial = bool(f) and not full and not pieces and s["kind"] != "card"
        state = "empty"
        if partial:
            state = "partial"
        elif pieces:                                                                 # S471: every piece must be read
            state = "read" if all(x["read_status"] == "read" for x in pieces) else "arrived"
        elif f:
            state = "arrived"
            if f["read_status"] == "read" or (s["kind"] == "card" and f["read_status"] in ("n/a", None, "")):
                state = "read"
            if s["kind"] == "card" and f:
                state = "read"
            if s["sanjeevni"] and f["read_status"] == "read":
                state = "matched"
        no_dec = bool(s["kind"] == "card" and not f and con.execute("SELECT 1 FROM stmt_file WHERE slot_id=? AND folder='cards' AND matched_status=? "
                                                            "AND period_to BETWEEN ? AND ? LIMIT 1", (s["id"], NO_TWIN, lo, hi)).fetchone())   # S417
        twin_missing = False
        if s["kind"] == "card":
            orig = [dict(r) for r in con.execute("SELECT * FROM stmt_file WHERE slot_id=? AND folder='cards' ORDER BY mtime DESC, id DESC LIMIT 1", (s["id"],))]
            dec = [dict(r) for r in con.execute("SELECT * FROM stmt_file WHERE slot_id=? AND folder='decrypted' ORDER BY mtime DESC, id DESC LIMIT 1", (s["id"],))]
            if orig and not (orig[0].get("unlocked_path") and orig[0].get("read_status") == "read") and (not dec or (dec[0]["mtime"] or "") < (orig[0]["mtime"] or "")):   # S503
                twin_missing = True
        out.append(dict(slot=s["key"], slot_id=s["id"], words=s["ident_words"], label=s["holder_label"], bank=s["bank"], kind=s["kind"], sanjeevni=s["sanjeevni"],
                        tail=(((s["ident_tail"] or "").replace(",", " / ") or None) if s["kind"] == "card" else s["ident_tail"]),   # S417: one card, every number
                        state=state, file=(dict(id=f["id"], name=f["name"], period_from=f["period_from"], period_to=f["period_to"],
                                                read_status=f["read_status"], matched_status=f["matched_status"], locked=bool(f.get("locked"))) if f else None),
                        twin_missing=twin_missing, no_decrypted=no_dec,
                        copies=([dict(id=x["id"], what=_copy_kind(x["name"], x["folder"]), read_status=x["read_status"], matched_status=x["matched_status"],
                                      period_from=x["period_from"], period_to=x["period_to"])
                                 for x in con.execute("SELECT * FROM stmt_file WHERE slot_id=? AND folder='bank' AND period_to IS NOT NULL AND period_from<=? "
                                                      "AND period_to>=? AND id<>? ORDER BY id", (s["id"], hi, lo, f["id"]))] if (f and s["kind"] != "card") else []),   # S503
                        pieces=[dict(id=x["id"], name=x["name"], period_from=x["period_from"], period_to=x["period_to"], read_status=x["read_status"]) for x in pieces],   # S471
                        partial_words=(_partial_words(f, lo, hi) if partial else "")))
    return out


def unplaced(con):
    ensure(con)
    return [dict(id=r["id"], name=r["name"], folder=r["folder"], subfolder=r["subfolder"], bank=r["bank"], kind=r["kind"], tail=r["tail"],
                 period_from=r["period_from"], period_to=r["period_to"], why=r["note"], fetched_at=r["fetched_at"], holder=r["holder"],
                 locked=str(r["note"] or "").startswith("password-protected"), no_password=(r["read_status"] == LOCKED_NO_PW))
            for r in con.execute("SELECT * FROM stmt_file WHERE slot_id IS NULL AND folder<>'all_txn' AND NOT (folder='cards' AND COALESCE(subfolder,'')='' "
                                 "AND LOWER(name) LIKE '%.xls%') AND COALESCE(read_status,'') NOT IN (?, 'set aside') "
                                 "AND NOT (?='1' AND COALESCE(locked,0)=1 AND (unlocked_path IS NULL OR unlocked_path='')) ORDER BY id DESC LIMIT 60",
                                 (OFF_STATUS, str(_setting(con, "packs.hide_locked", SETTINGS["packs.hide_locked"][0]))))]


def secret_keys(con):
    """S471: the keys a password may be stored under -- one per Yes Bank account on the shelf ('YES:<slot key>', the account's
    label) first, then the three banks. Data from stmt_slot; never a value."""
    out = []
    retired = _csv_setting(con, "packs.retired_slots")
    for r in con.execute("SELECT key, holder_label FROM stmt_slot WHERE bank='YES' AND kind<>'card' ORDER BY sort"):
        if r[0] not in retired:
            out.append(("YES:" + r[0], r[1], True))
    for r in con.execute("SELECT key, holder_label FROM stmt_slot WHERE bank='ICICI' AND kind<>'card' ORDER BY sort"):   # S503
        if r[0] not in retired:
            out.append(("ICICI:" + r[0], r[1], True))
    for r in con.execute("SELECT key, holder_label FROM stmt_slot WHERE kind='card' ORDER BY sort"):                     # S503
        if r[0] not in retired:
            out.append(("CARD:" + r[0], r[1], True))
    out += [(b, l + " (any other)", False) for b, l in SECRET_BANKS]
    return out


def secret_state(con):
    """S412: what the page may know of the passwords -- per key: set or not, by whom, when, how many candidates, how many files
    it opened. NEVER the value. S471: one key per Yes Bank account, then the three banks."""
    ensure(con)
    have = {r[0]: r for r in con.execute("SELECT bank, set_by, set_at, secret FROM stmt_secret")}
    opened_by = {r[0]: r[1] for r in con.execute("SELECT bank, COUNT(*) FROM stmt_secret_log WHERE action LIKE 'opened file %' GROUP BY bank")}
    out = []
    for bank, label, is_account in secret_keys(con):
        r = have.get(bank)
        n = 0
        if r:
            try:
                n = len(json.loads(r[3]))
            except (TypeError, ValueError):
                n = 1
        out.append(dict(bank=bank, label=label, account=is_account, set=bool(r), set_by=(r[1] if r else None), set_at=(r[2] if r else None), candidates=n,
                        opened=int(opened_by.get(bank, 0))))
    lk = con.execute("SELECT COUNT(*), SUM(CASE WHEN read_status=? THEN 1 ELSE 0 END) FROM stmt_file WHERE locked=1 AND (unlocked_path IS NULL OR unlocked_path='') "
                     "AND COALESCE(read_status,'')<>?", (LOCKED_NO_PW, SET_ASIDE)).fetchone()
    opened = con.execute("SELECT COUNT(*) FROM stmt_file WHERE locked=1 AND unlocked_path IS NOT NULL AND unlocked_path<>''").fetchone()[0]
    files, aside = _locked_list(con)                                  # S503: every closed file named, in words
    return dict(banks=out, locked_open=int(lk[0] or 0), no_password=int(lk[1] or 0), opened=int(opened or 0), locked_files=files, set_aside=aside)


def _copy_kind(name, folder=""):
    """S503: what a file IS, in words, from the shape of its name only (used to name a copy -- never to place or read a file)."""
    n = str(name or "")
    m = re.search(r"Statement_(\d{4})MTH(\d{2})_", n)
    if m:
        return "ICICI e-statement, %s" % _month_name("%s-%s" % (m.group(1), m.group(2)))
    if re.search(r"_YFB?_quarterly", n, re.I):
        return "Yes Bank quarterly statement"
    if re.search(r"\bCS_\d+_", n):
        return "Yes Bank consolidated e-statement"
    if "CASA_" in n:
        return "Yes Bank monthly e-statement"
    if folder == "cards":
        return "card statement (locked original)"
    if re.search(r"_0192\d{8}_\d{4}-\d{2}-\d{2}_\d{4}-\d{2}-\d{2}_", n):
        return "ICICI net-banking statement"
    if n.lower().endswith(".txt"):
        return "ICICI text statement"
    return "branch statement"


def _locked_list(con):
    """S503: the locked files still closed (and the ones the owner set aside), each named in words, its account number masked."""
    files, aside = [], []
    for r in con.execute("SELECT id, name, folder, subfolder, fetched_at, read_status, note FROM stmt_file WHERE locked=1 AND "
                         "(unlocked_path IS NULL OR unlocked_path='') ORDER BY id"):
        d = dict(id=r[0], name=re.sub(r"\d{6,}", lambda m: "x" + m.group(0)[-4:], r[1] or ""), what=_copy_kind(r[1], r[2]),
                 card=r[3] if r[2] == "cards" else "", arrived=(r[4] or "")[:10], why=(r[6] or "") if r[5] == SET_ASIDE else
                 ("no stored password opens it yet" if r[5] == LOCKED_NO_PW else "waiting for its first try"))
        (aside if r[5] == SET_ASIDE else files).append(d)
    return files, aside


def set_aside(con, fid, back, who, why=""):
    """S503: the owner's tap on a locked file -- 'set aside' (never tried, never counted) or put back (tried again on the next run)."""
    ensure(con)
    r = con.execute("SELECT id, locked, unlocked_path, read_status FROM stmt_file WHERE id=?", (fid,)).fetchone()
    if not r:
        return dict(ok=False, error="no_such"), 404
    if not r[1] or r[2]:
        return dict(ok=False, error="not_a_closed_locked_file"), 400
    if back:
        con.execute("UPDATE stmt_file SET read_status=NULL, note=NULL WHERE id=? AND read_status=?", (fid, SET_ASIDE))
    else:
        con.execute("UPDATE stmt_file SET read_status=?, note=? WHERE id=?", (SET_ASIDE, ("set aside by %s %s%s" % (who, _now()[:16], (" -- " + why[:120]) if why else "")), fid))
    con.commit()
    return dict(ok=True, file=fid, set_aside=not back), 200


def _retry_refused(con):
    """S503: a file a reader refused is offered again when any reader is newer than that refusal -- so a reader that learns a layout
    reads every file it refused before, by itself (the one-time correction carried by the design). Returns how many were re-offered."""
    try:
        newest = max(os.path.getmtime(os.path.join(HERE, x)) for x in READERS if os.path.exists(os.path.join(HERE, x)))
    except ValueError:
        return 0
    stamp = dt.datetime.fromtimestamp(newest).replace(microsecond=0).isoformat()
    cur = con.execute("UPDATE stmt_file SET read_status=NULL WHERE read_status LIKE 'refused%' AND COALESCE(identified_at,'')<?", (stamp,))
    con.commit()
    return cur.rowcount or 0


def _read_card_lines(con):
    """S504: every card statement on the shelf that is read (the decrypted twin, or the original the shelf opened) goes through
    card_lines.py -- every line, proved against its own printed totals. Idempotent (card_stmt_file remembers each shelf file); the
    same statement under two files is read once. Returns how many statements were read now. Never fatal."""
    try:
        import card_lines                                 # noqa: PLC0415
        card_lines.ensure(con)
    except Exception:                                    # noqa: BLE001
        return 0
    done = {r[0] for r in con.execute("SELECT file_id FROM card_stmt_file")}
    n = 0
    for fid, key in con.execute("SELECT f.id, s.key FROM stmt_file f JOIN stmt_slot s ON s.id=f.slot_id WHERE s.kind='card' "
                                "AND f.read_status='read' ORDER BY f.id").fetchall():
        if fid in done:
            continue
        p, _name = _best_path(con, fid)
        text, locked = _file_text(p or "")
        if not text or locked:
            continue
        try:
            _sid, _parsed, new = card_lines.ingest(con, fid, key, text)
        except Exception as ex:                          # noqa: BLE001  (a layout it does not know: said on the row, never fatal)
            con.execute("UPDATE stmt_file SET note=? WHERE id=?", (("card lines not read: %s" % ex)[:200], fid))
            continue
        n += 1 if new else 0
    con.commit()
    return n


def _card_labels(con):
    return {r[0]: re.sub(r"\s*\(card\)\s*$", "", r[1] or r[0]) for r in con.execute("SELECT key, holder_label FROM stmt_slot WHERE kind='card'")}


def card_spends(con, month):
    """S504: the month's card spends as the page shows them (card_lines.month_view); {} when the reader is not on this server."""
    try:
        import card_lines                                 # noqa: PLC0415
    except Exception:                                    # noqa: BLE001
        return {}
    return card_lines.month_view(con, month, _card_labels(con))


def build_card_sheet(con, month):
    """S504: the accountants' card workbook for the month -- (bytes, why_not). Built only when every card statement dated in the
    month is read and proved; otherwise (None, why) and the pack keeps the running All_Transactions.xlsx."""
    try:
        import card_lines                                 # noqa: PLC0415
    except Exception:                                    # noqa: BLE001
        return None, "the card reader is not on this server"
    v, sheets = card_lines.month_sheet_rows(con, month, _card_labels(con))
    cards = [r[0] for r in con.execute("SELECT key FROM stmt_slot WHERE kind='card' ORDER BY sort")]
    have = {s["slot"] for s in v["statements"]}
    miss = [k for k in cards if k not in have]
    if not v["statements"]:
        return None, "no card statement dated in %s has been read yet" % _month_name(month)
    if miss:
        return None, "not every card's statement of %s is read yet (%s)" % (_month_name(month), ", ".join(_card_labels(con).get(k, k) for k in miss))
    bad = [s["label"] for s in v["statements"] if not s["proof_ok"]]
    if bad:
        return None, "a statement does not add up to its own totals (%s) -- look before it goes" % ", ".join(bad)
    return _xlsx(sheets), None


def set_secret(con, bank, text, who):
    """S412: the owner's candidates for one bank (one per line, at most 10, 64 chars each) replace what was stored; the log says set /
    replaced without the value; the locked files still closed are retried at once. Returns (body, code) -- the body never carries the value."""
    ensure(con)
    bank = str(bank or "").strip()[:40]                                 # S471: 'YES:<slot key>' keeps its case; a bank name is upper
    bank = bank if ":" in bank else bank.upper()[:8]
    cands = []
    for ln in str(text or "").splitlines():
        ln = ln.strip()
        if ln and ln[:64] not in cands:
            cands.append(ln[:64])
    if bank not in {k for k, _l, _a in secret_keys(con)} or not cands or len(cands) > 10:
        return dict(ok=False, error="bad_request"), 400
    had = con.execute("SELECT 1 FROM stmt_secret WHERE bank=?", (bank,)).fetchone() is not None
    con.execute("INSERT INTO stmt_secret (bank, secret, set_by, set_at) VALUES (?,?,?,?) ON CONFLICT(bank) DO UPDATE SET secret=excluded.secret, "
                "set_by=excluded.set_by, set_at=excluded.set_at", (bank, json.dumps(cands), who, _now()))
    con.execute("INSERT INTO stmt_secret_log (bank, action, who, at) VALUES (?,?,?,?)", (bank, ("replaced" if had else "set"), who, _now()))
    con.execute("UPDATE stmt_file SET read_status=NULL WHERE locked=1 AND (unlocked_path IS NULL OR unlocked_path='') AND COALESCE(read_status,'')<>?", (SET_ASIDE,))
    con.commit()
    res = process_inbox(con)
    return dict(ok=True, bank=bank, action=("replaced" if had else "set"), candidates=len(cands), result=res, secrets=secret_state(con)), 200


# ---------------------------------------------------------------- the pack: builders
def _xlsx(sheets):
    """[(title, [row, ...]), ...] -> xlsx bytes (openpyxl, on both pythons)."""
    from openpyxl import Workbook                        # noqa: PLC0415
    from openpyxl.styles import Font                     # noqa: PLC0415
    wb = Workbook()
    first = True
    for title, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title[:30]
        for i, r in enumerate(rows):
            ws.append(list(r))
            if i == 0:
                for c in ws[1]:
                    c.font = Font(bold=True)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _tender(row):
    out = {"cash": 0, "online": 0, "card": 0, "split": 0}
    known = {"Cash": "cash", "Online Payment": "online", "Debit Card": "card", "Credit Card": "card", "Net Banking": "online", "Patient APP": "online", "Wallet": "online", "Split Payment": "split"}
    try:
        raw = json.loads(row["tender_json"] or "{}")
    except (ValueError, TypeError):
        return out
    for label, p in raw.items():
        b = known.get(str(label).strip())
        if b:
            out[b] += int(p)
    return out


def build_income(con, month):
    """(1) the doctor's monthly income sheet + the date-wise clinic ledger, physiotherapy excluded (it is its own unit)."""
    if not _has(con, "clinic_day_revenue"):
        return None, "the clinic day-revenue table is not on this server"
    rows = [dict(r) for r in con.execute("SELECT * FROM clinic_day_revenue WHERE business_date BETWEEN ? AND ? ORDER BY business_date", (month + "-01", _month_end(month)))]
    if not rows:
        return None, "no clinic day revenue rows for %s" % _month_name(month)
    ledger = [("Date", "Consultations (n)", "Consultations Rs", "X-ray (n)", "X-ray Rs", "Procedures (n)", "Procedures Rs", "Total (n)", "Total Rs", "Cash Rs", "Online Rs", "Card Rs", "Split Rs")]
    tot = [0] * 12
    for r in rows:
        t = _tender(r)
        vals = [r["cons_count"] or 0, (r["cons_amount_p"] or 0) / 100.0, r["xray_count"] or 0, (r["xray_amount_p"] or 0) / 100.0, r["proc_count"] or 0, (r["proc_amount_p"] or 0) / 100.0,
                r["total_count"] or 0, (r["total_amount_p"] or 0) / 100.0, t["cash"] / 100.0, t["online"] / 100.0, t["card"] / 100.0, t["split"] / 100.0]
        tot = [a + b for a, b in zip(tot, vals)]
        ledger.append([r["business_date"]] + vals)
    ledger.append(["TOTAL"] + tot)
    income = [("Dr Manoj Agarwal Clinic — income %s" % _month_name(month), ""), ("Days with takings", len(rows)), ("Consultations Rs", tot[1]), ("X-ray Rs", tot[3]),
              ("Procedures Rs", tot[5]), ("Total Rs", tot[7]), ("of which cash Rs", tot[8]), ("of which online Rs", tot[9]), ("of which card Rs", tot[10]), ("of which split Rs (part cash, part online)", tot[11]),
              ("Physiotherapy", "excluded (its own register)"), ("Paper OPD register", "the accountants' own add-on")]
    return _xlsx([("Income", income), ("Ledger date-wise", ledger)]), None


def build_upi_corrections(con, month):
    """(2) UPI totals per unit + the cash/UPI correction report (finance_app's own rows, filtered to the month)."""
    if not _has(con, "upi_txn"):
        return None, "no UPI feed on this server"
    upi = [("Unit", "Transactions", "UPI Rs")]
    for u, n, s in con.execute("SELECT unit, COUNT(*), COALESCE(SUM(amount_p),0) FROM upi_txn WHERE substr(txn_date,1,7)=? GROUP BY unit ORDER BY unit", (month,)):
        upi.append((u, n, (s or 0) / 100.0))
    if len(upi) == 1:
        return None, "no UPI transactions in %s" % _month_name(month)
    corr = [("Date", "Change", "Amount Rs", "What happened", "Likely bills", "Marg says UPI", "Bank says UPI", "Status")]
    try:
        import finance_app as fa                          # noqa: PLC0415
        for r in fa._correction_rows(con, 400, include_resolved=True):
            d = str(r.get("date") or r.get("business_date") or "")[:10]
            if d[:7] != month:
                continue
            corr.append((d, r.get("change") or "", abs(int(r.get("diff_p") or 0)) / 100.0, r.get("says") or "",   # S434: finance_app's keys
                         r.get("candidates") or "", r.get("declared_upi") or "", r.get("bank_upi") or "", r.get("status") or ""))
    except Exception as ex:                              # noqa: BLE001
        corr.append(("(the correction rows could not be read: %s)" % str(ex)[:80], "", "", "", "", "", ""))
    return _xlsx([("UPI totals", upi), ("Pharmacy cash-UPI corrections", corr)]), None


def build_digest(con, month):
    """(7) the monthly payment digest: the payment register's rows of the month (S234's derived table)."""
    if not _has(con, "payment_register"):
        return None, "the payment register is not on this server"
    rows = [("Date", "Vendor", "Description", "Amount Rs", "Attachment in Drive")]
    kept, dropped, total = digest_rows(con, month)
    for r in kept:
        rows.append((r["date"], r["vendor"], r["desc"], r["amount"] if r["amount"] is not None else "see the receipt", "yes" if r["in_drive"] else ""))
    if len(rows) == 1:
        return None, "no payments in the payment register for %s" % _month_name(month)
    rows.append(("", "", "TOTAL (rows with an amount)", total, ""))
    return _xlsx([("Payments %s" % month, rows)]), None


def digest_rows(con, month):
    """S434: the payment register's rows of the month that ARE payments. (kept, dropped, total Rs)."""
    drop = [w.strip().upper() for w in _setting(con, "packs.digest_drop_words", SETTINGS["packs.digest_drop_words"][0]).split(",") if w.strip()]
    kept, dropped, seen, total = [], [], set(), 0.0
    for r in con.execute("SELECT date_iso, vendor, description, amount_raw, amount_paise, in_drive, gmail_link FROM payment_register "
                         "WHERE substr(date_iso,1,7)=? ORDER BY date_iso, row_no", (month,)):
        desc = str(r[2] or "").split(" | ")[0].strip()
        up = (str(r[1] or "") + " " + desc).upper()
        amt = (r[4] / 100.0) if r[4] else None
        why = None
        if desc.startswith("[NOT A PAYMENT]"):                         # S510: the Janitor's own mark (v2.4, 10-Oct-2026)
            why = "marked not a payment in the sheet"
        elif desc.startswith("[DUPLICATE"):
            why = "the same mail twice (marked in the sheet)"
        elif "APPOINTMENT" in up:
            why = "a patient appointment (never to the accountants)"
        elif any(w in up for w in drop):
            why = "not a payment"
        elif not any(k in up for k in DIGEST_KEEP):
            why = "not a receipt, invoice or bill"
        key = (r[6] or "") or (r[0], r[1], desc, r[4])
        if not why and key in seen:
            why = "the same mail twice"
        if why:
            dropped.append(dict(date=r[0], vendor=r[1], why=why))
            continue
        seen.add(key)
        kept.append(dict(date=r[0], vendor=r[1], desc=desc, amount=amt, in_drive=r[5]))
        total += amt or 0
    return kept, dropped, round(total, 2)


def build_neft(con, month):
    """(6, Sanjeevni) the NEFT advice Excel + the covering letter, from purchase_app's own routes' code."""
    try:
        import purchase_app as pa                         # noqa: PLC0415
        st = pa._month_status(con, month)
        if st["status"] != "final":
            return None, "%s is not finalised on the pay sheet" % _month_name(month)
        blob, total, missing = pa._advice_xlsx_s267(con, month)
        letter = pa._letter_html_s266(con, month, total)
        return [("Sanjeevni NEFT advice - %s.xlsx" % _month_name(month), blob, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("Sanjeevni NEFT letter to the bank - %s.pdf" % _month_name(month), _letter_pdf(letter), "application/pdf")], None   # S434
    except Exception as ex:                              # noqa: BLE001
        return None, "the pay sheet could not be read (%s)" % str(ex)[:80]


def electricity_lines(con, month):
    """(5) the ICICI statement lines whose narration matches packs.electricity_words -> one line each."""
    words = [w.strip().upper() for w in _setting(con, "packs.electricity_words", SETTINGS["packs.electricity_words"][0]).split(",") if w.strip()]
    # the ICICI accounts = the tails the ICICI slots learned + the tails of every ICICI statement placed on the shelf (never a Yes Bank line)
    icici = {r[0] for r in con.execute("SELECT ident_tail FROM stmt_slot WHERE bank='ICICI' AND kind<>'card' AND ident_tail IS NOT NULL AND ident_tail<>''")}   # S417
    icici |= {r[0] for r in con.execute("SELECT f.tail FROM stmt_file f JOIN stmt_slot s ON s.id=f.slot_id WHERE s.bank='ICICI' AND s.kind<>'card' AND f.tail IS NOT NULL AND f.tail<>''")}
    if not _has(con, "icici_statement_line"):                       # S411: ICICI's own table (finance_icici v1.1), never the Yes Bank one
        return [], "no ICICI statement lines on this server yet"
    if not icici:
        return [], "no ICICI statement on the shelf yet"
    rows = [dict(r) for r in con.execute("SELECT account_ref, txn_date, description, withdrawal_p FROM icici_statement_line WHERE withdrawal_p>0 AND substr(txn_date,1,7)=? ORDER BY txn_date", (month,))]
    rows = [r for r in rows if r["account_ref"] in icici]
    hits = [r for r in rows if any(w in (r["description"] or "").upper() for w in words)]
    out = ["auto-paid %s on %s from account …%s (%s)" % (_inr(r["withdrawal_p"]), _dmy(r["txn_date"]), r["account_ref"], (r["description"] or "")[:40]) for r in hits]
    out += card_electricity(con, month, words)                          # S434: the bill the Amazon Pay card pays
    if not out:
        return [], "no electricity payment found for %s in the ICICI statements or the card sheet" % _month_name(month)
    return out, None


def card_electricity(con, month, words):
    """S434: the electricity bills paid by a card, from the card sheet on the shelf (All_Transactions.xlsx): Date · Card · Narration ·
    Amount · Dr/Cr -- a debit whose narration carries an electricity word and whose date falls in the month.
    S504: the card statements' own lines (card_line) first; the running Excel only when they hold nothing for the month.
    S505 (the owner, 10-Oct-2026: for September's pack "the 17-Aug payment" -- the one on the card statement dated 12-Sep): a card's
    bill belongs to the month its STATEMENT is dated in, as the card sheet and the attached statement do; the running Excel is read
    (by the line's own date, as before) only when no card statement dated in the month has been read."""
    if _has(con, "card_line") and _has(con, "card_stmt") and con.execute("SELECT 1 FROM card_stmt WHERE substr(statement_date,1,7)=? LIMIT 1", (month,)).fetchone():
        lab = _card_labels(con)
        hits = [r for r in con.execute("SELECT l.txn_date, l.slot_key, l.description, l.amount_p, s.statement_date FROM card_line l JOIN card_stmt s ON s.id=l.stmt_id "
                                       "WHERE l.credit=0 AND substr(s.statement_date,1,7)=? ORDER BY l.txn_date", (month,)).fetchall()
                if any(w in (r[2] or "").upper() for w in words)]
        return ["auto-paid %s on %s by the %s card, on its statement dated %s (%s)" % (_inr(r[3]), _dmy(r[0]), lab.get(r[1], r[1]), _dmy(r[4]),
                                                                                     re.sub(r"\d{6,}", "", r[2] or "")[:40]) for r in hits]
    r = con.execute("SELECT local_path FROM stmt_file WHERE folder='all_txn' ORDER BY mtime DESC, id DESC LIMIT 1").fetchone()
    if not r or not r[0] or not os.path.exists(r[0]):
        return []
    try:
        from openpyxl import load_workbook                  # noqa: PLC0415
        wb = load_workbook(r[0], read_only=True, data_only=True)
        ws = wb["ALL"] if "ALL" in wb.sheetnames else wb[wb.sheetnames[0]]
        out = []
        head = None
        for row in ws.iter_rows(values_only=True):
            if head is None:
                head = [str(c or "").strip().lower() for c in row]
                continue
            rec = dict(zip(head, row))
            d = rec.get("date")
            if isinstance(d, dt.datetime):
                iso = d.date().isoformat()
            elif isinstance(d, dt.date):
                iso = d.isoformat()
            else:
                s = str(d or "").strip()
                iso = ("%s-%s-%s" % (s[6:10], s[3:5], s[0:2])) if re.match(r"^\d{2}/\d{2}/\d{4}$", s) else ""
            nar = str(rec.get("narration") or "").upper()
            if iso[:7] != month or str(rec.get("dr/cr") or "").strip().upper() != "DR" or not any(w in nar for w in words):
                continue
            try:
                amt = int(round(float(rec.get("amount") or 0) * 100))
            except (TypeError, ValueError):
                continue
            nar_clean = re.sub(r"\d{6,}", "", str(rec.get("narration") or "")).strip()
            out.append("auto-paid %s on %s by the %s card (%s)" % (_inr(amt), _dmy(iso), rec.get("card") or "", nar_clean[:40]))
        wb.close()
        return out
    except Exception:                                    # noqa: BLE001
        return []


def _assets_con():
    try:
        import purchase_app as pa                         # noqa: PLC0415
        return pa._assets_con()
    except Exception:                                    # noqa: BLE001
        return None


def _text_pdf(title, lines):
    """A one-or-more-page A4 text PDF through clinic_day_pdf's dependency-free writer (the index page of a bundle)."""
    import clinic_day_pdf as cdp                          # noqa: PLC0415
    p = cdp._PDF(title)
    p.new_page()
    y = 800
    p.text(40, y, title, 13, bold=True)
    y -= 24
    for ln in lines:
        if y < 60:
            p.new_page()
            y = 800
        p.text(40, y, ln[:110], 9.5)
        y -= 13
    return p.build()


def _html_lines(html):
    """S434: an HTML letter as its lines of text (the styles dropped, one line per block)."""
    import html as _h                                     # noqa: PLC0415
    s = re.sub(r"(?is)<(style|script)[^>]*>.*?</\1>", "", html or "")
    s = re.sub(r"(?i)<br\s*/?>|</(div|p|tr|h\d|li)>", "\n", s)
    s = _h.unescape(re.sub(r"<[^>]+>", "", s))
    out = []
    for ln in s.splitlines():
        ln = " ".join(ln.split())
        if ln or (out and out[-1]):
            out.append(ln)
    return out


def _letter_pdf(html):
    """S434: the NEFT covering letter as an A4 PDF, laid out as it prints (the firm's name bold, the paragraphs wrapped)."""
    import textwrap                                       # noqa: PLC0415
    import clinic_day_pdf as cdp                          # noqa: PLC0415
    lines = _html_lines(html)
    p = cdp._PDF("Letter")
    p.new_page()
    y = 790
    for i, ln in enumerate(lines):
        if not ln:
            y -= 9
            continue
        bold = ("SANJEEVNI MEDICOS" in ln.upper() and len(ln) < 40) or ln.lower().startswith("sub")
        for part in (textwrap.wrap(ln, 92) or [""]):
            if y < 60:
                p.new_page()
                y = 790
            p.text(56, y, part, 13 if (bold and "SANJEEVNI" in ln.upper()) else 11, bold=bold)
            y -= 16
    return p.build()


def build_bundle(con, month, lane, count_only=False):
    """(8) every S409 bill of the lane with bill_date in the month, plus the lane's rows scanned in the month with an earlier
    bill_date under 'Late -- belongs to <month>' -> one PDF: the index page (stamp numbers), then each scan.
    count_only: (number of rows, why) without building -- the checklist's and Needs-you's cheap look."""
    acon = _assets_con()
    if acon is None:
        return None, "the asset app's database is not readable from here"
    try:
        cols = {r[1] for r in acon.execute("PRAGMA table_info(bills)")}
        if "lane" not in cols:
            return None, "the asset app has no lanes yet (S409 not installed)"
        mcol = "COALESCE(NULLIF(bill_month,''), substr(bill_date,1,7))" if "bill_month" in cols else "substr(bill_date,1,7)"      # S434/S435 (F-658)
        rows = [dict(r) for r in acon.execute(
            "SELECT id, stamp_no, vendor, bill_no, bill_date, total_amount, source_stored, late_for, created_at, %s AS bmonth FROM bills WHERE lane=? AND status<>'rejected' "
            "AND ((%s=?) OR (substr(created_at,1,7)=? AND late_for IS NOT NULL AND %s<>?)) ORDER BY late_for IS NOT NULL, bill_date, id" % (mcol, mcol, mcol), (lane, month, month, month))]
        sent = {r[0]: r[1] for r in con.execute("SELECT month, MIN(sent_at) FROM pack_send GROUP BY month")}
        rows = [r for r in rows if not (r["late_for"] and r["bmonth"] != month and sent.get(r["bmonth"]) and str(sent[r["bmonth"]]) >= str(r["created_at"] or ""))]
    finally:
        acon.close()
    if not rows:
        return None, "no %s scans for %s" % (lane.replace("_", " "), _month_name(month))
    if count_only:
        return len(rows), ""
    tmp = tempfile.mkdtemp(prefix="bundle_")
    try:
        idx = ["%s  %s  %s  %s  %s%s" % (r["stamp_no"] or "-", _dmy(r["bill_date"]) if (r["bill_date"] and str(r["bill_date"])[:7] == r["bmonth"]) else _month_name(r["bmonth"] or month),
                                        (r["vendor"] or "")[:30], (r["bill_no"] or "")[:16],
                                        ("Rs %.2f" % r["total_amount"]) if r["total_amount"] is not None else "",
                                        (" -- Late, belongs to %s" % _month_name(r["bmonth"])) if (r["bmonth"] and r["bmonth"] != month) else "")
               for r in rows]
        parts = [os.path.join(tmp, "00_index.pdf")]
        with open(parts[0], "wb") as fh:
            fh.write(_text_pdf("%s bills -- %s (%d)" % (lane.replace("_", " ").title(), _month_name(month), len(rows)), idx))
        for i, r in enumerate(rows, 1):
            src = os.path.join(UPLOADS, r["source_stored"] or "")
            if not r["source_stored"] or not os.path.exists(src):
                continue
            dst = os.path.join(tmp, "%02d.pdf" % i)
            if src.lower().endswith(".pdf"):
                shutil.copy(src, dst)
            else:
                subprocess.run(["convert", src, dst], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
            if os.path.exists(dst):
                parts.append(dst)
        out = os.path.join(tmp, "bundle.pdf")
        if len(parts) == 1:
            shutil.copy(parts[0], out)
        else:
            subprocess.run(["pdfunite"] + parts + [out], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        if not os.path.exists(out):
            return None, "pdfunite could not build the bundle"
        with open(out, "rb") as fh:
            return fh.read(), None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _nice(label):
    """'NK Pathology (Yes Bank current)' -> 'NK Pathology - Yes Bank current' (S434: a readable attachment name)."""
    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", label or "")
    s = ("%s - %s" % (m.group(1), m.group(2))) if m else (label or "")
    return re.sub(r"[\\/:*?\"<>|]", "-", s)


def clinic_upi_check(con, month):
    """S434, the owner's page only: the clinic's UPI as the bank saw it against the takings paid online (+ split)."""
    try:
        bank = con.execute("SELECT COALESCE(SUM(amount_p),0) FROM upi_txn WHERE unit='clinic' AND substr(txn_date,1,7)=?", (month,)).fetchone()[0] or 0
        on = sp = 0
        for r in con.execute("SELECT * FROM clinic_day_revenue WHERE business_date BETWEEN ? AND ?", (month + "-01", _month_end(month))):
            t = _tender(dict(r))
            on += t["online"]
            sp += t["split"]
        return "clinic: the bank's UPI %s · takings paid online %s + split %s (part online) · difference %s" % (_inr(bank), _inr(on), _inr(sp), _inr(bank - on - sp))
    except Exception:                                    # noqa: BLE001
        return ""


def pack_rows(con, month, light=False):
    """Every item of the owner's list as a row: ready / missing (why) / late; the attachments it would send.
    light: the bundles are counted, not built (the checklist and Needs-you read the statuses only)."""
    ensure(con)
    rows, att = [], []

    def row(key, title, status, why="", files=None, preview=None, san=False):
        rows.append(dict(key=key, title=title, status=status, why=why, files=[f[0] for f in (files or [])], preview=preview, sanjeevni=san))
        att.extend(files or [])
    b, why = build_income(con, month)
    row("income", "1 · Income sheet + date-wise clinic ledger (physiotherapy excluded)", "ready" if b else "missing", why or "", [("Clinic income and ledger - %s.xlsx" % _month_name(month), b, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")] if b else None, "/finance/clinic/day?m=" + month)
    b, why = build_upi_corrections(con, month)
    why = (why or "") or clinic_upi_check(con, month)                  # S434: the owner sees the check; the accountants get the sheet
    row("upi", "2 · UPI totals + the cash/UPI correction report", "ready" if b else "missing", why or "", [("UPI totals and pharmacy corrections - %s.xlsx" % _month_name(month), b, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")] if b else None, None)
    cl = cells(con, month)
    n3 = 0
    for c in cl:
        if c["kind"] == "card":
            continue
        n3 += 1
        f = c["file"]
        st = "ready" if (f and c["state"] in ("read", "matched", "arrived")) else ("partial" if c["state"] == "partial" else "missing")   # S471
        files = None
        if f and st == "ready":
            pieces = c.get("pieces") or []                                     # S471: the stitched month goes in as part 1 of 2
            srcs = [(p["id"], i + 1, len(pieces)) for i, p in enumerate(pieces)] if pieces else [(f["id"], 0, 0)]
            files = []
            for fid, i, n in srcs:
                fr = _best_path(con, fid)                              # S479: the copy the accountants can open
                if fr[0]:
                    ext = os.path.splitext(fr[0])[1].lower() or ".pdf"
                    with open(fr[0], "rb") as fh:
                        files.append(("%s statement - %s%s%s" % (_nice(c["label"]), _month_name(month), (" (part %d of %d)" % (i, n)) if n else "", ext), fh.read(), "application/pdf" if ext == ".pdf" else "text/plain"))   # S434
            files = files or None
        why_ = ""
        if not f:
            why_ = "not on the shelf"
        elif c["state"] == "partial":                                  # S411 · S471: which statement is still to come
            why_ = c.get("partial_words") or ("only %s → %s on the shelf, not the whole month" % (_dmy(f["period_from"]), _dmy(f["period_to"])))
        elif c.get("pieces"):                                          # S471
            why_ = "%d statements together cover the month: %s" % (len(c["pieces"]), " + ".join("%s → %s" % (_dmy(p["period_from"]), _dmy(p["period_to"])) for p in c["pieces"]))
            if st != "ready":
                why_ += "; read: " + ", ".join(p["read_status"] or "arrived" for p in c["pieces"])
        elif st != "ready":
            why_ = "read: " + (f["read_status"] or "arrived")
        elif f["read_status"] and f["read_status"] != "read":
            why_ = f["read_status"]
        row("stmt:" + c["slot"], "3.%d · %s statement" % (n3, c["label"]), st, why_, files, "/finance/packs/file/%d" % f["id"] if f else None, bool(c["sanjeevni"]))
    n4 = 0
    for c in cl:
        if c["kind"] != "card":
            continue
        f = c["file"]
        files = None
        if f:
            fr = _best_path(con, f["id"])                                   # S503: an original the shelf opened attaches its unlocked copy
            if fr and fr[0] and os.path.exists(fr[0]):
                with open(fr[0], "rb") as fh:
                    files = [("%s statement - %s.pdf" % (_nice(c["label"]), ("dated " + _dmy(f["period_to"])) if f.get("period_to") else _month_name(month)), fh.read(), "application/pdf")]   # S434
        why_ = ("statement dated %s" % _dmy(f["period_to"])) if f else (NO_TWIN if c.get("no_decrypted") else "no decrypted statement on the shelf")   # S417
        if c["twin_missing"]:
            why_ = (why_ + "; " if why_ else "") + "the newest original has no decrypted twin"
        n4 += 1
        row("card:" + c["slot"], "4.%d · %s statement (password removed)" % (n4, c["label"]), ("ready" if f else "missing") if not c["twin_missing"] else ("late" if f else "missing"), why_, files, "/finance/packs/file/%d" % f["id"] if f else None)
    cb, cwhy = build_card_sheet(con, month)                             # S504: the cards' own lines, by ledger, proved
    if cb:
        row("cards", "4.%d · Card spends by ledger (every line, proved against each statement)" % (n4 + 1), "ready", "",
            [("Card spends by ledger - %s.xlsx" % _month_name(month), cb, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")],
            "/finance/packs/preview/cards?month=" + month)
    else:
        at = con.execute("SELECT id, local_path FROM stmt_file WHERE folder='all_txn' ORDER BY mtime DESC, id DESC LIMIT 1").fetchone()
        files = None
        if at and at[1] and os.path.exists(at[1]):
            with open(at[1], "rb") as fh:
                files = [("All card transactions.xlsx", fh.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")]   # S434
        row("cards", "4.%d · All_Transactions.xlsx (the cards' sheet)" % (n4 + 1), "ready" if files else "missing",
            ("the running Excel goes, because " + cwhy) if files else "not fetched from Drive yet", files, "/finance/packs/file/%d" % at[0] if at else None)
    el, why = electricity_lines(con, month)
    try:
        want = int(_setting(con, "packs.electricity_expected", SETTINGS["packs.electricity_expected"][0]) or 2)
    except ValueError:
        want = 2
    if len(el) < want and not any(c["file"] for c in cl if c["bank"] == "ICICI" and c["kind"] != "card"):   # S471: the bills come with the statements
        why = "the ICICI statements for %s are not on the shelf yet -- the two bills are paid around the 17th and show there" % _month_name(month)
    row("electricity", "5 · Electricity: the auto-paid bills (ICICI statements and the Amazon Pay card)", "ready" if len(el) >= want else "missing",   # S434
        (why or "") if not el else ("%d of %d found: " % (len(el), want) if len(el) < want else "") + " · ".join(el), None, None)
    nf, why = build_neft(con, month)
    row("neft", "6 · Sanjeevni NEFT details: advice Excel + letter copy", "ready" if nf else "missing", why or "", nf, "/finance/purchase/page/pay/%s" % month, True)
    b, why = build_digest(con, month)
    row("digest", "7 · Monthly payment digest", "ready" if b else "missing", why or "", [("Payment digest - %s.xlsx" % _month_name(month), b, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")] if b else None, None)
    for n8, (lane, label) in enumerate((("lab_purchase", "Dr Bhawna's lab-purchase bills"), ("owner_expense", "Dr Manoj's expense-file bills")), 1):
        b, why = build_bundle(con, month, lane, count_only=light)
        row("bundle:" + lane, "8.%d · Scanned bills: %s (numbered bundle + index)" % (n8, label), "ready" if b else "missing", why or "",
            ([("%s - %s.pdf" % (label, _month_name(month)), b, "application/pdf")] if b else None) if not light else None, None)   # S434
    ticks = {r[0]: (r[1], r[2]) for r in con.execute("SELECT item, done_by, done_at FROM pack_tick WHERE month=?", (month,))}
    for k, t in (("lab_register", "Physical: lab register accounts"), ("lab_receipt_book", "Physical: lab receipt book")):
        rows.append(dict(key=k, title=t, status="ready" if k in ticks else "tick", why=("ticked by %s %s" % ticks[k]) if k in ticks else "hand it over, then tick", files=[], preview=None, tick=True))
    return rows, att


# ---------------------------------------------------------------- the send
def _env():
    out = {}
    try:
        with open(ENV_PATH, encoding="utf-8") as fh:
            for ln in fh:
                ln = ln.strip()
                if ln and not ln.startswith("#") and "=" in ln:
                    k, v = ln.split("=", 1)
                    out[k.strip()] = v.strip().strip('"').strip("'")
    except OSError:
        pass
    return out


def send_pack(con, u, month):
    ensure(con)
    to = [x.strip() for x in _setting(con, "packs.accountant_to", "").split(",") if x.strip()]
    if not to:
        return dict(ok=False, error="no_addresses", message="packs.accountant_to is empty -- the accountants' addresses are not on this server"), 409
    rows, att = pack_rows(con, month)
    att = [a for a in att if a[1]]
    if not att:
        return dict(ok=False, error="nothing_ready", message="no item of the pack is ready to send"), 409
    limit = int(float(_setting(con, "packs.mail_limit_mb", "20") or 20) * 1024 * 1024)
    parts, cur, size = [], [], 0
    for a in att:
        if cur and size + len(a[1]) > limit:
            parts.append(cur)
            cur, size = [], 0
        cur.append(a)
        size += len(a[1])
    if cur:
        parts.append(cur)
    prior = con.execute("SELECT COUNT(*) FROM pack_send WHERE month=?", (month,)).fetchone()[0]
    who = _who(u)
    env = _env()
    ids, total = [], 0
    for i, part in enumerate(parts, 1):
        from email.message import EmailMessage             # noqa: PLC0415
        msg = EmailMessage()
        subj = "Dr Manoj Agarwal Clinic - accounts pack %s" % _month_name(month) + ((" (part %d of %d)" % (i, len(parts))) if len(parts) > 1 else "") + (" - resend" if prior else "")
        msg["Subject"] = subj
        msg["From"] = env.get("WATCHDOG_SMTP_FROM") or env.get("WATCHDOG_SMTP_USER") or "clinic@localhost"
        msg["To"] = ", ".join(to)
        body = ["The month-end accounts pack for %s%s." % (_month_name(month), (", part %d of %d" % (i, len(parts))) if len(parts) > 1 else ""), ""]
        for r in rows:
            body.append("%s  [%s]%s" % (r["title"], r["status"], ("  " + r["why"]) if r["why"] else ""))
        body += ["", "Sent from the clinic server by %s on %s." % (who, _dmy(_today().isoformat()))]
        msg.set_content("\n".join(body))
        for name, blob, mime in part:
            mt, st = (mime.split("/", 1) + ["octet-stream"])[:2]
            msg.add_attachment(blob, maintype=mt, subtype=st, filename=name)
        raw = msg.as_bytes()
        total += len(raw)
        mid = hashlib.sha256(raw).hexdigest()[:16]
        stub = os.environ.get("PACKS_MAIL_STUB", "")
        if stub:
            os.makedirs(stub, exist_ok=True)
            with open(os.path.join(stub, "pack_%s_part%d.eml" % (month, i)), "wb") as fh:
                fh.write(raw)
        else:
            import smtplib                                    # noqa: PLC0415
            user, pw = env.get("WATCHDOG_SMTP_USER", ""), env.get("WATCHDOG_SMTP_PASS", "")
            if not (user and pw):
                return dict(ok=False, error="no_smtp", message="the box's SMTP credentials are not set (/root/wa/.env)"), 503
            s = smtplib.SMTP(env.get("WATCHDOG_SMTP_HOST", "smtp.gmail.com"), int(env.get("WATCHDOG_SMTP_PORT", "587")), timeout=60)
            try:
                s.starttls()
                s.login(user, pw)
                s.send_message(msg, from_addr=user, to_addrs=to)
            finally:
                s.quit()
        con.execute("INSERT INTO pack_send (month, sent_at, sent_by, what, message_id, to_list, parts, bytes, resend) VALUES (?,?,?,?,?,?,?,?,?)",
                    (month, _now(), who, json.dumps([n for n, _b, _m in part]), mid, ", ".join(to), len(parts), len(raw), 1 if prior else 0))
        ids.append(mid)
    con.commit()
    return dict(ok=True, parts=len(parts), attachments=len(att), bytes=total, resend=bool(prior), message_ids=ids), 200


# ---------------------------------------------------------------- Amir's pack (Sanjeevni)
def amir_pack(con, month):
    """Available when both Sanjeevni statements of the month are on the shelf; the paid NEFT sheet (S407's marks) + the two PDFs."""
    ensure(con)
    cl = {c["slot"]: c for c in cells(con, month)}
    ys, ic = cl.get("yes_cur_sanj"), cl.get("icici_sanj")
    ready = bool(ys and ys["file"] and ys["state"] != "partial" and ic and ic["file"] and ic["state"] != "partial")   # S411
    seen = con.execute("SELECT done_by, done_at FROM pack_tick WHERE month=? AND item='amir_seen'", (month,)).fetchone()
    return dict(month=month, ready=ready, yes=(ys["file"] if ys else None), icici=(ic["file"] if ic else None), seen=(dict(by=seen[0], at=seen[1]) if seen else None))


def build_neft_paid_sheet(con, month):
    """The pay sheet with every line marked paid and the bank date once confirmed (S407's state)."""
    import purchase_app as pa                             # noqa: PLC0415
    s, groups = pa._pay_rows(con, month)
    st = None
    try:
        import supplier_msg                               # noqa: PLC0415
        st = supplier_msg.state(con, month)
    except Exception:                                    # noqa: BLE001
        st = None
    ev = (st or {}).get("event") or {}
    bank = (st or {}).get("bank_line") or {}
    rows = [("Vendor", "Payable Rs", "Route", "Paid", "NEFT date", "Bank date")]
    for g in groups:
        paid = "paid" if (g["route"] == "NEFT" and ev) else ("cheque" if g["route"] != "NEFT" else "not yet")
        rows.append((g["name"], g["payable_p"] / 100.0, g["route"], paid, _dmy(ev.get("date")) if (g["route"] == "NEFT" and ev) else "", _dmy(bank.get("date")) if (g["route"] == "NEFT" and bank) else ""))
    return _xlsx([("NEFT paid %s" % month, rows)])


def amir_card(con):
    """'Pichle mahine ka pack' on Amir's board: tap-to-open each piece, a 'dekh liya' tick. Nothing when not ready."""
    try:
        m = _prev_month(_today().strftime("%Y-%m"))
        p = amir_pack(con, m)
        if not p["ready"]:
            return ""
        seen = p["seen"]
        return ("<div class=card><h2>Pichle mahine ka pack — %s</h2>"
                "<p><a class='btn plain' href='/finance/amir/pack/%s/neft'>Paid NEFT sheet (Excel)</a></p>"
                "<p><a class='btn plain' href='/finance/amir/pack/%s/yes'>Yes Bank Sanjeevni statement (PDF)</a></p>"
                "<p><a class='btn plain' href='/finance/amir/pack/%s/icici'>ICICI Sanjeevni statement (PDF)</a></p>"
                "%s</div>" % (_esc(_month_name(m)), m, m, m,
                              ("<p class=sub>dekh liya ✓ %s %s</p>" % (_esc(seen["by"]), _esc((seen["at"] or "")[:16]))) if seen else
                              ("<form method=post action='/finance/amir/pack/%s/seen'><button class=btn name=go value=1>Dekh liya</button></form>" % m)))
    except Exception:                                    # noqa: BLE001
        return ""


# ---------------------------------------------------------------- the checklist
def _auto_state(con, month, key, rows_by_key):
    """The system's own answer for an automatic item: (done?, words)."""
    if not key:
        return None, ""
    if key.startswith("row:"):
        r = rows_by_key.get(key[4:])
        if key == "row:stmts":
            st = [x for x in rows_by_key.values() if x["key"].startswith("stmt:")]
            n = sum(1 for x in st if x["status"] == "ready")
            return (n == len(st) and n > 0), "%d / %d shelf par" % (n, len(st))
        if key == "row:cards":
            st = [x for x in rows_by_key.values() if x["key"].startswith("card:") or x["key"] == "cards"]
            n = sum(1 for x in st if x["status"] == "ready")
            return (n == len(st) and n > 0), "%d / %d" % (n, len(st))
        if key == "row:bundles":
            st = [x for x in rows_by_key.values() if x["key"].startswith("bundle:")]
            n = sum(1 for x in st if x["status"] == "ready")
            return n > 0, "%d bundle" % n
        return (bool(r) and r["status"] == "ready"), (r["why"] if r else "")
    if key == "sent":
        r = con.execute("SELECT sent_at, sent_by FROM pack_send WHERE month=? ORDER BY id DESC LIMIT 1", (month,)).fetchone()
        return bool(r), ("%s %s" % (r[1], r[0][:16])) if r else ""
    if key == "pm_final":
        try:
            import purchase_app as pa                     # noqa: PLC0415
            return pa._month_status(con, month)["status"] == "final", ""
        except Exception:                                # noqa: BLE001
            return False, ""
    if key == "neft_done":
        r = con.execute("SELECT id FROM purchase_neft_event WHERE month=? AND kind<>'rejected'", (month,)).fetchone() if _has(con, "purchase_neft_event") else None
        return bool(r), ""
    if key == "scans_clear":
        try:
            import porders                                # noqa: PLC0415
            n = len([b for b in porders.unscanned_bills(con) if str(b.get("month")) == month])
            return n == 0, ("%d baaki" % n) if n else ""
        except Exception:                                # noqa: BLE001
            return False, ""
    if key == "sanj_stmts":
        cl = {c["slot"]: c for c in cells(con, month)}
        ok = all(cl[k].get("file") for k in ("yes_cur_sanj", "icici_sanj") if k in cl) and any(k in cl for k in ("yes_cur_sanj", "icici_sanj"))   # S433
        return ok, ""
    if key == "clinic_stmts":
        cl = {c["slot"]: c for c in cells(con, month)}
        ok = all(cl[k].get("file") for k in ("yes_cur_clinic", "icici_clinic") if k in cl) and any(k in cl for k in ("yes_cur_clinic", "icici_clinic"))   # S433
        return ok, ""
    if key.startswith("slot:"):
        cl = {c["slot"]: c for c in cells(con, month)}
        return bool(cl.get(key[5:], {}).get("file")), ""
    if key == "clinic_days":
        if not _has(con, "clinic_day_revenue"):
            return False, ""
        n = con.execute("SELECT COUNT(*) FROM clinic_day_revenue WHERE substr(business_date,1,7)=?", (month,)).fetchone()[0]
        return n > 0, "%d din" % n
    return None, ""


def _due_day(it):
    """S472: the day of the following month this item is due by -- the owner's own value, else the rule by kind."""
    try:
        if it.get("due_day"):
            return max(1, min(28, int(it["due_day"])))
    except (TypeError, ValueError):
        pass
    k = it.get("auto_key") or ""
    if k in DUE_BY_KEY:
        return DUE_BY_KEY[k]
    if k.startswith("row:") or k.startswith("slot:"):
        return 10 if k.startswith("row:") else 5
    return DUE_BY_GROUP.get(it.get("grp") or "", 10)


def _due_iso(month, day):
    """The due date of month's checklist: <day> of the FOLLOWING month."""
    y, m = int(month[:4]), int(month[5:7])
    y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return "%04d-%02d-%02d" % (y, m, day)


def _is_handover(it):
    s = (it.get("item") or "").lower()
    return not it.get("auto_key") and any(w in s for w in HANDOVER_WORDS)


def checklist(con, month, today=None):
    ensure(con)
    today = today or _today()
    rows, _att = pack_rows(con, month, light=True)
    by = {r["key"]: r for r in rows}
    done = {r[0]: (r[1], r[2], r[3]) for r in con.execute("SELECT item_id, done_by, done_at, note FROM packs_done WHERE month=?", (month,))}
    out = []
    for it in [dict(r) for r in con.execute("SELECT * FROM packs_item WHERE active=1 ORDER BY sort, id")]:
        auto, words = _auto_state(con, month, it["auto_key"], by)
        d = done.get(it["id"])
        dd = _due_day(it)
        due = _due_iso(month, dd)
        is_done = bool(d) or bool(auto)
        photo = (d[2] if (d and d[2] and str(d[2]).startswith("handover/")) else None)
        out.append(dict(id=it["id"], grp=it["grp"], item=it["item"], auto=(it["auto_key"] is not None), auto_done=auto, words=words,
                        done=is_done, done_by=(d[0] if d else ("system" if auto else None)), done_at=(d[1] if d else None),
                        due_day=dd, due=due, late=(not is_done and today.isoformat() > due), due_today=(not is_done and today.isoformat() == due),   # S472
                        handover=_is_handover(it), photo=photo))
    return out


def checklist_line(con, month=None, today=None):
    """S472: the one line for Shavez's tile: 'N kaam baaki · M late · aaj: <item>'. Hindi, counts only."""
    today = today or _today()
    month = month or _prev_month(today.strftime("%Y-%m"))
    items = checklist(con, month, today)
    open_ = [x for x in items if not x["done"]]
    late = [x for x in open_ if x["late"]]
    due_today = [x for x in open_ if x["due_today"]]
    if not open_:
        text = "%s: sab ho gaya ✓" % _month_name(month)
    else:
        text = "%d kaam baaki" % len(open_)
        if late:
            text += " · %d late" % len(late)
        if due_today:
            text += " · aaj: " + due_today[0]["item"][:40]
        elif late:
            text += " · pehle: " + late[0]["item"][:40]
    return dict(ok=True, month=month, open=len(open_), late=len(late), due_today=len(due_today), text_hi=text)


def open_items(con, month):
    return [x for x in checklist(con, month) if not x["done"]]


def needs_you_lines(con):
    """After the 10th: the previous month's empty cells and open checklist items."""
    try:
        ensure(con)
        t = _today()
        if t.day < DUE_DAY:
            return []
        m = _prev_month(t.strftime("%Y-%m"))
        empty = [c for c in cells(con, m) if c["state"] == "empty"]
        out = []
        if empty:
            out.append(dict(cls="warn", target="bank", text="%s: %d statement%s not on the shelf (%s)" % (_month_name(m), len(empty), "" if len(empty) == 1 else "s", ", ".join(c["label"].split(" (")[0] for c in empty[:4]) + (" …" if len(empty) > 4 else ""))))
        op = open_items(con, m)
        if op:
            out.append(dict(cls="warn", target="bank", text="Shavez's month-end checklist for %s: %d item%s open" % (_month_name(m), len(op), "" if len(op) == 1 else "s")))
        # S412: locked Yes Bank e-statements waiting -- only when the previous month's Yes Bank cell is empty on BOTH routes (no branch copy either)
        if any(c["bank"] == "YES" for c in empty):
            lk = con.execute("SELECT COUNT(*), SUM(CASE WHEN read_status=? THEN 1 ELSE 0 END) FROM stmt_file WHERE folder='bank' AND locked=1 AND (unlocked_path IS NULL OR unlocked_path='')", (LOCKED_NO_PW,)).fetchone()
            n_lk, n_np = int(lk[0] or 0), int(lk[1] or 0)
            if n_lk and not _secret_count(con):
                out.append(dict(cls="warn", target="bank", text="Yes Bank statements need their password on the packs page (%d locked file%s waiting)" % (n_lk, "" if n_lk == 1 else "s")))
            elif n_lk and n_np:
                out.append(dict(cls="warn", target="bank", text="the stored Yes Bank password opens none of %d locked statement%s -- replace it on the packs page" % (n_np, "" if n_np == 1 else "s")))
        return out
    except Exception:                                    # noqa: BLE001
        return []


# ---------------------------------------------------------------- S476 (F-723): the statement road, watched by what it produces
def _int_setting(con, key):
    try:
        return int(str(_setting(con, key, SETTINGS[key][0]) or SETTINGS[key][0]).strip())
    except (TypeError, ValueError):
        return int(SETTINGS[key][0])


def statement_road(con, today=None):
    """S476 (F-723): is the relay still filing the banks' statements into Drive? It runs under the owner's personal Google
    authorisation and its failure mail reaches his personal inbox only, so this looks at what it PRODUCES: the newest bank file the
    shelf fetched. Returns dict(state, text, hint, last, this_month, missing, waiting, month). 'warn' from road_from_day of a month
    on, while no bank statement has reached the shelf this month and last month's cells are still empty or partial (it stays
    until one arrives); 'info' from road_from_day to road_to_day when some came and then none for road_quiet_days while a cell is
    still empty; 'ok' otherwise. A partial cell alone is never a grey note: a cycle statement's second half is due after the 10th.
    Never raises."""
    try:
        ensure(con)
        t = today or _today()
        d_from, d_to, quiet = _int_setting(con, "packs.road_from_day"), _int_setting(con, "packs.road_to_day"), _int_setting(con, "packs.road_quiet_days")
        r = con.execute("SELECT MAX(fetched_at), SUM(CASE WHEN fetched_at>=? THEN 1 ELSE 0 END) FROM stmt_file WHERE folder='bank'",
                        (t.replace(day=1).isoformat(),)).fetchone()
        last, n_month = str(r[0] or "")[:10], int(r[1] or 0)
        quiet = max(1, quiet)
        m = _prev_month(t.strftime("%Y-%m"))
        bank_cells = [c for c in cells(con, m) if c["kind"] != "card"]
        empty = [c["label"].split(" (")[0] for c in bank_cells if c["state"] == "empty"]
        waiting = [c for c in bank_cells if c["state"] in ("empty", "partial")]
        out = dict(state="ok", text="", hint="", last=last, this_month=n_month, missing=len(empty), waiting=len(waiting), month=m)
        if not last:
            out.update(state="info", text="No bank statement has reached the shelf yet.")
            return out
        age = (t - dt.date.fromisoformat(last)).days
        when = "%s (%s)" % (_dmy(last), "today" if age <= 0 else ("yesterday" if age == 1 else "%d days ago" % age))
        in_window = d_from <= t.day <= d_to
        if t.day >= d_from and waiting and n_month == 0:
            out.update(state="warn",
                       text="No bank statement has reached Drive this month -- the last one came on %s; %d of %s's statements are still awaited."
                            % (when, len(waiting), _month_name(m)),
                       hint="By now the banks have mailed them. If they are sitting in your personal inbox, the relay's Google "
                            "authorisation has lapsed: open the Janitor script, run it once and allow it. The shelf picks the "
                            "files up at its next fetch (05:40 and 07:30).")
        elif in_window and empty and age >= quiet:
            out.update(state="info",
                       text="No bank statement has reached Drive for %d days (the last on %s); %d of %s's statements are still missing: %s."
                            % (age, _dmy(last), len(empty), _month_name(m), ", ".join(empty[:4]) + (" ..." if len(empty) > 4 else "")),
                       hint="Usual while a bank's own copy is late or a statement password is not set. If the statements are "
                            "sitting in your personal inbox, the relay needs allowing again.")
        else:
            out.update(text="The last bank statement reached the shelf on %s -- %d this month." % (when, n_month))
        return out
    except Exception as ex:                              # noqa: BLE001
        return dict(state="info", text="the statement road could not be read (%s)" % str(ex)[:120], hint="", last="", this_month=0,
                    missing=0, waiting=0, month="")


def statement_road_row(add, con):
    """S476 (F-723): one row on /finance/health, called by finance_app's _health_state. Never raises."""
    r = statement_road(con)
    add("stmtroad", "Bank statements reaching Drive", r["state"], r["text"], r.get("hint") or "")


# ---------------------------------------------------------------- routes
def _owner():
    return _require("checker", unit=_unit)


def _staff():
    return _require("maker", "checker", unit=_unit)


def _tpl(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        return fh.read()


@bp.route("/finance/packs")
def page_packs():
    u, err = _owner()
    if err:
        return err
    return render_template_string(_tpl("packs.html"), month=request.args.get("month") or _prev_month(_today().strftime("%Y-%m")), me=_who(u))


@bp.route("/finance/packs/checklist")
def page_checklist():
    u, err = _staff()
    if err:
        return err
    return render_template_string(_tpl("packs_checklist.html"), month=request.args.get("month") or _prev_month(_today().strftime("%Y-%m")), me=_who(u),
                                  owner=("checker" in set((u or {}).get("roles") or []) or (u or {}).get("role") == "checker"))


@bp.route("/finance/packs/api/state")
def api_state():
    u, err = _owner()
    if err:
        return err
    m = request.args.get("month") or _prev_month(_today().strftime("%Y-%m"))
    if not _good_month(m):
        return jsonify(ok=False, error="bad_month"), 400
    con = _db()
    rows, att = pack_rows(con, m)
    sends = [dict(r) for r in con.execute("SELECT id, sent_at, sent_by, parts, bytes, resend, to_list FROM pack_send WHERE month=? ORDER BY id", (m,))]
    return jsonify(ok=True, month=m, month_name=_month_name(m), cells=cells(con, m), unplaced=unplaced(con), slots=[dict(r) for r in con.execute("SELECT id, key, holder_label, bank, kind, ident_tail FROM stmt_slot ORDER BY sort")],
                   rows=[dict(r, files=r.get("files", [])) for r in rows], attachments=[dict(name=a[0], bytes=len(a[1])) for a in att if a[1]],
                   sends=sends, to=_setting(con, "packs.accountant_to", ""), limit_mb=_setting(con, "packs.mail_limit_mb", "20"),
                   amir=amir_pack(con, m), checklist=checklist(con, m), items=[dict(r) for r in con.execute("SELECT * FROM packs_item ORDER BY sort, id")],
                   inbox_dir=INBOX, today=_today().isoformat(), secrets=secret_state(con), road=statement_road(con), cards=card_spends(con, m))   # S476 S504


@bp.route("/finance/packs/api/secret", methods=["POST"])
def api_secret():
    """S412: the owner types a bank's statement password(s) once -- stored in finance.db only, never echoed, never logged; the locked
    files are tried at once. The answer carries counts and 'set on', never the value."""
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    body, code = set_secret(_db(), b.get("bank"), b.get("text"), _who(u))
    return jsonify(**body), code


@bp.route("/finance/packs/api/card-merchant", methods=["POST"])
def api_card_merchant():
    """S504: the owner names a card merchant once (a category from the list); every line of it, past and future, follows."""
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    try:
        import card_lines                                 # noqa: PLC0415
    except Exception:                                    # noqa: BLE001
        return jsonify(ok=False, error="no_reader"), 500
    con = _db()
    body, code = card_lines.name_merchant(con, str(b.get("merchant") or "").strip(), str(b.get("category") or "").strip(), _who(u))
    return jsonify(**body), code


@bp.route("/finance/packs/api/set-aside", methods=["POST"])
def api_set_aside():
    """S503: the owner sets a locked file aside (never tried, never counted) or puts it back."""
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    try:
        fid = int(b.get("file") or 0)
    except (TypeError, ValueError):
        fid = 0
    if not fid:
        return jsonify(ok=False, error="bad_request"), 400
    body, code = set_aside(_db(), fid, bool(b.get("back")), _who(u), str(b.get("why") or ""))
    return jsonify(**body), code


@bp.route("/finance/packs/api/assign", methods=["POST"])
def api_assign():
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    try:
        fid, sid = int(b.get("file") or 0), int(b.get("slot") or 0)
    except (TypeError, ValueError):
        fid = sid = 0
    if not fid or not sid:
        return jsonify(ok=False, error="bad_request"), 400
    body, code = assign(_db(), fid, sid, _who(u))
    return jsonify(**body), code


@bp.route("/finance/packs/api/slot-words", methods=["POST"])
def api_slot_words():
    """S411: the owner edits the words a slot's holder must print (data, never code); the unplaced files are re-identified."""
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    try:
        sid = int(b.get("slot") or 0)
    except (TypeError, ValueError):
        sid = 0
    words = " ".join(str(b.get("words") or "").upper().replace(",", " ").split())[:80]
    con = _db()
    ensure(con)
    r = con.execute("SELECT key, ident_words FROM stmt_slot WHERE id=?", (sid,)).fetchone()
    if not r or not words:
        return jsonify(ok=False, error="bad_request"), 400
    con.execute("UPDATE stmt_slot SET ident_words=?, owner_set=? WHERE id=?", (words, "words by %s %s" % (_who(u), _now()), sid))
    con.commit()
    return jsonify(ok=True, slot=r[0], before=r[1], words=words, reidentified=process_inbox(con))


@bp.route("/finance/packs/api/tick", methods=["POST"])
def api_tick():
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    m, item = str(b.get("month") or ""), str(b.get("item") or "")
    if not _good_month(m) or item not in ("lab_register", "lab_receipt_book"):
        return jsonify(ok=False, error="bad_request"), 400
    con = _db()
    ensure(con)
    con.execute("INSERT OR REPLACE INTO pack_tick (month, item, done_by, done_at) VALUES (?,?,?,?)", (m, item, _who(u), _now()))
    con.commit()
    return jsonify(ok=True)


@bp.route("/finance/packs/api/send", methods=["POST"])
def api_send():
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    m = str(b.get("month") or "")
    if not _good_month(m):
        return jsonify(ok=False, error="bad_month"), 400
    con = _db()
    ensure(con)
    last = con.execute("SELECT sent_at FROM pack_send WHERE month=? ORDER BY id DESC LIMIT 1", (m,)).fetchone()
    if last:
        try:
            if (dt.datetime.now() - dt.datetime.fromisoformat(last[0])).total_seconds() < REPEAT_MIN * 60:
                return jsonify(ok=True, already=True, message="sent %s -- a repeat within %d minutes is not sent again" % (last[0][:16], REPEAT_MIN)), 200
        except ValueError:
            pass
    body, code = send_pack(con, u, m)
    return jsonify(**body), code


@bp.route("/finance/packs/api/refresh", methods=["POST"])
def api_refresh():
    """On demand: identify what the inbox holds (the fetch itself is the venv cron's; here the identification runs)."""
    u, err = _owner()
    if err:
        return err
    return jsonify(ok=True, **process_inbox(_db()))


@bp.route("/finance/packs/api/item", methods=["POST"])
def api_item():
    """The owner edits the checklist: add / rename / retire an item."""
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    con = _db()
    ensure(con)
    act = str(b.get("action") or "")
    if act == "add":
        grp, item = str(b.get("grp") or "").strip()[:40], str(b.get("item") or "").strip()[:160]
        if not grp or not item:
            return jsonify(ok=False, error="bad_request"), 400
        n = con.execute("SELECT COALESCE(MAX(sort),0)+1 FROM packs_item").fetchone()[0]
        con.execute("INSERT INTO packs_item (grp, item, auto_key, sort, edited_by, edited_at, source) VALUES (?,?,NULL,?,?,?,?)", (grp, item, n, _who(u), _now(), "owner"))
    elif act in ("rename", "retire", "restore"):
        try:
            iid = int(b.get("id") or 0)
        except (TypeError, ValueError):
            iid = 0
        if not con.execute("SELECT 1 FROM packs_item WHERE id=?", (iid,)).fetchone():
            return jsonify(ok=False, error="no_such_item"), 404
        if act == "rename":
            con.execute("UPDATE packs_item SET item=?, edited_by=?, edited_at=? WHERE id=?", (str(b.get("item") or "").strip()[:160], _who(u), _now(), iid))
            try:                                                                              # S472: the due day rides the same tap
                dd = int(b.get("due_day")) if b.get("due_day") not in (None, "") else None
                con.execute("UPDATE packs_item SET due_day=? WHERE id=?", ((max(1, min(28, dd)) if dd else None), iid))
            except (TypeError, ValueError):
                pass
        else:
            con.execute("UPDATE packs_item SET active=?, edited_by=?, edited_at=? WHERE id=?", (0 if act == "retire" else 1, _who(u), _now(), iid))
    else:
        return jsonify(ok=False, error="bad_action"), 400
    con.commit()
    return jsonify(ok=True)


@bp.route("/finance/packs/api/checklist")
def api_checklist():
    u, err = _staff()
    if err:
        return err
    m = request.args.get("month") or _prev_month(_today().strftime("%Y-%m"))
    if not _good_month(m):
        return jsonify(ok=False, error="bad_month"), 400
    con = _db()
    return jsonify(ok=True, month=m, month_name=_month_name(m), items=checklist(con, m), me=_who(u))


@bp.route("/finance/packs/api/checklist/tick", methods=["POST"])
def api_checklist_tick():
    u, err = _staff()
    if err:
        return err
    b = request.get_json(silent=True) or (dict(request.form) if request.form else {})      # S472: multipart when a photo comes
    m = str(b.get("month") or "")
    try:
        iid = int(b.get("id") or 0)
    except (TypeError, ValueError):
        iid = 0
    if not _good_month(m) or not iid:
        return jsonify(ok=False, error="bad_request"), 400
    con = _db()
    ensure(con)
    it = con.execute("SELECT id, auto_key FROM packs_item WHERE id=? AND active=1", (iid,)).fetchone()
    if not it:
        return jsonify(ok=False, error="no_such_item"), 404
    if it[1]:
        return jsonify(ok=False, error="automatic", message="yeh item system khud bharta hai"), 409
    d = con.execute("SELECT done_by, done_at FROM packs_done WHERE month=? AND item_id=?", (m, iid)).fetchone()
    if d:
        try:
            if (dt.datetime.now() - dt.datetime.fromisoformat(d[1])).total_seconds() < REPEAT_MIN * 60:
                return jsonify(ok=True, already=True, done_by=d[0], done_at=d[1]), 200
        except ValueError:
            pass
        return jsonify(ok=True, already=True, done_by=d[0], done_at=d[1]), 200
    note = str(b.get("note") or "")[:120]
    photo_rel = _save_handover_photo(m, iid)                      # S472: the photo of what was handed over, when one came
    if photo_rel:
        note = photo_rel
    con.execute("INSERT INTO packs_done (month, item_id, done_by, done_at, note) VALUES (?,?,?,?,?)", (m, iid, _who(u), _now(), note))
    con.commit()
    return jsonify(ok=True, done_by=_who(u), done_at=_now(), photo=bool(photo_rel)), 200


def _save_handover_photo(month, iid):
    """S472: the optional photo on a tick (multipart field 'photo'; jpeg/png; 8 MB) -> handover/<month>/<item>_<stamp>.<ext>
    under the packs folder; returns the relative path or None. A bad file is refused silently (the tick still stands)."""
    f = request.files.get("photo") if request.files else None
    if not f or not f.filename:
        return None
    blob = f.read(8 * 1024 * 1024 + 1)
    if not blob or len(blob) > 8 * 1024 * 1024:
        return None
    ext = "jpg" if blob[:3] == b"\xff\xd8\xff" else ("png" if blob[:8] == b"\x89PNG\r\n\x1a\n" else None)
    if not ext:
        return None
    rel = "handover/%s/%d_%s.%s" % (month, iid, _now().replace(":", "").replace("-", "")[:15], ext)
    path = os.path.join(PACK_DIR, rel)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as fh:
        fh.write(blob)
    return rel


@bp.route("/finance/packs/api/checklist/line")
def api_checklist_line():
    """S472: the tile's line (staff and the owner). Counts and one item name; nothing else."""
    u, err = _staff()
    if err:
        return err
    try:
        return jsonify(**checklist_line(_db()))
    except Exception:                                          # noqa: BLE001 -- the tile never waits on this
        return jsonify(ok=False), 200


@bp.route("/finance/packs/handover/<month>/<name>")
def page_handover(month, name):
    """S472: a hand-over photo, to staff and the owner."""
    u, err = _staff()
    if err:
        return err
    if not _good_month(month) or not re.fullmatch(r"[0-9]+_[0-9T]+\.(jpg|png)", name or ""):
        return "no such photo", 404
    p = os.path.join(PACK_DIR, "handover", month, name)
    if not os.path.exists(p):
        return "no such photo", 404
    return send_file(p, as_attachment=False)


@bp.route("/finance/packs/file/<int:fid>")
def page_file(fid):
    """A shelf file, to the owner (preview)."""
    u, err = _owner()
    if err:
        return err
    r = _best_path(_db(), fid)                                         # S479: the unlocked copy when the shelf opened it
    if not r[0]:
        return "not on the shelf", 404
    return send_file(r[0], as_attachment=False, download_name=re.sub(r"[^A-Za-z0-9._-]", "_", r[1] or "statement")[:80])


@bp.route("/finance/packs/api/text/<int:fid>")
def api_text(fid):
    """S476 (F-725): the text the readers see of one shelf file -- pdftotext -layout of the unlocked copy when the shelf opened
    it, else of the file as fetched -- to the owner only. A statement is read in its own bytes before its reader is judged
    (F-720). Read-only: nothing is stored, and the answer is never cached."""
    u, err = _owner()
    if err:
        return err
    r = _db().execute("SELECT id, name, folder, local_path, unlocked_path, locked, read_status, period_from, period_to FROM stmt_file WHERE id=?",
                      (fid,)).fetchone()
    if not r:
        return jsonify(ok=False, error="not_on_shelf"), 404
    lp = r[4] or r[3] or ""
    text, still_locked = _file_text(lp)
    resp = jsonify(ok=True, id=r[0], name=r[1], folder=r[2], source=("the unlocked copy" if r[4] else "the file as fetched"),
                   on_disk=bool(lp and os.path.exists(lp)), locked=bool(r[5]), still_locked=bool(still_locked), read_status=r[6],
                   period_from=r[7], period_to=r[8], chars=len(text), lines=text.count("\n"), text=text)
    resp.headers["Cache-Control"] = "no-store"
    return resp


@bp.route("/finance/packs/preview/<key>")
def page_preview(key):
    """A built attachment, to the owner: income / upi / digest / neft / bundle_lab_purchase / bundle_owner_expense."""
    u, err = _owner()
    if err:
        return err
    m = request.args.get("month") or _prev_month(_today().strftime("%Y-%m"))
    if not _good_month(m):
        return "bad month", 400
    con = _db()
    if key == "income":
        b, why = build_income(con, m)
        name, mime = "Clinic income and ledger - %s.xlsx" % _month_name(m), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif key == "upi":
        b, why = build_upi_corrections(con, m)
        name, mime = "UPI totals and pharmacy corrections - %s.xlsx" % _month_name(m), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif key == "digest":
        b, why = build_digest(con, m)
        name, mime = "Payment digest - %s.xlsx" % _month_name(m), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif key == "neft_letter":                                         # S434
        nf, why = build_neft(con, m)
        pdf = [x for x in (nf or []) if x[2] == "application/pdf"]
        b, name, mime = (pdf[0][1] if pdf else None), (pdf[0][0] if pdf else "letter.pdf"), "application/pdf"
    elif key == "cards":                                               # S504
        b, why = build_card_sheet(con, m)
        name, mime = "Card spends by ledger - %s.xlsx" % _month_name(m), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif key.startswith("bundle_"):
        b, why = build_bundle(con, m, key[7:])
        name, mime = "Bills_%s_%s.pdf" % (key[7:], m), "application/pdf"
    else:
        return "no such preview", 404
    if not b:
        return "not ready: %s" % why, 404
    return send_file(io.BytesIO(b), mimetype=mime, as_attachment=(not mime.endswith("pdf")), download_name=name)


# Amir's pack (Sanjeevni): under /finance/amir/... so the front gate's medical unit lets him in
@bp.route("/finance/amir/pack/<month>/<what>")
def page_amir_pack(month, what):
    u, err = _require("maker", "checker", "viewer", unit="medical")
    if err:
        return err
    if not _good_month(month):
        return "bad month", 400
    con = _db()
    if what == "neft":                                   # S452 (the owner, 02-Oct): a PDF, and only once the month's NEFT is confirmed
        import supplier_msg                              # noqa: PLC0415 -- Sanjeevni's rule decides (the bank SMS read, or the owner's entry)
        return supplier_msg.s452_amir_neft_pdf(con, month)
    p = amir_pack(con, month)
    if not p["ready"]:
        return "pack abhi tayyar nahi (dono statement shelf par nahi)", 404
    fid = (p["yes"] if what == "yes" else p["icici"] if what == "icici" else None)
    if not fid:
        return "no such piece", 404
    r = _best_path(con, fid["id"])                                     # S479: the copy Amir can open
    if not r[0]:
        return "not on the shelf", 404
    return send_file(r[0], as_attachment=False, download_name=re.sub(r"[^A-Za-z0-9._-]", "_", r[1] or "statement")[:80])


@bp.route("/finance/amir/pack/<month>/seen", methods=["POST"])
def api_amir_seen(month):
    u, err = _require("maker", "checker", "viewer", unit="medical")
    if err:
        return err
    if not _good_month(month):
        return "bad month", 400
    con = _db()
    ensure(con)
    con.execute("INSERT OR IGNORE INTO pack_tick (month, item, done_by, done_at) VALUES (?,?,?,?)", (month, "amir_seen", _who(u), _now()))
    con.commit()
    from flask import redirect                            # noqa: PLC0415
    return redirect("/finance/amir", code=303)
