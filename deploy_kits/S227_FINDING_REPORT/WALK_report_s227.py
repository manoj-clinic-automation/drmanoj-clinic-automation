"""LIVE-SHAPE WALK -- S227 FINDING REPORT.

The owner, 06-Sep-2026: an analytics layer for the stock-check report -- match
the count against everything accumulated (purchases, sales, sale returns, Marg's
own figures), flag unexplained movement, and on every major difference an inline
expandable list of the item's purchases this financial year (vendor, date, bill
number, item name, quantity) so the Marg operator reconciles from one screen.

This builds the shop with purchase bills (three vendors, one spelling with a
trailing dot, one bill before the financial year, one purchase return), sale
lines, two Marg export days; records a count with differences through the page;
then OPENS THE REPORT at phone width and reads it the way the doctor would: the
head, the seal, the mismatch, the sections; taps a major difference and checks
every purchase row and the residue against its own arithmetic.
"""
import os, shutil, subprocess, sys, threading, time
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import WALK_readiness_s226 as R
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

OK, BAD = [], []
def chk(n, c, d=""):
    (OK if c else BAD).append(n + (("  -- " + str(d)) if d and not c else ""))
    print(("  ok   " if c else "  FAIL ") + n + (("   " + str(d)) if d else ""))

import random; random.seed(2271)
import re as _re
def sale_key(s): return _re.sub(r"\s+", " ", _re.sub(r"[^A-Z0-9 ]+", " ", s.upper())).strip()
FORMS = [("TAB", "1*10", 10), ("TAB", "1*15", 15), ("CAP", "1*10", 10), ("SYP 100ML", "1*1", 1), ("INJ", "1*1", 1), ("TAB", "1*30", 30)]
ITEMS = []
for i in range(1, 121):
    f = FORMS[i % len(FORMS)]
    ITEMS.append(dict(item="WALK-%03d %s %s" % (i, random.choice(["ZINC", "CALPOL", "PANTO", "DOLO", "MOX", "ORTHO"]), f[0]), packing=f[1], pack_size=f[2]))
names = [r["item"] for r in ITEMS]; PS = {r["item"]: r["pack_size"] for r in ITEMS}

R.DB = os.path.join(HERE, "_walkreport.db"); con = R.build(); R.fill(con)
con.execute("DELETE FROM stock_snapshot"); con.execute("DELETE FROM purchase_bill"); con.execute("DELETE FROM purchase_line"); con.execute("DELETE FROM sale_line_item")
MARG, MARG_AUG, STRIP_P = {}, {}, {}
for r in ITEMS:
    q = random.choice([4, 13, 40, 137, 405]); MARG[r["item"]] = q
    if r is ITEMS[0]: q = 137; MARG[r["item"]] = q                  # X: enough on the shelf to be short by 22
    con.execute("INSERT OR REPLACE INTO stock_snapshot (as_on,item,qty,packing,pack_size,source,loaded_at) VALUES (?,?,?,?,?,?,?)",
                ("06-09-2026", r["item"], q, r["packing"], PS[r["item"]], "p", "2026-09-06T09:00:59"))
for r in ITEMS[:50]:
    con.execute("INSERT OR REPLACE INTO stock_rate (item, rate_p, as_of, source) VALUES (?,?,?,?)", (r["item"], 2500, "2026-09-06", "walk"))
seq = 500
for n, r in enumerate(ITEMS[:30] + ITEMS[100:110]):
    STRIP_P[r["item"]] = 1000 + n * 350
    for k in range(3):
        seq += 1
        con.execute("INSERT INTO sale_line_item (day_entry_id, unit, business_date, bill_no, is_return, seq, item_name, item_key, qty_raw, pack, amount_p) "
                    "VALUES (1,'medical',?,?,0,?,?,?,?,?,?)", ("2026-09-0%d" % (k + 1), "A0035%02d" % (seq % 100), seq, r["item"], sale_key(r["item"]), "1:0", r["packing"], STRIP_P[r["item"]]))
con.commit()

# ---- THE ITEM UNDER THE GLASS: names[0], pack 15 (i=1 -> FORMS[1]); its life this financial year
X = names[0]; PSX = PS[X]
MARG_AUG[X] = 100
con.execute("INSERT OR REPLACE INTO stock_snapshot (as_on,item,qty,packing,pack_size,source,loaded_at) VALUES (?,?,?,?,?,?,?)",
            ("01-08-2026", X, 100, "1*15", PSX, "p", "2026-08-01T09:00:00"))
for sup, sn in (("ANMOL AGENCIES", "anmol agencies"), ("BHARAT MEDICAL", "bharat medical"), ("PRIME ORTHO", "prime ortho")):
    con.execute("INSERT INTO purchase_bill (supplier_norm, supplier, bill_no, bill_date, month, amount_p) VALUES (?,?,?,?,?,?)", (sn, sup, "X/1", "2026-08-10", "2026-08", 100))
PL = [  # (supplier_norm, bill_no, bill_date, item as written, qty, free, loose, rate_p, direction)
    ("anmol agencies", "AN/1188", "2026-03-20", X, 10, 0, 0, 9000, "PURCHASE"),           # LAST financial year -- must not show
    ("anmol agencies", "AN/1301", "2026-04-15", X, 4, 1, 0, 9100, "PURCHASE"),            # 5 strips = 75
    ("bharat medical", "BM-7741", "2026-08-12", X + ".", 6, 0, 3, 9250, "PURCHASE"),     # trailing dot; 6 strips + 3 loose = 93
    ("prime ortho", "PO/2210", "2026-08-20", X, 2, 0, 0, 9300, "RETURN"),                # a purchase return: -30
    ("prime ortho", "PO/2299", "2026-09-02", X, 1, 0, 0, 9300, "PURCHASE"),               # 15
]
for (sn, bn, bd, it, qty, free, loose, rate, dr) in PL:
    con.execute("INSERT INTO purchase_line (supplier_norm, bill_no, bill_date, month, item, packing, qty, free, loose_qty, rate_p, direction, source_md5) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,'walk')", (sn, bn, bd, bd[:7], it, "1*15", qty, free, loose, rate, dr))
# sales of X between the two export days: 3 bills of 1 strip 2 tabs (17 each = 51), one return of 5
for k, day in enumerate(("2026-08-05", "2026-08-18", "2026-09-04")):
    seq += 1
    con.execute("INSERT INTO sale_line_item (day_entry_id, unit, business_date, bill_no, is_return, seq, item_name, item_key, qty_raw, pack, amount_p) "
                "VALUES (1,'medical',?,?,0,?,?,?,?,?,?)", (day, "A0036%02d" % k, seq, X, sale_key(X), "1:2", "1*15", STRIP_P[X]))
seq += 1
con.execute("INSERT INTO sale_line_item (day_entry_id, unit, business_date, bill_no, is_return, seq, item_name, item_key, qty_raw, pack, amount_p) "
            "VALUES (1,'medical','2026-08-25','CN00170',1,?,?,?,'0:5','1*15',?)", (seq, X, sale_key(X), STRIP_P[X]))
STRIP_P[X] = 150000                                             # Rs 1,500 a strip: X's shortage is MAJOR at MRP
con.execute("UPDATE sale_line_item SET amount_p=? WHERE item_key=?", (STRIP_P[X], sale_key(X)))
con.commit()
# the arithmetic the report must reproduce, for X between 01-08 and 06-09:
FY_P_UNITS = 75 + 93 - 30 + 15            # this FY, all four dated lines (the March one excluded)
WIN_P_UNITS = 93 - 30 + 15                # after 01-08 up to 06-09
WIN_SOLD = 51 + 3 * 15 + 0                # 1:0 sale lines from the generic loop: 3 bills x 15 units on 01..03-Sep, plus the three 1:2 bills
WIN_SOLD = 3 * (1 * PSX + 2) + 3 * (1 * PSX)   # = 51 + 45 = 96
WIN_RET = 5
MARG_DELTA = MARG[X] - 100
DOCS_DELTA = WIN_P_UNITS - WIN_SOLD + WIN_RET
RESIDUE = MARG_DELTA - DOCS_DELTA

ARCH = os.path.join(HERE, "pad_uploads")
if os.path.isdir(ARCH): shutil.rmtree(ARCH)
app = R.app_for(con)
threading.Thread(target=lambda: app.run(port=8837, threaded=False), daemon=True).start()
time.sleep(1.5)
HOST = "http://127.0.0.1:8837"; URL = HOST + "/finance/stock/page/count"
errs = []


def fill_pad(src, dst, entries):
    shutil.copy(src, dst); wb = load_workbook(dst); ws = wb.worksheets[0]; where = {}
    for r in range(9, 400):
        nm = ws.cell(row=r, column=2).value
        if isinstance(nm, str) and nm.strip(): where[nm.strip()] = r
    for it, (st, lo) in entries.items():
        r = where[it]
        if st is not None: ws.cell(row=r, column=5, value=st)
        if lo is not None: ws.cell(row=r, column=6, value=lo)
    wb.save(dst); return dst


good, EXPECT = {}, {}
for i, nm in enumerate(names[:40]):
    ps = PS[nm]; q = MARG[nm]
    tgt = q if i % 3 else (max(q - 7, 0) if i % 2 else q + 5)
    if nm == X: tgt = max(q - 22, 0)                          # X is short by 22: 1 strip 7 tabs -> a MAJOR line at its MRP
    good[nm] = (tgt // ps if ps > 1 else None, tgt % ps if ps > 1 else tgt); EXPECT[nm] = tgt
DIFF_ITEMS = [nm for nm in good if EXPECT[nm] != MARG[nm]]


def gate(pg):
    pg.locator('#whoC button[data-u="Darpan"]').click(); pg.locator('#whoE button[data-u="Amir"]').click()
    pg.fill("#bill", "A003425"); pg.fill("#billdate", "2026-09-06"); pg.wait_for_timeout(150)


with sync_playwright() as pw:
    b = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium")
    ctx = b.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True)
    pg = ctx.new_page()
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.on("console", lambda m: errs.append(m.text) if m.type == "error" and "Failed to load resource" not in m.text else None)
    pg.on("dialog", lambda d: d.accept())

    # ------------------------------------------------------------ 1 a count through the page
    pg.goto(URL); pg.wait_for_timeout(700); gate(pg)
    with pg.expect_download() as dl:
        pg.locator("#padget").click()
    PAD = os.path.join(HERE, "pad_prefilled.xlsx"); dl.value.save_as(PAD)
    pg.set_input_files("#padfile", fill_pad(PAD, os.path.join(HERE, "pad_filled.xlsx"), good)); pg.wait_for_timeout(3000)
    out = pg.inner_text("#padout")
    chk("the count is recorded through the page", "Recorded as count #1" in out, out[:120])
    chk("THE BOX OFFERS THE STOCK CHECK REPORT", pg.locator("#padout a.dl.report").count() == 1 and "Open the STOCK CHECK REPORT" in out
        and pg.locator("#padout a.dl.report").get_attribute("href") == "/finance/stock/page/report?count=1")
    chk("...and the recent list links it", "Open the report" in pg.inner_text("#recent"))

    # ------------------------------------------------------------ 2 the report JSON
    j = pg.request.get(HOST + "/finance/stock/api/pad/report/1").json()
    chk("THE REPORT PAYLOAD: count #1, one sheet, sealed and unchanged, the four readiness lines frozen",
        j["ok"] and j["count_id"] == 1 and j["parts"] == 0 and j["seal_ok"] and len(j["seals"]) == 1 and j["seals"][0]["finding_no"]
        and len(j["readiness_lines"]) == 4 and any(x.startswith("Sale report:") for x in j["readiness_lines"]), (j.get("seals"), j.get("readiness_lines")))
    chk("...the numbers: 40 counted of 120, %d differ, 80 not counted, the mismatch at MRP present" % len(DIFF_ITEMS),
        j["counted"] == 40 and j["items_in_shop"] == 120 and j["differed"] == len(DIFF_ITEMS) and j["not_counted"] == 80 and j["mismatch"]["lines"] == len(DIFF_ITEMS))
    chk("...the financial year window: from 01-04-2026 to the count's as-on 06-09-2026", j["fy_from"] == "2026-04-01" and j["upto"] == "2026-09-06", (j["fy_from"], j["upto"]))
    rx = [d for d in j["differences"] if d["item"] == X][0]
    chk("...X is a MAJOR difference: short 22 at its MRP, over Rs 1,000", rx["major"] and rx["diff"] == -22 and rx["mrp_p"] is not None and abs(rx["mrp_p"]) >= 100000, rx)
    chk("...every difference carries its life link, its answer / cause / decision slots, and MAJOR is counted",
        all(d["life"].startswith("/finance/stock/api/pad/item/1/") and "answer" in d and "cause" in d and "decision" in d for d in j["differences"])
        and j["major"] == sum(1 for d in j["differences"] if d["major"]))
    chk("...NOT COUNTED carries Marg's stock and its value at MRP sitting unchecked, and a total",
        len(j["not_counted_rows"]) == 80 and j["unchecked_p"] > 0 and all("marg" in r and "mrp_p" in r for r in j["not_counted_rows"]))
    chk("...MATCHED is the list of exact matches", len(j["matched"]) == 40 - len(DIFF_ITEMS))

    # ------------------------------------------------------------ 3 the life of X
    life = pg.request.get(HOST + "/finance/stock/api/pad/item/1/" + X).json()
    P = life["purchases"]
    chk("THE LIFE OF X: four purchase lines THIS financial year -- the March bill is not there",
        life["ok"] and len(P) == 4 and not any(p["bill_no"] == "AN/1188" for p in P), [p["bill_no"] for p in P])
    chk("...each with vendor, date, bill number, item as written, qty + free + loose, units, rate",
        P[0] == dict(vendor="ANMOL AGENCIES", bill_date="2026-04-15", bill_date_text="15-04-2026", bill_no="AN/1301", item=X, packing="1*15",
                     batch="", expiry="", qty=4.0, free=1.0, loose=0.0, units=75, rate_p=9100, purchase_rate_p=None, direction="PURCHASE"), P[0])
    chk("...THE TRAILING-DOT SPELLING ON THE BILL IS MATCHED (BM-7741, 6 strips + 3 loose = 93 units)",
        any(p["bill_no"] == "BM-7741" and p["item"] == X + "." and p["units"] == 93 and p["vendor"] == "BHARAT MEDICAL" for p in P)
        and sorted(life["purchase_names"]) == sorted([X, X + "."]), [(p["bill_no"], p["units"]) for p in P])
    chk("...a purchase RETURN is negative units and named", any(p["bill_no"] == "PO/2210" and p["units"] == -30 and p["direction"] == "RETURN" for p in P))
    chk("...bought this FY = %d units, in date order" % FY_P_UNITS, life["purchased_units"] == FY_P_UNITS and [p["bill_date"] for p in P] == sorted(p["bill_date"] for p in P))
    chk("...sold and returned, by the sale lines: %d sold in 6 bills, 5 returned in 1 credit note" % WIN_SOLD,
        life["sold_units"] == WIN_SOLD and life["sale_bills"] == 6 and life["returned_units"] == 5 and life["credit_notes"] == 1, (life["sold_units"], life["sale_bills"], life["returned_units"]))
    chk("...Marg's own figure by export day: 01-08 = 100 -> 06-09 = %d" % MARG[X],
        [(m["text"], m["qty"]) for m in life["marg"]] == [("01-08-2026", 100), ("06-09-2026", MARG[X])], life["marg"])
    t = life["detector"]
    chk("THE DETECTOR: Marg moved %+d, the documents explain %+d (bought %d, sold %d, returned %d), residue %+d" % (MARG_DELTA, DOCS_DELTA, WIN_P_UNITS, WIN_SOLD, WIN_RET, RESIDUE),
        t and t["marg_delta"] == MARG_DELTA and t["purchased"] == WIN_P_UNITS and t["sold"] == WIN_SOLD and t["returned"] == WIN_RET
        and t["docs_delta"] == DOCS_DELTA and t["residue"] == RESIDUE and ("NO document" in t["sentence"] if RESIDUE else "every unit is accounted" in t["sentence"]), t)
    y = pg.request.get(HOST + "/finance/stock/api/pad/item/1/" + names[115]).json()
    chk("an item with no purchase, no sale and one export day says so plainly, no detector",
        y["ok"] and y["purchases"] == [] and y["sold_units"] == 0 and y["detector"] is None and len(y["marg"]) == 1)
    chk("a life for a count that does not exist is 404", pg.request.get(HOST + "/finance/stock/api/pad/item/99/" + X).status == 404)

    # ------------------------------------------------------------ 4 THE REPORT PAGE, on the phone
    pg.goto(HOST + "/finance/stock/page/report?count=1"); pg.wait_for_timeout(1200)
    head = pg.inner_text("#head")
    chk("THE REPORT OPENS: count #1, the clinic, when, status OPEN, who, Marg as-on, after bill, the finding",
        "Stock check report — count #1" in head and "Advanced Orthopaedic Surgery Centre" in head and "OPEN" in head
        and "Darpan" in head and "Amir" in head and "06-09-2026" in head and "A003425" in head and j["seals"][0]["finding_no"] in head, head[:400])
    chk("...the seal line says unchanged since sealing", "unchanged since sealing" in head)
    chk("...the four readiness lines, frozen", "Sale report:" in head and "Purchases:" in head and "frozen at the moment of sealing" in head)
    chk("...the KPI tiles: counted, agree, differ + major, not counted, explained so far",
        "items counted" in head and "agree with Marg" in head and ("%d major" % j["major"]) in head and "not counted" in head and "explained so far" in head)
    chk("...THE VALUE OF THE MISMATCH AT MRP, staff and doctor seeing the same figures", "Value of the mismatch at MRP: short Rs" in head and "over Rs" in head and "same figures" in head)
    chk("...the tools: print, the result sheet, the hand sheet, the mismatch data, the proof", "Print / save as PDF" in head and "Result sheet (Excel)" in head and "Differences to check by hand (PDF)" in head and "Proof of sheet 1" in head)
    body = pg.inner_text("#body")
    chk("SECTION 1 -- Differences, %d items, largest value first, with the FY window in words" % len(DIFF_ITEMS),
        ("1 · Differences — %d items" % len(DIFF_ITEMS)) in body and "from 01-04-2026 to 06-09-2026" in body)
    chk("...the filter chips: All, Major, Short, Over, Unexplained", all(x in body for x in ("All (", "Major (", "Short (", "Over (", "Unexplained (")))
    chk("...%d rows, each with Marg -> counted, the difference in strips and tabs, MAJOR where it is, 'not yet explained'" % len(DIFF_ITEMS),
        pg.locator("#rows .row").count() == len(DIFF_ITEMS) and pg.locator("#rows .badge.major").count() == j["major"]
        and pg.locator("#rows .badge", has_text="not yet explained").count() == len(DIFF_ITEMS))
    rowX = pg.locator('#rows .row .h[data-item="%s"]' % X)
    chk("...X reads 'short 1 strip 7 tabs' in the doctor's convention, with its MRP value", "short 1 strip 7 tabs" in rowX.inner_text() and "at MRP" in rowX.inner_text(), rowX.inner_text())
    def shot(loc, name):
        try: loc.screenshot(path=os.path.join(HERE, name), timeout=5000)
        except Exception as e: print("  (no screenshot %s: %s)" % (name, e.__class__.__name__))
    shot(pg.locator("#s1"), "_shot_report_s1.png")
    # tap X: the life comes inline
    rowX.click(); pg.wait_for_timeout(1200)
    life_el = rowX.locator("xpath=..").locator(".life")
    lt = life_el.inner_text()
    chk("TAP X: THE INLINE SECTION OPENS with the financial year and the four purchase rows -- vendor, date, bill no, item, qty",
        life_el.is_visible() and "THIS FINANCIAL YEAR — 01-04-2026 to 06-09-2026" in lt and "ANMOL AGENCIES" in lt and "15-04-2026" in lt and "AN/1301" in lt
        and "4 + 1 free" in lt and "BHARAT MEDICAL" in lt and "BM-7741" in lt and "6 + 3 loose" in lt and "PO/2210" in lt and "(return)" in lt and "AN/1188" not in lt, lt[:500])
    chk("...bought / sold / returned totals in strips and tabs, Marg's figures by export day",
        ("Bought: " in lt and "(%d units) in 4 bill lines" % FY_P_UNITS in lt) and "Sold:" in lt and "in 6 bills" in lt and "Returned by patients:" in lt and "01-08-2026 = 100" in lt and "06-09-2026 = %d" % MARG[X] in lt, lt[500:1000])
    chk("...THE RESIDUE, in one sentence, and what a residue means",
        ("%+d units moved with NO document behind them" % RESIDUE in lt if RESIDUE else "every unit is accounted for" in lt) and "A residue is a stock change no document explains" in lt)
    chk("...the spelling note: the bill writes it two ways", "On purchase bills this item is written as:" in lt and X + "." in lt)
    pg.screenshot(path=os.path.join(HERE, "_shot_report_life.png"), full_page=True)
    rowX.click(); pg.wait_for_timeout(200)
    chk("...tap again: it folds away (and does not re-fetch)", not life_el.is_visible())
    pg.locator('.filters button[data-f="major"]').click(); pg.wait_for_timeout(400)
    chk("THE MAJOR FILTER shows only the major lines", pg.locator("#rows .row").count() == j["major"] and pg.locator("#rows .badge.major").count() == j["major"])
    pg.locator('.filters button[data-f="short"]').click(); pg.wait_for_timeout(400)
    chk("...the Short filter", pg.locator("#rows .row").count() == sum(1 for d in j["differences"] if d["diff"] < 0))
    pg.locator('.filters button[data-f="all"]').click(); pg.wait_for_timeout(400)
    chk("SECTION 2 -- Not counted: 80 items, Marg's stock sitting unchecked valued at MRP, the sentence about a partial count",
        "2 · Not counted — 80 items" in body and "sitting unchecked:" in body and "at MRP" in body and "partial count" in body)
    chk("SECTION 3 -- Matched, collapsed to one line", ("3 · %d items matched exactly" % (40 - len(DIFF_ITEMS))) in body and not pg.locator("#s3 details").get_attribute("open"))
    pg.locator("#s3 summary").click(); pg.wait_for_timeout(200)
    chk("...opens on a tap to the list", pg.locator("#s3 details").get_attribute("open") is not None and pg.locator("#s3 table tr").count() == 41 - len(DIFF_ITEMS))
    chk("no section 4 when nothing was sent back", pg.locator("#s4").count() == 0)
    pg.evaluate("()=>window.scrollTo(0,0)"); pg.wait_for_timeout(200)
    pg.screenshot(path=os.path.join(HERE, "_shot_report_top.png"), full_page=False)

    # ------------------------------------------------------------ 5 explanations arrive, the report shows them
    d_id = con.execute("SELECT id FROM stock_diff WHERE count_id=1 AND item=?", (X,)).fetchone()[0]
    con.execute("INSERT INTO stock_diff_answer (diff_id, reason, note, answered_by, answered_at) VALUES (?,?,?,?,?)", (d_id, "breakage", "one strip crushed", "amir", "2026-09-06T16:00:00"))
    con.execute("INSERT INTO stock_diff_decision (diff_id, decision, recover_from, recover_p, recovery_state, note, decided_by, decided_at) VALUES (?,?,?,?,?,?,?,?)",
                (d_id, "WRITE_OFF", None, None, "none", "", "manoj", "2026-09-06T16:05:00"))
    con.commit()
    pg.reload(); pg.wait_for_timeout(1200)
    rt = pg.locator('#rows .row .h[data-item="%s"]' % X).inner_text()
    chk("A STAFF REASON AND A DECISION SHOW ON THE LINE, and 'explained so far' counts it",
        "staff: breakage — one strip crushed" in rt and "written off" in rt and "not yet explained" not in rt and "1 / %d" % len(DIFF_ITEMS) in pg.inner_text("#head").replace("\n", " "), rt)
    pg.locator('.filters button[data-f="open"]').click(); pg.wait_for_timeout(400)
    chk("...and the Unexplained filter leaves it out", pg.locator("#rows .row").count() == len(DIFF_ITEMS) - 1)

    # ------------------------------------------------------------ 6 no count given: the newest; a closed count; the seal check
    pg.goto(HOST + "/finance/stock/page/report"); pg.wait_for_timeout(1200)
    chk("/page/report with no count opens the newest", "count #1" in pg.inner_text("#head"))
    r = pg.request.post(HOST + "/finance/stock/api/pad/close/1").json()
    pg.reload(); pg.wait_for_timeout(1200)
    chk("a closed count reads CLOSED in the head", "CLOSED" in pg.inner_text("#head") and "by the doctor as it stood" in pg.inner_text("#head"))
    con.execute("UPDATE stock_diff SET counted_qty=counted_qty+1 WHERE count_id=1 AND item=?", (X,)); con.commit()
    pg.reload(); pg.wait_for_timeout(1200)
    chk("A SEALED FIGURE TAMPERED WITH (a stock_diff row edited): the report says so, in red, before anything else", "SEALED FIGURES HAVE CHANGED" in pg.inner_text("#head") and pg.locator("#head .seal.bad").count() == 1)
    con.execute("UPDATE stock_diff SET counted_qty=counted_qty-1 WHERE count_id=1 AND item=?", (X,)); con.commit()
    chk("a report for a count that does not exist says so", pg.request.get(HOST + "/finance/stock/api/pad/report/99").status == 404)
    chk("THE PAGE THREW NO ERROR THROUGH ALL OF IT", not errs, errs)
    b.close()

print("\n%d ok, %d FAIL" % (len(OK), len(BAD)))
for x in BAD:
    print("  FAIL " + x)
sys.exit(1 if BAD else 0)
