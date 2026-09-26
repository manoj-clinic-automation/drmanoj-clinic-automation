#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""packs.py -- S408 (26-Sep-2026, D624). Month-end packs: the statement shelf, the accountant pack, Amir's pack, Shavez's
monthly checklist. OWNER: PARENT (clinic). The three Sanjeevni pieces (Amir's pack, the NEFT details, the Sanjeevni statement
matching) are marked below. Everything is server-side; the accountants receive nothing until the owner taps Send.

  the shelf     stmt_slot (one per account/card the owner listed; the tail is LEARNED by his one tap) · stmt_file (every file
                stmt_shelf.py fetched, identified BY CONTENT through pdftotext: the bank's name, the holder's words, the
                account / card tail, the period). Bank statements are read into bank_statement_period/_line by finance_yesbank
                (Yes Bank) or finance_icici (ICICI, new); a card statement is filed as arrived. Every account's month is one
                cell: arrived / read / matched. A month with any cell empty by the 10th -> a Needs-you line + a checklist line.
  the pack      /finance/packs (owner, English; unit 'packs', the doctor checker): every item of the owner's formal list as a
                row -- ready / missing (why) / late (belongs to <month>); previews; the two paper items as ticks; ONE button
                'Send to accountants' -> one mail (part 2 over packs.mail_limit_mb) to setting packs.accountant_to through the
                box's SMTP (the health report's .env), receipt in pack_send; a second send is a re-send, recorded as such.
  Amir's pack   (Sanjeevni) when both Sanjeevni statements of a month are on the shelf: the paid NEFT sheet (S407's marks) +
                the two PDFs, on his board as 'Pichle mahine ka pack' with tap-to-open and a 'dekh liya' tick. No email to Amir.
  the checklist /finance/packs/checklist (staff, Hindi; unit 'packs', shavez maker): packs_item rows -- 'system se ho gaya'
                where automatic (linked to the pack row), a tick for the manual ones (who, when, 10-minute guard); the Lab
                items ticks only; what is still open on the 10th -> the owner's Needs you.
NO patient data, no account number beyond a tail, no password anywhere in this file. The mail credentials are read from
/root/wa/.env at send time and never printed. PACKS_TODAY / PACKS_MAIL_STUB / PACKS_ENV / STMT_INBOX serve the walk.
"""
import datetime as dt
import glob
import hashlib
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile

from flask import Blueprint, jsonify, render_template_string, request, send_file

bp = Blueprint("packs", __name__)
_db = _require = None
_unit = "packs"
HERE = os.path.dirname(os.path.abspath(__file__))
INBOX = os.environ.get("STMT_INBOX", os.path.join(HERE, "statements", "inbox"))
PACK_DIR = os.environ.get("PACKS_DIR", os.path.join(HERE, "statements", "packs"))
UPLOADS = os.environ.get("ASSETS_UPLOADS", "/root/assetapp/uploads")
ENV_PATH = os.environ.get("PACKS_ENV", "/root/wa/.env")
REPEAT_MIN = 10
DUE_DAY = 10

DDL = (
    "CREATE TABLE IF NOT EXISTS stmt_slot (id INTEGER PRIMARY KEY, key TEXT NOT NULL UNIQUE, bank TEXT NOT NULL, holder_label TEXT NOT NULL,"
    " kind TEXT NOT NULL CHECK (kind IN ('current','savings','card')), ident_tail TEXT, ident_words TEXT, owner_set TEXT, sanjeevni INTEGER NOT NULL DEFAULT 0,"
    " sort INTEGER NOT NULL DEFAULT 0)",
    "CREATE TABLE IF NOT EXISTS stmt_file ("
    " id INTEGER PRIMARY KEY, drive_id TEXT NOT NULL UNIQUE, name TEXT NOT NULL, mtime TEXT, size INTEGER, folder TEXT NOT NULL,"
    " subfolder TEXT, slot_id INTEGER, period_from TEXT, period_to TEXT, read_status TEXT, matched_status TEXT, sha256 TEXT,"
    " fetched_at TEXT, local_path TEXT, bank TEXT, holder TEXT, tail TEXT, kind TEXT, ident_how TEXT, note TEXT, identified_at TEXT)",
    "CREATE TABLE IF NOT EXISTS pack_send (id INTEGER PRIMARY KEY, month TEXT NOT NULL, sent_at TEXT NOT NULL, sent_by TEXT, what TEXT,"
    " message_id TEXT, to_list TEXT, parts INTEGER, bytes INTEGER, resend INTEGER NOT NULL DEFAULT 0)",
    "CREATE TABLE IF NOT EXISTS pack_tick (month TEXT NOT NULL, item TEXT NOT NULL, done_by TEXT, done_at TEXT, PRIMARY KEY (month, item))",
    "CREATE TABLE IF NOT EXISTS packs_item (id INTEGER PRIMARY KEY, grp TEXT NOT NULL, item TEXT NOT NULL, auto_key TEXT, sort INTEGER NOT NULL DEFAULT 0,"
    " active INTEGER NOT NULL DEFAULT 1, edited_by TEXT, edited_at TEXT, source TEXT)",
    "CREATE TABLE IF NOT EXISTS packs_done (month TEXT NOT NULL, item_id INTEGER NOT NULL, done_by TEXT, done_at TEXT, note TEXT, PRIMARY KEY (month, item_id))",
)

# the owner's §1 list: the slots. bank · holder label · kind · the words the statement must carry (ALL of them, upper-case) · Sanjeevni?
SLOTS = [
    ("yes_cur_nk",       "YES",   "NK Pathology (Yes Bank current)",           "current", "NK PATHOLOGY",                 0),
    ("yes_cur_clinic",   "YES",   "Dr Manoj Agarwal Clinic (Yes Bank current)", "current", "MANOJ AGARWAL CLINIC",         0),
    ("yes_cur_sanj",     "YES",   "Sanjeevni Medicos (Yes Bank current)",      "current", "SANJEEVNI",                    1),
    ("yes_sav_manoj",    "YES",   "Dr Manoj Agarwal (Yes Bank savings)",       "savings", "MANOJ AGARWAL",                0),
    ("yes_sav_bhawna",   "YES",   "Dr Bhawna Agarwal (Yes Bank savings)",      "savings", "BHAWNA",                       0),
    ("yes_sav_huf",      "YES",   "Manoj Kumar Agarwal HUF (Yes Bank savings)", "savings", "HUF",                          0),
    ("icici_sanj",       "ICICI", "Sanjeevni Medicos (ICICI)",                 "current", "SANJEEVNI",                    1),
    ("icici_clinic",     "ICICI", "Dr Manoj Agarwal Clinic (ICICI)",           "current", "MANOJ AGARWAL CLINIC",         0),
    ("icici_personal",   "ICICI", "Dr Manoj Agarwal (ICICI)",                  "savings", "MANOJ AGARWAL",                0),
    ("card_hdfc_regalia", "HDFC", "HDFC Business Regalia (card)",              "card",    "REGALIA",                      0),
    ("card_icici_amazon", "ICICI", "ICICI Amazon Pay (card)",                  "card",    "AMAZON",                       0),
    ("card_icici_coral",  "ICICI", "ICICI Coral (card)",                       "card",    "CORAL",                        0),
]
CARD_FOLDERS = {"HDFC Business Regalia": "card_hdfc_regalia", "ICICI Amazon Pay": "card_icici_amazon", "ICICI Card 5007": "card_icici_coral", "ICICI Coral": "card_icici_coral"}
SETTINGS = {
    "packs.accountant_to": ("", "S408 D624 -- the accountants' addresses (read from the filer script at install); a data edit changes it"),
    "packs.electricity_words": ("PVVNL,UPPCL,ELECTRICITY,BIJLI,PASCHIMANCHAL,POWER CORP", "S408 D624 -- ICICI narration words that mean an electricity auto-pay; seeded from typical words (no real ICICI statement was on the box) -- edit when the real one shows"),
    "packs.mail_limit_mb": ("20", "S408 D624 -- a pack heavier than this goes as part 1 / part 2"),
}
# Shavez's checklist: the workbook could not be read by the box (Drive 404 to the service account), so the four groups are seeded
# from the brief's own words; the owner edits them on his page. auto_key ties an item to a pack row / a system fact.
ITEMS = [
    ("Sanjeevni", "Vendor payments sheet lock (mahine ka purchase final)", "pm_final"),
    ("Sanjeevni", "NEFT done aur supplier ko bata diya", "neft_done"),
    ("Sanjeevni", "Cheque register update (handed over)", None),
    ("Sanjeevni", "Pharmacy bill scan poore (Bill scan karo khaali)", "scans_clear"),
    ("Sanjeevni", "Yes Bank aur ICICI Sanjeevni statement shelf par", "sanj_stmts"),
    ("Lab", "Lab register accounts (kaagaz) accountant ko", None),
    ("Lab", "Lab receipt book (kaagaz) accountant ko", None),
    ("Lab", "NK Pathology ka Yes Bank statement shelf par", "slot:yes_cur_nk"),
    ("Clinic", "Docterz day revenue har din aa gaya", "clinic_days"),
    ("Clinic", "Clinic ka Yes Bank / ICICI statement shelf par", "clinic_stmts"),
    ("Clinic", "Petty book aur attendance close", None),
    ("Accountant ka samaan", "Income sheet + clinic ledger tayyar", "row:income"),
    ("Accountant ka samaan", "UPI totals + correction report tayyar", "row:upi"),
    ("Accountant ka samaan", "Saare bank statement shelf par", "row:stmts"),
    ("Accountant ka samaan", "Card statements (password hataye hue) + All_Transactions", "row:cards"),
    ("Accountant ka samaan", "Bijli ke do bill (ICICI se)", "row:electricity"),
    ("Accountant ka samaan", "Sanjeevni NEFT advice + letter", "row:neft"),
    ("Accountant ka samaan", "Payment digest", "row:digest"),
    ("Accountant ka samaan", "Scan kiye bill bundle (lab + expense)", "row:bundles"),
    ("Accountant ka samaan", "Pack accountant ko bheja", "sent"),
]


def init(app, db_getter, require_fn, unit="packs"):
    global _db, _require, _unit
    _db, _require, _unit = db_getter, require_fn, unit
    app.register_blueprint(bp)
    return bp


# ---------------------------------------------------------------- small things
def _today():
    v = os.environ.get("PACKS_TODAY", "")
    try:
        return dt.date.fromisoformat(v) if v else dt.date.today()
    except ValueError:
        return dt.date.today()


def _now():
    return dt.datetime.now().replace(microsecond=0).isoformat()


def _month_name(ym):
    try:
        return dt.date(int(ym[:4]), int(ym[5:7]), 1).strftime("%B %Y")
    except (TypeError, ValueError):
        return ym


def _prev_month(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    return "%04d-%02d" % ((y - 1, 12) if m == 1 else (y, m - 1))


def _month_end(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    nxt = dt.date(y + (m == 12), (m % 12) + 1, 1)
    return (nxt - dt.timedelta(days=1)).isoformat()


def _good_month(m):
    return bool(re.match(r"^\d{4}-\d{2}$", str(m or "")))


def _dmy(iso):
    try:
        return dt.date.fromisoformat(str(iso)[:10]).strftime("%d-%b-%Y")
    except (TypeError, ValueError):
        return str(iso or "")


def _inr(p):
    v = str(int(round(abs(int(p or 0)) / 100.0)))
    if len(v) > 3:
        head, tail = v[:-3], v[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        v = ",".join(parts) + "," + tail
    return "₹" + v


def _esc(s):
    return (str("" if s is None else s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


def _has(con, t):
    return con.execute("SELECT 1 FROM sqlite_master WHERE name=?", (t,)).fetchone() is not None


def _setting(con, key, default=""):
    try:
        r = con.execute("SELECT value FROM setting WHERE key=?", (key,)).fetchone()
        return (r[0] if r and r[0] is not None else default)
    except Exception:                                    # noqa: BLE001
        return default


def ensure(con):
    for d in DDL:
        con.execute(d)
    for i, (key, bank, label, kind, words, sanj) in enumerate(SLOTS):
        con.execute("INSERT OR IGNORE INTO stmt_slot (key, bank, holder_label, kind, ident_words, sanjeevni, sort) VALUES (?,?,?,?,?,?,?)",
                    (key, bank, label, kind, words, sanj, i))
    if not con.execute("SELECT 1 FROM packs_item LIMIT 1").fetchone():
        for i, (grp, item, key) in enumerate(ITEMS):
            con.execute("INSERT INTO packs_item (grp, item, auto_key, sort, source) VALUES (?,?,?,?,?)", (grp, item, key, i, "seed S408 (the brief's four groups)"))
    con.commit()


def _who(u):
    return (u or {}).get("user") or (u or {}).get("username") or ""


def _con():
    con = sqlite3.connect(os.environ.get("FINANCE_DB", os.path.join(HERE, "finance.db")), timeout=60)
    con.row_factory = sqlite3.Row
    return con


# ---------------------------------------------------------------- the shelf: identification by content
def pdf_text(path):
    try:
        r = subprocess.run(["pdftotext", "-layout", path, "-"], stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=60)
        return r.stdout.decode("utf-8", "replace")
    except Exception:                                    # noqa: BLE001
        return ""


TAIL_RES = (re.compile(r"(?:card|account|a/c)\s*(?:no\.?|number|#)?\s*[:\-]?\s*(?:[Xx*]{2,}[\s\-]*)+(\d{4})\b", re.I),
            re.compile(r"(?:account|a/c)\s*(?:no\.?|number|#)?\s*[:\-]?\s*(\d{9,18})\b", re.I),
            re.compile(r"\b\d{4}[\s\-]?[Xx*]{4}[\s\-]?[Xx*]{4}[\s\-]?(\d{4})\b"),
            re.compile(r"\b(?:[Xx*]{4}[\s\-]?){3}(\d{4})\b"))
PERIOD_RES = (re.compile(r"(?:statement\s+period|period|from)\s*[:\-]?\s*(\d{2}[/-]\d{2}[/-]\d{4})\s*(?:to|-|–)\s*(\d{2}[/-]\d{2}[/-]\d{4})", re.I),
              re.compile(r"statement\s+date\s*[:\-]?\s*(\d{2}[/-]\d{2}[/-]\d{4})", re.I))


def _iso_any(s):
    for fmt in ("%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def identify_text(text):
    """The bank, the kind, the tail, the period and the holder-word hits, from the text alone."""
    up = " ".join(text.upper().split())
    bank = "YES" if ("YES BANK" in up or "YESBANK" in up) else ("ICICI" if "ICICI" in up else ("HDFC" if "HDFC" in up else ""))
    kind = "card" if ("CREDIT CARD" in up or "MINIMUM AMOUNT DUE" in up or "CARD NUMBER" in up) else ("savings" if "SAVINGS" in up else ("current" if "CURRENT" in up else ""))
    tail = ""
    for rx in TAIL_RES:
        m = rx.search(text)
        if m:
            tail = m.group(1)[-4:]
            break
    pf = pt = None
    m = PERIOD_RES[0].search(text)
    if m:
        pf, pt = _iso_any(m.group(1)), _iso_any(m.group(2))
    else:
        m = PERIOD_RES[1].search(text)
        if m:
            d = _iso_any(m.group(1))
            if d:
                pt = d
                pf = d[:8] + "01"
    return dict(bank=bank, kind=kind, tail=tail, period_from=pf, period_to=pt, up=up)


def place(con, ident):
    """(slot row, how) or (None, why). By the learned tail first; else by the holder's words + bank (+ kind when the text says);
    a tie is unplaced -- the owner's one tap decides."""
    slots = [dict(r) for r in con.execute("SELECT * FROM stmt_slot ORDER BY sort")]
    if ident["tail"]:
        for s in slots:
            if s["ident_tail"] and s["ident_tail"] == ident["tail"] and (not ident["bank"] or s["bank"] == ident["bank"]):
                return s, "by tail"
    cands = []
    for s in slots:
        if ident["bank"] and s["bank"] != ident["bank"]:
            continue
        if ident["kind"] and s["kind"] != ident["kind"] and not (ident["kind"] in ("current", "savings") and s["kind"] in ("current", "savings") and ident["kind"] == s["kind"]):
            if ident["kind"] == "card" or s["kind"] == "card":
                continue
        words = [w.strip() for w in (s["ident_words"] or "").split(",") if w.strip()]
        if words and all(w in ident["up"] for w in words):
            cands.append((len(" ".join(words)), s))
    if not cands:
        return None, "no slot's words are in the text"
    cands.sort(key=lambda x: -x[0])
    if len(cands) > 1 and cands[0][0] == cands[1][0]:
        return None, "two slots fit the words (%s / %s)" % (cands[0][1]["holder_label"], cands[1][1]["holder_label"])
    best = cands[0][1]
    if best["ident_tail"] and ident["tail"] and best["ident_tail"] != ident["tail"]:
        return None, "the words say %s but the tail %s is not that account's (%s)" % (best["holder_label"], ident["tail"], best["ident_tail"])
    return best, "by words"


def read_file(con, f, slot):
    """Read a placed bank statement into the bank tables; a card is filed as arrived. Returns (read_status, matched_status)."""
    if not slot or slot["kind"] == "card" or not f["local_path"] or not os.path.exists(f["local_path"]):
        return ("n/a" if slot and slot["kind"] == "card" else "no file"), ""
    with open(f["local_path"], "rb") as fh:
        blob = fh.read()
    try:
        if slot["bank"] == "YES":
            import finance_yesbank                        # noqa: PLC0415
            res = finance_yesbank.ingest_statement(con, f["name"], blob, None, now=_now())
        elif slot["bank"] == "ICICI":
            import finance_icici                          # noqa: PLC0415
            res = finance_icici.ingest_statement(con, f["name"], blob, now=_now())
            if slot["key"] == "icici_sanj":
                _anchor_refresh(con, res)
        else:
            return "n/a", ""
    except Exception as ex:                              # noqa: BLE001
        return "refused: %s" % str(ex)[:160], ""
    matched = "ingested %d lines" % res.get("lines", 0)
    return "read", matched


def _anchor_refresh(con, res):
    """(Sanjeevni) bank_anchor for ICICI Sanjeevni follows the newest statement's closing balance (S377's row, refreshed)."""
    if not _has(con, "bank_anchor"):
        return
    r = con.execute("SELECT as_on FROM bank_anchor WHERE unit='medical' AND account='icici'").fetchone()
    if r and str(r[0] or "") >= res["period"][1]:
        return
    con.execute("INSERT INTO bank_anchor (unit, account, as_on, balance_p, source, entered_by, entered_at) VALUES ('medical','icici',?,?,?,?,?) "
                "ON CONFLICT(unit, account) DO UPDATE SET as_on=excluded.as_on, balance_p=excluded.balance_p, source=excluded.source, "
                "entered_by=excluded.entered_by, entered_at=excluded.entered_at",
                (res["period"][1], res["closing_p"], "ICICI Sanjeevni statement %s..%s, the bank's own closing balance (S408 shelf)" % tuple(res["period"]), "shelf", _now()))


def process_inbox(con=None, verbose=False):
    """Identify every fetched file not yet identified (or re-identify the unplaced ones after the owner's tap). Idempotent."""
    own = con is None
    con = con or _con()
    ensure(con)
    out = dict(placed=0, unplaced=0, read=0, refused=0, skipped=0)
    for f in [dict(r) for r in con.execute("SELECT * FROM stmt_file WHERE slot_id IS NULL OR read_status IS NULL ORDER BY id")]:
        if f["folder"] == "all_txn":
            con.execute("UPDATE stmt_file SET read_status='n/a', matched_status='', identified_at=? WHERE id=?", (_now(), f["id"]))
            out["skipped"] += 1
            continue
        lp = f["local_path"] or ""
        text = pdf_text(lp) if lp.lower().endswith(".pdf") and os.path.exists(lp) else ""
        ident = identify_text(text) if text else dict(bank="", kind="", tail="", period_from=None, period_to=None, up="")
        slot, how = place(con, ident) if text else (None, "not a readable PDF" if lp else "no file")
        if f["folder"] in ("cards", "decrypted") and not slot and f["subfolder"] in CARD_FOLDERS:
            slot = dict(con.execute("SELECT * FROM stmt_slot WHERE key=?", (CARD_FOLDERS[f["subfolder"]],)).fetchone() or {}) or None
            how = "by the card folder" if slot else how
        con.execute("UPDATE stmt_file SET bank=?, kind=?, tail=?, period_from=?, period_to=?, holder=?, slot_id=?, ident_how=?, note=?, identified_at=? WHERE id=?",
                    (ident["bank"], ident["kind"] or (slot["kind"] if slot else ""), ident["tail"], ident["period_from"], ident["period_to"],
                     (slot["holder_label"] if slot else ""), (slot["id"] if slot else None), (how if slot else None), (None if slot else how), _now(), f["id"]))
        if slot:
            out["placed"] += 1
            if f["folder"] == "cards":
                rs, ms = "locked original", ""
            else:
                rs, ms = read_file(con, f, slot)
            con.execute("UPDATE stmt_file SET read_status=?, matched_status=? WHERE id=?", (rs, ms, f["id"]))
            out["read" if rs == "read" else ("refused" if rs.startswith("refused") else "skipped")] += 1
        else:
            out["unplaced"] += 1
            con.execute("UPDATE stmt_file SET read_status=NULL WHERE id=?", (f["id"],))
    con.commit()
    if own:
        con.close()
    return out


def assign(con, file_id, slot_id, who):
    """The owner's one tap: this unplaced file is that account -> the slot learns the tail for ever; the file is read."""
    ensure(con)
    f = con.execute("SELECT * FROM stmt_file WHERE id=?", (file_id,)).fetchone()
    s = con.execute("SELECT * FROM stmt_slot WHERE id=?", (slot_id,)).fetchone()
    if not f or not s:
        return dict(ok=False, error="no_such"), 404
    f, s = dict(f), dict(s)
    if f["tail"]:
        con.execute("UPDATE stmt_slot SET ident_tail=?, owner_set=? WHERE id=?", (f["tail"], "%s %s" % (who, _now()), s["id"]))
    con.execute("UPDATE stmt_file SET slot_id=?, holder=?, ident_how=?, note=NULL, identified_at=? WHERE id=?", (s["id"], s["holder_label"], "the owner's tap", _now(), f["id"]))
    rs, ms = ("locked original", "") if f["folder"] == "cards" else read_file(con, f, s)
    con.execute("UPDATE stmt_file SET read_status=?, matched_status=? WHERE id=?", (rs, ms, f["id"]))
    con.commit()
    return dict(ok=True, file=f["id"], slot=s["key"], tail_learned=f["tail"] or None, read_status=rs), 200


# ---------------------------------------------------------------- the month cells
def cells(con, month):
    """One cell per slot for the month: the file (the one whose period covers the month, or falls in it) and its state."""
    ensure(con)
    lo, hi = month + "-01", _month_end(month)
    out = []
    for s in [dict(r) for r in con.execute("SELECT * FROM stmt_slot ORDER BY sort")]:
        fs = [dict(r) for r in con.execute(
            "SELECT * FROM stmt_file WHERE slot_id=? AND folder<>'cards' AND period_to IS NOT NULL AND period_from<=? AND period_to>=? "
            "ORDER BY period_to DESC, id DESC", (s["id"], hi, lo))]
        f = fs[0] if fs else None
        state = "empty"
        if f:
            state = "arrived"
            if f["read_status"] == "read" or (s["kind"] == "card" and f["read_status"] in ("n/a", None, "")):
                state = "read"
            if s["kind"] == "card" and f:
                state = "read"
            if s["sanjeevni"] and f["read_status"] == "read":
                state = "matched"
        twin_missing = False
        if s["kind"] == "card":
            orig = [dict(r) for r in con.execute("SELECT * FROM stmt_file WHERE slot_id=? AND folder='cards' ORDER BY mtime DESC, id DESC LIMIT 1", (s["id"],))]
            dec = [dict(r) for r in con.execute("SELECT * FROM stmt_file WHERE slot_id=? AND folder='decrypted' ORDER BY mtime DESC, id DESC LIMIT 1", (s["id"],))]
            if orig and (not dec or (dec[0]["mtime"] or "") < (orig[0]["mtime"] or "")):
                twin_missing = True
        out.append(dict(slot=s["key"], label=s["holder_label"], bank=s["bank"], kind=s["kind"], tail=s["ident_tail"], sanjeevni=s["sanjeevni"],
                        state=state, file=(dict(id=f["id"], name=f["name"], period_from=f["period_from"], period_to=f["period_to"],
                                                read_status=f["read_status"], matched_status=f["matched_status"]) if f else None),
                        twin_missing=twin_missing))
    return out


def unplaced(con):
    ensure(con)
    return [dict(id=r["id"], name=r["name"], folder=r["folder"], subfolder=r["subfolder"], bank=r["bank"], kind=r["kind"], tail=r["tail"],
                 period_from=r["period_from"], period_to=r["period_to"], why=r["note"], fetched_at=r["fetched_at"])
            for r in con.execute("SELECT * FROM stmt_file WHERE slot_id IS NULL AND folder<>'all_txn' ORDER BY id DESC LIMIT 60")]


# ---------------------------------------------------------------- the pack: builders
def _xlsx(sheets):
    """[(title, [row, ...]), ...] -> xlsx bytes (openpyxl, on both pythons)."""
    from openpyxl import Workbook                        # noqa: PLC0415
    from openpyxl.styles import Font                     # noqa: PLC0415
    wb = Workbook()
    first = True
    for title, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title[:30]
        for i, r in enumerate(rows):
            ws.append(list(r))
            if i == 0:
                for c in ws[1]:
                    c.font = Font(bold=True)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _tender(row):
    out = {"cash": 0, "online": 0, "card": 0, "split": 0}
    known = {"Cash": "cash", "Online Payment": "online", "Debit Card": "card", "Credit Card": "card", "Net Banking": "online", "Patient APP": "online", "Wallet": "online", "Split Payment": "split"}
    try:
        raw = json.loads(row["tender_json"] or "{}")
    except (ValueError, TypeError):
        return out
    for label, p in raw.items():
        b = known.get(str(label).strip())
        if b:
            out[b] += int(p)
    return out


def build_income(con, month):
    """(1) the doctor's monthly income sheet + the date-wise clinic ledger, physiotherapy excluded (it is its own unit)."""
    if not _has(con, "clinic_day_revenue"):
        return None, "the clinic day-revenue table is not on this server"
    rows = [dict(r) for r in con.execute("SELECT * FROM clinic_day_revenue WHERE business_date BETWEEN ? AND ? ORDER BY business_date", (month + "-01", _month_end(month)))]
    if not rows:
        return None, "no clinic day revenue rows for %s" % _month_name(month)
    ledger = [("Date", "Consultations (n)", "Consultations Rs", "X-ray (n)", "X-ray Rs", "Procedures (n)", "Procedures Rs", "Total (n)", "Total Rs", "Cash Rs", "Online Rs", "Card Rs", "Split Rs")]
    tot = [0] * 12
    for r in rows:
        t = _tender(r)
        vals = [r["cons_count"] or 0, (r["cons_amount_p"] or 0) / 100.0, r["xray_count"] or 0, (r["xray_amount_p"] or 0) / 100.0, r["proc_count"] or 0, (r["proc_amount_p"] or 0) / 100.0,
                r["total_count"] or 0, (r["total_amount_p"] or 0) / 100.0, t["cash"] / 100.0, t["online"] / 100.0, t["card"] / 100.0, t["split"] / 100.0]
        tot = [a + b for a, b in zip(tot, vals)]
        ledger.append([r["business_date"]] + vals)
    ledger.append(["TOTAL"] + tot)
    income = [("Dr Manoj Agarwal Clinic — income %s" % _month_name(month), ""), ("Days with takings", len(rows)), ("Consultations Rs", tot[1]), ("X-ray Rs", tot[3]),
              ("Procedures Rs", tot[5]), ("Total Rs", tot[7]), ("of which cash Rs", tot[8]), ("of which online Rs", tot[9]), ("of which card Rs", tot[10]),
              ("Physiotherapy", "excluded (its own register)"), ("Paper OPD register", "the accountants' own add-on")]
    return _xlsx([("Income", income), ("Ledger date-wise", ledger)]), None


def build_upi_corrections(con, month):
    """(2) UPI totals per unit + the cash/UPI correction report (finance_app's own rows, filtered to the month)."""
    if not _has(con, "upi_txn"):
        return None, "no UPI feed on this server"
    upi = [("Unit", "Transactions", "UPI Rs")]
    for u, n, s in con.execute("SELECT unit, COUNT(*), COALESCE(SUM(amount_p),0) FROM upi_txn WHERE substr(txn_date,1,7)=? GROUP BY unit ORDER BY unit", (month,)):
        upi.append((u, n, (s or 0) / 100.0))
    if len(upi) == 1:
        return None, "no UPI transactions in %s" % _month_name(month)
    corr = [("Date", "Change", "Amount Rs", "Likely bill", "Books say UPI", "Bank says UPI", "Status")]
    try:
        import finance_app as fa                          # noqa: PLC0415
        for r in fa._correction_rows(con, 400, include_resolved=True):
            d = str(r.get("date") or r.get("business_date") or "")[:10]
            if d[:7] != month:
                continue
            corr.append((d, r.get("change") or r.get("instruction") or "", (r.get("amount_p") or 0) / 100.0, r.get("bill") or r.get("likely_bill") or "",
                         (r.get("books_upi_p") or 0) / 100.0, (r.get("bank_upi_p") or 0) / 100.0, r.get("status") or ""))
    except Exception as ex:                              # noqa: BLE001
        corr.append(("(the correction rows could not be read: %s)" % str(ex)[:80], "", "", "", "", "", ""))
    return _xlsx([("UPI totals", upi), ("Cash-UPI corrections", corr)]), None


def build_digest(con, month):
    """(7) the monthly payment digest: the payment register's rows of the month (S234's derived table)."""
    if not _has(con, "payment_register"):
        return None, "the payment register is not on this server"
    rows = [("Date", "Vendor", "Description", "Amount Rs", "Attachment in Drive")]
    for r in con.execute("SELECT date_iso, vendor, description, amount_paise, in_drive FROM payment_register WHERE substr(date_iso,1,7)=? ORDER BY date_iso, row_no", (month,)):
        rows.append((r[0], r[1], r[2], (r[3] or 0) / 100.0, "yes" if r[4] else ""))
    if len(rows) == 1:
        return None, "no payment-register rows for %s" % _month_name(month)
    return _xlsx([("Payments %s" % month, rows)]), None


def build_neft(con, month):
    """(6, Sanjeevni) the NEFT advice Excel + the covering letter, from purchase_app's own routes' code."""
    try:
        import purchase_app as pa                         # noqa: PLC0415
        st = pa._month_status(con, month)
        if st["status"] != "final":
            return None, "%s is not finalised on the pay sheet" % _month_name(month)
        blob, total, missing = pa._advice_xlsx_s267(con, month)
        letter = pa._letter_html_s266(con, month, total)
        return [("Sanjeevni_NEFT_advice_%s.xlsx" % month, blob, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("Sanjeevni_NEFT_letter_%s.html" % month, letter.encode("utf-8"), "text/html")], None
    except Exception as ex:                              # noqa: BLE001
        return None, "the pay sheet could not be read (%s)" % str(ex)[:80]


def electricity_lines(con, month):
    """(5) the ICICI statement lines whose narration matches packs.electricity_words -> one line each."""
    words = [w.strip().upper() for w in _setting(con, "packs.electricity_words", SETTINGS["packs.electricity_words"][0]).split(",") if w.strip()]
    # the ICICI accounts = the tails the ICICI slots learned + the tails of every ICICI statement placed on the shelf (never a Yes Bank line)
    icici = {r[0] for r in con.execute("SELECT ident_tail FROM stmt_slot WHERE bank='ICICI' AND ident_tail IS NOT NULL AND ident_tail<>''")}
    icici |= {r[0] for r in con.execute("SELECT f.tail FROM stmt_file f JOIN stmt_slot s ON s.id=f.slot_id WHERE s.bank='ICICI' AND f.tail IS NOT NULL AND f.tail<>''")}
    if not _has(con, "bank_statement_line"):
        return [], "no statement lines on this server"
    if not icici:
        return [], "no ICICI statement on the shelf yet"
    rows = [dict(r) for r in con.execute("SELECT account_ref, txn_date, description, withdrawal_p FROM bank_statement_line WHERE withdrawal_p>0 AND substr(txn_date,1,7)=? ORDER BY txn_date", (month,))]
    rows = [r for r in rows if r["account_ref"] in icici]
    hits = [r for r in rows if any(w in (r["description"] or "").upper() for w in words)]
    if not rows:
        return [], "no ICICI statement lines on the shelf for %s" % _month_name(month)
    if not hits:
        return [], "no line matched packs.electricity_words (%s) in the ICICI lines of %s" % (", ".join(words[:4]), _month_name(month))
    return ["auto-paid %s on %s from account …%s (%s)" % (_inr(r["withdrawal_p"]), _dmy(r["txn_date"]), r["account_ref"], (r["description"] or "")[:40]) for r in hits], None


def _assets_con():
    try:
        import purchase_app as pa                         # noqa: PLC0415
        return pa._assets_con()
    except Exception:                                    # noqa: BLE001
        return None


def _text_pdf(title, lines):
    """A one-or-more-page A4 text PDF through clinic_day_pdf's dependency-free writer (the index page of a bundle)."""
    import clinic_day_pdf as cdp                          # noqa: PLC0415
    p = cdp._PDF(title)
    p.new_page()
    y = 800
    p.text(40, y, title, 13, bold=True)
    y -= 24
    for ln in lines:
        if y < 60:
            p.new_page()
            y = 800
        p.text(40, y, ln[:110], 9.5)
        y -= 13
    return p.build()


def build_bundle(con, month, lane, count_only=False):
    """(8) every S409 bill of the lane with bill_date in the month, plus the lane's rows scanned in the month with an earlier
    bill_date under 'Late -- belongs to <month>' -> one PDF: the index page (stamp numbers), then each scan.
    count_only: (number of rows, why) without building -- the checklist's and Needs-you's cheap look."""
    acon = _assets_con()
    if acon is None:
        return None, "the asset app's database is not readable from here"
    try:
        cols = {r[1] for r in acon.execute("PRAGMA table_info(bills)")}
        if "lane" not in cols:
            return None, "the asset app has no lanes yet (S409 not installed)"
        rows = [dict(r) for r in acon.execute(
            "SELECT id, stamp_no, vendor, bill_no, bill_date, total_amount, source_stored, late_for, created_at FROM bills WHERE lane=? AND status<>'rejected' "
            "AND ((substr(bill_date,1,7)=? ) OR (substr(created_at,1,7)=? AND late_for IS NOT NULL)) ORDER BY late_for IS NOT NULL, bill_date, id", (lane, month, month))]
    finally:
        acon.close()
    if not rows:
        return None, "no %s scans for %s" % (lane.replace("_", " "), _month_name(month))
    if count_only:
        return len(rows), ""
    tmp = tempfile.mkdtemp(prefix="bundle_")
    try:
        idx = ["%s  %s  %s  %s  %s%s" % (r["stamp_no"] or "-", _dmy(r["bill_date"]) if r["bill_date"] else "no date", (r["vendor"] or "")[:30], (r["bill_no"] or "")[:16],
                                        ("Rs %.2f" % r["total_amount"]) if r["total_amount"] is not None else "", (" -- Late, belongs to %s" % r["late_for"]) if r["late_for"] else "")
               for r in rows]
        parts = [os.path.join(tmp, "00_index.pdf")]
        with open(parts[0], "wb") as fh:
            fh.write(_text_pdf("%s bills -- %s (%d)" % (lane.replace("_", " ").title(), _month_name(month), len(rows)), idx))
        for i, r in enumerate(rows, 1):
            src = os.path.join(UPLOADS, r["source_stored"] or "")
            if not r["source_stored"] or not os.path.exists(src):
                continue
            dst = os.path.join(tmp, "%02d.pdf" % i)
            if src.lower().endswith(".pdf"):
                shutil.copy(src, dst)
            else:
                subprocess.run(["convert", src, dst], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
            if os.path.exists(dst):
                parts.append(dst)
        out = os.path.join(tmp, "bundle.pdf")
        if len(parts) == 1:
            shutil.copy(parts[0], out)
        else:
            subprocess.run(["pdfunite"] + parts + [out], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        if not os.path.exists(out):
            return None, "pdfunite could not build the bundle"
        with open(out, "rb") as fh:
            return fh.read(), None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def pack_rows(con, month, light=False):
    """Every item of the owner's list as a row: ready / missing (why) / late; the attachments it would send.
    light: the bundles are counted, not built (the checklist and Needs-you read the statuses only)."""
    ensure(con)
    rows, att = [], []

    def row(key, title, status, why="", files=None, preview=None, san=False):
        rows.append(dict(key=key, title=title, status=status, why=why, files=[f[0] for f in (files or [])], preview=preview, sanjeevni=san))
        att.extend(files or [])
    b, why = build_income(con, month)
    row("income", "1 · Income sheet + date-wise clinic ledger (physiotherapy excluded)", "ready" if b else "missing", why or "", [("Clinic_income_%s.xlsx" % month, b, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")] if b else None, "/finance/clinic/day?m=" + month)
    b, why = build_upi_corrections(con, month)
    row("upi", "2 · UPI totals + the cash/UPI correction report", "ready" if b else "missing", why or "", [("UPI_and_corrections_%s.xlsx" % month, b, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")] if b else None, "/finance/api/marg-corrections.xlsx?all=1")
    cl = cells(con, month)
    for c in cl:
        if c["kind"] == "card":
            continue
        f = c["file"]
        st = "ready" if (f and c["state"] in ("read", "matched", "arrived")) else "missing"
        files = None
        if f:
            fr = con.execute("SELECT local_path, name FROM stmt_file WHERE id=?", (f["id"],)).fetchone()
            if fr and fr[0] and os.path.exists(fr[0]):
                with open(fr[0], "rb") as fh:
                    files = [("Statement_%s_%s.pdf" % (c["slot"], month), fh.read(), "application/pdf")]
        row("stmt:" + c["slot"], "3 · %s statement" % c["label"], st, ("not on the shelf" if not f else ("read: " + (f["read_status"] or "arrived"))) if st != "ready" or f else "", files, "/finance/packs/file/%d" % f["id"] if f else None, bool(c["sanjeevni"]))
    for c in cl:
        if c["kind"] != "card":
            continue
        f = c["file"]
        files = None
        if f:
            fr = con.execute("SELECT local_path FROM stmt_file WHERE id=?", (f["id"],)).fetchone()
            if fr and fr[0] and os.path.exists(fr[0]):
                with open(fr[0], "rb") as fh:
                    files = [("Card_%s_%s.pdf" % (c["slot"], month), fh.read(), "application/pdf")]
        why_ = "" if f else "no decrypted statement on the shelf"
        if c["twin_missing"]:
            why_ = (why_ + "; " if why_ else "") + "the newest original has no decrypted twin"
        row("card:" + c["slot"], "4 · %s statement (password removed)" % c["label"], ("ready" if f else "missing") if not c["twin_missing"] else ("late" if f else "missing"), why_, files, "/finance/packs/file/%d" % f["id"] if f else None)
    at = con.execute("SELECT id, local_path FROM stmt_file WHERE folder='all_txn' ORDER BY mtime DESC, id DESC LIMIT 1").fetchone()
    files = None
    if at and at[1] and os.path.exists(at[1]):
        with open(at[1], "rb") as fh:
            files = [("All_Transactions.xlsx", fh.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")]
    row("cards", "4 · All_Transactions.xlsx (the cards' sheet)", "ready" if files else "missing", "" if files else "not fetched from Drive yet", files, "/finance/packs/file/%d" % at[0] if at else None)
    el, why = electricity_lines(con, month)
    row("electricity", "5 · Electricity: the two auto-paid bills (from the ICICI statements)", "ready" if el else "missing", (why or "") if not el else " · ".join(el), None, None)
    nf, why = build_neft(con, month)
    row("neft", "6 · Sanjeevni NEFT details: advice Excel + letter copy", "ready" if nf else "missing", why or "", nf, "/finance/purchase/page/pay/%s" % month, True)
    b, why = build_digest(con, month)
    row("digest", "7 · Monthly payment digest", "ready" if b else "missing", why or "", [("Payment_digest_%s.xlsx" % month, b, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")] if b else None, None)
    for lane, label in (("lab_purchase", "Dr Bhawna's lab-purchase bills"), ("owner_expense", "Dr Manoj's expense-file bills")):
        b, why = build_bundle(con, month, lane, count_only=light)
        row("bundle:" + lane, "8 · Scanned bills: %s (numbered bundle + index)" % label, "ready" if b else "missing", why or "",
            ([("Bills_%s_%s.pdf" % (lane, month), b, "application/pdf")] if b else None) if not light else None, None)
    ticks = {r[0]: (r[1], r[2]) for r in con.execute("SELECT item, done_by, done_at FROM pack_tick WHERE month=?", (month,))}
    for k, t in (("lab_register", "Physical: lab register accounts"), ("lab_receipt_book", "Physical: lab receipt book")):
        rows.append(dict(key=k, title=t, status="ready" if k in ticks else "tick", why=("ticked by %s %s" % ticks[k]) if k in ticks else "hand it over, then tick", files=[], preview=None, tick=True))
    return rows, att


# ---------------------------------------------------------------- the send
def _env():
    out = {}
    try:
        with open(ENV_PATH, encoding="utf-8") as fh:
            for ln in fh:
                ln = ln.strip()
                if ln and not ln.startswith("#") and "=" in ln:
                    k, v = ln.split("=", 1)
                    out[k.strip()] = v.strip().strip('"').strip("'")
    except OSError:
        pass
    return out


def send_pack(con, u, month):
    ensure(con)
    to = [x.strip() for x in _setting(con, "packs.accountant_to", "").split(",") if x.strip()]
    if not to:
        return dict(ok=False, error="no_addresses", message="packs.accountant_to is empty -- the accountants' addresses are not on this server"), 409
    rows, att = pack_rows(con, month)
    att = [a for a in att if a[1]]
    if not att:
        return dict(ok=False, error="nothing_ready", message="no item of the pack is ready to send"), 409
    limit = int(float(_setting(con, "packs.mail_limit_mb", "20") or 20) * 1024 * 1024)
    parts, cur, size = [], [], 0
    for a in att:
        if cur and size + len(a[1]) > limit:
            parts.append(cur)
            cur, size = [], 0
        cur.append(a)
        size += len(a[1])
    if cur:
        parts.append(cur)
    prior = con.execute("SELECT COUNT(*) FROM pack_send WHERE month=?", (month,)).fetchone()[0]
    who = _who(u)
    env = _env()
    ids, total = [], 0
    for i, part in enumerate(parts, 1):
        from email.message import EmailMessage             # noqa: PLC0415
        msg = EmailMessage()
        subj = "Dr Manoj Agarwal Clinic - accounts pack %s" % _month_name(month) + ((" (part %d of %d)" % (i, len(parts))) if len(parts) > 1 else "") + (" - resend" if prior else "")
        msg["Subject"] = subj
        msg["From"] = env.get("WATCHDOG_SMTP_FROM") or env.get("WATCHDOG_SMTP_USER") or "clinic@localhost"
        msg["To"] = ", ".join(to)
        body = ["The month-end accounts pack for %s%s." % (_month_name(month), (", part %d of %d" % (i, len(parts))) if len(parts) > 1 else ""), ""]
        for r in rows:
            body.append("%s  [%s]%s" % (r["title"], r["status"], ("  " + r["why"]) if r["why"] else ""))
        body += ["", "Sent from the clinic server by %s on %s." % (who, _dmy(_today().isoformat()))]
        msg.set_content("\n".join(body))
        for name, blob, mime in part:
            mt, st = (mime.split("/", 1) + ["octet-stream"])[:2]
            msg.add_attachment(blob, maintype=mt, subtype=st, filename=name)
        raw = msg.as_bytes()
        total += len(raw)
        mid = hashlib.sha256(raw).hexdigest()[:16]
        stub = os.environ.get("PACKS_MAIL_STUB", "")
        if stub:
            os.makedirs(stub, exist_ok=True)
            with open(os.path.join(stub, "pack_%s_part%d.eml" % (month, i)), "wb") as fh:
                fh.write(raw)
        else:
            import smtplib                                    # noqa: PLC0415
            user, pw = env.get("WATCHDOG_SMTP_USER", ""), env.get("WATCHDOG_SMTP_PASS", "")
            if not (user and pw):
                return dict(ok=False, error="no_smtp", message="the box's SMTP credentials are not set (/root/wa/.env)"), 503
            s = smtplib.SMTP(env.get("WATCHDOG_SMTP_HOST", "smtp.gmail.com"), int(env.get("WATCHDOG_SMTP_PORT", "587")), timeout=60)
            try:
                s.starttls()
                s.login(user, pw)
                s.send_message(msg, from_addr=user, to_addrs=to)
            finally:
                s.quit()
        con.execute("INSERT INTO pack_send (month, sent_at, sent_by, what, message_id, to_list, parts, bytes, resend) VALUES (?,?,?,?,?,?,?,?,?)",
                    (month, _now(), who, json.dumps([n for n, _b, _m in part]), mid, ", ".join(to), len(parts), len(raw), 1 if prior else 0))
        ids.append(mid)
    con.commit()
    return dict(ok=True, parts=len(parts), attachments=len(att), bytes=total, resend=bool(prior), message_ids=ids), 200


# ---------------------------------------------------------------- Amir's pack (Sanjeevni)
def amir_pack(con, month):
    """Available when both Sanjeevni statements of the month are on the shelf; the paid NEFT sheet (S407's marks) + the two PDFs."""
    ensure(con)
    cl = {c["slot"]: c for c in cells(con, month)}
    ys, ic = cl.get("yes_cur_sanj"), cl.get("icici_sanj")
    ready = bool(ys and ys["file"] and ic and ic["file"])
    seen = con.execute("SELECT done_by, done_at FROM pack_tick WHERE month=? AND item='amir_seen'", (month,)).fetchone()
    return dict(month=month, ready=ready, yes=(ys["file"] if ys else None), icici=(ic["file"] if ic else None), seen=(dict(by=seen[0], at=seen[1]) if seen else None))


def build_neft_paid_sheet(con, month):
    """The pay sheet with every line marked paid and the bank date once confirmed (S407's state)."""
    import purchase_app as pa                             # noqa: PLC0415
    s, groups = pa._pay_rows(con, month)
    st = None
    try:
        import supplier_msg                               # noqa: PLC0415
        st = supplier_msg.state(con, month)
    except Exception:                                    # noqa: BLE001
        st = None
    ev = (st or {}).get("event") or {}
    bank = (st or {}).get("bank_line") or {}
    rows = [("Vendor", "Payable Rs", "Route", "Paid", "NEFT date", "Bank date")]
    for g in groups:
        paid = "paid" if (g["route"] == "NEFT" and ev) else ("cheque" if g["route"] != "NEFT" else "not yet")
        rows.append((g["name"], g["payable_p"] / 100.0, g["route"], paid, _dmy(ev.get("date")) if (g["route"] == "NEFT" and ev) else "", _dmy(bank.get("date")) if (g["route"] == "NEFT" and bank) else ""))
    return _xlsx([("NEFT paid %s" % month, rows)])


def amir_card(con):
    """'Pichle mahine ka pack' on Amir's board: tap-to-open each piece, a 'dekh liya' tick. Nothing when not ready."""
    try:
        m = _prev_month(_today().strftime("%Y-%m"))
        p = amir_pack(con, m)
        if not p["ready"]:
            return ""
        seen = p["seen"]
        return ("<div class=card><h2>Pichle mahine ka pack — %s</h2>"
                "<p><a class='btn plain' href='/finance/amir/pack/%s/neft'>Paid NEFT sheet (Excel)</a></p>"
                "<p><a class='btn plain' href='/finance/amir/pack/%s/yes'>Yes Bank Sanjeevni statement (PDF)</a></p>"
                "<p><a class='btn plain' href='/finance/amir/pack/%s/icici'>ICICI Sanjeevni statement (PDF)</a></p>"
                "%s</div>" % (_esc(_month_name(m)), m, m, m,
                              ("<p class=sub>dekh liya ✓ %s %s</p>" % (_esc(seen["by"]), _esc((seen["at"] or "")[:16]))) if seen else
                              ("<form method=post action='/finance/amir/pack/%s/seen'><button class=btn name=go value=1>Dekh liya</button></form>" % m)))
    except Exception:                                    # noqa: BLE001
        return ""


# ---------------------------------------------------------------- the checklist
def _auto_state(con, month, key, rows_by_key):
    """The system's own answer for an automatic item: (done?, words)."""
    if not key:
        return None, ""
    if key.startswith("row:"):
        r = rows_by_key.get(key[4:])
        if key == "row:stmts":
            st = [x for x in rows_by_key.values() if x["key"].startswith("stmt:")]
            n = sum(1 for x in st if x["status"] == "ready")
            return (n == len(st) and n > 0), "%d / %d shelf par" % (n, len(st))
        if key == "row:cards":
            st = [x for x in rows_by_key.values() if x["key"].startswith("card:") or x["key"] == "cards"]
            n = sum(1 for x in st if x["status"] == "ready")
            return (n == len(st) and n > 0), "%d / %d" % (n, len(st))
        if key == "row:bundles":
            st = [x for x in rows_by_key.values() if x["key"].startswith("bundle:")]
            n = sum(1 for x in st if x["status"] == "ready")
            return n > 0, "%d bundle" % n
        return (bool(r) and r["status"] == "ready"), (r["why"] if r else "")
    if key == "sent":
        r = con.execute("SELECT sent_at, sent_by FROM pack_send WHERE month=? ORDER BY id DESC LIMIT 1", (month,)).fetchone()
        return bool(r), ("%s %s" % (r[1], r[0][:16])) if r else ""
    if key == "pm_final":
        try:
            import purchase_app as pa                     # noqa: PLC0415
            return pa._month_status(con, month)["status"] == "final", ""
        except Exception:                                # noqa: BLE001
            return False, ""
    if key == "neft_done":
        r = con.execute("SELECT id FROM purchase_neft_event WHERE month=? AND kind<>'rejected'", (month,)).fetchone() if _has(con, "purchase_neft_event") else None
        return bool(r), ""
    if key == "scans_clear":
        try:
            import porders                                # noqa: PLC0415
            n = len([b for b in porders.unscanned_bills(con) if str(b.get("month")) == month])
            return n == 0, ("%d baaki" % n) if n else ""
        except Exception:                                # noqa: BLE001
            return False, ""
    if key == "sanj_stmts":
        cl = {c["slot"]: c for c in cells(con, month)}
        ok = all(cl.get(k, {}).get("file") for k in ("yes_cur_sanj", "icici_sanj"))
        return ok, ""
    if key == "clinic_stmts":
        cl = {c["slot"]: c for c in cells(con, month)}
        ok = all(cl.get(k, {}).get("file") for k in ("yes_cur_clinic", "icici_clinic"))
        return ok, ""
    if key.startswith("slot:"):
        cl = {c["slot"]: c for c in cells(con, month)}
        return bool(cl.get(key[5:], {}).get("file")), ""
    if key == "clinic_days":
        if not _has(con, "clinic_day_revenue"):
            return False, ""
        n = con.execute("SELECT COUNT(*) FROM clinic_day_revenue WHERE substr(business_date,1,7)=?", (month,)).fetchone()[0]
        return n > 0, "%d din" % n
    return None, ""


def checklist(con, month):
    ensure(con)
    rows, _att = pack_rows(con, month, light=True)
    by = {r["key"]: r for r in rows}
    done = {r[0]: (r[1], r[2]) for r in con.execute("SELECT item_id, done_by, done_at FROM packs_done WHERE month=?", (month,))}
    out = []
    for it in [dict(r) for r in con.execute("SELECT * FROM packs_item WHERE active=1 ORDER BY sort, id")]:
        auto, words = _auto_state(con, month, it["auto_key"], by)
        d = done.get(it["id"])
        out.append(dict(id=it["id"], grp=it["grp"], item=it["item"], auto=(it["auto_key"] is not None), auto_done=auto, words=words,
                        done=bool(d) or bool(auto), done_by=(d[0] if d else ("system" if auto else None)), done_at=(d[1] if d else None)))
    return out


def open_items(con, month):
    return [x for x in checklist(con, month) if not x["done"]]


def needs_you_lines(con):
    """After the 10th: the previous month's empty cells and open checklist items."""
    try:
        ensure(con)
        t = _today()
        if t.day < DUE_DAY:
            return []
        m = _prev_month(t.strftime("%Y-%m"))
        empty = [c for c in cells(con, m) if c["state"] == "empty"]
        out = []
        if empty:
            out.append(dict(cls="warn", target="bank", text="%s: %d statement%s not on the shelf (%s)" % (_month_name(m), len(empty), "" if len(empty) == 1 else "s", ", ".join(c["label"].split(" (")[0] for c in empty[:4]) + (" …" if len(empty) > 4 else ""))))
        op = open_items(con, m)
        if op:
            out.append(dict(cls="warn", target="bank", text="Shavez's month-end checklist for %s: %d item%s open" % (_month_name(m), len(op), "" if len(op) == 1 else "s")))
        return out
    except Exception:                                    # noqa: BLE001
        return []


# ---------------------------------------------------------------- routes
def _owner():
    return _require("checker", unit=_unit)


def _staff():
    return _require("maker", "checker", unit=_unit)


def _tpl(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        return fh.read()


@bp.route("/finance/packs")
def page_packs():
    u, err = _owner()
    if err:
        return err
    return render_template_string(_tpl("packs.html"), month=request.args.get("month") or _prev_month(_today().strftime("%Y-%m")), me=_who(u))


@bp.route("/finance/packs/checklist")
def page_checklist():
    u, err = _staff()
    if err:
        return err
    return render_template_string(_tpl("packs_checklist.html"), month=request.args.get("month") or _prev_month(_today().strftime("%Y-%m")), me=_who(u),
                                  owner=("checker" in set((u or {}).get("roles") or []) or (u or {}).get("role") == "checker"))


@bp.route("/finance/packs/api/state")
def api_state():
    u, err = _owner()
    if err:
        return err
    m = request.args.get("month") or _prev_month(_today().strftime("%Y-%m"))
    if not _good_month(m):
        return jsonify(ok=False, error="bad_month"), 400
    con = _db()
    rows, att = pack_rows(con, m)
    sends = [dict(r) for r in con.execute("SELECT id, sent_at, sent_by, parts, bytes, resend, to_list FROM pack_send WHERE month=? ORDER BY id", (m,))]
    return jsonify(ok=True, month=m, month_name=_month_name(m), cells=cells(con, m), unplaced=unplaced(con), slots=[dict(r) for r in con.execute("SELECT id, key, holder_label, bank, kind, ident_tail FROM stmt_slot ORDER BY sort")],
                   rows=[dict(r, files=r.get("files", [])) for r in rows], attachments=[dict(name=a[0], bytes=len(a[1])) for a in att if a[1]],
                   sends=sends, to=_setting(con, "packs.accountant_to", ""), limit_mb=_setting(con, "packs.mail_limit_mb", "20"),
                   amir=amir_pack(con, m), checklist=checklist(con, m), items=[dict(r) for r in con.execute("SELECT * FROM packs_item ORDER BY sort, id")],
                   inbox_dir=INBOX, today=_today().isoformat())


@bp.route("/finance/packs/api/assign", methods=["POST"])
def api_assign():
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    try:
        fid, sid = int(b.get("file") or 0), int(b.get("slot") or 0)
    except (TypeError, ValueError):
        fid = sid = 0
    if not fid or not sid:
        return jsonify(ok=False, error="bad_request"), 400
    body, code = assign(_db(), fid, sid, _who(u))
    return jsonify(**body), code


@bp.route("/finance/packs/api/tick", methods=["POST"])
def api_tick():
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    m, item = str(b.get("month") or ""), str(b.get("item") or "")
    if not _good_month(m) or item not in ("lab_register", "lab_receipt_book"):
        return jsonify(ok=False, error="bad_request"), 400
    con = _db()
    ensure(con)
    con.execute("INSERT OR REPLACE INTO pack_tick (month, item, done_by, done_at) VALUES (?,?,?,?)", (m, item, _who(u), _now()))
    con.commit()
    return jsonify(ok=True)


@bp.route("/finance/packs/api/send", methods=["POST"])
def api_send():
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    m = str(b.get("month") or "")
    if not _good_month(m):
        return jsonify(ok=False, error="bad_month"), 400
    con = _db()
    ensure(con)
    last = con.execute("SELECT sent_at FROM pack_send WHERE month=? ORDER BY id DESC LIMIT 1", (m,)).fetchone()
    if last:
        try:
            if (dt.datetime.now() - dt.datetime.fromisoformat(last[0])).total_seconds() < REPEAT_MIN * 60:
                return jsonify(ok=True, already=True, message="sent %s -- a repeat within %d minutes is not sent again" % (last[0][:16], REPEAT_MIN)), 200
        except ValueError:
            pass
    body, code = send_pack(con, u, m)
    return jsonify(**body), code


@bp.route("/finance/packs/api/refresh", methods=["POST"])
def api_refresh():
    """On demand: identify what the inbox holds (the fetch itself is the venv cron's; here the identification runs)."""
    u, err = _owner()
    if err:
        return err
    return jsonify(ok=True, **process_inbox(_db()))


@bp.route("/finance/packs/api/item", methods=["POST"])
def api_item():
    """The owner edits the checklist: add / rename / retire an item."""
    u, err = _owner()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    con = _db()
    ensure(con)
    act = str(b.get("action") or "")
    if act == "add":
        grp, item = str(b.get("grp") or "").strip()[:40], str(b.get("item") or "").strip()[:160]
        if not grp or not item:
            return jsonify(ok=False, error="bad_request"), 400
        n = con.execute("SELECT COALESCE(MAX(sort),0)+1 FROM packs_item").fetchone()[0]
        con.execute("INSERT INTO packs_item (grp, item, auto_key, sort, edited_by, edited_at, source) VALUES (?,?,NULL,?,?,?,?)", (grp, item, n, _who(u), _now(), "owner"))
    elif act in ("rename", "retire", "restore"):
        try:
            iid = int(b.get("id") or 0)
        except (TypeError, ValueError):
            iid = 0
        if not con.execute("SELECT 1 FROM packs_item WHERE id=?", (iid,)).fetchone():
            return jsonify(ok=False, error="no_such_item"), 404
        if act == "rename":
            con.execute("UPDATE packs_item SET item=?, edited_by=?, edited_at=? WHERE id=?", (str(b.get("item") or "").strip()[:160], _who(u), _now(), iid))
        else:
            con.execute("UPDATE packs_item SET active=?, edited_by=?, edited_at=? WHERE id=?", (0 if act == "retire" else 1, _who(u), _now(), iid))
    else:
        return jsonify(ok=False, error="bad_action"), 400
    con.commit()
    return jsonify(ok=True)


@bp.route("/finance/packs/api/checklist")
def api_checklist():
    u, err = _staff()
    if err:
        return err
    m = request.args.get("month") or _prev_month(_today().strftime("%Y-%m"))
    if not _good_month(m):
        return jsonify(ok=False, error="bad_month"), 400
    con = _db()
    return jsonify(ok=True, month=m, month_name=_month_name(m), items=checklist(con, m), me=_who(u))


@bp.route("/finance/packs/api/checklist/tick", methods=["POST"])
def api_checklist_tick():
    u, err = _staff()
    if err:
        return err
    b = request.get_json(silent=True) or {}
    m = str(b.get("month") or "")
    try:
        iid = int(b.get("id") or 0)
    except (TypeError, ValueError):
        iid = 0
    if not _good_month(m) or not iid:
        return jsonify(ok=False, error="bad_request"), 400
    con = _db()
    ensure(con)
    it = con.execute("SELECT id, auto_key FROM packs_item WHERE id=? AND active=1", (iid,)).fetchone()
    if not it:
        return jsonify(ok=False, error="no_such_item"), 404
    if it[1]:
        return jsonify(ok=False, error="automatic", message="yeh item system khud bharta hai"), 409
    d = con.execute("SELECT done_by, done_at FROM packs_done WHERE month=? AND item_id=?", (m, iid)).fetchone()
    if d:
        try:
            if (dt.datetime.now() - dt.datetime.fromisoformat(d[1])).total_seconds() < REPEAT_MIN * 60:
                return jsonify(ok=True, already=True, done_by=d[0], done_at=d[1]), 200
        except ValueError:
            pass
        return jsonify(ok=True, already=True, done_by=d[0], done_at=d[1]), 200
    con.execute("INSERT INTO packs_done (month, item_id, done_by, done_at, note) VALUES (?,?,?,?,?)", (m, iid, _who(u), _now(), str(b.get("note") or "")[:120]))
    con.commit()
    return jsonify(ok=True, done_by=_who(u), done_at=_now()), 200


@bp.route("/finance/packs/file/<int:fid>")
def page_file(fid):
    """A shelf file, to the owner (preview)."""
    u, err = _owner()
    if err:
        return err
    r = _db().execute("SELECT local_path, name FROM stmt_file WHERE id=?", (fid,)).fetchone()
    if not r or not r[0] or not os.path.exists(r[0]):
        return "not on the shelf", 404
    return send_file(r[0], as_attachment=False, download_name=re.sub(r"[^A-Za-z0-9._-]", "_", r[1] or "statement")[:80])


@bp.route("/finance/packs/preview/<key>")
def page_preview(key):
    """A built attachment, to the owner: income / upi / digest / neft / bundle_lab_purchase / bundle_owner_expense."""
    u, err = _owner()
    if err:
        return err
    m = request.args.get("month") or _prev_month(_today().strftime("%Y-%m"))
    if not _good_month(m):
        return "bad month", 400
    con = _db()
    if key == "income":
        b, why = build_income(con, m)
        name, mime = "Clinic_income_%s.xlsx" % m, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif key == "upi":
        b, why = build_upi_corrections(con, m)
        name, mime = "UPI_and_corrections_%s.xlsx" % m, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif key == "digest":
        b, why = build_digest(con, m)
        name, mime = "Payment_digest_%s.xlsx" % m, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif key.startswith("bundle_"):
        b, why = build_bundle(con, m, key[7:])
        name, mime = "Bills_%s_%s.pdf" % (key[7:], m), "application/pdf"
    else:
        return "no such preview", 404
    if not b:
        return "not ready: %s" % why, 404
    return send_file(io.BytesIO(b), mimetype=mime, as_attachment=(not mime.endswith("pdf")), download_name=name)


# Amir's pack (Sanjeevni): under /finance/amir/... so the front gate's medical unit lets him in
@bp.route("/finance/amir/pack/<month>/<what>")
def page_amir_pack(month, what):
    u, err = _require("maker", "checker", "viewer", unit="medical")
    if err:
        return err
    if not _good_month(month):
        return "bad month", 400
    con = _db()
    p = amir_pack(con, month)
    if not p["ready"]:
        return "pack abhi tayyar nahi (dono statement shelf par nahi)", 404
    if what == "neft":
        try:
            b = build_neft_paid_sheet(con, month)
        except Exception as ex:                          # noqa: BLE001
            return "sheet nahi ban saki: %s" % _esc(str(ex)[:80]), 500
        return send_file(io.BytesIO(b), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="NEFT_paid_%s.xlsx" % month)
    fid = (p["yes"] if what == "yes" else p["icici"] if what == "icici" else None)
    if not fid:
        return "no such piece", 404
    r = con.execute("SELECT local_path, name FROM stmt_file WHERE id=?", (fid["id"],)).fetchone()
    if not r or not r[0] or not os.path.exists(r[0]):
        return "not on the shelf", 404
    return send_file(r[0], as_attachment=False, download_name=re.sub(r"[^A-Za-z0-9._-]", "_", r[1] or "statement")[:80])


@bp.route("/finance/amir/pack/<month>/seen", methods=["POST"])
def api_amir_seen(month):
    u, err = _require("maker", "checker", "viewer", unit="medical")
    if err:
        return err
    if not _good_month(month):
        return "bad month", 400
    con = _db()
    ensure(con)
    con.execute("INSERT OR IGNORE INTO pack_tick (month, item, done_by, done_at) VALUES (?,?,?,?)", (month, "amir_seen", _who(u), _now()))
    con.commit()
    from flask import redirect                            # noqa: PLC0415
    return redirect("/finance/amir", code=303)
