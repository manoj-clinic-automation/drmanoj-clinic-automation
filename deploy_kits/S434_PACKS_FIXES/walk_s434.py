#!/usr/bin/env python3
"""walk_s434.py -- S434 on the box, BEFORE anything is placed: the new packs.py + packs.html on a COPY of the live finance.db (the live
modules -- purchase_app, finance_app, clinic_day_pdf -- are only imported, from the finance folder, against the copy). Usage:
walk_s434.py <kit dir> <scratch dir> <db copy> <finance dir>. Last line 'WALK OK ...' or 'WALK RED ...'."""
import io
import os
import re
import shutil
import sys

KIT, W, DB, FIN = sys.argv[1:5]
for f in ("packs.py", "packs.html"):
    shutil.copy(os.path.join(KIT, f), W)
for f in ("yes_branch.py", "packs_checklist.html"):
    shutil.copy(os.path.join(FIN, f), W)
os.environ.update(FINANCE_DB=DB, STMT_INBOX=os.path.join(FIN, "statements", "inbox"), STMT_VENV_PYTHON="/nonexistent")
sys.path.insert(0, W)
sys.path.append(FIN)
import packs  # noqa: E402

bad = []
def check(name, cond, info=""):
    print(("  PASS " if cond else "  FAIL ") + name + ("" if cond else "  -- %s" % (info,)))
    if not cond:
        bad.append(name)

M = "2026-08"
con = packs._con()
packs.ensure(con)
con.execute("UPDATE setting SET value='' WHERE key='packs.retired_slots' AND value='yes_cur_clinic'")
for k in ("packs.electricity_expected", "packs.digest_drop_words"):
    con.execute("INSERT OR IGNORE INTO setting (key, value, note) VALUES (?,?,?)", (k, packs.SETTINGS[k][0], packs.SETTINGS[k][1]))
con.commit()
cl = {c["slot"]: c for c in packs.cells(con, M)}
check("the clinic's Yes Bank current row is back", "yes_cur_clinic" in cl)
check("... and reads empty until the branch sends it", (cl.get("yes_cur_clinic") or {}).get("state") == "empty", cl.get("yes_cur_clinic"))
rows, att = packs.pack_rows(con, M)
by = {r["key"]: r for r in rows}
kept, dropped, total = packs.digest_rows(con, M)
check("the digest keeps payments (%d kept, %d set aside, total %.2f)" % (len(kept), len(dropped), total), len(kept) > 0 and total > 0)
check("no appointment line (a patient's name) in the digest", not any("APPOINT" in (r["desc"] or "").upper() for r in kept))
check("no OTP / sign-in / failed / trial line in the digest", not any(re.search(r"OTP|SECURE LINK|UNSUCCESSFUL|TRIAL|EXPIRING", (r["desc"] or "").upper()) for r in kept))
el = by.get("electricity") or {}
print("  electricity: %s" % (el.get("why") or "")[:200])
check("electricity: both bills found for August (ICICI + the Amazon Pay card)", el.get("status") == "ready", el.get("why"))
names = [a[0] for a in att]
print("  attachments: %s" % " | ".join(names))
check("every attachment name is readable (no internal slot codes)", not any(re.search(r"_(yes|icici|card)_|_2026-\d\d\.", n) for n in names), names)
neft = [a for a in att if a[0].startswith("Sanjeevni NEFT letter")]
check("the NEFT letter is a PDF", bool(neft) and neft[0][1][:5] == b"%PDF-", [a[0] for a in att if "NEFT" in a[0]])
inc = [a for a in att if a[0].startswith("Clinic income")]
if inc:
    from openpyxl import load_workbook
    vals = [r for r in load_workbook(io.BytesIO(inc[0][1]))["Income"].iter_rows(values_only=True)]
    lab = {str(r[0]): r[1] for r in vals}
    parts = sum(float(lab.get(k) or 0) for k in lab if str(k).startswith("of which"))
    check("the income summary adds up (cash + online + card + split = total)", abs(parts - float(lab.get("Total Rs") or 0)) < 0.5, lab)
upi = [a for a in att if a[0].startswith("UPI totals")]
if upi:
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(upi[0][1]))["Pharmacy cash-UPI corrections"]
    cr = list(ws.iter_rows(values_only=True))[1:]
    print("  pharmacy corrections in August: %d" % len(cr))
    check("no correction row reads Rs 0", all((r[2] or 0) > 0 for r in cr if r and str(r[0] or "").startswith("20")), cr[:3])
yes = [r for r in rows if r["key"].startswith("stmt:yes_")]
check("the five Yes Bank statements the branch sent are still ready", sum(1 for r in yes if r["status"] == "ready") == 5, [(r["key"], r["status"]) for r in yes])
# the page renders
from flask import Flask
app = Flask("walk434")
packs._owner = lambda: ({"user": "walk"}, None)
packs.init(app, packs._con, None)
c = app.test_client()
p = c.get("/finance/packs?month=%s" % M)
check("the page renders with its folded sections", p.status_code == 200 and b'id="d_pack"' in p.data and b'id="topcard"' in p.data, p.status_code)
s = c.get("/finance/packs/api/state?month=%s" % M)
check("the page's data answers", s.status_code == 200 and s.get_json().get("ok"), s.status_code)
lp = c.get("/finance/packs/preview/neft_letter?month=%s" % M)
check("the letter preview is a PDF", lp.status_code == 200 and lp.data[:5] == b"%PDF-", lp.status_code)
ready = sum(1 for r in rows if r["status"] == "ready")
con.close()
print("WALK OK -- August %d of %d ready; digest %d payments; electricity %s" % (ready, len(rows), len(kept), el.get("status"))
      if not bad else "WALK RED -- %d check(s) failed: %s" % (len(bad), "; ".join(bad)))
sys.exit(1 if bad else 0)
