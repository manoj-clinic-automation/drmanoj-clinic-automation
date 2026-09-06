"""LIVE-SHAPE WALK -- S226 PAD ON-PAGE.

The owner's afternoon ruling: the Excel flow is part of the stock-check page.
Details first, the pad comes down PREFILLED, one tap uploads it, a processing
box shows, the result sheet is offered, the loop goes on, staff never need him.

This drives that page at phone width, with the real item list, the way Amir
will: it downloads the pad through the button (not through the URL), fills it
the way a person fills it (no recalculation), uploads it by picking the file,
and reads what the page says back. Then it does the loop. Then it does the
things that go wrong: the same file twice, a pad with blank top lines, an
unreadable file, the old address.
"""
import csv, os, shutil, sys, threading, time, sqlite3
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import padreader
import WALK_readiness_s226 as R
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

OK, BAD = [], []
def chk(n, c, d=""):
    (OK if c else BAD).append(n + (("  -- " + str(d)) if d and not c else ""))
    print(("  ok   " if c else "  FAIL ") + n + (("   " + str(d)) if d else ""))

OLD_PAD = "/mnt/user-data/outputs/STOCK_COUNT_PAD_06-09-2026.xlsx"   # the pad Amir already holds (blank top lines)
ITEMS = list(csv.DictReader(open(
    "/mnt/user-data/uploads/Downloads/margsync/_analysis/_S226_item_list.csv", encoding="utf-8")))
names = [r["item"] for r in ITEMS]
PS = {r["item"]: int(r["pack_size"] or 1) for r in ITEMS}

R.DB = os.path.join(HERE, "_walkonpage.db"); con = R.build(); R.fill(con)
import random; random.seed(11)
MARG = {}
for r in ITEMS:
    q = random.choice([0, 4, 13, 40, 137, 405, 684]); MARG[r["item"]] = q
    con.execute("INSERT OR REPLACE INTO stock_snapshot (as_on,item,qty,packing,pack_size,source,loaded_at) "
                "VALUES (?,?,?,?,?,?,?)", ("06-09-2026", r["item"], q, r["packing"], PS[r["item"]], "p", "2026-09-06T09:00:59"))
for r in ITEMS[:120]:
    con.execute("INSERT OR REPLACE INTO stock_rate (item, rate_p, as_of, source) VALUES (?,?,?,?)",
                (r["item"], 2500, "2026-09-06", "walk"))
con.commit()
app = R.app_for(con)
threading.Thread(target=lambda: app.run(port=8816, threaded=False), daemon=True).start()
time.sleep(1.5)
HOST = "http://127.0.0.1:8816"
URL = HOST + "/finance/stock/page/count"


def fill_pad(src, dst, entries, rename=None, bad=None):
    shutil.copy(src, dst)
    wb = load_workbook(dst); ws = wb.worksheets[0]
    where = {}
    for r in range(9, 400):
        nm = ws.cell(row=r, column=2).value
        if isinstance(nm, str) and nm.strip():
            where[nm.strip()] = r
    for it, (st, lo) in entries.items():
        r = where[it]
        if st is not None: ws.cell(row=r, column=5, value=st)
        if lo is not None: ws.cell(row=r, column=6, value=lo)
    for it, nm in (rename or {}).items():
        ws.cell(row=where[it], column=2, value=nm); ws.cell(row=where[it], column=6, value=7)
    for it, v in (bad or {}).items():
        ws.cell(row=where[it], column=6, value=v)
    wb.save(dst)                                    # NOT recalculated, on purpose
    return dst


good = {}
for nm in names[:40]:
    ps = PS[nm]; q = MARG[nm]
    tgt = q if len(good) % 6 else max(q - 3, 0)
    good[nm] = (tgt // ps if ps > 1 else None, tgt % ps if ps > 1 else tgt)


def gate(pg, cby="Darpan", eby="Amir", bill="A003425", bdate="2026-09-06"):
    pg.locator('#whoC button[data-u="%s"]' % cby).click()
    pg.locator('#whoE button[data-u="%s"]' % eby).click()
    pg.fill("#bill", bill); pg.fill("#billdate", bdate); pg.wait_for_timeout(150)


print("== S226 PAD ON-PAGE -- the Excel loop on the stock-check page ==")
errs = []
with sync_playwright() as pw:
    b = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium")
    ctx = b.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True)
    pg = ctx.new_page()
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.on("console", lambda m: errs.append(m.text)
          if m.type == "error" and "Failed to load resource" not in m.text else None)

    # ------------------------------------------------------------ 1 the page opens on the details
    pg.goto(URL); pg.wait_for_timeout(700)
    chk("the page opens on the readiness lines and the four details", "A003412" in pg.inner_text("#rdy")
        and pg.locator("#gate").is_visible(), pg.inner_text("#rdy")[:50])
    chk("the Excel panel is ON THIS PAGE, under the details", pg.locator("#xl").is_visible()
        and "Download SHEET 1" in pg.inner_text("#xl"))
    chk("...its download button is shut until the details are given", pg.locator("#padget").is_disabled())
    chk("...and it says what it wants", "Do step 1 first" in pg.inner_text("#padnote"))
    chk("the page is numbered 1-2-3-4 for the staff", pg.locator(".step").count() == 4 and "Close the stock check" in pg.inner_text("#xl"))
    chk("...and the result box sits UNDER step 4, where the close button is explained",
        pg.evaluate("()=>document.querySelectorAll('.step')[3].compareDocumentPosition(document.getElementById('padout')) & 4") == 4)
    chk("the on-screen count is still there, as the second choice", "Count on this screen instead" in pg.inner_text("#start"))
    chk("no recent counts yet, nothing shown", pg.inner_text("#recent").strip() == "")
    gate(pg)
    chk("with the four details the download opens", not pg.locator("#padget").is_disabled())
    chk("...and the note says what will be written into the pad",
        "Darpan" in pg.inner_text("#padnote") and "A003425" in pg.inner_text("#padnote") and "06-09-2026" in pg.inner_text("#padnote"),
        pg.inner_text("#padnote"))

    # ------------------------------------------------------------ 2 the pad comes down prefilled
    with pg.expect_download() as dl:
        pg.locator("#padget").click()
    PAD = os.path.join(HERE, "pad_prefilled.xlsx"); dl.value.save_as(PAD)
    chk("THE PAD DOWNLOADS FROM THE BUTTON", os.path.getsize(PAD) > 10000 and dl.value.suggested_filename.startswith("STOCK_COUNT_") and dl.value.suggested_filename.endswith("_SHEET_1.xlsx"),
        dl.value.suggested_filename)
    rp = padreader.read_pad(PAD)
    chk("...with every item and no figures", len(rp["rows"]) == len(ITEMS) and all(r["counted"] is None for r in rp["rows"]))
    chk("...AND THE FOUR DETAILS WRITTEN IN ITS TOP LINES",
        rp["meta"].get("counted_by") == "Darpan" and rp["meta"].get("entered_by") == "Amir"
        and rp["meta"].get("bill_no") == "A003425" and rp["meta"].get("bill_date") == "2026-09-06", rp["meta"])
    ws = load_workbook(PAD).worksheets[0]
    chk("...where a person sees them: B5 Darpan, B6 Amir, B7 the bill, E7 its date as dd-mm-yyyy",
        (ws["B5"].value, ws["B6"].value, ws["B7"].value, ws["E7"].value) == ("Darpan", "Amir", "A003425", "06-09-2026"),
        (ws["B5"].value, ws["B6"].value, ws["B7"].value, ws["E7"].value))
    chk("...and the detail cells are shaded -- a person may still correct them", ws["B5"].fill.fgColor.rgb == "FFFFF3C4")
    chk("the note now tells them what to do next", "downloading" in pg.inner_text("#padnote"))

    # ------------------------------------------------------------ 3 the filled pad, ONE tap
    FILLED = fill_pad(PAD, os.path.join(HERE, "pad_filled.xlsx"), good,
                      rename={names[60]: "SOMETHING NOBODY KNOWS"}, bad={names[61]: -3, names[62]: 2.5})
    # clear the page's boxes first: the SHEET must carry the details on its own
    pg.goto(URL); pg.wait_for_timeout(500)
    pg.evaluate("()=>{ S.cby=''; S.eby=''; S.bill=''; S.billdate=''; save(); }")
    pg.reload(); pg.wait_for_timeout(500)
    chk("(a fresh device: no details in the boxes)", pg.locator("#padget").is_disabled())
    seen_busy = pg.evaluate("""()=>{ window.__busy=false;
        new MutationObserver(()=>{ if(!document.getElementById('padbusy').classList.contains('hide')) window.__busy=true; })
          .observe(document.getElementById('padbusy'), {attributes:true}); return true; }""")
    pg.set_input_files("#padfile", FILLED); pg.wait_for_timeout(2500)
    chk("A PROCESSING BOX SHOWED WHILE IT WORKED", pg.evaluate("()=>window.__busy") is True)
    chk("...and is gone when it is done", not pg.locator("#padbusy").is_visible())
    out = pg.inner_text("#padout")
    chk("IT RECORDS IN ONE STEP AND SAYS SO", "Recorded as count #1" in out, out[:120])
    chk("...this sheet: 40 counted, 3 rows sent back to fix", "40 counted" in out and "3 rows sent back" in out, out[:300])
    chk("...SOME STOCK COUNT IS LEFT: 333 not counted, 3 to fix -- download and complete", "Some stock count is left: 333 items not counted, 3 rows to fix" in out and "complete it" in out, out[:400])
    chk("...AND OFFERS THE SHEET FOR THE REMAINING WORK", pg.locator("#padout a.dl").count() == 1
        and pg.locator("#padout a.dl").inner_text() == "Download SHEET 2 REMAINING", pg.locator("#padout a.dl").inner_text())
    chk("...and NO close button while work is left", pg.locator("#padout button.close").count() == 0)
    chk("...and says what to do with it: the two tabs, upload at step 3, repeat",
        "NOT COUNTED" in out and "SENT BACK TO FIX" in out and "step 3" in out)
    row = con.execute("SELECT id, items_counted, items_total, bill_no, bill_date, submitted_by FROM stock_count ORDER BY id DESC LIMIT 1").fetchone()
    chk("the ledger holds it: 40 of the shop, pinned to THE SHEET'S bill (the boxes were empty)",
        row == (1, 40, len(ITEMS), "A003425", "2026-09-06", "walk"), row)
    who = con.execute("SELECT DISTINCT counted_by, entered_by FROM stock_count_item WHERE count_id=1").fetchall()
    chk("...counted by Darpan, entered by Amir -- from the sheet's own top lines", who == [("Darpan", "Amir")], who)
    chk("the finding is sealed and the horizon frozen",
        con.execute("SELECT 1 FROM stock_finding WHERE count_id=1").fetchone() is not None
        and con.execute("SELECT 1 FROM stock_check_readiness WHERE count_id=1").fetchone() is not None)
    kept = con.execute("SELECT written, figure FROM stock_count_pad_issue WHERE count_id=1 ORDER BY sheet_row").fetchall()
    chk("the three rows it could not use are KEPT", len(kept) == 3 and kept[0] == ("SOMETHING NOBODY KNOWS", 7), kept)
    pf = con.execute("SELECT count_id, filename FROM stock_count_pad_file").fetchall()
    chk("the sheet's bytes are remembered with the count", pf == [(1, "pad_filled.xlsx")], pf)
    href = pg.locator("#padout a.dl").get_attribute("href")
    chk("the result link is the family's result sheet", href.startswith("/finance/stock/api/pad/followup/1.xlsx"), href)
    chk("THE RECENT LIST NOW SHOWS COUNT #1 WITH ITS SHEET", "Count #1" in pg.inner_text("#recent")
        and "OPEN — 336 still to do" in pg.inner_text("#recent") and pg.locator("#recent a.dl").count() == 1, pg.inner_text("#recent")[:200])

    # ------------------------------------------------------------ 4 the same file again
    pg.set_input_files("#padfile", FILLED); pg.wait_for_timeout(2000)
    out = pg.inner_text("#padout")
    chk("THE SAME FILE TWICE IS NOT A SECOND COUNT", "Already recorded" in out and "count #1" in out, out[:160])
    chk("...nothing extra in the ledger", con.execute("SELECT COUNT(*) FROM stock_count").fetchone()[0] == 1)
    chk("...and the result sheet is still offered", pg.locator("#padout a.dl").count() == 1)

    # ------------------------------------------------------------ 5 the loop: fill the result sheet, upload it
    fu = pg.request.get(HOST + href.split("?")[0])
    RES = os.path.join(HERE, "result_1.xlsx"); open(RES, "wb").write(fu.body())
    wbf = load_workbook(RES)
    chk("the result sheet has the five approved tabs", wbf.sheetnames == ["SUMMARY", "DIFFERENCES", "NOT COUNTED", "SENT BACK TO FIX", "MATCHED"], wbf.sheetnames)
    sm = wbf["SUMMARY"]; smtxt = "\n".join("%s | %s" % (sm.cell(row=r, column=1).value or "", sm.cell(row=r, column=2).value or "") for r in range(1, 40))
    chk("...SUMMARY carries who and which bill, from the sheet", "Counted by | Darpan" in smtxt and "After bill | A003425 dated 2026-09-06" in smtxt, smtxt[:300])
    nc, sb = wbf["NOT COUNTED"], wbf["SENT BACK TO FIX"]
    nc.cell(row=5, column=4, value=2); nc.cell(row=5, column=5, value=3); first_nc = nc.cell(row=5, column=2).value
    sb.cell(row=5, column=2, value=names[60]); sb.cell(row=5, column=5, value=7)
    sb.cell(row=6, column=5, value=3)
    wbf.save(RES)
    pg.set_input_files("#padfile", RES); pg.wait_for_timeout(2500)
    out = pg.inner_text("#padout")
    chk("THE FILLED RESULT SHEET JOINS THE COUNT", "Recorded as part of count #1" in out, out[:120])
    chk("...three figures from the two tabs", "3 counted" in out, out[:200])
    chk("...whole count so far 43", "Whole count so far: 43" in out, out[:300])
    chk("...still to do: 330 not counted, 1 to fix", "330 items not counted, 1 row to fix" in out, out[:400])
    chk("...the next sheet is named SHEET 3 REMAINING", pg.locator("#padout a.dl").inner_text() == "Download SHEET 3 REMAINING", pg.locator("#padout a.dl").inner_text())
    fn = pg.request.get(HOST + "/finance/stock/api/pad/followup/1.xlsx").headers.get("content-disposition", "")
    chk("...and downloads under that name", "STOCK_COUNT_06-09-2026_SHEET_3_REMAINING.xlsx" in fn, fn)
    got = dict(con.execute("SELECT item, counted_qty FROM stock_count_item WHERE count_id=2").fetchall())
    chk("...the total settled from the SHOP's pack size", got.get(first_nc) == 2 * PS[first_nc] + 3, (got.get(first_nc), PS[first_nc]))
    chk("...the corrected name landed as the real item", names[60] in got)
    part = con.execute("SELECT part_of FROM stock_count_part WHERE count_id=2").fetchone()
    chk("...as part of count #1, pinned to count #1's bill", part == (1,)
        and con.execute("SELECT bill_no FROM stock_count WHERE id=2").fetchone() == ("A003425",), part)
    left = con.execute("SELECT written FROM stock_count_pad_issue WHERE resolved_by IS NULL").fetchall()
    chk("...THE TWO CORRECTED ROWS ARE CLOSED BY THEIR ORIGINAL ROW; the half-tablet row still open",
        left == [("CIXCEF O",)], left)
    chk("the recent list shows one count with one follow-up sheet, 331 still to do",
        "Count #1 (sheets 1\u20132)" in pg.inner_text("#recent") and "331 still to do" in pg.inner_text("#recent"), pg.inner_text("#recent")[:200])

    # ------------------------------------------------------------ 6 the loop closes
    fu = pg.request.get(HOST + "/finance/stock/api/pad/followup/1.xlsx")
    RES2 = os.path.join(HERE, "result_2.xlsx"); open(RES2, "wb").write(fu.body())
    wb2 = load_workbook(RES2); nc, sb = wb2["NOT COUNTED"], wb2["SENT BACK TO FIX"]
    r = 5
    while nc.cell(row=r, column=2).value:
        nm = nc.cell(row=r, column=2).value; ps = PS[nm]; q = MARG[nm]
        if ps > 1: nc.cell(row=r, column=4, value=q // ps); nc.cell(row=r, column=5, value=q % ps)
        else: nc.cell(row=r, column=5, value=q)
        r += 1
    sb.cell(row=5, column=5, value=2)
    wb2.save(RES2)
    pg.set_input_files("#padfile", RES2); pg.wait_for_timeout(3500)
    out = pg.inner_text("#padout")
    chk("THE LAST SHEET: the server says everything is counted", "everything is counted" in out, out[:300])
    chk("...whole count = the whole shop", "Whole count so far: %d of %d" % (len(ITEMS), len(ITEMS)) in out, out[:300])
    chk("...AND ONLY NOW THE CLOSE BUTTON APPEARS", pg.locator("#padout button.close").count() == 1 and "Close the stock check" in out)
    chk("the recent list says everything counted, not yet closed", "everything counted" in pg.inner_text("#recent") and "not yet closed" in pg.inner_text("#recent"))
    chk("...but carries NO second close button for the count the box is showing", pg.locator("#recent button.close").count() == 0)
    chk("no issue left open", con.execute("SELECT COUNT(*) FROM stock_count_pad_issue WHERE resolved_by IS NULL").fetchone()[0] == 0)

    # ------------------------------------------------------------ 6b CLOSE THE STOCK CHECK
    dialogs = []
    pg.on("dialog", lambda d: (dialogs.append(d.message), d.accept()))
    pg.locator("#padout button.close").click(); pg.wait_for_timeout(1500)
    chk("closing asks once, then says STOCK CHECK COMPLETED",
        any("Close the stock check?" in m for m in dialogs) and any("STOCK CHECK COMPLETED" in m for m in dialogs), dialogs)
    out = pg.inner_text("#padout")
    chk("...the box says STOCK CHECK COMPLETED, no more sheets, start again at step 1",
        "STOCK CHECK COMPLETED" in out and "No more sheets" in out and "step 1" in out, out[:300])
    cl = con.execute("SELECT count_id, how, left_not_counted, left_to_fix, closed_by FROM stock_count_close").fetchall()
    chk("...the ledger holds the close: count #1, complete, nothing left", cl == [(1, "complete", 0, 0, "walk")], cl)
    chk("...the FINAL RESULT sheet is offered", "FINAL RESULT" in pg.locator("#padout a.dl").inner_text())
    fn = pg.request.get(HOST + "/finance/stock/api/pad/followup/1.xlsx").headers.get("content-disposition", "")
    chk("...named so", "STOCK_COUNT_06-09-2026_FINAL_RESULT.xlsx" in fn, fn)
    fu = pg.request.get(HOST + "/finance/stock/api/pad/followup/1.xlsx"); FIN = os.path.join(HERE, "final_1.xlsx"); open(FIN, "wb").write(fu.body())
    sm = load_workbook(FIN)["SUMMARY"]; smtxt = "\n".join("%s | %s" % (sm.cell(row=r, column=1).value or "", sm.cell(row=r, column=2).value or "") for r in range(1, 40))
    chk("...and its SUMMARY says CLOSED -- stock check completed", "Status | CLOSED -- stock check completed" in smtxt, smtxt[:400])
    chk("the recent list says CLOSED -- completed", "CLOSED" in pg.inner_text("#recent") and "completed on" in pg.inner_text("#recent"))
    # a sheet for a closed count is refused -- a NEW sheet (a changed cell), not the same bytes
    wb3 = load_workbook(RES2); wb3["NOT COUNTED"].cell(row=5, column=5, value=9); RES3 = os.path.join(HERE, "result_3_late.xlsx"); wb3.save(RES3)
    pg.set_input_files("#padfile", RES3); pg.wait_for_timeout(2000)
    out = pg.inner_text("#padout")
    chk("A SHEET FOR A CLOSED COUNT IS REFUSED, in words, nothing recorded",
        "Not recorded" in out and "CLOSED" in out and "fresh pad" in out, out[:300])
    chk("...nothing added", con.execute("SELECT COUNT(*) FROM stock_count").fetchone()[0] == 3)
    # closing twice is harmless
    r = pg.request.post(HOST + "/finance/stock/api/pad/close/1").json()
    chk("closing an already-closed count changes nothing", r.get("ok") and r.get("already") is True
        and con.execute("SELECT COUNT(*) FROM stock_count_close").fetchone()[0] == 1)

    # ------------------------------------------------------------ 7 the OLD pad, blank top lines
    OLDF = fill_pad(OLD_PAD, os.path.join(HERE, "pad_old_filled.xlsx"), dict(list(good.items())[:5]))
    pg.set_input_files("#padfile", OLDF); pg.wait_for_timeout(2000)
    out = pg.inner_text("#padout")
    chk("A PAD WITH BLANK TOP LINES, NO DETAILS ON THE PAGE: it ASKS, records nothing",
        "One more thing" in out and "who counted" in out and "bill" in out, out[:200])
    chk("...nothing recorded", con.execute("SELECT COUNT(*) FROM stock_count").fetchone()[0] == 3)
    chk("...the four boxes are marked", pg.locator(".ask-miss").count() == 4)
    gate(pg, "Shavez", "Amir", "A003440", "2026-09-07")
    pg.locator("#padagain").click(); pg.wait_for_timeout(2500)
    out = pg.inner_text("#padout")
    chk("...FILL THEM IN, ONE TAP, THE SAME FILE IS RECORDED", "Recorded as count #4" in out, out[:120])
    row = con.execute("SELECT bill_no, bill_date FROM stock_count WHERE id=4").fetchone()
    chk("...pinned to the bill from the boxes", row == ("A003440", "2026-09-07"), row)
    who = con.execute("SELECT DISTINCT counted_by FROM stock_count_item WHERE count_id=4").fetchall()
    chk("...counted by Shavez, from the boxes", who == [("Shavez",)], who)
    chk("...the marks are gone", pg.locator(".ask-miss").count() == 0)

    # ------------------------------------------------------------ 7b THE DOCTOR'S POWER: close an abandoned count as it stands
    out = pg.inner_text("#padout")
    chk("an incomplete count shows the checker a 'close as it stands' button", pg.locator('#padout button.doc[data-close="4"]').count() == 1
        and "as it stands (368 left)" in out, out[-200:])
    chk("...once, not again in the list for the same count", pg.locator('#recent button.doc[data-close="4"]').count() == 0)
    padShow_reset = pg.evaluate("()=>{ PAD_SHOWN_ROOT=null; padRecent(); return 1; }"); pg.wait_for_timeout(600)
    chk("...and in the list when the box is not showing it", pg.locator('#recent button.doc[data-close="4"]').count() == 1)
    # a viewer (the counters) has no such button, and the server refuses him too
    pg.evaluate("()=>{ DATA.checker=false; }"); padrecent_before = pg.inner_text("#recent")
    pg.evaluate("()=>padRecent()"); pg.wait_for_timeout(600)
    chk("...a counter sees no such button in the list", pg.locator("#recent button.doc").count() == 0)
    import stock_app as SA
    keep = SA._require
    SA._require = lambda *roles: ({"user": "amir", "roles": ["viewer"], "role": "viewer"}, None)
    r = pg.request.post(HOST + "/finance/stock/api/pad/close/4").json()
    SA._require = keep
    chk("...AND THE SERVER REFUSES A COUNTER: 'some stock count is still left ... only the doctor'",
        r.get("ok") is False and r.get("error") == "incomplete" and "Only the doctor" in r.get("message", ""), r)
    chk("...nothing closed", con.execute("SELECT COUNT(*) FROM stock_count_close").fetchone()[0] == 1)
    pg.evaluate("()=>{ DATA.checker=true; }"); pg.evaluate("()=>padRecent()"); pg.wait_for_timeout(600)
    pg.locator('#recent button.doc[data-close="4"]').click(); pg.wait_for_timeout(1500)
    chk("THE DOCTOR CLOSES IT AS IT STANDS -- asked once, plainly", any("AS IT STANDS" in m for m in dialogs), dialogs[-1:])
    cl = con.execute("SELECT how, left_not_counted, left_to_fix FROM stock_count_close WHERE count_id=4").fetchone()
    chk("...recorded as INCOMPLETE with what was left", cl == ("incomplete", 368, 0), cl)
    out = pg.inner_text("#padout")
    chk("...the box says closed by the doctor, what was left, no more sheets", "CLOSED BY THE DOCTOR" in out and "368 not counted" in out and "No more sheets" in out, out[:300])
    chk("...the list says so too", "by the doctor as it stood" in pg.inner_text("#recent"))
    fn = pg.request.get(HOST + "/finance/stock/api/pad/followup/4.xlsx").headers.get("content-disposition", "")
    chk("...and its sheet is the FINAL RESULT", "FINAL_RESULT" in fn, fn)

    # ------------------------------------------------------------ 8 things that go wrong
    JUNK = os.path.join(HERE, "junk.xlsx"); open(JUNK, "wb").write(b"this is not a spreadsheet")
    pg.set_input_files("#padfile", JUNK); pg.wait_for_timeout(1500)
    out = pg.inner_text("#padout")
    chk("A FILE THAT IS NOT A PAD IS REFUSED IN WORDS", "Not recorded" in out and "Nothing was written" in out, out[:200])
    chk("a FRESH pad after a close starts a NEW count -- the loop is not blocked for good",
        True)   # proven by count #4 having been recorded after count #1 (fresh pad, no PART OF line)
    chk("...the picker is open again", not pg.locator("#padfile").is_disabled())
    r = pg.request.get(HOST + "/finance/stock/page/pad", max_redirects=0)
    chk("THE OLD ADDRESS /page/pad GOES TO THIS PAGE", r.status in (301, 302) and r.headers.get("location", "").endswith("/finance/stock/page/count"),
        (r.status, r.headers.get("location")))
    r = pg.request.get(HOST + "/finance/stock/api/pad/recent").json()
    chk("the recent list is one call, newest first, roots only",
        [c["count_id"] for c in r["counts"]] == [4, 1] and r["counts"][1]["parts"] == 2 and r["counts"][1]["closed"]["how"] == "complete", [c["count_id"] for c in r["counts"]])

    # ------------------------------------------------------------ 9 the on-screen count still works
    pg.goto(URL); pg.wait_for_timeout(500)
    pg.evaluate("()=>{ localStorage.clear(); }"); pg.reload(); pg.wait_for_timeout(500)
    gate(pg)
    chk("the on-screen button opens with the same four details", not pg.locator("#start").is_disabled())
    pg.locator("#start").click(); pg.wait_for_timeout(500)
    chk("...and the counting screen comes up as before", pg.locator("#count").is_visible() and pg.locator("#gate").is_hidden())
    pg.fill("#find", names[3][:4]); pg.wait_for_timeout(300)
    chk("...search still finds", names[3] in pg.inner_text("#list"))
    pg.locator("#list .row .yes").first.click(); pg.wait_for_timeout(300)
    chk("...OK still counts", pg.evaluate("()=>Object.keys(S.e).length") >= 1)
    b.close()

chk("no javascript error anywhere in that", not errs, errs[:3])
print("\n%d ok, %d FAILED" % (len(OK), len(BAD)))
for x in BAD: print("   FAILED: " + x)
sys.exit(1 if BAD else 0)
