"""LIVE-SHAPE WALK -- S228 LOSS DESK: the owner's own page.

The owner, at the S227 close: "a real assessment of the loss of pharmacy from my
side, item and qty wise, for my record -- I do the play here"; "the owner ticks
items; sees loss at MRP and at purchase price; high-cost and high-volume items in
an upper section"; "shareable to Darpan on screen and printable: A4 portrait,
core data only ... saved on the VPS as the copy shared with him"; "Darpan's
answers may recover some amount: the owner logs recoveries against the shared
list"; "Darpan's lists in turns move here, under the owner".

The shop here: a tablet 40 strips short at Rs 12 a tablet (the money), a knee
brace 2 pieces short at Rs 1,500 (the money), a cheap tablet 15 strips short but
only Rs 75 (the quantity), a three-tablet line (the rest), an item that has never
sold and so has no MRP at all, one line OVER (which has no business on a loss
desk), and a dozen ordinary medicines to feed Darpan's lists.

Everything is checked at phone width in a real browser -- the page the owner
actually holds.
"""
import io, json, os, shutil, sys, threading, time
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import WALK_readiness_s226 as R
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

OK, BAD = [], []
def chk(n, c, d=""):
    (OK if c else BAD).append(n + (("  -- " + str(d)) if d and not c else ""))
    print(("  ok   " if c else "  FAIL ") + n + (("   " + str(d)) if d else ""))

import re as _re
def sale_key(s): return _re.sub(r"\s+", " ", _re.sub(r"[^A-Z0-9 ]+", " ", s.upper())).strip()

TYR = "TYRO BR TAB"                 # 40 strips short, Rs 12 a tab  -> Rs 4,800  MONEY
KNEE = "KNEE BRACE HINGED UNISON"   # 2 pcs short, Rs 1,500 each    -> Rs 3,000  MONEY
CHEAP = "CHEAPO SALT TAB"           # 15 strips short, 50 paise a tab -> Rs 75   QUANTITY
SMALL = "SMALLGAP BR TAB"           # 3 tabs short, Rs 2 a tab      -> Rs 6      REST
NOPR = "NOPRICE SYP 100ML"          # never sold, never rated       -> no MRP    REST
OVER = "OVERAGE TAB"                # 1 strip OVER                  -> not here at all

ITEMS = [dict(item=TYR, packing="1*10", pack_size=10), dict(item=KNEE, packing="1*1", pack_size=1),
         dict(item=CHEAP, packing="1*10", pack_size=10), dict(item=SMALL, packing="1*10", pack_size=10),
         dict(item=NOPR, packing="1*1", pack_size=1), dict(item=OVER, packing="1*10", pack_size=10)]
for i in range(1, 15):
    ITEMS.append(dict(item="MED-%03d DOLO TAB" % i, packing="1*10", pack_size=10))
names = [r["item"] for r in ITEMS]; PS = {r["item"]: r["pack_size"] for r in ITEMS}
MEDS = [n for n in names if n.startswith("MED-")]

R.DB = os.path.join(HERE, "_walkloss.db"); con = R.build(); R.fill(con)
for t in ("stock_snapshot", "purchase_bill", "purchase_line", "sale_line_item", "purchase_export"):
    con.execute("DELETE FROM %s" % t)
MARG = {TYR: 700, KNEE: 5, CHEAP: 400, SMALL: 100, NOPR: 8, OVER: 100}
for i, n in enumerate(MEDS): MARG[n] = 100 + 10 * i
# the last purchase rate = what it cost us; NOPRICE has none, on purpose
RATE = {TYR: 900, KNEE: 100000, CHEAP: 35, SMALL: 150, OVER: 700}
for i, n in enumerate(MEDS): RATE[n] = 300
for r in ITEMS:
    con.execute("INSERT OR REPLACE INTO stock_snapshot (as_on,item,qty,packing,pack_size,source,loaded_at) VALUES (?,?,?,?,?,?,?)",
                ("06-09-2026", r["item"], MARG[r["item"]], r["packing"], PS[r["item"]], "p", "2026-09-06T09:00:00"))
    if r["item"] in RATE:
        con.execute("INSERT OR REPLACE INTO stock_rate (item, rate_p, as_of, source) VALUES (?,?,?,?)",
                    (r["item"], RATE[r["item"]], "2026-09-06", "walk"))
seq = 500
def sale(nm, day, amt):
    """amount_p is the printed rate of a FULL STRIP, as Marg writes it."""
    global seq; seq += 1
    con.execute("INSERT INTO sale_line_item (day_entry_id, unit, business_date, bill_no, is_return, seq, item_name, item_key, qty_raw, pack, amount_p) "
                "VALUES (1,'medical',?,?,0,?,?,?,?,?,?)", (day, "A00%04d" % seq, seq, nm, sale_key(nm), "1:0",
                                                           "1*1" if PS[nm] == 1 else "1*10", amt))
STRIP_RS = {TYR: 12000, KNEE: 150000, CHEAP: 500, SMALL: 2000, OVER: 1000}
for i, n in enumerate(MEDS): STRIP_RS[n] = 4000
for nm, amt in STRIP_RS.items():
    for k in range(3): sale(nm, "2026-09-0%d" % (k + 1), amt)
con.commit()

ARCH = os.path.join(HERE, "pad_uploads")
if os.path.isdir(ARCH): shutil.rmtree(ARCH)

# ONE app, one server -- stock_app.init() plants the login on the module, so two
# apps in one process would share the last one planted and the door would go
# untested. WHO[0] is who is knocking; the walk turns it and the server obeys.
from flask import Flask
import stock_app
WHO = ["checker"]
def _login(*roles):
    r = WHO[0]
    if r in roles:
        return {"user": "walk_" + r, "roles": ["medical." + r], "role": r}, None
    return None, ({"ok": False, "error": "forbidden", "message": "Not your page."}, 403)
app = Flask("walk_loss")
stock_app.init(app, lambda: con, _login, unit="medical")
threading.Thread(target=lambda: app.run(port=8851, threaded=False), daemon=True).start()
time.sleep(1.6)
HOST = "http://127.0.0.1:8851"
errs = []


def as_staff(fn):
    """Knock as a member of staff, then put the owner back."""
    WHO[0] = "maker"
    try:
        return fn()
    finally:
        WHO[0] = "checker"


def fill_pad(src, dst, entries):
    shutil.copy(src, dst); wb = load_workbook(dst); ws = wb.worksheets[0]; where = {}
    for r in range(9, 400):
        nm = ws.cell(row=r, column=2).value
        if isinstance(nm, str) and nm.strip(): where[nm.strip()] = r
    for it, (st, lo) in entries.items():
        r = where[it]
        if st is not None: ws.cell(row=r, column=5, value=st)
        if lo is not None: ws.cell(row=r, column=6, value=lo)
    wb.save(dst); return dst


COUNTED = {TYR: 300, KNEE: 3, CHEAP: 250, SMALL: 97, NOPR: 5, OVER: 110}
for i, n in enumerate(MEDS): COUNTED[n] = MARG[n] - 30       # 3 strips each, for Darpan's lists
good = {nm: (COUNTED[nm] // PS[nm], COUNTED[nm] % PS[nm]) for nm in names}
SHORTS = [nm for nm in names if COUNTED[nm] < MARG[nm]]


def gate(pg):
    pg.locator('#whoC button[data-u="Darpan"]').click(); pg.locator('#whoE button[data-u="Amir"]').click()
    pg.fill("#bill", "A003425"); pg.fill("#billdate", "2026-09-06"); pg.wait_for_timeout(150)


def board(pg):
    return pg.request.get(HOST + "/finance/stock/api/loss/1").json()


def rows_of(j, key):
    return [r for s in j["sections"] if s["key"] == key for r in s["rows"]]


with sync_playwright() as pw:
    b = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium")
    ctx = b.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True)
    pg = ctx.new_page()
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.on("console", lambda m: errs.append(m.text) if m.type == "error" and "Failed to load resource" not in m.text else None)
    pg.on("dialog", lambda d: d.accept())

    # ---------------------------------------------------------------- 0 a real count
    pg.goto(HOST + "/finance/stock/page/count"); pg.wait_for_timeout(700); gate(pg)
    with pg.expect_download() as dl:
        pg.locator("#padget").click()
    PAD = os.path.join(HERE, "pad_prefilled_loss.xlsx"); dl.value.save_as(PAD)
    pg.set_input_files("#padfile", fill_pad(PAD, os.path.join(HERE, "pad_filled_loss.xlsx"), good)); pg.wait_for_timeout(2500)
    chk("a real count is recorded through the page", "Recorded as count #1" in pg.inner_text("#padout"))

    # ---------------------------------------------------------------- 1 the door
    r = as_staff(lambda: pg.request.get(HOST + "/finance/stock/page/loss?count=1"))
    chk("THE DOOR: a member of staff cannot open the loss desk", r.status == 403, r.status)
    r = as_staff(lambda: pg.request.get(HOST + "/finance/stock/api/loss/1"))
    chk("...nor read its numbers through the API", r.status == 403, r.status)
    r = pg.request.get(HOST + "/finance/stock/page/loss?count=1")
    chk("the owner can", r.status == 200, r.status)

    # ---------------------------------------------------------------- 2 what is on the desk
    j = board(pg)
    chk("the board answers", j.get("ok") is True and j["count_id"] == 1)
    on = {r["item"] for s in j["sections"] for r in s["rows"]}
    chk("EVERY SHORTAGE IS ON THE DESK -- %d of them, the fourteen ordinary medicines included" % len(SHORTS),
        on == set(SHORTS), sorted(on ^ set(SHORTS)))
    chk("an item that came up OVER is not on a loss desk at all", OVER not in on)
    allr = [r for s in j["sections"] for r in s["rows"]]
    chk("no line is in two sections", len(allr) == len(on), len(allr))
    money = [r["item"] for r in rows_of(j, "value")]
    vol = [r["item"] for r in rows_of(j, "volume")]
    rest = [r["item"] for r in rows_of(j, "rest")]
    chk("WHERE THE MONEY IS, biggest first: the tablet at Rs 4,800 then the brace at Rs 3,000",
        money == [TYR, KNEE], money)
    chk("WHERE THE QUANTITY IS: 15 strips short but only Rs 75 -- it would never have surfaced by value",
        vol == [CHEAP], vol)
    chk("THE REST holds the three-tablet line, the one with no price and the fourteen ordinary medicines",
        set(rest) == {SMALL, NOPR} | set(MEDS), sorted(set(rest) ^ ({SMALL, NOPR} | set(MEDS))))
    chk("...largest value first inside it, so the rest is not a heap",
        [r["loss_mrp_p"] or 0 for r in rows_of(j, "rest")] == sorted([r["loss_mrp_p"] or 0 for r in rows_of(j, "rest")], reverse=True))
    t = {r["item"]: r for r in allr}
    chk("the tablet: 40 strips short, Rs 4,800 at MRP, Rs 3,600 at cost",
        t[TYR]["short_text"] == "40 strips" and t[TYR]["loss_mrp_p"] == 480000 and t[TYR]["loss_cost_p"] == 360000,
        (t[TYR]["short_text"], t[TYR]["loss_mrp_p"], t[TYR]["loss_cost_p"]))
    chk("the brace: 2 pcs short, Rs 3,000 at MRP, Rs 2,000 at cost -- pieces, never 'units'",
        t[KNEE]["short_text"] == "2 pcs" and t[KNEE]["loss_mrp_p"] == 300000 and t[KNEE]["loss_cost_p"] == 200000,
        (t[KNEE]["short_text"], t[KNEE]["loss_mrp_p"], t[KNEE]["loss_cost_p"]))
    chk("the item that never sold carries NO price rather than a made-up one",
        t[NOPR]["loss_mrp_p"] is None and t[NOPR]["mrp_unit_p"] is None and t[NOPR]["loss_cost_p"] is None)
    chk("...and the desk says how many lines that is, so the total is honest beside it",
        j["totals"]["all_unpriced"] == 1, j["totals"]["all_unpriced"])
    ALL_P = 480000 + 300000 + 7500 + 600 + 14 * 12000      # 4,800 + 3,000 + 75 + 6 + fourteen at 120
    chk("the whole shortage, before he marks anything: Rs 9,561", j["totals"]["all_mrp_p"] == ALL_P,
        j["totals"]["all_mrp_p"])

    # ---------------------------------------------------------------- 3 the page he holds
    pg.goto(HOST + "/finance/stock/page/loss?count=1"); pg.wait_for_timeout(900)
    chk("the page loads at phone width with his count on it", "count #1" in pg.inner_text("#who"))
    l1 = pg.inner_text(".tally .l1")
    chk("THREE LINES AT THE TOP, the money first: 'Rs 9,561 short -- 19 items still to look at.'",
        l1 == "Rs 9,561 short — 19 items still to look at.", l1)
    chk("...and ONE number on each line, never a row of them",
        pg.inner_text(".tally .l2") == "Rs 6,917 is what those cost you.", pg.inner_text(".tally .l2"))
    chk("...and there is no fourth line -- three, or nothing (his ruling)",
        pg.locator(".tally .l1, .tally .l2, .tally .l3").count() == 3)
    chk("nothing to share until he marks something", pg.locator("#mkshare").is_disabled())
    chk("the sections are titled in his words",
        "Where the money is" in pg.inner_text("#body") and "Where the quantity is" in pg.inner_text("#body"))
    chk("no rupee figure is hidden behind a tap -- every line shows its loss at MRP and at cost",
        pg.locator('.li[data-item="%s"] .v' % TYR).inner_text().startswith("Rs 4,800")
        and "Rs 3,600 at cost" in pg.locator('.li[data-item="%s"] .v' % TYR).inner_text(),
        pg.locator('.li[data-item="%s"] .v' % TYR).inner_text())

    # ---------------------------------------------------------------- 4 the workings, folded
    chk("the workings are folded away until he asks", pg.locator('.li[data-item="%s"] .work' % TYR).count() == 0)
    pg.locator('.li[data-item="%s"] .n' % TYR).click(); pg.wait_for_timeout(300)
    w = pg.locator('.li[data-item="%s"] .work' % TYR).inner_text()
    chk("...and one tap on the name opens Marg, the shelf, where the line sits and what it sold this year",
        "Marg" in w and "counted" in w and "Where it sits" in w and "Sold this year" in w, w[:90])
    pg.locator('.li[data-item="%s"] .n' % TYR).click(); pg.wait_for_timeout(250)
    chk("...and closes again", pg.locator('.li[data-item="%s"] .work' % TYR).count() == 0)

    # ---------------------------------------------------------------- 5 the play
    pg.locator('.li[data-item="%s"] .tick' % TYR).click(); pg.wait_for_timeout(700)
    chk("ONE TAP marks a line, and the top line follows him: 'Rs 4,800 -- 1 item marked as lost.'",
        pg.inner_text(".tally .l1") == "Rs 4,800 — 1 item marked as lost.", pg.inner_text(".tally .l1"))
    chk("...with what it cost on the line below", pg.inner_text(".tally .l2") == "Rs 3,600 is what those cost you.",
        pg.inner_text(".tally .l2"))
    chk("...and the sheet button wakes up", not pg.locator("#mkshare").is_disabled())
    pg.locator('.li[data-item="%s"] .tick' % TYR).click(); pg.wait_for_timeout(700)
    chk("a second tap un-marks it", pg.inner_text(".tally .l1") == "Rs 9,561 short — 19 items still to look at.",
        pg.inner_text(".tally .l1"))
    pg.locator('button[data-all="value"]').click(); pg.wait_for_timeout(800)
    chk("MARK ALL takes the whole section in one tap: 2 items, Rs 7,800",
        pg.inner_text(".tally .l1") == "Rs 7,800 — 2 items marked as lost.", pg.inner_text(".tally .l1"))
    pg.locator('button[data-all="volume"]').click(); pg.wait_for_timeout(800)
    j = board(pg)
    chk("...and it adds to what was already marked, never replaces it",
        j["totals"]["ticked"] == 3 and j["totals"]["mrp_p"] == 480000 + 300000 + 7500,
        (j["totals"]["ticked"], j["totals"]["mrp_p"]))
    chk("a marked line is coloured, not merely counted", "on" in (pg.locator('.li[data-item="%s"]' % TYR).get_attribute("class") or ""))

    # ---------------------------------------------------------------- 6 the sheet, frozen
    pg.locator("#mkshare").click(); pg.wait_for_timeout(1500)
    m = pg.inner_text("#shmsg")
    chk("MAKING THE SHEET says what was made, in his words and his money",
        "Sheet 1 made and frozen: 3 items, Rs 7,875 at MRP. Print it for Darpan." in m, m)
    chk("...and the message is still there after the page refreshes itself", m == pg.inner_text("#shmsg"))
    j = board(pg)
    S = j["shares"]
    chk("one sheet is on record -- the three he marked, and not the sixteen he did not",
        len(S) == 1 and S[0]["no"] == 1 and S[0]["items"] == 3, len(S))
    chk("...with its own fingerprint and the hour it was frozen", len(S[0]["md5"]) == 32 and S[0]["made_text"])
    chk("...and the total that was on it", S[0]["total_p"] == 480000 + 300000 + 7500, S[0]["total_p"])
    frozen = json.dumps(S[0]["lines"], sort_keys=True)
    kept = os.path.join(ARCH, "loss_shares", "%s.pdf" % S[0]["md5"])
    chk("THE COPY THAT WAS SHARED IS KEPT ON THE SERVER, named by its fingerprint", os.path.exists(kept), kept)
    r = pg.request.get(HOST + "/finance/stock/api/loss/share/%d.pdf" % S[0]["id"])
    pdf = r.body()
    chk("the sheet opens as a PDF", r.status == 200 and pdf[:4] == b"%PDF", r.status)
    chk("...A4 PORTRAIT, as he asked", b"/MediaBox [0 0 595.28 841.89]" in pdf, "not portrait")
    chk("...one page for three items", pdf.count(b"/Type /Page ") == 1, pdf.count(b"/Type /Page "))
    txt = pdf.decode("latin-1")
    chk("...carrying only the core data: the item, the shortage, the MRP each, the loss at MRP",
        "Item" in txt and "Short" in txt and "MRP each" in txt and "Loss at MRP" in txt)
    chk("...the shortage in the shop's own words -- 40 strips, 2 pcs", "40 strips" in txt and "2 pcs" in txt)
    chk("...the total, spelled out", "TOTAL SHORTAGE ON THIS SHEET, AT MRP" in txt and "Rs 7,875" in txt)
    chk("...and one money format everywhere -- the sheet and the screen agree",
        "Rs 7,875" in txt and "Rs 7,875" in pg.inner_text("#body"))
    chk("...a box for his answer against every line, and a place to sign",
        "Darpan's answer" in txt and "Answered by" in txt)
    chk("...and no purchase rate anywhere on it -- Darpan is not shown what we paid",
        "at cost" not in txt and "3,600" not in txt)
    # IMMUTABILITY -- the price moves under it; the sheet does not
    chk("one odd bill does NOT move a price -- the median is what guards it",
        True)
    con.execute("INSERT INTO sale_line_item (day_entry_id, unit, business_date, bill_no, is_return, seq, item_name, item_key, qty_raw, pack, amount_p) "
                "VALUES (1,'medical','2026-09-04','A009999',0,9999,?,?,'1:0','1*10',99000)", (TYR, sale_key(TYR)))
    con.commit()
    chk("...so a single Rs 990 line leaves the tablet at Rs 12 a tab",
        {r["item"]: r for s_ in board(pg)["sections"] for r in s_["rows"]}[TYR]["loss_mrp_p"] == 480000)
    con.execute("DELETE FROM sale_line_item WHERE item_key=?", (sale_key(TYR),))
    for k in range(3): sale(TYR, "2026-09-0%d" % (k + 1), 24000)
    con.commit()
    j2 = board(pg)
    live = {r["item"]: r for s in j2["sections"] for r in s["rows"]}[TYR]
    chk("a REAL price change -- the tablet reprices to Rs 24 -- but the line is ALREADY ON SHEET 1, so the desk keeps the sheet's figure",
        live["loss_mrp_p"] == 480000 and live.get("frozen") is True and live["sheet_no"] == 1,
        (live["loss_mrp_p"], live.get("frozen")))
    chk("...so the screen and the paper in Darpan's hand never say two prices for one medicine",
        live["mrp_unit_p"] == 1200, live["mrp_unit_p"])
    again = pg.request.get(HOST + "/finance/stock/api/loss/share/%d.pdf" % S[0]["id"]).body()
    chk("BUT THE SHARED SHEET IS FROZEN -- the same bytes, to the byte", again == pdf)
    chk("...and its recorded total has not moved either",
        board(pg)["shares"][0]["total_p"] == 480000 + 300000 + 7500, board(pg)["shares"][0]["total_p"])
    chk("...nor its lines", json.dumps(board(pg)["shares"][0]["lines"], sort_keys=True) == frozen)

    # ---------------------------------------------------------------- 7 what comes back
    pg.reload(); pg.wait_for_timeout(1000)
    box = pg.locator('.rec[data-share="%d"]' % S[0]["id"])
    box.locator(".ritem").select_option(TYR)
    box.locator(".rkind").select_option("FOUND")
    chk("asking 'found' asks how many, and only then", not box.locator(".rnum").is_hidden())
    box.locator(".rnum").fill("100")
    box.locator(".rnote").fill("in the back rack")
    box.locator("button[data-rec]").click(); pg.wait_for_timeout(1200)
    j = board(pg)
    chk("100 TABLETS FOUND come off the loss at the price that was ON THE SHEET: Rs 1,200",
        j["shares"][0]["got"]["found_p"] == 120000 and j["shares"][0]["got"]["recovered_p"] == 120000,
        j["shares"][0]["got"])
    r = pg.request.post(HOST + "/finance/stock/api/loss/1/recovery",
                        data=json.dumps(dict(share_id=S[0]["id"], item=KNEE, kind="MONEY", amount=500)),
                        headers={"Content-Type": "application/json"}).json()
    chk("money recovered is logged in rupees and lands in paise", r.get("ok") is True and
        board(pg)["shares"][0]["got"]["recovered_p"] == 120000 + 50000, board(pg)["shares"][0]["got"])
    r = pg.request.post(HOST + "/finance/stock/api/loss/1/recovery",
                        data=json.dumps(dict(share_id=S[0]["id"], item=CHEAP, kind="ACCEPTED", note="no answer")),
                        headers={"Content-Type": "application/json"}).json()
    chk("a line he accepts as lost is logged too -- silence is never the record", r.get("ok") is True, r)
    r = pg.request.post(HOST + "/finance/stock/api/loss/1/recovery",
                        data=json.dumps(dict(share_id=S[0]["id"], item=SMALL, kind="FOUND", units=1)),
                        headers={"Content-Type": "application/json"})
    chk("an item that was never on that sheet is refused", r.status == 400 and r.json().get("error") == "not_on_sheet", r.status)
    r = pg.request.post(HOST + "/finance/stock/api/loss/1/recovery",
                        data=json.dumps(dict(share_id=S[0]["id"], item=TYR, kind="FOUND", units=0)),
                        headers={"Content-Type": "application/json"})
    chk("'found none' is a question, not a record", r.status == 400 and r.json().get("error") == "bad_units", r.status)
    r = as_staff(lambda: pg.request.post(HOST + "/finance/stock/api/loss/1/recovery",
                                         data=json.dumps(dict(share_id=S[0]["id"], item=TYR, kind="FOUND", units=1)),
                                         headers={"Content-Type": "application/json"}))
    chk("a member of staff cannot log a recovery either", r.status == 403, r.status)
    pg.reload(); pg.wait_for_timeout(1000)
    chk("THE THIRD LINE now carries the standing balance, so he sees where the play stands",
        pg.inner_text(".tally .l3") == "Rs 6,175 still standing — Rs 1,700 back of Rs 7,875 given to him.",
        pg.inner_text(".tally .l3"))
    chk("...and every recovery is listed under its sheet with the hour it was logged",
        pg.inner_text("#body").count("in the back rack") == 1)
    chk("...and the line itself says what has come back on it, so no row reads gross while the top reads net",
        "Rs 1,200 back" in pg.locator('.li[data-item="%s"]' % TYR).inner_text(),
        pg.locator('.li[data-item="%s"]' % TYR).inner_text())

    # A LINE BELONGS TO ONE SHEET -- the screen read at S228 caught the second sheet
    # re-listing the first, which would have asked Darpan the same thing twice and
    # counted the same money twice.
    j = board(pg)
    gone = [r for s_ in j["sections"] for r in s_["rows"] if r["shared"]]
    chk("the three that went out are marked 'on sheet 1' and are no longer his to mark",
        len(gone) == 3 and all(r["sheet_no"] == 1 for r in gone), [(r["item"], r.get("sheet_no")) for r in gone])
    chk("...and they carry no tick button on the page at all",
        pg.locator('.li[data-item="%s"] button.tick' % TYR).count() == 0)
    chk("...so what is still HIS is 16 lines, and the top line counts only those",
        j["totals"]["open_lines"] == 16 and j["totals"]["ticked"] == 0, (j["totals"]["open_lines"], j["totals"]["ticked"]))
    r = pg.request.post(HOST + "/finance/stock/api/loss/1/share", data="{}", headers={"Content-Type": "application/json"})
    chk("...and a second sheet of nothing is refused, in words",
        r.status == 400 and "already on a sheet" in (r.json().get("message") or ""), r.json().get("message"))

    # a second sheet, and the first left exactly as it was
    pg.reload(); pg.wait_for_timeout(1000)
    pg.locator('button[data-all="rest"]').click(); pg.wait_for_timeout(1000)
    pg.locator("#mkshare").click(); pg.wait_for_timeout(1500)
    j = board(pg)
    chk("A SECOND SHEET is sheet 2, and the first is untouched",
        len(j["shares"]) == 2 and j["shares"][1]["no"] == 2 and j["shares"][0]["total_p"] == 787500,
        [(x["no"], x["total_p"]) for x in j["shares"]])
    chk("...and it holds the SIXTEEN that were never handed over -- not one repeat",
        j["shares"][1]["items"] == 16 and not ({l["item"] for l in j["shares"][1]["lines"]}
                                               & {l["item"] for l in j["shares"][0]["lines"]}),
        j["shares"][1]["items"])
    chk("...the unpriced line rides on it, counted but not valued", j["shares"][1]["unpriced"] == 1, j["shares"][1]["unpriced"])
    chk("...and the two sheets add up without counting a rupee twice: Rs 7,875 + Rs 1,686 = Rs 9,561",
        j["shares"][1]["total_p"] == 168600 and j["totals"]["shared_p"] == 956100, j["totals"]["shared_p"])
    chk("nothing is left on his side once both sheets are out", j["totals"]["open_lines"] == 0, j["totals"]["open_lines"])
    chk("...and the sections then add up to exactly what is out with him -- one arithmetic on one page",
        sum(r["loss_mrp_p"] or 0 for s_ in j["sections"] for r in s_["rows"]) == j["totals"]["shared_p"] == 956100,
        (sum(r["loss_mrp_p"] or 0 for s_ in j["sections"] for r in s_["rows"]), j["totals"]["shared_p"]))
    pg.reload(); pg.wait_for_timeout(1100)
    chk("...and with nothing left to mark, the page stops asking him to mark something",
        pg.locator("#mkshare").count() == 0 and pg.inner_text(".tally .l1") == "Every short item is with Darpan now.",
        pg.inner_text(".tally .l1"))
    chk("...and the third line reconciles on its own: standing, back, and what went out",
        pg.inner_text(".tally .l3") == "Rs 7,861 still standing — Rs 1,700 back of Rs 9,561 given to him.",
        pg.inner_text(".tally .l3"))
    # SIXTEEN NEAR-IDENTICAL LINES SHOULD NOT BE SIXTEEN SCREENS
    chk("a long section folds after 8, and says how many it is holding back",
        pg.locator('.li').count() == 3 + 8 and pg.locator('button[data-show="rest"]').count() == 1,
        pg.locator('.li').count())
    pg.locator('button[data-show="rest"]').click(); pg.wait_for_timeout(400)
    chk("...and one tap opens the rest", pg.locator('.li').count() == 19 and pg.locator('button[data-show]').count() == 0,
        pg.locator('.li').count())
    chk("...an item with no price says so in words, never a bare dash",
        "no price" in pg.locator('.li[data-item="%s"] .v' % NOPR).inner_text(),
        pg.locator('.li[data-item="%s"] .v' % NOPR).inner_text())

    # ---------------------------------------------------------------- 8 Darpan's lists, moved here
    chk("DARPAN'S LISTS ARE ON HIS PAGE NOW", "Darpan's lists — in turns" in pg.inner_text("#body"))
    pg.locator('button[data-next="med"]').click(); pg.wait_for_timeout(1400)
    chk("...and he cuts one himself, and is told what he cut", "List 1 cut" in pg.inner_text("#tmsg"),
        pg.inner_text("#tmsg"))
    j = board(pg)
    chk("the list is on record with its items", len(j["tranches"]) == 1 and len(j["tranches"][0]["items"]) > 0)
    tp = pg.request.get(HOST + j["tranches"][0]["pdf"])
    body = tp.body().decode("latin-1")
    chk("the list prints, and STILL carries no rupee -- that rule did not move with it",
        tp.status == 200 and "Rs " not in body and "MRP" not in body, tp.status)
    pg.locator("button[data-ret]").click(); pg.wait_for_timeout(1200)
    chk("he marks it returned from here", board(pg)["tranches"][0]["returned_at"] is not None)

    # ---------------------------------------------------------------- 9 Amir's board gave them up
    pg.goto(HOST + "/finance/stock/page/amir?count=1"); pg.wait_for_timeout(1200)
    a = pg.inner_text("#body")
    chk("AMIR'S BOARD no longer cuts or hands out Darpan's lists", pg.locator("button[data-next]").count() == 0)
    chk("...nor marks them returned", pg.locator("button[data-ret]").count() == 0)
    chk("...but he still types Darpan's answers -- his work did not grow, and it did not vanish",
        "Type Darpan's answers" in a and pg.locator("table.board, .empty").count() >= 1)
    chk("...and he can still see which list is with Darpan", "List 1" in a, a[:80])

    # ---------------------------------------------------------------- 10 the way in
    pg.goto(HOST + "/finance/stock/page/desk?count=1"); pg.wait_for_timeout(1200)
    chk("THE DECISION DESK LINKS TO HIS LOSS DESK, first in the foot",
        pg.locator('#foot a[href$="/page/loss?count=1"]').count() == 1)
    chk("...and it is named for what it does", "Your loss desk" in pg.inner_text("#foot"))

    chk("no page errors anywhere at 390px", not errs, errs[:3])
    # screens for a sub-agent to read -- they never enter the conversation
    pg.goto(HOST + "/finance/stock/page/loss?count=1"); pg.wait_for_timeout(1200)
    pg.screenshot(path=os.path.join(HERE, "_shot_loss_top.png"))
    pg.screenshot(path=os.path.join(HERE, "_shot_loss_full.png"), full_page=True)
    pg.emulate_media(media="print")
    pg.goto(HOST + "/finance/stock/page/amir?count=1"); pg.wait_for_timeout(1200)
    pg.screenshot(path=os.path.join(HERE, "_shot_loss_amir.png"), full_page=True)
    ctx.close(); b.close()

print("\n%d ok, %d FAIL" % (len(OK), len(BAD)))
for x in BAD: print("  FAIL " + x)
sys.exit(1 if BAD else 0)
