"""LIVE-SHAPE WALK -- S227 PAD PROOF.

The owner, 06-Sep-2026: save the uploaded Excel sheets, give the stock checkers
a PDF printout as proof of what the VPS ingested, and preserve the sheets for any
dispute during the reconciliation of a stock mismatch.

This drives the stock-check page at phone width the way Amir does -- the pad
comes down from the button, is filled without recalculation, uploaded by picking
the file -- and then LOOKS: at the box, at the receipt it offers, at the folder
on disk, at the ledger. Then the loop, the repeat, an unreadable file, the
checker's archive, the counter refused, the close. The S226 loop is walked again
underneath, because the upload route was rewritten around it.

Run from inside the kit folder. Needs playwright + openpyxl (readers that are
not ours) and pdftotext (poppler) to read the receipt back as a person would.
"""
import hashlib, os, shutil, subprocess, sys, threading, time, zipfile, io
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import padreader
import WALK_readiness_s226 as R
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

OK, BAD = [], []
def chk(n, c, d=""):
    (OK if c else BAD).append(n + (("  -- " + str(d)) if d and not c else ""))
    print(("  ok   " if c else "  FAIL ") + n + (("   " + str(d)) if d else ""))

# ---- a shop of 120 items, the shape of the real list (names, packings, pack sizes)
import random; random.seed(227)
FORMS = [("TAB", "1*10", 10), ("TAB", "1*15", 15), ("CAP", "1*10", 10), ("SYP 100ML", "1*1", 1),
         ("INJ", "1*1", 1), ("OINT 30G", "1*1", 1), ("TAB", "1*30", 30)]
ITEMS = []
for i in range(1, 121):
    f = FORMS[i % len(FORMS)]
    ITEMS.append(dict(item="WALK-%03d %s %s" % (i, random.choice(["ZINC", "CALPOL", "PANTO", "DOLO", "MOX", "ORTHO"]), f[0]),
                      packing=f[1], pack_size=f[2]))
names = [r["item"] for r in ITEMS]
PS = {r["item"]: r["pack_size"] for r in ITEMS}

R.DB = os.path.join(HERE, "_walkproof.db"); con = R.build(); R.fill(con)
con.execute("DELETE FROM stock_snapshot")
MARG = {}
for r in ITEMS:
    q = random.choice([0, 4, 13, 40, 137, 405]); MARG[r["item"]] = q
    con.execute("INSERT OR REPLACE INTO stock_snapshot (as_on,item,qty,packing,pack_size,source,loaded_at) "
                "VALUES (?,?,?,?,?,?,?)", ("06-09-2026", r["item"], q, r["packing"], PS[r["item"]], "p", "2026-09-06T09:00:59"))
for r in ITEMS[:50]:
    con.execute("INSERT OR REPLACE INTO stock_rate (item, rate_p, as_of, source) VALUES (?,?,?,?)",
                (r["item"], 2500, "2026-09-06", "walk"))
con.commit()
ARCH = os.path.join(HERE, "pad_uploads")
if os.path.isdir(ARCH):
    shutil.rmtree(ARCH)
app = R.app_for(con)
threading.Thread(target=lambda: app.run(port=8827, threaded=False), daemon=True).start()
time.sleep(1.5)
HOST = "http://127.0.0.1:8827"
URL = HOST + "/finance/stock/page/count"
errs = []


def fill_pad(src, dst, entries, rename=None, bad=None):
    shutil.copy(src, dst)
    wb = load_workbook(dst); ws = wb.worksheets[0]
    where = {}
    for r in range(9, 400):
        nm = ws.cell(row=r, column=2).value
        if isinstance(nm, str) and nm.strip():
            where[nm.strip()] = r
    for it, (st, lo) in entries.items():
        r = where[it]
        if st is not None: ws.cell(row=r, column=5, value=st)
        if lo is not None: ws.cell(row=r, column=6, value=lo)
    for it, nm in (rename or {}).items():
        ws.cell(row=where[it], column=2, value=nm); ws.cell(row=where[it], column=6, value=7)
    for it, v in (bad or {}).items():
        ws.cell(row=where[it], column=6, value=v)
    wb.save(dst)                                    # NOT recalculated, on purpose
    return dst


def pdf_text(b):
    p = subprocess.run(["pdftotext", "-layout", "-", "-"], input=b, stdout=subprocess.PIPE)
    return p.stdout.decode("utf-8", "replace")


def md5f(p):
    return hashlib.md5(open(p, "rb").read()).hexdigest()


good = {}
for nm in names[:40]:
    ps = PS[nm]; q = MARG[nm]
    tgt = q if len(good) % 6 else max(q - 3, 0)
    good[nm] = (tgt // ps if ps > 1 else None, tgt % ps if ps > 1 else tgt)
EXPECT = {nm: (v[0] or 0) * PS[nm] + (v[1] or 0) for nm, v in good.items()}


def gate(pg, cby="Darpan", eby="Amir", bill="A003425", bdate="2026-09-06"):
    pg.locator('#whoC button[data-u="%s"]' % cby).click()
    pg.locator('#whoE button[data-u="%s"]' % eby).click()
    pg.fill("#bill", bill); pg.fill("#billdate", bdate); pg.wait_for_timeout(150)


with sync_playwright() as pw:
    b = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium")
    ctx = b.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True)
    pg = ctx.new_page()
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.on("console", lambda m: errs.append(m.text)
          if m.type == "error" and "Failed to load resource" not in m.text else None)
    dialogs = []
    pg.on("dialog", lambda d: (dialogs.append(d.message), d.accept()))

    # ------------------------------------------------------------ 1 the page, the pad
    pg.goto(URL); pg.wait_for_timeout(700)
    chk("the page opens: four steps, the Excel panel, no recent counts", pg.locator(".step").count() == 4
        and pg.locator("#xl").is_visible() and pg.inner_text("#recent").strip() == "")
    gate(pg)
    with pg.expect_download() as dl:
        pg.locator("#padget").click()
    PAD = os.path.join(HERE, "pad_prefilled.xlsx"); dl.value.save_as(PAD)
    rp = padreader.read_pad(PAD)
    chk("SHEET 1 comes down from the button with every item and the four details", len(rp["rows"]) == len(ITEMS)
        and rp["meta"].get("counted_by") == "Darpan" and rp["meta"].get("bill_no") == "A003425", rp["meta"])

    # ------------------------------------------------------------ 2 the filled pad, ONE tap -- and the PROOF
    FILLED = fill_pad(PAD, os.path.join(HERE, "pad_filled.xlsx"), good,
                      rename={names[60]: "SOMETHING NOBODY KNOWS"}, bad={names[61]: -3, names[62]: 2.5})
    FILLED_MD5 = md5f(FILLED)
    pg.set_input_files("#padfile", FILLED); pg.wait_for_timeout(3000)
    out = pg.inner_text("#padout")
    chk("IT RECORDS IN ONE STEP AND SAYS SO", "Recorded as count #1" in out, out[:120])
    chk("...this sheet: 40 counted, 3 rows sent back to fix", "40 counted" in out and "3 rows sent back" in out, out[:300])
    chk("...some stock count is left: 80 not counted, 3 to fix", "Some stock count is left: 80 items not counted, 3 rows to fix" in out, out[:400])
    chk("THE BOX OFFERS THE PROOF OF THIS SHEET, a PDF, right under the numbers",
        pg.locator("#padout a.dl.proof").count() == 1 and "Download the PROOF of this sheet (PDF)" in out, out[:600])
    chk("...and says the Excel file itself is kept on the server, with its fingerprint",
        "kept on the server, unchanged" in out and FILLED_MD5[:8] in out, out[:700])
    chk("...the proof sits ABOVE the remaining-work sheet link, the result sheet still offered",
        pg.locator("#padout a.dl").count() == 2 and pg.locator("#padout a.dl").nth(0).get_attribute("class").endswith("proof")
        and pg.locator("#padout a.dl").nth(1).inner_text() == "Download SHEET 2 REMAINING")
    pg.locator("#padout").screenshot(path=os.path.join(HERE, "_shot_box_after_upload.png"))
    href = pg.locator("#padout a.dl.proof").get_attribute("href")
    chk("...the proof link is the receipt for count #1", href.startswith("/finance/stock/api/pad/receipt/1.pdf"), href)
    r = pg.request.get(HOST + href.split("?")[0])
    RECEIPT1 = r.body()
    chk("THE RECEIPT IS A PDF, opened inline under the sheet's name",
        r.status == 200 and RECEIPT1[:5] == b"%PDF-" and "application/pdf" in r.headers.get("content-type", "")
        and 'inline; filename="STOCK_COUNT_' in r.headers.get("content-disposition", "")
        and "_SHEET_1_PROOF.pdf" in r.headers.get("content-disposition", ""), r.headers.get("content-disposition"))
    txt = pdf_text(RECEIPT1)
    chk("...it names the count and the sheet", "Count #1 - sheet 1" in txt, txt[:200])
    chk("...who: uploaded by walk, counted by Darpan, entered by Amir, bill A003425 dated 06-09-2026",
        "Uploaded by (login)" in txt and "walk" in txt and "Darpan" in txt and "Amir" in txt and "A003425 dated 06-09-2026" in txt)
    chk("...THE FILE: its name, its size, its md5, KEPT on the server, verified",
        "pad_filled.xlsx" in txt and ("%d bytes" % os.path.getsize(FILLED)) in txt and FILLED_MD5 in txt
        and "Kept on the server" in txt and "YES - unchanged" in txt and "the md5 matched" in txt, txt[txt.find("THE FILE"):][:400])
    chk("...what this sheet did: 43 rows read, 40 counted, 3 sent back",
        "43 rows with figures - 40 items counted" in txt and "3 rows sent back to fix" in txt, txt[txt.find("WHAT THIS SHEET DID"):][:200])
    chk("...the whole count after it: 40 of 120, 80 not yet counted, 3 to fix",
        "40 of 120 items counted" in txt and "80 not yet counted, 3 to fix" in txt)
    ndiff = sum(1 for nm in good if EXPECT[nm] != MARG[nm])
    chk("...DIFFERENCES FROM THIS SHEET, first, with the count", ("DIFFERENCES FROM THIS SHEET (%d)" % ndiff) in txt, ndiff)
    chk("...the three rows sent back, as written, with why",
        "ROWS SENT BACK TO FIX (3)" in txt and "SOMETHING NOBODY KNOWS" in txt and "not in the shop list" in txt, txt[txt.find("ROWS SENT"):][:400])
    chk("...EVERY ITEM RECORDED FROM THIS SHEET (40), each with Marg, counted, strips, loose, diff",
        "EVERY ITEM RECORDED FROM THIS SHEET (40)" in txt and "40 items on this sheet - %d differ" % ndiff in txt)
    # every figure on the receipt is the figure the sheet carried
    sec = txt[txt.find("EVERY ITEM RECORDED"):]
    allrows = True
    for nm in good:
        line = [l for l in sec.splitlines() if nm in l]
        if not line:
            allrows = False; break
        cols = line[0].split()
        # ... Marg Counted Strips Loose Diff  -> the last five tokens
        marg, counted, diff = cols[-5], cols[-4], cols[-1]
        if int(marg) != MARG[nm] or int(counted) != EXPECT[nm] or int(diff) != EXPECT[nm] - MARG[nm]:
            allrows = False; break
    chk("...AND EVERY ONE OF THE 40 FIGURES ON THE RECEIPT IS THE FIGURE THE SHEET CARRIED", allrows, nm)
    chk("...the foot names the count, the full md5, and the time in IST, unclipped",
        "Count #1 sheet 1 - md5 %s - PDF made " % FILLED_MD5 in txt and " IST" in txt.split("PDF made")[1][:30], txt[txt.find("Count #1 sheet 1 - md5"):][:120])
    pages = RECEIPT1.count(b"/Type /Page ")
    chk("...on %d page(s) of A4" % pages, pages >= 1 and b"/MediaBox [0 0 595.28 841.89]" in RECEIPT1)

    # ------------------------------------------------------------ 3 the folder on disk, the ledger
    kept = os.path.join(ARCH, FILLED_MD5 + ".xlsx")
    chk("THE UPLOADED BYTES ARE ON DISK, beside the database, named by their md5", os.path.exists(kept), ARCH)
    chk("...BYTE FOR BYTE the file that was uploaded", open(kept, "rb").read() == open(FILLED, "rb").read())
    chk("...and a second reader opens the kept copy", load_workbook(kept).worksheets[0]["B5"].value == "Darpan")
    row = con.execute("SELECT path, filename, bytes, times, outcome, count_id, verified, seen_by FROM stock_count_pad_archive WHERE md5=?", (FILLED_MD5,)).fetchone()
    chk("the archive ledger holds it: recorded, count #1, verified, once, by walk",
        row == (FILLED_MD5 + ".xlsx", "pad_filled.xlsx", os.path.getsize(FILLED), 1, "recorded", 1, 1, "walk"), row)
    idx = open(os.path.join(ARCH, "INDEX.txt"), encoding="utf-8").read()
    chk("INDEX.txt, a line a person can read: IST, md5, recorded, count 1, by walk, the name, verified",
        " IST | %s | recorded" % FILLED_MD5 in idx and "count 1" in idx and "by walk" in idx and "pad_filled.xlsx" in idx and "verified" in idx, idx[:200])
    stored = [f for f in os.listdir(ARCH) if f.endswith("_PROOF.pdf")]
    chk("A FROZEN COPY OF THE RECEIPT WAS WRITTEN BESIDE THE SHEET the moment it was recorded",
        len(stored) == 1 and stored[0].startswith(FILLED_MD5[:8] + "_count1_") and stored[0].endswith("_SHEET_1_PROOF.pdf"), stored)
    stxt = pdf_text(open(os.path.join(ARCH, stored[0]), "rb").read())
    chk("...and it says what the route says (the same sealed rows)", stxt.split("PDF made")[0] == txt.split("PDF made")[0])
    rj = pg.request.get(HOST + "/finance/stock/api/pad/recent").json()
    chk("THE RECENT LIST CARRIES THE SHEETS WITH THEIR RECEIPTS", rj["counts"][0]["sheets"] == [dict(count_id=1, n=1, when=rj["counts"][0]["sheets"][0]["when"],
        receipt="/finance/stock/api/pad/receipt/1.pdf", md5=FILLED_MD5, filename="pad_filled.xlsx", kept=True, from_sheet=True)], rj["counts"][0]["sheets"])
    chk("...and the page shows 'Proof of each sheet (PDF): sheet 1' under count #1",
        "Proof of each sheet (PDF): sheet 1" in pg.inner_text("#recent"), pg.inner_text("#recent")[:300])

    # ------------------------------------------------------------ 4 the same file again
    pg.set_input_files("#padfile", FILLED); pg.wait_for_timeout(2000)
    out = pg.inner_text("#padout")
    chk("THE SAME FILE TWICE IS NOT A SECOND COUNT -- and the proof is still offered",
        "Already recorded" in out and "count #1" in out and pg.locator("#padout a.dl.proof").count() == 1, out[:200])
    row = con.execute("SELECT times, outcome, count_id FROM stock_count_pad_archive WHERE md5=?", (FILLED_MD5,)).fetchone()
    chk("...the archive counts it as seen twice, STILL recorded (a recorded sheet never regresses), count #1, ONE file on disk",
        row == (2, "recorded", 1) and len([f for f in os.listdir(ARCH) if f.endswith(".xlsx")]) == 1, row)
    idx = open(os.path.join(ARCH, "INDEX.txt"), encoding="utf-8").read()
    chk("...INDEX.txt logs the repeat as an event", "| repeat" in idx and idx.count(FILLED_MD5) == 2)
    chk("...nothing extra in the ledger", con.execute("SELECT COUNT(*) FROM stock_count").fetchone()[0] == 1)

    # ------------------------------------------------------------ 5 the loop: the result sheet, filled, joins -- with ITS proof
    fu = pg.request.get(HOST + "/finance/stock/api/pad/followup/1.xlsx")
    RES = os.path.join(HERE, "result_1.xlsx"); open(RES, "wb").write(fu.body())
    wbf = load_workbook(RES); nc, sb = wbf["NOT COUNTED"], wbf["SENT BACK TO FIX"]
    nc.cell(row=5, column=4, value=2); nc.cell(row=5, column=5, value=3); first_nc = nc.cell(row=5, column=2).value
    sb.cell(row=5, column=2, value=names[60]); sb.cell(row=5, column=5, value=7)
    sb.cell(row=6, column=5, value=3)
    wbf.save(RES); RES_MD5 = md5f(RES)
    pg.set_input_files("#padfile", RES); pg.wait_for_timeout(3000)
    out = pg.inner_text("#padout")
    chk("THE FILLED RESULT SHEET JOINS THE COUNT", "Recorded as part of count #1" in out and "3 counted" in out, out[:200])
    chk("...with its own proof link, for count #2", pg.locator("#padout a.dl.proof").get_attribute("href").startswith("/finance/stock/api/pad/receipt/2.pdf"))
    r = pg.request.get(HOST + "/finance/stock/api/pad/receipt/2.pdf"); txt2 = pdf_text(r.body())
    chk("THE PART'S RECEIPT: count #1 sheet 2, recorded as count #2 part of #1, named SHEET_2_PROOF",
        "Count #1 - sheet 2" in txt2 and "recorded as count #2, part of #1" in txt2 and "_SHEET_2_PROOF.pdf" in r.headers.get("content-disposition", ""), txt2[:200])
    chk("...3 items counted, the corrected name among them, the settled total from the shop's pack size",
        "3 items counted" in txt2 and names[60] in txt2 and (" %d " % (2 * PS[first_nc] + 3)) in [l for l in txt2.splitlines() if first_nc in l][0], [l for l in txt2.splitlines() if first_nc in l])
    chk("...the whole count after it: 43 of 120", "43 of 120 items counted" in txt2)
    chk("...the sheet's own md5 on it, kept", RES_MD5 in txt2 and "YES - unchanged" in txt2)
    chk("both sheets on disk, both in the ledger", os.path.exists(os.path.join(ARCH, RES_MD5 + ".xlsx"))
        and con.execute("SELECT COUNT(*) FROM stock_count_pad_archive WHERE outcome='recorded'").fetchone()[0] == 2)
    chk("the recent list: 'Proof of each sheet (PDF): sheet 1 · sheet 2'",
        "Proof of each sheet (PDF): sheet 1 · sheet 2" in pg.inner_text("#recent"), pg.inner_text("#recent")[:400])
    chk("...and the page has thrown no error so far", not errs, errs)

    # ------------------------------------------------------------ 6 a file that is not a pad is refused -- AND KEPT
    JUNK = os.path.join(HERE, "junk.xlsx"); open(JUNK, "wb").write(b"this is not a spreadsheet, but it is what was sent")
    JUNK_MD5 = md5f(JUNK)
    pg.set_input_files("#padfile", JUNK); pg.wait_for_timeout(1500)
    out = pg.inner_text("#padout")
    chk("A FILE THAT IS NOT A PAD IS REFUSED IN WORDS, nothing recorded", "Not recorded" in out and "Nothing was written" in out, out[:200])
    row = con.execute("SELECT outcome, count_id, verified, filename FROM stock_count_pad_archive WHERE md5=?", (JUNK_MD5,)).fetchone()
    chk("...BUT THE BYTES ARE KEPT ALL THE SAME -- refused-unreadable, no count, verified",
        row == ("refused-unreadable", None, 1, "junk.xlsx") and open(os.path.join(ARCH, JUNK_MD5 + ".xlsx"), "rb").read() == open(JUNK, "rb").read(), row)
    chk("...nothing in the count ledger for it", con.execute("SELECT COUNT(*) FROM stock_count").fetchone()[0] == 2)
    # a pad with blank top lines and no details: HELD, and kept
    BLANK = os.path.join(HERE, "pad_blank_top.xlsx")
    import padwriter
    open(BLANK, "wb").write(padwriter.fresh_pad([(r["item"], r["packing"], r["pack_size"]) for r in ITEMS], "06-09-2026", "06-09-2026"))
    HELD = fill_pad(BLANK, os.path.join(HERE, "pad_held.xlsx"), dict(list(good.items())[40:] or [(names[70], (None, 5))]))
    pg.evaluate("()=>{ S.cby=''; S.eby=''; S.bill=''; S.billdate=''; save(); }"); pg.reload(); pg.wait_for_timeout(500)
    pg.set_input_files("#padfile", HELD); pg.wait_for_timeout(2000)
    out = pg.inner_text("#padout")
    row = con.execute("SELECT outcome, count_id FROM stock_count_pad_archive WHERE md5=?", (md5f(HELD),)).fetchone()
    chk("A PAD HELD FOR DETAILS is kept as held-need_details, no count yet", "One more thing" in out and row == ("held-need_details", None), (out[:100], row))
    gate(pg, "Shavez", "Amir", "A003440", "2026-09-07")
    pg.locator("#padagain").click(); pg.wait_for_timeout(2500)
    out = pg.inner_text("#padout")
    row = con.execute("SELECT outcome, count_id, times FROM stock_count_pad_archive WHERE md5=?", (md5f(HELD),)).fetchone()
    chk("...filled in, the same file is recorded as count #3 -- and the SAME archive row becomes recorded, count 3, seen twice",
        "Recorded as count #3" in out and row == ("recorded", 3, 2), (out[:80], row))

    # ------------------------------------------------------------ 7 the checker's archive
    r = pg.request.get(HOST + "/finance/stock/api/pad/file/%s.xlsx" % FILLED_MD5)
    chk("THE CHECKER TAKES THE ORIGINAL BACK, byte for byte, under a name that says which",
        r.status == 200 and r.body() == open(FILLED, "rb").read() and FILLED_MD5[:8] + "_pad_filled.xlsx" in r.headers.get("content-disposition", ""), r.headers.get("content-disposition"))
    r = pg.request.get(HOST + "/finance/stock/api/pad/file/deadbeef.xlsx")
    chk("...a bad md5 is refused", r.status == 400)
    r = pg.request.get(HOST + "/finance/stock/api/pad/file/%s.xlsx" % ("0" * 32))
    chk("...an unknown md5 is 404", r.status == 404)
    j = pg.request.get(HOST + "/finance/stock/api/pad/archive").json()
    chk("THE ARCHIVE LIST: four files, newest first, each on disk and verified, with its outcome and receipt",
        j["ok"] and len(j["files"]) == 4 and all(f["on_disk"] and f["verified"] for f in j["files"])
        and sorted(f["outcome"] for f in j["files"]) == ["recorded", "recorded", "recorded", "refused-unreadable"]
        and [f["receipt"] for f in j["files"] if f["count_id"] == 1] == ["/finance/stock/api/pad/receipt/1.pdf"], [(f["outcome"], f["count_id"]) for f in j["files"]])
    r = pg.request.get(HOST + "/finance/stock/api/pad/archive.zip")
    z = zipfile.ZipFile(io.BytesIO(r.body()))
    namesz = sorted(z.namelist())
    chk("THE WHOLE FOLDER AS ONE ZIP: the four sheets, three receipts, INDEX.txt, the two CSVs",
        r.status == 200 and 'attachment; filename="pad_uploads_' in r.headers.get("content-disposition", "")
        and sum(1 for n in namesz if n.endswith(".xlsx")) == 4 and sum(1 for n in namesz if n.endswith("_PROOF.pdf")) == 3
        and "pad_uploads/INDEX.txt" in namesz and "pad_uploads/ARCHIVE_LEDGER.csv" in namesz and "pad_uploads/SHEETS_RECORDED.csv" in namesz, namesz)
    chk("...and the sheet inside the zip is the sheet that was uploaded", z.read("pad_uploads/" + FILLED_MD5 + ".xlsx") == open(FILLED, "rb").read())
    chk("...the ledger CSV inside says recorded / count 1 / verified for it",
        (",recorded,1,1" in z.read("pad_uploads/ARCHIVE_LEDGER.csv").decode()) and FILLED_MD5 in z.read("pad_uploads/SHEETS_RECORDED.csv").decode())
    chk("the checker sees the archive line under the recent list", "download every uploaded sheet (zip)" in pg.inner_text("#recent"))

    # ------------------------------------------------------------ 8 a counter is refused the originals, but gets the receipt
    import stock_app as SA
    keep = SA._require
    def as_counter(*roles):
        if "checker" in roles and len(roles) == 1:
            from flask import jsonify
            return None, (jsonify(ok=False, error="forbidden"), 403)
        return {"user": "amir", "roles": ["viewer"], "role": "viewer"}, None
    SA._require = as_counter
    r1 = pg.request.get(HOST + "/finance/stock/api/pad/file/%s.xlsx" % FILLED_MD5)
    r2 = pg.request.get(HOST + "/finance/stock/api/pad/archive.zip")
    r3 = pg.request.get(HOST + "/finance/stock/api/pad/archive")
    r4 = pg.request.get(HOST + "/finance/stock/api/pad/receipt/1.pdf")
    SA._require = keep
    chk("A COUNTER IS REFUSED the originals, the zip and the list (403, 403, 403)", (r1.status, r2.status, r3.status) == (403, 403, 403), (r1.status, r2.status, r3.status))
    chk("...but gets the receipt of a recorded sheet", r4.status == 200 and r4.body()[:5] == b"%PDF-")
    pg.evaluate("()=>{ DATA.checker=false; padRecent(); }"); pg.wait_for_timeout(600)
    chk("...and sees no archive line on the page", "download every uploaded sheet" not in pg.inner_text("#recent"))
    pg.evaluate("()=>{ DATA.checker=true; padRecent(); }"); pg.wait_for_timeout(600)

    # ------------------------------------------------------------ 9 the close still works; the proofs stay
    fu = pg.request.get(HOST + "/finance/stock/api/pad/followup/1.xlsx")
    RES2 = os.path.join(HERE, "result_2.xlsx"); open(RES2, "wb").write(fu.body())
    wb2 = load_workbook(RES2); nc, sb = wb2["NOT COUNTED"], wb2["SENT BACK TO FIX"]
    r = 5
    while nc.cell(row=r, column=2).value:
        nm = nc.cell(row=r, column=2).value; ps = PS[nm]; q = MARG[nm]
        if ps > 1: nc.cell(row=r, column=4, value=q // ps); nc.cell(row=r, column=5, value=q % ps)
        else: nc.cell(row=r, column=5, value=q)
        r += 1
    sb.cell(row=5, column=5, value=2)
    wb2.save(RES2)
    pg.set_input_files("#padfile", RES2); pg.wait_for_timeout(3500)
    out = pg.inner_text("#padout")
    chk("THE LAST SHEET: everything counted, the close button appears, its proof offered",
        "everything is counted" in out and pg.locator("#padout button.close").count() == 1 and pg.locator("#padout a.dl.proof").count() == 1, out[:300])
    pg.locator("#padout button.close").click(); pg.wait_for_timeout(1500)
    out = pg.inner_text("#padout")
    pg.locator("#padout").screenshot(path=os.path.join(HERE, "_shot_box_closed.png"))
    pg.locator("#recent").screenshot(path=os.path.join(HERE, "_shot_recent.png"))
    chk("CLOSED: STOCK CHECK COMPLETED, the FINAL RESULT sheet, AND the proof of each of the three sheets",
        "STOCK CHECK COMPLETED" in out and "FINAL RESULT" in out and "Proof of each sheet (PDF): sheet 1 · sheet 2 · sheet 3" in out, out[:400])
    r = pg.request.get(HOST + "/finance/stock/api/pad/receipt/1.pdf"); t = pdf_text(r.body())
    chk("...sheet 1's receipt, made again after the close, carries the same 40 figures and now says the count is CLOSED",
        "EVERY ITEM RECORDED FROM THIS SHEET (40)" in t and "Count #1 is CLOSED - completed on" in t, t[t.find("Whole count"):][:300])
    chk("a receipt for a count that does not exist is 404", pg.request.get(HOST + "/finance/stock/api/pad/receipt/99.pdf").status == 404)
    r = pg.request.get(HOST + "/finance/stock/page/pad", max_redirects=0)
    chk("the old address still goes to this page", r.status in (301, 302))
    chk("THE PAGE THREW NO ERROR THROUGH ALL OF IT", not errs, errs)
    b.close()

print("\n%d ok, %d FAIL" % (len(OK), len(BAD)))
for x in BAD:
    print("  FAIL " + x)
sys.exit(1 if BAD else 0)
