"""LIVE-SHAPE WALK -- S228 THREE-WAY: the shelf, Marg, and our own system.

The owner, 06-Sep-2026: "today's stock check having three columns, one for the
physical, second for the Marg, and third for our own system. Then we compare our
system with the Marg export -- we should match -- and then only we should populate
the screens." And, on why he asked at all: "in the internal calculations on
prompting and searching, errors are continuously coming up."

The thing this walk exists to prove is not the third column. It is the LINE ABOVE
IT. stock_snapshot is keyed (as_on, item) and last-write-wins, and two different
pushes land on it -- Marg's export (push_snapshot) and our own computed figure
(push_expected). Whichever arrived later IS what the count was measured against,
and until now no screen said which. A count could be measured against our own
arithmetic while every page called it Marg. The walk drives all three cases.

The shop here: a tablet where our system and Marg disagree by 40 strips, a cheap
tablet where they disagree by 3, a brace where they agree, an item our system
never computed at all, one item only our system knows and one only Marg knows.
"""
import io, json, os, shutil, sys, threading, time
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import WALK_readiness_s226 as R
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

OK, BAD = [], []
def chk(n, c, d=""):
    (OK if c else BAD).append(n + (("  -- " + str(d)) if d and not c else ""))
    print(("  ok   " if c else "  FAIL ") + n + (("   " + str(d)) if d else ""))

import re as _re
def sale_key(s): return _re.sub(r"\s+", " ", _re.sub(r"[^A-Z0-9 ]+", " ", s.upper())).strip()

AS_ON = "06-09-2026"
TYR = "TYRO BR TAB"          # ours and Marg differ by 40 strips
CHEAP = "CHEAPO SALT TAB"    # ours and Marg differ by 3 strips
KNEE = "KNEE BRACE UNISON"   # they agree
NOEXP = "NOTCOMPUTED SYP"    # our system never computed it
ONLYO = "ONLYOURS TAB"       # only our system knows it
ITEMS = [dict(item=TYR, packing="1*10", pack_size=10), dict(item=CHEAP, packing="1*10", pack_size=10),
         dict(item=KNEE, packing="1*1", pack_size=1), dict(item=NOEXP, packing="1*1", pack_size=1)]
for i in range(1, 7):
    ITEMS.append(dict(item="MED-%03d DOLO TAB" % i, packing="1*10", pack_size=10))
names = [r["item"] for r in ITEMS]; PS = {r["item"]: r["pack_size"] for r in ITEMS}

R.DB = os.path.join(HERE, "_walktw.db"); con = R.build(); R.fill(con)
for t in ("stock_snapshot", "purchase_bill", "purchase_line", "sale_line_item", "purchase_export", "stock_feed"):
    try: con.execute("DELETE FROM %s" % t)
    except Exception: pass
MARG = {TYR: 700, CHEAP: 400, KNEE: 5, NOEXP: 8}
for i, n in enumerate(names):
    if n.startswith("MED-"): MARG[n] = 100 + 10 * i
# OUR system's own figure for the same day: two items disagree, one is absent
OURS = dict(MARG); OURS[TYR] = 700 + 400; OURS[CHEAP] = 400 - 30
del OURS[NOEXP]
OURS[ONLYO] = 250

def snap(source):
    """Write stock_snapshot exactly as /api/snapshot does, from one sender."""
    for r in ITEMS:
        con.execute("INSERT OR REPLACE INTO stock_snapshot (as_on,item,qty,packing,pack_size,source,loaded_at) "
                    "VALUES (?,?,?,?,?,?,?)",
                    (AS_ON, r["item"], MARG[r["item"]], r["packing"], PS[r["item"]], source, "2026-09-06T09:00:00"))
def feed(source, items, at):
    for it, q in items.items():
        con.execute("INSERT INTO stock_feed (as_on, source, item, qty, received_at) VALUES (?,?,?,?,?)",
                    (AS_ON, source, it, int(q), at))

snap("push_snapshot.py v3")                       # Marg's export wrote the figures
feed("push_snapshot.py v3", MARG, "2026-09-06T08:52:00")
feed("push_expected.py v2", OURS, "2026-09-06T09:10:00")
for r in ITEMS:
    con.execute("INSERT OR REPLACE INTO stock_rate (item, rate_p, as_of, source) VALUES (?,?,?,?)",
                (r["item"], 900, "2026-09-06", "walk"))
seq = 500
for nm in names:
    for k in range(3):
        seq += 1
        con.execute("INSERT INTO sale_line_item (day_entry_id, unit, business_date, bill_no, is_return, seq, "
                    "item_name, item_key, qty_raw, pack, amount_p) VALUES (1,'medical',?,?,0,?,?,?,'1:0',?,?)",
                    ("2026-09-0%d" % (k + 1), "A00%04d" % seq, seq, nm, sale_key(nm),
                     "1*1" if PS[nm] == 1 else "1*10", 12000))
con.commit()

ARCH = os.path.join(HERE, "pad_uploads")
if os.path.isdir(ARCH): shutil.rmtree(ARCH)
from flask import Flask
import stock_app
app = Flask("walk_tw")
stock_app.init(app, lambda: con,
               lambda *roles: ({"user": "walk", "roles": ["medical.checker"], "role": "checker"}, None),
               unit="medical")
threading.Thread(target=lambda: app.run(port=8861, threaded=False), daemon=True).start()
time.sleep(1.6)
HOST = "http://127.0.0.1:8861"
errs = []


def fill_pad(src, dst, entries):
    shutil.copy(src, dst); wb = load_workbook(dst); ws = wb.worksheets[0]; where = {}
    for r in range(9, 400):
        nm = ws.cell(row=r, column=2).value
        if isinstance(nm, str) and nm.strip(): where[nm.strip()] = r
    for it, (st, lo) in entries.items():
        r = where[it]
        if st is not None: ws.cell(row=r, column=5, value=st)
        if lo is not None: ws.cell(row=r, column=6, value=lo)
    wb.save(dst); return dst


COUNTED = dict(MARG); COUNTED[TYR] = 300; COUNTED[CHEAP] = 250; COUNTED[KNEE] = 3
good = {nm: (COUNTED[nm] // PS[nm], COUNTED[nm] % PS[nm]) for nm in names}


def report(pg):
    return pg.request.get(HOST + "/finance/stock/api/pad/report/1").json()


with sync_playwright() as pw:
    b = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium")
    ctx = b.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True)
    pg = ctx.new_page()
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.on("console", lambda m: errs.append(m.text) if m.type == "error" and "Failed to load resource" not in m.text else None)
    pg.on("dialog", lambda d: d.accept())

    pg.goto(HOST + "/finance/stock/page/count"); pg.wait_for_timeout(700)
    pg.locator('#whoC button[data-u="Darpan"]').click(); pg.locator('#whoE button[data-u="Amir"]').click()
    pg.fill("#bill", "A003425"); pg.fill("#billdate", "2026-09-06"); pg.wait_for_timeout(150)
    with pg.expect_download() as dl:
        pg.locator("#padget").click()
    PAD = os.path.join(HERE, "pad_tw.xlsx"); dl.value.save_as(PAD)
    pg.set_input_files("#padfile", fill_pad(PAD, os.path.join(HERE, "pad_tw_filled.xlsx"), good)); pg.wait_for_timeout(2500)
    chk("a real count is recorded", "Recorded as count #1" in pg.inner_text("#padout"))

    # ------------------------------------------------------------ 1 the three feeds
    j = report(pg); T = j["three_way"]
    chk("the report carries the three-way block", isinstance(T, dict) and T["as_on"] == AS_ON)
    chk("it knows BOTH feeds landed for that day", T["have_ours"] and T["have_marg"])
    chk("THE PROVENANCE: the count's own figures came from Marg's export, and the page says so",
        T["snapshot_source"] == "marg" and T["measured_against"].startswith("Marg's own closing-stock export"),
        (T["snapshot_source"], T["measured_against"][:40]))
    chk("...with the hour Marg's export was pushed", T["marg_at_text"] == "06-09-2026 08:52", T["marg_at_text"])

    # ------------------------------------------------------------ 2 the verdict
    chk("OUR SYSTEM AND MARG DIFFER on exactly the two items built to differ",
        T["differ"] == 2 and {g["item"] for g in T["gaps"]} == {TYR, CHEAP}, [g["item"] for g in T["gaps"]])
    chk("...the biggest gap is first, and it is said in STRIPS, never in units (D384)",
        T["gaps"][0]["item"] == TYR and T["gaps"][0]["gap_text"] == "40 strips", T["gaps"][0])
    chk("...the smaller one is 3 strips", T["gaps"][1]["gap_text"] == "3 strips", T["gaps"][1]["gap_text"])
    chk("...the agreeing items are counted, not listed", T["agree"] == T["compared"] - 2 and T["agree"] > 0,
        (T["agree"], T["compared"]))
    chk("...and the headline NAMES ITS DENOMINATOR, so it cannot be read against the count's own tiles",
        T["line"].startswith("Our system and Marg's export differ on 2 of the 9 items both of them carry")
        and TYR in T["line"], T["line"])
    chk("an item only OUR system knows is named, not silently dropped",
        T["only_ours_n"] == 1 and T["only_ours"] == ["ONLYOURS TAB"], (T["only_ours_n"], T["only_ours"]))
    chk("an item only MARG knows is named too", T["only_marg_n"] == 1 and T["only_marg"] == [NOEXP],
        (T["only_marg_n"], T["only_marg"]))
    chk("the count is NOT called trustworthy while the two disagree", T["trustworthy"] is False)

    # ------------------------------------------------------------ 3 the third figure on the row
    D = {x["item"]: x for x in j["differences"]}
    chk("THE THIRD COLUMN: the tablet carries Marg 700, ours 1100, counted 300",
        D[TYR]["marg"] == 700 and D[TYR]["ours"] == 1100 and D[TYR]["counted"] == 300 and D[TYR]["ours_gap"] == 400,
        (D[TYR]["marg"], D[TYR]["ours"], D[TYR]["ours_gap"]))
    chk("an item our system never computed shows NO third figure -- never a made-up one",
        NOEXP not in D or D[NOEXP]["ours"] is None,
        (D.get(NOEXP) or {}).get("ours"))
    M = {x["item"]: x for x in j["matched"]}
    chk("the matched rows carry it as well", any(x.get("ours") is not None for x in M.values()))

    # ------------------------------------------------------------ 4 the page itself
    pg.goto(HOST + "/finance/stock/page/report?count=1"); pg.wait_for_timeout(1400)
    body = pg.inner_text("body")
    chk("the page shows the provenance line", "Measured against" in body and "Marg's own closing-stock export" in body)
    chk("...and the verdict card", "our system vs marg" in body.lower() and "differ on 2 of the 9" in body)
    chk("...with NO banner while the figures really are Marg's", pg.locator(".tw .bang").count() == 0)
    chk("...with the two differing lines in a small table, Marg beside ours",
        pg.locator(".tw table tr").count() == 3, pg.locator(".tw table tr").count())
    chk("...and it always says WHAT it compares, so it cannot be read against the count's own tiles",
        "two COMPUTER figures only" in pg.inner_text(".tw"))
    chk("ROUTINE CONFIRMATION SITS AFTER THE NUMBERS -- the tiles come first",
        pg.evaluate("document.querySelector('.kpis').compareDocumentPosition(document.querySelector('.tw')) & Node.DOCUMENT_POSITION_FOLLOWING") > 0,
        "the card is above the tiles")
    pg.screenshot(path=os.path.join(HERE, "_shot_tw_report.png"), full_page=True)
    # OPEN EVERY LANE -- a collapsed page proves nothing about the line itself.
    for i in range(pg.locator('[data-tg]').count()):
        try: pg.locator('[data-tg]').nth(i).click(); pg.wait_for_timeout(60)
        except Exception: pass
    pg.wait_for_timeout(400)
    row = pg.locator('[data-item="%s"]' % TYR).first
    txt = " ".join(row.inner_text().split())
    chk("THE COUNT'S OWN LINE IS UNCHANGED -- Marg, then counted, then the shortage",
        "Marg 70 strips \u2192 counted 30 strips" in txt and "short 40 strips" in txt, txt[:120])
    chk("...and the third figure sits on its OWN line, so nothing is squeezed",
        "our own system said 110 strips" in txt and "40 strips more than Marg" in txt, txt[:160])
    chk("...and NO figure anywhere on an item line is split from its unit by a line break",
        pg.evaluate("""(() => {const bad=[];document.querySelectorAll('.q .nb').forEach(e=>{if(e.getClientRects().length>1)bad.push(e.textContent);});return bad;})()""") == [],
        pg.evaluate("""(() => {const bad=[];document.querySelectorAll('.q .nb').forEach(e=>{if(e.getClientRects().length>1)bad.push(e.textContent);});return bad;})()"""))
    chk("...where our system agrees with Marg the line is three words, not a sentence",
        "our own system: same" in " ".join(pg.locator('[data-item="%s"]' % KNEE).first.inner_text().split()),
        " ".join(pg.locator('[data-item="%s"]' % KNEE).first.inner_text().split())[:120])
    chk("...the page still does not scroll sideways at 390px with every lane open",
        pg.evaluate("document.documentElement.scrollWidth") <= 390,
        pg.evaluate("document.documentElement.scrollWidth"))
    pg.screenshot(path=os.path.join(HERE, "_shot_tw_open.png"), full_page=True)
    chk("...the card is marked as a disagreement, not as agreement",
        "bad" in (pg.locator(".tw").first.get_attribute("class") or ""),
        pg.locator(".tw").first.get_attribute("class"))
    chk("...and it says what a gap usually means, so nobody recounts the shelf for it",
        "keyed after the export" in body)
    chk("no page errors at 390px", not errs, errs[:3])

    # ------------------------------------------------------------ 5 THE CASE THE OWNER ASKED ABOUT
    # our own push lands LAST on the same day: from now on the count's "Marg"
    # column is our own arithmetic, and the page must say so in as many words.
    con.execute("UPDATE stock_snapshot SET source='push_expected.py v2' WHERE as_on=?", (AS_ON,))
    con.commit()
    T2 = report(pg)["three_way"]
    chk("WHEN OUR OWN PUSH WROTE THE FIGURES, the page stops calling them Marg's",
        T2["snapshot_source"] == "expected" and "OUR OWN COMPUTED FIGURE, not Marg's export" in T2["measured_against"],
        T2["measured_against"][:60])
    chk("...and the CARD carries its own banner, so a reader who skips the header cannot be misled",
        "came from OUR SYSTEM, not from Marg's export" in (T2["bang"] or ""), T2["bang"][:60])
    chk("...and such a count is never called trustworthy", T2["trustworthy"] is False)
    pg.goto(HOST + "/finance/stock/page/report?count=1"); pg.wait_for_timeout(1200)
    chk("...and it is flagged on the page, not buried",
        "OUR OWN COMPUTED FIGURE" in pg.inner_text("body") and pg.locator(".warn").count() >= 1)
    chk("...as a filled banner across the card, not a red word in a table",
        pg.locator(".tw .bang").count() == 1 and "OUR SYSTEM" in pg.locator(".tw .bang").inner_text(),
        pg.locator(".tw .bang").inner_text()[:60] if pg.locator(".tw .bang").count() else "no banner")
    chk("...and the card is red even though the two feeds themselves are comparable",
        "bad" in (pg.locator(".tw").first.get_attribute("class") or ""))
    chk("...and THIS time the card JUMPS the queue and sits above the numbers",
        pg.evaluate("document.querySelector('.tw').compareDocumentPosition(document.querySelector('.kpis')) & Node.DOCUMENT_POSITION_FOLLOWING") > 0,
        "the alarm is below the tiles")
    pg.screenshot(path=os.path.join(HERE, "_shot_tw_warn.png"), full_page=True)

    # both senders on one day -- the figures are then from no single source
    con.execute("UPDATE stock_snapshot SET source='push_snapshot.py v3' WHERE as_on=? AND item=?", (AS_ON, TYR))
    con.commit()
    T3 = report(pg)["three_way"]
    chk("when TWO senders wrote one day, the page says the figures are not from one source",
        T3["snapshot_source"] == "mixed" and "more than one sender" in T3["measured_against"],
        T3["measured_against"][:50])
    chk("...and that state gets its own banner too", "more than one sender" in (T3["bang"] or ""), T3["bang"][:50])

    # ------------------------------------------------------------ 6 nothing computed at all
    con.execute("UPDATE stock_snapshot SET source='push_snapshot.py v3' WHERE as_on=?", (AS_ON,))
    con.execute("DELETE FROM stock_feed WHERE source LIKE 'push_expected%'")
    con.commit()
    T4 = report(pg)["three_way"]
    chk("WITH NO COMPUTED FIGURE the page says so plainly and compares nothing",
        T4["have_ours"] is False and T4["differ"] == 0 and "has not computed a figure" in T4["line"], T4["line"])
    j4 = report(pg)
    chk("...and every third figure is absent rather than zero",
        all(x["ours"] is None for x in j4["differences"]),
        [(x["item"], x["ours"]) for x in j4["differences"] if x["ours"] is not None][:3])
    pg.goto(HOST + "/finance/stock/page/report?count=1"); pg.wait_for_timeout(1200)
    chk("...the card is neither green nor red -- there is nothing to judge",
        (pg.locator(".tw").first.get_attribute("class") or "").strip() == "tw",
        pg.locator(".tw").first.get_attribute("class"))
    chk("...and the difference lines simply do not mention our system at all",
        "our own system" not in pg.locator('[data-item="%s"]' % TYR).first.inner_text())

    # ------------------------------------------------------------ 7 the agreeing shop
    feed("push_expected.py v2", MARG, "2026-09-06T09:10:00")
    con.commit()
    T5 = report(pg)["three_way"]
    chk("WHEN THEY AGREE the page says so in one line, and the count is trustworthy",
        T5["differ"] == 0 and T5["trustworthy"] is True
        and T5["line"].startswith("Our system and Marg's export AGREE on all"), T5["line"])
    chk("...and no banner at all", not T5["bang"])
    pg.goto(HOST + "/finance/stock/page/report?count=1"); pg.wait_for_timeout(1200)
    chk("...and the card turns green, with a tick, so agreement LOOKS like agreement",
        "ok" in (pg.locator(".tw").first.get_attribute("class") or "")
        and "\u2713" in pg.inner_text(".tw"), pg.inner_text(".tw")[:60])
    chk("...and it does NOT print the gap explanation when there is no gap",
        "keyed after the export" not in pg.inner_text(".tw"), pg.inner_text(".tw")[-70:])
    chk("no page errors in any of the five states", not errs, errs[:3])
    pg.screenshot(path=os.path.join(HERE, "_shot_tw_agree.png"), full_page=True)
    ctx.close(); b.close()

print("\n%d ok, %d FAIL" % (len(OK), len(BAD)))
for x in BAD: print("  FAIL " + x)
sys.exit(1 if BAD else 0)
