#!/usr/bin/env python3
"""walk_s538.py -- kit S538 (session 311, 12-Oct-2026; D714; F-867, F-868, F-869). READ-ONLY.
    python3 -B walk_s538.py <folder holding packs.py> [--db <a COPY of finance.db> --month 2026-09] [--expect-red]
Part A builds a made-up database in memory (no person, no number) and holds the new packs.py to every rule of D714:
splits by their parts, the Razorpay booking as the consultation fee only, a 'Wallet' with no booking number decided against the
bank, the bank check day by day (D+1), the digest one row per charge. --expect-red runs Part A against the OLD packs.py: it must
fail (the negative control). Part B (--db) runs the same builders on a copy of the real database and holds September to what
S305's test measured: the nine days 11, 14, 17, 19, 22, 23, 24, 25, 29 equal to the rupee, 12-Sep Rs 2,950 short.
Prints 'WALK OK ...' on success; exit 1 on any failure."""
import importlib.util
import io
import os
import sqlite3
import sys
import types

FAILS = []
RUN = [0]


def ok(cond, what):
    RUN[0] += 1
    if not cond:
        FAILS.append(what)


def load_packs(folder):
    try:
        import flask                                     # noqa: F401
    except ImportError:                                  # the PC has no flask: a stand-in, enough to import the module
        fl = types.ModuleType("flask")

        class _BP:
            def __init__(self, *a, **k):
                pass

            def route(self, *a, **k):
                return lambda f: f
        fl.Blueprint = _BP
        for n in ("jsonify", "render_template_string", "request", "send_file"):
            setattr(fl, n, None)
        sys.modules["flask"] = fl
    spec = importlib.util.spec_from_file_location("packs_walk", os.path.join(folder, "packs.py"))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def made_up():
    con = sqlite3.connect(":memory:")
    con.row_factory = sqlite3.Row
    con.executescript("""
    CREATE TABLE setting (key TEXT PRIMARY KEY, value TEXT);
    CREATE TABLE clinic_day_revenue (business_date TEXT PRIMARY KEY, cons_count INT, cons_amount_p INT, xray_count INT, xray_amount_p INT,
      proc_count INT, proc_amount_p INT, total_count INT, total_amount_p INT, tender_json TEXT);
    CREATE TABLE clinic_day_line (business_date TEXT, section TEXT, sn INT, patient TEXT DEFAULT '', clinic_id TEXT DEFAULT '', amount_p INT,
      mode TEXT, shift TEXT DEFAULT '', gateway_ref TEXT NOT NULL DEFAULT '');
    CREATE TABLE clinic_day_tender (business_date TEXT, clinic_id TEXT DEFAULT '', invoice_no TEXT DEFAULT '', tender TEXT, amount_p INT, source_file TEXT DEFAULT '');
    CREATE TABLE icici_statement_line (id INTEGER PRIMARY KEY, account_ref TEXT, txn_date TEXT, value_date TEXT, description TEXT, reference TEXT,
      withdrawal_p INT DEFAULT 0, deposit_p INT DEFAULT 0, balance_p INT, is_cash_deposit INT DEFAULT 0, source_file TEXT, sha256 TEXT, ingested_at TEXT);
    CREATE TABLE icici_statement_period (id INTEGER PRIMARY KEY, account_ref TEXT, period_from TEXT, period_to TEXT, opening_p INT, closing_p INT,
      source_file TEXT, sha256 TEXT, ingested_at TEXT, layout TEXT, closing_printed INT);
    CREATE TABLE upi_txn (id INTEGER PRIMARY KEY, merchant_id TEXT, unit TEXT, txn_date TEXT, amount_p INT, rrn TEXT, mode TEXT);
    INSERT INTO upi_txn (merchant_id, unit, txn_date, amount_p, rrn, mode) VALUES ('m', 'clinic', '2031-03-02', 10000, 'r1', 'UPI');
    CREATE TABLE payment_register (row_no INTEGER PRIMARY KEY, date_raw TEXT DEFAULT '', date_iso TEXT, vendor TEXT DEFAULT '', description TEXT DEFAULT '',
      amount_raw TEXT DEFAULT '', amount_paise INT, in_drive INT, gmail_link TEXT DEFAULT '', content_md5 TEXT DEFAULT 'x', first_seen TEXT DEFAULT '', last_seen TEXT DEFAULT '');
    """)
    R = 100

    def day(d, lines, tenders=()):
        tot = sum(a for _s, _m, a, _g in lines)
        cons = sum(a for s, _m, a, _g in lines if s == "consult")
        proc = sum(a for s, _m, a, _g in lines if s == "proc")
        con.execute("INSERT INTO clinic_day_revenue VALUES (?,?,?,?,?,?,?,?,?,?)", (d, 1, cons, 0, 0, 1, proc, len(lines), tot, "{}"))
        for i, (s, m, a, g) in enumerate(lines):
            con.execute("INSERT INTO clinic_day_line (business_date, section, sn, amount_p, mode, gateway_ref) VALUES (?,?,?,?,?,?)", (d, s, i, a, m, g))
        for t, a in tenders:
            con.execute("INSERT INTO clinic_day_tender (business_date, tender, amount_p) VALUES (?,?,?)", (d, t, a))

    def settle(d, a):
        con.execute("INSERT INTO icici_statement_line (account_ref, txn_date, description, deposit_p) VALUES ('7777',?,?,?)", (d, "EZY/ICICIPOS_SET_10XX/000", a))
    # 01: cash 1000, online 500, a split 1100 = cash 600 + online 500 -> ICICI expects 1000; the bank settles 1000 on 02 -> exact
    day("2031-03-01", [("consult", "Cash", 1000 * R, ""), ("consult", "Online Payment", 500 * R, ""), ("xray", "Split Payment", 1100 * R, "")],
        [("Cash", 600 * R), ("Online Payment", 500 * R)])
    settle("2031-03-02", 1000 * R)
    # 02: a Razorpay booking (consultation 600) and the same visit's procedure 800 carrying the same number (F-869): portal 600 only,
    #     the 800 is expected in ICICI; the bank settles 800 on 03 -> exact
    day("2031-03-02", [("consult", "Online Payment", 600 * R, "pay_MADEUP01"), ("proc", "Online Payment", 800 * R, "pay_MADEUP01")])
    settle("2031-03-03", 800 * R)
    # 03: a 'Wallet' consultation with no booking number (600); the bank settles online 300 + 600 extra -> the Wallet was online
    day("2031-03-03", [("consult", "Online Payment", 300 * R, ""), ("consult", "Wallet", 600 * R, "")])
    settle("2031-03-04", 900 * R)
    # 04: a 'Wallet' split part 500 the bank never saw -> cash; online 200 settled 200
    day("2031-03-04", [("xray", "Split Payment", 700 * R, ""), ("consult", "Cash", 400 * R, "")], [("Online Payment", 200 * R), ("Wallet", 500 * R)])
    settle("2031-03-05", 200 * R)
    # 05: a split whose parts the export did not carry -> named, never guessed; card 300 settled 250 -> Rs 50 less
    day("2031-03-05", [("consult", "Split Payment", 1000 * R, ""), ("xray", "Debit Card", 300 * R, "")])
    settle("2031-03-06", 250 * R)
    # 06: no Docterz day; the bank settles 300 on 07 -> 'no Docterz entry'
    settle("2031-03-07", 300 * R)
    # 31: online 400, its D+1 (01-Apr) is past the statement -> waiting
    day("2031-03-31", [("consult", "Online Payment", 400 * R, ""), ("consult", "Wallet", 100 * R, "")])
    con.execute("INSERT INTO icici_statement_period (account_ref, period_from, period_to) VALUES ('7777','2031-03-01','2031-03-31')")
    # the digest: a renewal and its receipt 6 days apart (one charge); two receipts of one amount (two charges); a renewal and a
    # receipt 9 days apart (two papers); a confirmation and its receipt the next day (one charge)
    for i, (d, v, desc, a) in enumerate([("2031-03-09", "Fruitco", "Your Subscription Renewal", 24900), ("2031-03-15", "Fruitco", "Your receipt from Fruitco.", 24900),
                                         ("2031-03-10", "Playstore", "Your Order Receipt", 299), ("2031-03-12", "Playstore", "Your Order Receipt", 299),
                                         ("2031-03-01", "Fruitco", "Your Subscription Renewal", 1999), ("2031-03-10", "Fruitco", "Your receipt from Fruitco.", 1999),
                                         ("2031-03-11", "Fruitco", "Subscription is Confirmed", 699),
                                         ("2031-03-12", "Fruitco", "Your receipt from Fruitco.", 699)], 1):
        con.execute("INSERT INTO payment_register (row_no, date_iso, vendor, description, amount_paise) VALUES (?,?,?,?,?)", (i, d, v, desc, a * R))
    return con


def part_a(P):
    con = made_up()
    P.ensure(con)
    con.execute("UPDATE stmt_slot SET ident_tail='7777' WHERE key='icici_clinic'")
    m = "2031-03"
    f = getattr(P, "day_parts", None)
    ok(f is not None, "A1 packs.day_parts exists")
    if f:
        d1, d2, d3, d4, d5 = (f(con, "2031-03-0%d" % i) for i in range(1, 6))
        ok(d1 and d1["cash"] == 160000 and d1["online"] == 100000 and d1["split_open"] == 0, "A2 a split counts by its parts (cash 1600, online 1000)")
        ok(d2 and d2["portal"] == 60000 and d2["online"] == 80000, "A3 the booking is the consultation fee only; the procedure is ICICI's (F-869)")
        ok(d3 and d3["wallet"] == 60000 and d3["portal"] == 0, "A4 a 'Wallet' with no booking number is a wallet entry, not a booking")
        ok(d4 and d4["wallet"] == 50000 and d4["online"] == 20000, "A5 a 'Wallet' split part is a wallet entry")
        ok(d5 and d5["split_open"] == 100000 and d5["card"] == 30000, "A6 a split without parts is named, never guessed")
    if f:                                                 # the review's case: a booked consultation in a split + another split without parts
        con.execute("INSERT INTO clinic_day_revenue VALUES ('2031-04-02',1,50000,0,0,1,180000,3,230000,'{}')")
        for i, (s_, m_, a_, g_, c_) in enumerate([("consult", "Split Payment", 50000, "pay_MADEUP02", "A"), ("proc", "Split Payment", 100000, "", "A"),
                                                  ("proc", "Split Payment", 80000, "", "B")]):
            con.execute("INSERT INTO clinic_day_line (business_date, section, sn, amount_p, mode, gateway_ref, clinic_id) VALUES ('2031-04-02',?,?,?,?,?,?)",
                        (s_, i, a_, m_, g_, c_))
        for t_, a_ in (("Wallet", 50000), ("Cash", 100000)):
            con.execute("INSERT INTO clinic_day_tender (business_date, clinic_id, tender, amount_p) VALUES ('2031-04-02','A',?,?)", (t_, a_))
        d6 = f(con, "2031-04-02")
        ok(d6 and d6["portal"] == 50000 and d6["wallet"] == 0 and d6["cash"] == 100000 and d6["split_open"] == 80000,
           "A6b a booked split is netted from ITS patient's parts only (got %s)" % (d6,))
    bc = {x["day"]: x for x in P.bank_check(con, m)} if hasattr(P, "bank_check") else {}
    ok(bool(bc), "A7 packs.bank_check exists and answers")
    if bc:
        ok(bc["2031-03-01"]["exact"] and bc["2031-03-02"]["exact"], "A8 the settlement of D+1 equals Docterz's online + card -> exact")
        ok(bc["2031-03-03"]["wallet_as"] == "online" and "Wallet" in bc["2031-03-03"]["verdict"] and not bc["2031-03-03"]["differs"],
           "A9 the bank's extra equals the Wallet entry -> online, nothing to ask")
        ok(bc["2031-03-04"]["wallet_as"] == "cash", "A10 the bank never saw the Wallet part -> cash")
        ok(bc["2031-03-05"]["diff"] == -5000 and bc["2031-03-05"]["differs"] and "reception" not in bc["2031-03-05"]["verdict"],
           "A11 Rs 50 less -> a day that differs, never put to reception")
        ok("2031-03-06" in bc and "no Docterz entry" in bc["2031-03-06"]["verdict"], "A12 a bank day with no Docterz day is listed")
        ok(bc["2031-03-31"]["bank"] is None and bc["2031-03-31"]["wallet_as"] == "waiting for the bank", "A13 past the statement -> waiting, undecided")
    b, why = P.build_income(con, m)
    ok(b is not None, "A14 the income workbook builds (%s)" % why)
    if b:
        from openpyxl import load_workbook               # noqa: PLC0415
        wb = load_workbook(io.BytesIO(b), data_only=True)
        ok(wb.sheetnames == ["Income", "Ledger date-wise"], "A15 the accountants' workbook has no bank check (the owner's only)")
        ok(not any("ask reception" in str(c or "").lower() or "bank check" in str(c or "").lower() for r in wb["Income"].iter_rows(values_only=True) for c in r),
           "A15b the accountants' income sheet names no bank check and no reception question")
        head = [c.value for c in wb["Ledger date-wise"][1]]
        ok(not any(str(h or "").strip() == "Split Rs" for h in head), "A16 no 'Split' lump column")
        rows = list(wb["Ledger date-wise"].iter_rows(values_only=True))[1:-1]
        ok(all((r[-1] or "") == "" for r in rows), "A17 every day's parts add to its total")
        inc = {str(r[0]): r[1] for r in wb["Income"].iter_rows(values_only=True)}
        k = [x for x in inc if x.startswith("of which booked online")]
        ok(k and inc[k[0]] == 600, "A18 the income sheet shows the booking apart (Rs 600)")
    kept, dropped, total = P.digest_rows(con, m)
    vend = [(r["vendor"], r["amount"]) for r in kept]
    ok(vend.count(("Fruitco", 24900.0)) == 1, "A19 a renewal and its receipt 6 days apart are one row (F-868)")
    ok(vend.count(("Playstore", 299.0)) == 2, "A20 two receipts of one amount stay two rows")
    ok(vend.count(("Fruitco", 1999.0)) == 2, "A21 a renewal and a receipt 9 days apart stay two rows")
    ok(vend.count(("Fruitco", 699.0)) == 1, "A22 a confirmation and its receipt are one row")
    ok(abs(total - (24900 + 299 * 2 + 1999 * 2 + 699)) < 0.01, "A23 the digest total counts each charge once")
    bd, _w = P.build_digest(con, m)
    if bd:
        from openpyxl import load_workbook               # noqa: PLC0415
        first = load_workbook(io.BytesIO(bd)).active["A1"].value or ""
        ok("Payment papers found in mail" in first and "overlap the card spends" in first, "A24 the digest says it is payment papers that overlap the card spends")
    line = P.clinic_upi_check(con, m)
    ok("differ" in line and "same to the rupee" in line and "not sent" in line and "reception" not in line, "A25 the owner's line: the month's totals and counts, his only")
    bb, _w = P.build_bank_check(con, m) if hasattr(P, "build_bank_check") else (None, "")
    ok(bb is not None, "A26 the owner's bank check builds on its own")
    if bb:
        from openpyxl import load_workbook               # noqa: PLC0415
        cells = [str(c or "") for r in load_workbook(io.BytesIO(bb)).active.iter_rows(values_only=True) for c in r]
        ok(any(c.startswith("Month: Docterz online + card") for c in cells) and any(c.startswith("Month: the ICICI card machine") for c in cells),
           "A27 the bank check opens with the month's two totals")
    rows_, _att = P.pack_rows(con, m, light=True)
    t = {r["key"]: r["title"] for r in rows_}
    ok("Payment papers" in t.get("digest", "") and "bank check" not in " ".join(t.values()).lower(), "A28 the pack rows: payment papers; no bank-check row")
    import tempfile                                       # noqa: PLC0415
    stub = tempfile.mkdtemp(prefix="walk538_mail_")
    os.environ["PACKS_MAIL_STUB"] = stub
    con.execute("INSERT OR REPLACE INTO setting (key, value) VALUES ('packs.accountant_to', 'accounts@example.invalid')")
    body, code = P.send_pack(con, {"user": "walk"}, m)
    import email                                          # noqa: PLC0415
    from email import policy                              # noqa: PLC0415
    head = ""
    for f in os.listdir(stub):
        with open(os.path.join(stub, f), "rb") as fh:
            msg = email.message_from_binary_file(fh, policy=policy.default)
        part = msg.get_body(preferencelist=("plain",))
        head += part.get_content() if part else ""
    ok(code == 200 and "2 · UPI totals" in head and "Docterz online" not in head and "clinic: " not in head and "differ" not in head,
       "A29 the accountants' mail carries no clinic bank line")
    for f in os.listdir(stub):
        os.remove(os.path.join(stub, f))
    os.rmdir(stub)
    cl = {x["item"]: x for x in P.checklist(con, m)}
    up = [x for x in cl.values() if "UPI totals" in x["item"]]
    ok(up and "Docterz" not in (up[0]["words"] or "") and "differ" not in (up[0]["words"] or ""), "A31 Shavez's checklist does not carry the owner's bank line")
    hp = os.path.join(FOLDER[0], "packs.html")
    if os.path.exists(hp):
        ok("/finance/packs/preview/bankcheck?month=" in open(hp, encoding="utf-8").read(), "A30 the packs page links the owner's bank check")


def part_b(P, db, month):
    con = sqlite3.connect("file:%s?mode=ro" % db, uri=True)
    con.row_factory = sqlite3.Row
    bc = P.bank_check(con, month)
    exact = sorted(int(x["day"][8:]) for x in bc if x["exact"])
    by = {x["day"]: x for x in bc}
    if month == "2026-09":
        want = [11, 14, 17, 19, 22, 23, 24, 25, 29]
        ok(all(d in exact for d in want), "B1 September: the nine exact days of S305's test (got %s)" % exact)
        ok(set(exact) - set(want) <= {1, 2}, "B2 September: no exact day beyond S305's nine but 1-2 Sep (outside its exports) (got %s)" % exact)
        ok(by.get("2026-09-12", {}).get("diff") == -295000, "B3 September: 12-Sep the bank Rs 2,950 short")
    rows = [dict(r) for r in con.execute("SELECT * FROM clinic_day_revenue WHERE substr(business_date,1,7)=?", (month,))]
    bad = []
    for r in rows:
        p = P.day_parts(con, r["business_date"]) or P._parts_from_tender(r)
        if sum(p[k] for k in P.S538_PARTS) != int(r["total_amount_p"] or 0):
            bad.append(r["business_date"])
    ok(not bad, "B4 every day's parts add to its total (%s)" % ", ".join(bad))
    b, why = P.build_income(con, month)
    ok(b is not None, "B5 the income workbook builds on the real data (%s)" % why)
    kept, _d, total = P.digest_rows(con, month)
    if month == "2026-09":
        ok(sum(1 for r in kept if (r["amount"] or 0) == 24900) == 1, "B6 September: the Rs 24,900 charge is one row")
    n_ask = sum(1 for x in bc if x["differs"])
    return "%s: %d days in the bank check, %d exact (%s), %d differ; digest %d rows Rs %.2f" % (month, len(bc), len(exact), ",".join(map(str, exact)), n_ask, len(kept), total)


FOLDER = [""]


def main(argv):
    folder = argv[1]
    FOLDER[0] = folder
    P = load_packs(folder)
    part_a(P)
    note = ""
    if "--db" in argv:
        db = argv[argv.index("--db") + 1]
        month = argv[argv.index("--month") + 1] if "--month" in argv else "2026-09"
        try:
            note = part_b(P, db, month)
        except Exception as ex:                          # noqa: BLE001
            FAILS.append("B0 the real-data part raised %s" % ex)
    red = "--expect-red" in argv
    if red:
        if FAILS:
            print("WALK OK -- the negative control is RED as it must be (%d checks fail on the old file, e.g. %s)" % (len(FAILS), FAILS[0]))
            return 0
        print("WALK RED -- the old file passed the checks: the walk proves nothing")
        return 1
    if FAILS:
        for x in FAILS:
            print("FAIL", x)
        return 1
    print("WALK OK -- %d checks green%s" % (RUN[0], ("; " + note) if note else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
