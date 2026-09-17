#!/usr/bin/env python3
"""selftest_s301.py -- proves S301_MARG_VOUCHERS offline.
Bytes: stock_app.py and stock_amir.html at their S299 pins, stock_hub.html at S299's; anchors once;
refusal; idempotence; backups; compile; the Amir and hub scripts parse. Behaviour over the S299
fixture: confirmed swaps wait as STOCK ISSUE / STOCK RECEIVE lines; a round freezes them in
batches (at most six, or the setting), ISSUE and RECEIVE apart; a second press finds nothing; a
remainder decided later goes on the next round from Marg-after-swap; an answer taken back after
its round becomes a reversal on the next; every item's lines, keyed in order, end where the lines
now say; a voucher recorded as entered and taken back; the workbook's two sheets.
    python3 -B selftest_s301.py --base <dir with stock_app.py, stock_amir.html, stock_hub.html (S299) and padwriter.py>"""
import argparse, hashlib, importlib.util, io, os, re, shutil, sqlite3, subprocess, sys, tempfile, types, zipfile
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import patch_marg_vouchers_s301 as P  # noqa: E402
PINS = dict(stock_app="75ffe86168e94c98fd1d724d0226a373", stock_amir="3181d9de5133afd51fe46dd2f61e9133")
HUB_FROM = "d0bec54f682cbe41c5bec70b30bbefb9"
EXT = dict(stock_app=".py", stock_amir=".html")
FAILS = []
def check(n, ok, d=""):
    print(("ok   " if ok else "FAIL ") + n + ("" if ok else "  -- " + str(d)[:500]))
    if not ok: FAILS.append(n)
def md5(p): return hashlib.md5(io.open(p, "rb").read()).hexdigest()
def stub_flask():
    f = types.ModuleType("flask")
    class BP:
        def __init__(s, *a, **k): pass
        def route(s, *a, **k): return lambda fn: fn
    f.Blueprint = BP; f.jsonify = lambda *a, **k: (a, k); f.request = types.SimpleNamespace(args={}, form={}, get_json=lambda *a, **k: {})
    f.Response = object; f.redirect = lambda *a, **k: None
    sys.modules.setdefault("flask", f)
def load(path, name):
    stub_flask(); spec = importlib.util.spec_from_file_location(name, path); m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m
def fixture():
    def x(item, lane, marg, counted, mrp_unit=100, pack=1, word=None, consumable=False):
        return dict(item=item, lane=lane, marg=marg, counted=counted, diff=counted - marg, pack=pack, packing="1*%d" % pack,
                    mrp_p=(counted - marg) * mrp_unit, cost_p=(counted - marg) * mrp_unit // 2, consumable=consumable, why="fixture",
                    word=(dict(action=word, label=word.lower().replace("_", " ")) if word else None))
    diffs = [x("ANKLE BINDER BAMBOO L", "ortho", 3, 0, 27520), x("ANKLE BINDER L TYNOR", "ortho", 0, 2, 30960),
             x("L S BELT CONT GRAY UNISON M", "ortho", 2, 1, 97300), x("L S BELT CONT GRAY UNISON L", "ortho", 2, 3, 139000),
             x("KNEE SUPPORT HINGED XXL UNISO", "ortho", 2, 1, 80500), x("KNEE SUPPORT HINGED M", "ortho", 3, 4, 80500),
             x("RIB BELT S TYNOR", "ortho", 1, 0, 52000), x("DISPO SYRINGE NIPRO 3ML", "consume", 62, 0, 1000, consumable=True),
             x("CERVICAL COLLAR SOFT HOPE L", "ortho", -1, 0, 19600), x("CERVICAL COLLAR SOFT HOPE M", "ortho", 1, 0, 19600),
             x("ETOZOX 90", "loss", 100, 80, 20, 10, word="PARKED"), x("PARI CR 12.5", "recount", 84, 367, 15, 10),
             x("MOTIF FORTE", "loss", 60, 55, 468, 10), x("ASTOFEN P", "recount", 10, 30, 550, 10),
             x("ROSIKA FORTE", "loss", 723, 553, 30, 10)]
    pairs = [dict(salt="ETORICOXIB 90", short=[dict(item="ETOZOX 90")], over=[dict(item="PARI CR 12.5")]),
             dict(salt="ACECLOFENAC 100 + PCM", short=[dict(item="MOTIF FORTE")], over=[dict(item="ASTOFEN P")])]
    return diffs, pairs
def build(S, con, words_extra=None):
    diffs, pairs = fixture(); by = {x["item"]: x for x in diffs}
    for k, v in (words_extra or {}).items(): by[k]["word"] = v
    words = {x["item"]: x["word"] for x in diffs if x["word"]}
    d = dict(differences=diffs, matches=S._match_apply(con, 1, diffs, pairs, words), count_id=1, day="06-09-2026")
    return d, {x["item"]: x for x in diffs}
def main(argv=None):
    ap = argparse.ArgumentParser(); ap.add_argument("--base", required=True); a = ap.parse_args(argv)
    src = {k: os.path.join(a.base, k + EXT[k]) for k in PINS}
    for k, p in src.items(): check("%s is the S299 pin" % k, md5(p) == PINS[k], md5(p))
    check("stock_hub.html (live) is the S299 pin", md5(os.path.join(a.base, "stock_hub.html")) == HUB_FROM)
    tmp = tempfile.mkdtemp(prefix="s301_"); w = {k: os.path.join(tmp, os.path.basename(p)) for k, p in src.items()}
    for k in src: shutil.copy2(src[k], w[k])
    shutil.copy2(os.path.join(a.base, "padwriter.py"), os.path.join(tmp, "padwriter.py"))
    for k, edits in (("stock_app", P.APP_EDITS), ("stock_amir", P.AMIR_EDITS)):
        t = io.open(w[k], encoding="utf-8").read()
        for old, _ in edits: check("%s anchor once: %s" % (k, old.strip()[:40].replace("\n", " ")), t.count(old) == 1, t.count(old))
    cmd = [sys.executable, "-B", os.path.join(HERE, "patch_marg_vouchers_s301.py"), "--app", w["stock_app"], "--amir", w["stock_amir"]]
    r = subprocess.run(cmd + ["--app-from", "0" * 32], capture_output=True, text=True); check("refuses a wrong --from", r.returncode != 0 and "REFUSING" in r.stdout + r.stderr)
    r = subprocess.run(cmd + ["--app-from", PINS["stock_app"], "--amir-from", PINS["stock_amir"]], capture_output=True, text=True)
    check("patches both from the S299 pins", r.returncode == 0, r.stdout + r.stderr)
    h = {k: md5(p) for k, p in w.items()}
    r = subprocess.run(cmd, capture_output=True, text=True); check("idempotent", {k: md5(p) for k, p in w.items()} == h and r.stdout.count("ALREADY") == 2, r.stdout)
    for k, p in w.items(): check("%s backup is the S299 bytes" % k, md5(p + ".bak_S301_" + PINS[k][:8]) == PINS[k])
    r = subprocess.run([sys.executable, "-m", "py_compile", w["stock_app"]], capture_output=True, text=True); check("stock_app.py compiles", r.returncode == 0, r.stderr)
    hub_p = os.path.join(HERE, "stock_hub.html")
    for name, path in (("amir", w["stock_amir"]), ("hub", hub_p)):
        sc = re.findall(r"<script[^>]*>(.*?)</script>", io.open(path, encoding="utf-8").read(), re.S); js = os.path.join(tmp, name + ".js"); io.open(js, "w", encoding="utf-8").write(sc[-1])
        r = subprocess.run(["node", "--check", js], capture_output=True, text=True); check("%s script parses" % name, r.returncode == 0, r.stderr[:300])
    hub = io.open(hub_p, encoding="utf-8").read(); am = io.open(w["stock_amir"], encoding="utf-8").read(); app = io.open(w["stock_app"], encoding="utf-8").read()
    check("hub step 6: not yet on a voucher · entered · the button to Amir's board", "Not yet on a voucher" in hub and "#vouchers\">The vouchers" in hub and "Marg vouchers entered" in hub)
    check("hub keeps S299's swap wording and the answers button", "STOCK ISSUE on the short item" in hub and "by Amir or by you" in hub)
    check("Amir's board: make, entered, undo, the Excel, the S299 lists kept underneath", 'id="vmake"' in am and "data-ent=" in am and "data-unent=" in am and "vouchers (Excel)" in am and "Swaps confirmed — stock issue and stock receive" in am and "<details" in am)
    check("routes present", all(s in app for s in ('"/api/pad/vouchers/<int:cid>/make"', '"/api/pad/vouchers/<int:cid>/entered"', '"/api/pad/vouchers/<int:cid>/<int:rno>.xlsx"')))
    print("predicted: stock_app.py %s · stock_amir.html %s · stock_hub.html %s" % (h["stock_app"], h["stock_amir"], md5(hub_p)))
    sys.path.insert(0, tmp)
    S = load(w["stock_app"], "sa301")
    con = sqlite3.connect(":memory:"); con.executescript(S.MATCH_SCHEMA)
    ins = lambda s, o, q, a_: con.execute("INSERT INTO stock_match (count_id, short_item, over_item, qty, answer, by_user, at) VALUES (1,?,?,?,?, 'owner', '2026-09-17T21:00:00')", (s, o, q, a_))  # noqa: E731
    ins("ANKLE BINDER BAMBOO L", "ANKLE BINDER L TYNOR", 2, "YES"); ins("L S BELT CONT GRAY UNISON M", "L S BELT CONT GRAY UNISON L", 1, "YES")
    ins("KNEE SUPPORT HINGED XXL UNISO", "KNEE SUPPORT HINGED M", 1, "NO"); ins("MOTIF FORTE", "ASTOFEN P", 5, "NO"); con.commit()
    d, by = build(S, con)
    pend = S._voucher_pending(con, d); pk = [(p["kind"], p["item"], p["marg_from"], p["change"], p["marg_to"]) for p in pend]
    print("   waiting:", pk)
    check("confirmed swaps wait as voucher lines at once (a part swap too)", pk == [("ISSUE", "ANKLE BINDER BAMBOO L", 3, -2, 1), ("ISSUE", "L S BELT CONT GRAY UNISON M", 2, -1, 1),
                                                                               ("RECEIVE", "ANKLE BINDER L TYNOR", 0, 2, 2), ("RECEIVE", "L S BELT CONT GRAY UNISON L", 2, 1, 3)], pk)
    check("orthotics carry their family", all(p["section"] == "Orthotics" for p in pend) and pend[0]["family"] == "Ankle binder", [(p["section"], p["family"]) for p in pend])
    check("value at the count day's rate", pend[0]["rate_p"] == 27520 and pend[0]["value_p"] == -55040, (pend[0]["rate_p"], pend[0]["value_p"]))
    check("a No and an undecided line are not waiting", not any(p["item"] in ("KNEE SUPPORT HINGED XXL UNISO", "MOTIF FORTE", "RIB BELT S TYNOR", "ROSIKA FORTE") for p in pend))
    rno, n = S._voucher_make(con, d, {"user": "amir"})
    st = S._voucher_state(con, d)
    check("round 1 made: 4 lines, ISSUE and RECEIVE apart, one voucher each", rno == 1 and n == 4 and [(b["kind"], b["batch_no"], b["n"]) for b in st["rounds"][0]["batches"]] == [("ISSUE", 1, 2), ("RECEIVE", 1, 2)], st["rounds"])
    check("nothing waiting after the round", st["pending"] == [] and st["batches_total"] == 2 and st["batches_entered"] == 0)
    check("a second press finds nothing", S._voucher_make(con, d, {"user": "amir"}) == (None, 0))
    d, by = build(S, con, {"ANKLE BINDER BAMBOO L": dict(action="WRITE_OFF", label="written off", at_text="x")})
    pend = S._voucher_pending(con, d); pk = [(p["kind"], p["item"], p["marg_from"], p["change"], p["marg_to"], p["reason"]) for p in pend]
    check("a remainder decided later waits from Marg-after-swap (1) to the shelf (0)", pk == [("ISSUE", "ANKLE BINDER BAMBOO L", 1, -1, 0, "the rest of the line -- written off")], pk)
    con.execute("INSERT INTO stock_match (count_id, short_item, over_item, qty, answer, by_user, at) VALUES (1,'L S BELT CONT GRAY UNISON M','L S BELT CONT GRAY UNISON L',1,'OPEN','owner','x')"); con.commit()
    d, by = build(S, con, {"ANKLE BINDER BAMBOO L": dict(action="WRITE_OFF", label="written off", at_text="x")})
    pend = S._voucher_pending(con, d); pk = {p["item"]: (p["kind"], p["marg_from"], p["change"], p["marg_to"]) for p in pend}
    print("   after an answer taken back:", pk)
    check("an answer taken back after its round becomes a reversal", pk.get("L S BELT CONT GRAY UNISON M") == ("RECEIVE", 1, 1, 2) and pk.get("L S BELT CONT GRAY UNISON L") == ("ISSUE", 3, -1, 2), pk)
    check("and says so", all(p["reason"].startswith("corrects an earlier voucher") for p in pend if p["item"].startswith("L S BELT")), [p["reason"] for p in pend])
    con.execute("CREATE TABLE IF NOT EXISTS setting (key TEXT PRIMARY KEY, value TEXT)"); con.execute("INSERT INTO setting VALUES ('stock.voucher_lines','1')"); con.commit()
    rno2, n2 = S._voucher_make(con, d, {"user": "amir"})
    st = S._voucher_state(con, d); r2 = st["rounds"][1]
    check("round 2 with the setting at one line a voucher: 2 ISSUE vouchers, 1 RECEIVE", rno2 == 2 and n2 == 3 and [(b["kind"], b["batch_no"], b["batches_n"]) for b in r2["batches"]] == [("ISSUE", 1, 2), ("ISSUE", 2, 2), ("RECEIVE", 1, 1)], r2["batches"])
    moved = {}
    for r in con.execute("SELECT item, change, marg_from, marg_to FROM stock_voucher_line ORDER BY round_no, id"):
        m0 = by[r[0]]["marg"] + moved.get(r[0], 0)
        check("chain: %s line starts where the last ended" % r[0], r[2] == m0 and r[3] == m0 + r[1], (r, m0)); moved[r[0]] = moved.get(r[0], 0) + r[1]
    check("keyed in order, each item ends where its lines now say", moved == {"ANKLE BINDER BAMBOO L": -3, "ANKLE BINDER L TYNOR": 2, "L S BELT CONT GRAY UNISON M": 0, "L S BELT CONT GRAY UNISON L": 0}, moved)
    check("nothing waiting after round 2", S._voucher_pending(con, d) == [])
    S._voucher_ensure(con)
    con.execute("INSERT INTO stock_voucher_entered (count_id, round_no, kind, batch_no, marg_voucher_no, entered_on, by_user, at) VALUES (1,1,'ISSUE',1,'SI-12','17-09-2026','amir','2026-09-17T21:30:00')"); con.commit()
    st = S._voucher_state(con, d); b = st["rounds"][0]["batches"][0]
    check("a voucher recorded as entered", b["entered"] and b["entered"]["no"] == "SI-12" and st["batches_entered"] == 1 and not st["rounds"][0]["closed"], b["entered"])
    con.execute("INSERT INTO stock_voucher_entered (count_id, round_no, kind, batch_no, marg_voucher_no, entered_on, by_user, at) VALUES (1,1,'ISSUE',1,'',NULL,'amir','2026-09-17T21:31:00')"); con.commit()
    check("and taken back", S._voucher_state(con, d)["batches_entered"] == 0)
    for kk, bb in (("ISSUE", 1), ("RECEIVE", 1)):
        con.execute("INSERT INTO stock_voucher_entered (count_id, round_no, kind, batch_no, marg_voucher_no, entered_on, by_user, at) VALUES (1,1,?,?,'V9','17-09-2026','amir','2026-09-17T21:32:00')", (kk, bb))
    con.commit()
    check("a round with every voucher entered is closed", S._voucher_state(con, d)["rounds"][0]["closed"])
    data = S._voucher_workbook(con, d, 1)
    z = zipfile.ZipFile(io.BytesIO(data)); wbx = z.read("xl/workbook.xml").decode()
    check("workbook: VOUCHER 1 - SHORT and VOUCHER 2 - EXCESS", "VOUCHER 1 - SHORT" in wbx and "VOUCHER 2 - EXCESS" in wbx, wbx[:300])
    s1 = z.read("xl/worksheets/sheet1.xml").decode()
    check("sheet 1: the batch header, the columns, the lines, the entered number", "MARG VOUCHER  1  of  1" in s1 and "CORRECTION" in s1 and "Qty in Marg" in s1 and "Counted 06-09-2026" in s1 and "ANKLE BINDER BAMBOO L" in s1 and "V9" in s1, s1[:400])
    check("no round 9", S._voucher_workbook(con, d, 9) is None)
    try:
        import openpyxl  # noqa: PLC0415
        wb = openpyxl.load_workbook(io.BytesIO(data)); check("the workbook opens (openpyxl)", wb.sheetnames == ["VOUCHER 1 - SHORT", "VOUCHER 2 - EXCESS"], wb.sheetnames)
    except ImportError:
        print("   (openpyxl not here -- the zip checks above stand)")
    print("\n%d failed" % len(FAILS)); return 1 if FAILS else 0
if __name__ == "__main__": sys.exit(main())
